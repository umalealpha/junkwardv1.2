"""bonu/legal_reports.py — the two reports the Claims Legal Office asked for.

The Claims Legal Office sent a working demo (18 Aug 2026) as the guide for how
the BONU Legal section should be structured. Two of its tabs are reports the
office generates from the day's captured work:

  * **Monthly Fee Note** — the BONU-portfolio matter billing for a month, the
    external-vs-internal comparison, and the estimated monthly performance bonus.
  * **Quarterly Savings & Bonus Report** — the contractual quarterly bonus
    (2% of the quarter's invoice saving once the trigger is met), the invoices
    reduced, the per-client in-house saving, and the value delivered.

Both are built HERE, server-side, from the same registers the screens read
(`legal_calc`), so there is one source of numbers and the Word/PDF copies agree
with the dashboard. The layout, wording and figure basis follow the demo
exactly — this file is the demo's two report renderers ported to Python, not a
new report.

Downloadable as Word (.docx, python-docx) and PDF (reportlab), the two formats
the office works in. No figure here posts to the ledger or pays anybody.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from . import legal_calc as calc
from django.utils import timezone

# The office and entity, from the demo's seed. Fixed here rather than typed on
# every report; the officer is the Claims Legal Officer whose two bonuses these
# reports measure.
OFFICER = 'Bernard Balikani Jr'
COMPANY = 'Alpha Direct Insurance Co. (Pty) Ltd'
TARIFF_CODE = 'ADI-HC-LEGAL-TARIFF-2026-002'

ZERO = Decimal('0')
_MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July',
           'August', 'September', 'October', 'November', 'December']


# --------------------------------------------------------------------------
# Formatting — matches the demo: money as "P1,234.56" / "-P1,234.56", a
# percentage to one decimal place, month as "July 2026", quarter as "Q3 2026".
# --------------------------------------------------------------------------

def money(n) -> str:
    v = Decimal(n or 0)
    sign = '-' if v < 0 else ''
    return f'{sign}P{abs(v):,.2f}'


def pct(n) -> str:
    if n is None:
        return '—'
    return f'{Decimal(n) * 100:.1f}%'


def month_label(mkey: str) -> str:
    y, m = mkey.split('-')
    return f'{_MONTHS[int(m) - 1]} {y}'


def quarter_label(qkey: str) -> str:
    y, q = qkey.split('-Q')
    return f'Q{q} {y}'


def _today_label() -> str:
    return timezone.localdate().strftime('%d/%m/%Y')


# --------------------------------------------------------------------------
# Report models — the figures, computed once from the registers.
# These mirror the demo's computeFeeNoteModel / computeQuarterlyModel exactly,
# including which total uses all lines and which uses mapped lines only.
# --------------------------------------------------------------------------

def feenote_model(month: str, fees, savings, mappings, entered_bonus, settings) -> dict:
    index = calc.mapping_index(mappings)
    lines = sorted([f for f in fees if f.date and calc.month_key(f.date) == month],
                   key=lambda f: (f.date or date.min), reverse=True)

    matter_rows = []
    diff_rows = []
    matter_total = ZERO
    ext_total = ZERO
    for f in lines:
        internal = calc.internal_amount(f)
        external = calc.external_equivalent(f, index)
        m = index.get((f.description or '').strip().lower())
        matter_total += internal
        if external is not None:
            ext_total += external
        matter_rows.append({
            'date': f.date.isoformat() if f.date else '',
            'client': f.client, 'description': f.description, 'unit': f.unit,
            'rate': f.rate, 'qty': f.qty, 'amount': internal,
        })
        diff_rows.append({
            'description': f.description,
            'basis': 'Disbursement (at cost)' if (m and m.is_disbursement) else 'Legal task (flat panel rate)',
            'internal': internal,
            'external': external,
            'saving': None if external is None else external - internal,
        })

    # In-house saving on the demo's basis: mapped external cost less ALL matter
    # billing for the month (monthlyExternalTotal - monthlyFeeTotal).
    in_house_saving = ext_total - matter_total
    pct_saved = (in_house_saving / ext_total) if ext_total else None
    bonus = calc.monthly_bonus(entered_bonus, settings)

    return {
        'kind': 'feenote', 'month': month,
        'title': COMPANY,
        'subtitle': f'Claims Legal Officer — Compliance · Monthly Fee Note '
                    f'(Internal Legal Services Tariff {TARIFF_CODE})',
        'officer': OFFICER, 'month_label': month_label(month), 'generated': _today_label(),
        'matter_rows': matter_rows, 'matter_total': matter_total,
        'diff_rows': diff_rows, 'ext_total': ext_total,
        'in_house_saving': in_house_saving, 'pct_saved': pct_saved,
        'bonus': bonus, 'bonus_cap': settings.monthly_bonus_cap,
        'external_rate': settings.external_hourly_rate,
        'filename': f'Monthly_Fee_Note_{month_label(month).replace(" ", "_")}',
    }


def quarterly_model(quarter: str, fees, savings, advisory, mappings, settings) -> dict:
    index = calc.mapping_index(mappings)
    months = calc.quarter_months(quarter)

    month_savings = []
    for mk in months:
        s = calc.savings_total([r for r in savings if r.date_reviewed
                                and calc.month_key(r.date_reviewed) == mk])
        month_savings.append({'month': mk, 'saving': s})
    q_total = sum((r['saving'] for r in month_savings), ZERO)
    q = calc.quarterly_bonus(q_total, settings)

    q_fees = [f for f in fees if f.date and calc.month_key(f.date) in months]
    # Illustrative in-house saving: per mapped line (external - internal),
    # unmapped lines skipped — the demo's quarterlyInHouseSaving.
    illustrative = ZERO
    for f in q_fees:
        ext = calc.external_equivalent(f, index)
        if ext is not None:
            illustrative += ext - calc.internal_amount(f)

    q_adv = [a for a in advisory if a.date and calc.month_key(a.date) in months]
    adv_hours = sum((a.hours or ZERO for a in q_adv), ZERO)
    adv_value = calc.advisory_value(adv_hours, settings)

    invoices = sorted([r for r in savings if r.date_reviewed
                       and calc.month_key(r.date_reviewed) in months],
                      key=lambda r: (r.date_reviewed or date.min), reverse=True)

    by_client: dict[str, dict] = {}
    for f in q_fees:
        ext = calc.external_equivalent(f, index)
        if ext is None:
            continue
        e = by_client.setdefault(f.client, {'internal': ZERO, 'external': ZERO})
        e['internal'] += calc.internal_amount(f)
        e['external'] += ext
    client_rows = [{'client': c, 'internal': v['internal'], 'external': v['external'],
                    'saving': v['external'] - v['internal']}
                   for c, v in sorted(by_client.items())]

    internal_fee_total = sum((calc.internal_amount(f) for f in q_fees), ZERO)
    external_total_q = ZERO
    for f in q_fees:
        ext = calc.external_equivalent(f, index)
        if ext is not None:
            external_total_q += ext
    invoices_original_total = sum((r.original_amount or ZERO for r in invoices), ZERO)

    return {
        'kind': 'quarterly', 'quarter': quarter,
        'officer': OFFICER, 'quarter_label': quarter_label(quarter),
        'months_label': ', '.join(month_label(m) for m in months),
        'generated': _today_label(),
        'months': months, 'month_savings': month_savings,
        'q_total': q_total, 'q_met': q['met'], 'q_bonus': q['bonus'],
        'quarterly_threshold': settings.quarterly_threshold,
        'quarterly_bonus_pct': settings.quarterly_bonus_pct,
        'illustrative': illustrative, 'adv_value': adv_value, 'adv_hours': adv_hours,
        'invoices': invoices, 'client_rows': client_rows,
        'internal_fee_total': internal_fee_total, 'external_total_q': external_total_q,
        'invoices_original_total': invoices_original_total,
        'total_delivered': q_total + illustrative + adv_value,
        'filename': f'Quarterly_Savings_Bonus_Report_{quarter_label(quarter).replace(" ", "_")}',
    }


# --------------------------------------------------------------------------
# Word (.docx) — python-docx, the pattern used by underwriting/quote_export.py.
# --------------------------------------------------------------------------

_NAVY = (0x0D, 0x1B, 0x2A)


def _docx_doc():
    from docx import Document
    return Document()


def _docx_title(doc, text, size=15):
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(*_NAVY)
    return p


def _docx_meta(doc, text):
    from docx.shared import Pt
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    return p


def _docx_heading(doc, text):
    from docx.shared import Pt, RGBColor
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(*_NAVY)
    return p


def _docx_table(doc, head, rows, aligns, foot=None, empty=None, bold_rows=None):
    """`bold_rows` = indices into `rows` to render bold (the demo's group headers
    inside a table, e.g. "Internal costs")."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    bold_rows = bold_rows or set()
    align_map = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT}
    table = doc.add_table(rows=1, cols=len(head))
    table.style = 'Table Grid'
    for i, h in enumerate(head):
        cell = table.rows[0].cells[i]
        cell.text = str(h)
        for para in cell.paragraphs:
            para.alignment = align_map.get(aligns[i], WD_ALIGN_PARAGRAPH.LEFT)
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(8.5)
    body = rows if rows else ([[empty] + [''] * (len(head) - 1)] if empty else [])
    for ri, r in enumerate(body):
        cells = table.add_row().cells
        for c, val in enumerate(r):
            cells[c].text = str(val)
            for para in cells[c].paragraphs:
                para.alignment = align_map.get(aligns[c], WD_ALIGN_PARAGRAPH.LEFT)
                for run in para.runs:
                    run.font.size = Pt(8.5)
                    if ri in bold_rows:
                        run.bold = True
    if foot:
        cells = table.add_row().cells
        for c, val in enumerate(foot):
            cells[c].text = str(val)
            for para in cells[c].paragraphs:
                para.alignment = align_map.get(aligns[c], WD_ALIGN_PARAGRAPH.LEFT)
                for run in para.runs:
                    run.bold = True
                    run.font.size = Pt(8.5)
    return table


def feenote_docx(m: dict) -> bytes:
    doc = _docx_doc()
    _docx_title(doc, m['title'])
    _docx_meta(doc, m['subtitle'])
    _docx_meta(doc, f"Officer: {m['officer']}     Billing month: {m['month_label']}"
                    f"     Generated: {m['generated']}")

    _docx_heading(doc, 'Matter billing')
    _docx_table(
        doc,
        ['Date', 'Client / BONU member', 'Service', 'Unit', 'Rate', 'Qty/Hrs', 'Amount'],
        [[r['date'], r['client'], r['description'], r['unit'], money(r['rate']),
          f"{r['qty']}", money(r['amount'])] for r in m['matter_rows']],
        ['left', 'left', 'left', 'left', 'right', 'right', 'right'],
        foot=['Section 1 total — total matter billing', '', '', '', '', '', money(m['matter_total'])],
        empty='No matter lines logged this month.')

    _docx_heading(doc, 'External Finances vs. Internal Costs & Savings')
    _docx_meta(doc, f"Internal costs are what Alpha Direct actually paid, handling these BONU "
                    f"matters in-house. External finances are the illustrative cost exposure had "
                    f"the same matters gone to an external panel attorney, at a flat panel rate of "
                    f"{money(calc.FIXED_EXTERNAL_TASK_RATE)} per task (disbursement items — email, "
                    f"telephone, couriers — are passed through at cost, so they carry no saving).")
    _docx_table(
        doc,
        ['Service', 'Basis', 'Internal cost', 'External finance equiv.', 'Saving'],
        [[r['description'], r['basis'], money(r['internal']),
          '—' if r['external'] is None else money(r['external']),
          '—' if r['saving'] is None else money(r['saving'])] for r in m['diff_rows']],
        ['left', 'left', 'right', 'right', 'right'],
        foot=['Totals', '', money(m['matter_total']), money(m['ext_total']), money(m['in_house_saving'])],
        empty='No matter lines logged this month.')
    _docx_meta(doc, f"Percentage saved vs. external finance exposure this month: {pct(m['pct_saved'])} "
                    f"(illustrative — does not affect the contractual bonuses below).")

    _docx_heading(doc, 'Bonus estimate')
    from docx.shared import Pt
    p = doc.add_paragraph()
    r = p.add_run(f"Estimated Monthly Performance Bonus (capped {money(m['bonus_cap'])}): {money(m['bonus'])}")
    r.bold = True
    r.font.size = Pt(9.5)
    _docx_meta(doc, 'Indicative only, per Tariff §7.1 threshold table. Final bonus is confirmed by '
                    'CFO approval. Invoice review savings (quarterly 2% bonus) are logged separately.')
    doc.add_paragraph()
    _docx_meta(doc, f"Submitted by: {m['officer']}     Signature: __________________     "
                    f"Date: __________________")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def quarterly_docx(m: dict) -> bytes:
    doc = _docx_doc()
    _docx_title(doc, 'Quarterly Savings & Bonus Report')
    _docx_meta(doc, f"{COMPANY} — Claims Legal Office · Officer: {m['officer']}")
    _docx_meta(doc, f"Quarter: {m['quarter_label']} ({m['months_label']})     Generated: {m['generated']}")

    _docx_heading(doc, 'Contractual quarterly bonus (Sec. 2 — Invoice Savings Log)')
    _docx_table(
        doc, ['Month', 'Saving (BWP)'],
        [[month_label(r['month']), money(r['saving'])] for r in m['month_savings']],
        ['left', 'right'],
        foot=['Quarterly total saving', money(m['q_total'])])
    _docx_meta(doc, f"Minimum trigger ({money(m['quarterly_threshold'])}) met: "
                    f"{'YES' if m['q_met'] else 'NO'}")
    _docx_meta(doc, f"Quarterly bonus payable (2% of total): {money(m['q_bonus'])}")

    _docx_heading(doc, 'Invoices reviewed this quarter')
    _docx_table(
        doc, ['Date', 'Ref', 'External attorney', 'Original', 'Agreed', 'Saving'],
        [[r.date_reviewed.isoformat() if r.date_reviewed else '', r.invoice_ref,
          r.external_attorney, money(r.original_amount), money(r.agreed_amount),
          money(r.saving)] for r in m['invoices']],
        ['left', 'left', 'left', 'right', 'right', 'right'],
        empty='No invoices logged this quarter.')

    _docx_heading(doc, 'Illustrative value — in-house handling vs external attorney')
    _docx_meta(doc, f"Total saving from handling BONU matters in-house this quarter: "
                    f"{money(m['illustrative'])} (does not count toward the bonus above)")
    _docx_table(
        doc, ['Client / BONU member', 'Internal', 'External equiv.', 'Saving'],
        [[r['client'], money(r['internal']), money(r['external']), money(r['saving'])]
         for r in m['client_rows']],
        ['left', 'right', 'right', 'right'],
        empty='No mapped fee note lines this quarter.')

    _docx_heading(doc, 'Internal ADI advisory (illustrative)')
    _docx_meta(doc, f"Hours logged: {m['adv_hours']:.2f}     Estimated value: {money(m['adv_value'])}")

    _docx_heading(doc, 'External Finances vs. Internal Costs & Savings — Quarter Summary')
    _docx_meta(doc, 'A clear split between what the business actually paid internally, and the '
                    'illustrative cost had these matters been referred externally at the Standard '
                    'Rate (disbursement items keep their own separate rates and are unaffected by '
                    'the Standard Rate).')
    _docx_table(
        doc, ['', 'Amount (BWP)'],
        [['Internal costs', ''],
         ['Internal matter billing, in-house handling (Sec. 3 basis)', money(m['internal_fee_total'])],
         ['Internal ADI advisory hours logged', f"{m['adv_hours']:.2f} hrs"],
         ['External finances (illustrative exposure)', ''],
         ['External equivalent, had matters gone to panel attorney', money(m['external_total_q'])],
         ['External-rate value of ADI advisory hours', money(m['adv_value'])],
         ['Invoices reviewed — original external amount billed', money(m['invoices_original_total'])],
         ['Net savings', ''],
         ['Contractual invoice review saving (2% bonus basis)', money(m['q_total'])],
         ['In-house handling saving vs. external (illustrative)', money(m['illustrative'])]],
        ['left', 'right'], bold_rows={0, 3, 7})

    _docx_heading(doc, 'Total value delivered to the business this quarter')
    from docx.shared import Pt
    p = doc.add_paragraph()
    r = p.add_run(f"{money(m['total_delivered'])} (contractual invoice savings + in-house handling "
                  f"saving + ADI advisory value)")
    r.bold = True
    r.font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF — reportlab platypus, so the layout survives a page break.
# --------------------------------------------------------------------------

def _pdf_styles():
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.colors import HexColor
    ss = getSampleStyleSheet()
    navy = HexColor('#0D1B2A')
    title = ParagraphStyle('t', parent=ss['Title'], fontSize=15, textColor=navy,
                           alignment=0, spaceAfter=6)
    head = ParagraphStyle('h', parent=ss['Heading3'], fontSize=11, textColor=navy,
                          spaceBefore=12, spaceAfter=4)
    meta = ParagraphStyle('m', parent=ss['Normal'], fontSize=8.5, leading=12, spaceAfter=3)
    return title, head, meta


def _pdf_table(head, rows, aligns, foot=None, empty=None, bold_rows=None):
    """`bold_rows` = indices into `rows` to render bold (the demo's in-table group
    headers, e.g. "Internal costs")."""
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.colors import HexColor
    bold_rows = bold_rows or set()
    data = [head]
    if rows:
        data += rows
    elif empty:
        data.append([empty] + [''] * (len(head) - 1))
    if foot:
        data.append(foot)
    t = Table(data, hAlign='LEFT', repeatRows=1)
    style = [
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('TEXTCOLOR', (0, 0), (-1, 0), HexColor('#0D1B2A')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('LINEBELOW', (0, 0), (-1, 0), 0.6, colors.HexColor('#999999')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#DDDDDD')),
    ]
    for i, a in enumerate(aligns):
        style.append(('ALIGN', (i, 0), (i, -1), a.upper()))
    for ri in bold_rows:
        style.append(('FONTNAME', (0, ri + 1), (-1, ri + 1), 'Helvetica-Bold'))  # +1 for the head row
    if foot:
        style.append(('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'))
        style.append(('LINEABOVE', (0, -1), (-1, -1), 1, colors.HexColor('#111111')))
    t.setStyle(TableStyle(style))
    return t


def _pdf_bytes(flow) -> bytes:
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm)
    doc.build(flow)
    return buf.getvalue()


def feenote_pdf(m: dict) -> bytes:
    from reportlab.platypus import Paragraph, Spacer
    from reportlab.lib.units import mm
    title, head, meta = _pdf_styles()
    flow = [
        Paragraph(m['title'], title),
        Paragraph(m['subtitle'], meta),
        Paragraph(f"Officer: {m['officer']} &nbsp;&nbsp; Billing month: {m['month_label']} "
                  f"&nbsp;&nbsp; Generated: {m['generated']}", meta),
        Paragraph('Matter billing', head),
        _pdf_table(
            ['Date', 'Client / BONU member', 'Service', 'Unit', 'Rate', 'Qty/Hrs', 'Amount'],
            [[r['date'], r['client'], r['description'], r['unit'], money(r['rate']),
              f"{r['qty']}", money(r['amount'])] for r in m['matter_rows']],
            ['left', 'left', 'left', 'left', 'right', 'right', 'right'],
            foot=['Section 1 total — total matter billing', '', '', '', '', '', money(m['matter_total'])],
            empty='No matter lines logged this month.'),
        Paragraph('External Finances vs. Internal Costs &amp; Savings', head),
        Paragraph(f"Internal costs are what Alpha Direct actually paid, handling these BONU matters "
                  f"in-house. External finances are the illustrative cost exposure had the same "
                  f"matters gone to an external panel attorney, at a flat panel rate of "
                  f"{money(calc.FIXED_EXTERNAL_TASK_RATE)} per task (disbursement items — email, "
                  f"telephone, couriers — are passed through at cost, so they carry no saving).", meta),
        _pdf_table(
            ['Service', 'Basis', 'Internal cost', 'External finance equiv.', 'Saving'],
            [[r['description'], r['basis'], money(r['internal']),
              '—' if r['external'] is None else money(r['external']),
              '—' if r['saving'] is None else money(r['saving'])] for r in m['diff_rows']],
            ['left', 'left', 'right', 'right', 'right'],
            foot=['Totals', '', money(m['matter_total']), money(m['ext_total']), money(m['in_house_saving'])],
            empty='No matter lines logged this month.'),
        Paragraph(f"Percentage saved vs. external finance exposure this month: <b>{pct(m['pct_saved'])}</b> "
                  f"(illustrative — does not affect the contractual bonuses below).", meta),
        Paragraph('Bonus estimate', head),
        Paragraph(f"<b>Estimated Monthly Performance Bonus (capped {money(m['bonus_cap'])}): "
                  f"{money(m['bonus'])}</b>", meta),
        Paragraph('Indicative only, per Tariff §7.1 threshold table. Final bonus is confirmed by CFO '
                  'approval. Invoice review savings (quarterly 2% bonus) are logged separately.', meta),
        Spacer(1, 8 * mm),
        Paragraph(f"Submitted by: {m['officer']} &nbsp;&nbsp;&nbsp; Signature: __________________ "
                  f"&nbsp;&nbsp;&nbsp; Date: __________________", meta),
    ]
    return _pdf_bytes(flow)


def quarterly_pdf(m: dict) -> bytes:
    from reportlab.platypus import Paragraph
    title, head, meta = _pdf_styles()
    flow = [
        Paragraph('Quarterly Savings &amp; Bonus Report', title),
        Paragraph(f"{COMPANY} — Claims Legal Office · Officer: {m['officer']}", meta),
        Paragraph(f"Quarter: {m['quarter_label']} ({m['months_label']}) &nbsp;&nbsp; "
                  f"Generated: {m['generated']}", meta),
        Paragraph('Contractual quarterly bonus (Sec. 2 — Invoice Savings Log)', head),
        _pdf_table(
            ['Month', 'Saving (BWP)'],
            [[month_label(r['month']), money(r['saving'])] for r in m['month_savings']],
            ['left', 'right'],
            foot=['Quarterly total saving', money(m['q_total'])]),
        Paragraph(f"Minimum trigger ({money(m['quarterly_threshold'])}) met: "
                  f"<b>{'YES' if m['q_met'] else 'NO'}</b><br/>"
                  f"Quarterly bonus payable (2% of total): <b>{money(m['q_bonus'])}</b>", meta),
        Paragraph('Invoices reviewed this quarter', head),
        _pdf_table(
            ['Date', 'Ref', 'External attorney', 'Original', 'Agreed', 'Saving'],
            [[r.date_reviewed.isoformat() if r.date_reviewed else '', r.invoice_ref,
              r.external_attorney, money(r.original_amount), money(r.agreed_amount),
              money(r.saving)] for r in m['invoices']],
            ['left', 'left', 'left', 'right', 'right', 'right'],
            empty='No invoices logged this quarter.'),
        Paragraph('Illustrative value — in-house handling vs external attorney', head),
        Paragraph(f"Total saving from handling BONU matters in-house this quarter: "
                  f"<b>{money(m['illustrative'])}</b> (does not count toward the bonus above)", meta),
        _pdf_table(
            ['Client / BONU member', 'Internal', 'External equiv.', 'Saving'],
            [[r['client'], money(r['internal']), money(r['external']), money(r['saving'])]
             for r in m['client_rows']],
            ['left', 'right', 'right', 'right'],
            empty='No mapped fee note lines this quarter.'),
        Paragraph('Internal ADI advisory (illustrative)', head),
        Paragraph(f"Hours logged: <b>{m['adv_hours']:.2f}</b> &nbsp;&nbsp; "
                  f"Estimated value: <b>{money(m['adv_value'])}</b>", meta),
        Paragraph('External Finances vs. Internal Costs &amp; Savings — Quarter Summary', head),
        _pdf_table(
            ['', 'Amount (BWP)'],
            [['Internal costs', ''],
             ['Internal matter billing, in-house handling (Sec. 3 basis)', money(m['internal_fee_total'])],
             ['Internal ADI advisory hours logged', f"{m['adv_hours']:.2f} hrs"],
             ['External finances (illustrative exposure)', ''],
             ['External equivalent, had matters gone to panel attorney', money(m['external_total_q'])],
             ['External-rate value of ADI advisory hours', money(m['adv_value'])],
             ['Invoices reviewed — original external amount billed', money(m['invoices_original_total'])],
             ['Net savings', ''],
             ['Contractual invoice review saving (2% bonus basis)', money(m['q_total'])],
             ['In-house handling saving vs. external (illustrative)', money(m['illustrative'])]],
            ['left', 'right'], bold_rows={0, 3, 7}),
        Paragraph('Total value delivered to the business this quarter', head),
        Paragraph(f"<b>{money(m['total_delivered'])}</b> (contractual invoice savings + in-house "
                  f"handling saving + ADI advisory value)", meta),
    ]
    return _pdf_bytes(flow)
