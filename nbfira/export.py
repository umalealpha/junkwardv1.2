"""
nbfira/export.py — XLSX export of NBFIRA returns.

CFO directive (NBFIRA_MODULE_BLUEPRINT.md Phase 4):
  Filename: ADIC_NBFIRA_<TYPE>_<PERIOD>_v<N>.xlsx
  Layout must match the NBFIRA prescribed templates exactly. Phase 4
  delivers a readable multi-sheet export — one tab per schedule with
  line code / label / value columns. Full pixel-perfect prescribed
  layout matches lands in a follow-up once Finance maps every line.

Reuses openpyxl (already in deps).
"""

from __future__ import annotations

import io
import hashlib
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import NBFIRAReturn


NAVY  = 'FF0D1B2A'
ORANGE= 'FFF07F00'
WHITE = 'FFFFFFFF'
LIGHT = 'FFF8F9FA'


def _header(ws, title: str):
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=4)
    c = ws.cell(1, 1, value=title)
    c.font = Font(bold=True, color=WHITE, size=13)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 22


def _column_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(ord('A') + i - 1)].width = w


def _table_header(ws, row, columns):
    thin = Side(border_style='thin', color='FF888888')
    for i, h in enumerate(columns, start=1):
        c = ws.cell(row, i, value=h)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = PatternFill('solid', fgColor=ORANGE)
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.row_dimensions[row].height = 28


def _line_row(ws, row, line):
    ws.cell(row, 1, value=line.line_code).font = Font(name='Courier New', size=9)
    ws.cell(row, 2, value=line.section)
    ws.cell(row, 3, value=line.label)
    ws.cell(row, 4, value=float(line.value)).number_format = '#,##0.0000'


def build_workbook(ret: NBFIRAReturn) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    # Cover sheet
    cover = wb.create_sheet('Cover', 0)
    _header(cover, f'NBFIRA Return — {ret.type.upper()} {ret.period_label}')
    _column_widths(cover, [22, 50])
    cover.cell(3, 1, value='Company').font = Font(bold=True)
    cover.cell(3, 2, value=(ret.company.code if ret.company_id else 'ADIC'))
    cover.cell(4, 1, value='Period start').font = Font(bold=True)
    cover.cell(4, 2, value=str(ret.period_start))
    cover.cell(5, 1, value='Period end').font = Font(bold=True)
    cover.cell(5, 2, value=str(ret.period_end))
    cover.cell(6, 1, value='Status').font = Font(bold=True)
    cover.cell(6, 2, value=ret.status.upper())
    cover.cell(7, 1, value='Initiated by').font = Font(bold=True)
    cover.cell(7, 2, value=(ret.initiated_by.get_full_name() if ret.initiated_by_id else '—'))
    cover.cell(8, 1, value='Reviewed by').font = Font(bold=True)
    cover.cell(8, 2, value=(ret.reviewed_by.get_full_name() if ret.reviewed_by_id else '—'))
    cover.cell(9, 1, value='Approved by').font = Font(bold=True)
    cover.cell(9, 2, value=(ret.approved_by.get_full_name() if ret.approved_by_id else '—'))
    cover.cell(10, 1, value='Locked by').font = Font(bold=True)
    cover.cell(10, 2, value=(ret.locked_by.get_full_name() if ret.locked_by_id else '—'))
    cover.cell(11, 1, value='Submitted by').font = Font(bold=True)
    cover.cell(11, 2, value=(ret.submitted_by.get_full_name() if ret.submitted_by_id else '—'))
    cover.cell(13, 1, value='Generated').font = Font(bold=True)
    cover.cell(13, 2, value=datetime.utcnow().isoformat() + 'Z')

    # One sheet per schedule
    by_schedule: dict[str, list] = {}
    for line in ret.lines.order_by('schedule', 'sort_order', 'line_code'):
        by_schedule.setdefault(line.schedule, []).append(line)

    for schedule, lines in by_schedule.items():
        # openpyxl sheet name limit 31 chars; keep schedule code as-is
        sn = schedule.replace('/', '-')[:31]
        ws = wb.create_sheet(sn)
        _header(ws, f'Schedule {schedule}')
        _column_widths(ws, [14, 26, 60, 18])
        _table_header(ws, 3, ['Line code', 'Section', 'Label', 'Value (P\'000)'])
        for i, ln in enumerate(lines, start=4):
            _line_row(ws, i, ln)
            if i % 2 == 0:
                for c in range(1, 5):
                    ws.cell(i, c).fill = PatternFill('solid', fgColor=LIGHT)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def filename_for(ret: NBFIRAReturn, version: int = 1) -> str:
    t = 'QUARTERLY' if ret.type == 'quarterly' else 'ANNUAL'
    return f'ADIC_NBFIRA_{t}_{ret.period_label}_v{version}.xlsx'
