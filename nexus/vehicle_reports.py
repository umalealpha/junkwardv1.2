"""nexus/vehicle_reports.py — Vehicle Register analytics + branded export.

Read-only aggregations over the trip register that surface misuse:
  * trips per driver,
  * purpose breakdown,
  * km per vehicle + utilisation,
  * flagged / abuse cases (damage, "Other" purpose, overdue).

    GET  vehicle-register/reports/          JSON aggregations (?from=&to=)
    GET  vehicle-register/reports/export/   Alpha Direct-branded .xlsx

Brand (per Finance standard): Book Antiqua, Dark Navy #0D1B2A, Orange #F4A623.
"""
from __future__ import annotations

import io
import logging
from collections import defaultdict
from pathlib import Path

from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .models import FleetVehicle, TripPurpose, TripState, VehicleStatus, VehicleTrip

logger = logging.getLogger(__name__)

NAVY = '0D1B2A'
ORANGE = 'F4A623'
BRAND_FONT = 'Book Antiqua'

# Logo artwork — same candidate order as payroll/pdf.py and procurement/pdf.py.
# logo-clean.png is the no-reflection/no-shadow cut; the shadowed logo.png was
# retired 2026-07-09 because the shadow renders as a grey box on a white sheet.
_REPO_ROOT = Path(__file__).resolve().parent.parent
_LOGO_CANDIDATES = [
    _REPO_ROOT / 'procurement' / 'pdf_assets' / 'logo-clean.png',
    _REPO_ROOT / 'alpha-direct-design-system' / 'assets' / 'logo-full-color.png',
    _REPO_ROOT / 'frontend' / 'public' / 'brand' / 'logo-full-color.png',
]


def _logo_path() -> str | None:
    for p in _LOGO_CANDIDATES:
        if p.exists():
            return str(p)
    return None


def _add_logo(ws, height_px: int = 34) -> None:
    """Anchor the Alpha Direct logo at A1, scaled to `height_px` at its true
    aspect ratio. The ratio MUST be read before the first assignment — setting
    .height first and then reading .width returns the already-mutated value and
    the logo comes out stretched.

    A missing/unreadable logo must NOT 500 a report the CFO is waiting on, but
    it must not vanish quietly either — it is logged at WARNING so a broken
    export shows up in the logs instead of just shipping unbranded.
    """
    path = _logo_path()
    if not path:
        logger.warning('vehicle-register export: no logo artwork found in %s',
                       [str(p) for p in _LOGO_CANDIDATES])
        return
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(path)
        ratio = (img.width / img.height) if img.height else 2.0
        img.height = height_px
        img.width = int(height_px * ratio)
        ws.add_image(img, 'A1')
    except Exception:   # noqa: BLE001 — never block the export on artwork
        logger.warning('vehicle-register export: could not embed logo %s on sheet %s',
                       path, getattr(ws, 'title', '?'), exc_info=True)


def _range(request):
    """Optional ?from=YYYY-MM-DD&to=YYYY-MM-DD (inclusive), applied to checkout."""
    d_from = parse_date(request.query_params.get('from', '') or '')
    d_to = parse_date(request.query_params.get('to', '') or '')
    return d_from, d_to


def _filtered(request):
    qs = VehicleTrip.objects.select_related('vehicle')
    d_from, d_to = _range(request)
    if d_from:
        qs = qs.filter(checkout_at__date__gte=d_from)
    if d_to:
        qs = qs.filter(checkout_at__date__lte=d_to)
    return qs


def _aggregate(request) -> dict:
    trips = list(_filtered(request))
    vehicles = list(FleetVehicle.objects.all())

    per_driver: dict[str, dict] = defaultdict(lambda: {'trips': 0, 'km': 0, 'flagged': 0})
    per_purpose: dict[str, int] = defaultdict(int)
    per_vehicle: dict[str, dict] = {}
    purpose_label = dict(TripPurpose.choices)

    for v in vehicles:
        per_vehicle[str(v.id)] = {
            'registration': v.registration,
            'make_model': ' '.join(x for x in [v.make, v.model] if x),
            'status': v.get_status_display(),
            'trips': 0, 'km': 0, 'odometer_km': v.odometer_km,
        }

    flagged = []
    for t in trips:
        km = t.distance_km or 0
        pd = per_driver[t.driver_name or '(unnamed)']
        pd['trips'] += 1
        pd['km'] += km
        if t.flagged:
            pd['flagged'] += 1
        per_purpose[t.purpose] += 1
        pv = per_vehicle.get(str(t.vehicle_id))
        if pv:
            pv['trips'] += 1
            pv['km'] += km
        if t.flagged or t.damage_on_return or t.is_overdue:
            reasons = []
            if t.damage_on_return:
                reasons.append('damage')
            if t.purpose == TripPurpose.OTHER:
                reasons.append('purpose: other')
            if t.is_overdue:
                reasons.append('overdue')
            if t.flag_reason and not reasons:
                reasons.append(t.flag_reason)
            flagged.append({
                'trip_id': str(t.id),
                'registration': t.vehicle.registration if t.vehicle_id else '',
                'driver_name': t.driver_name,
                'purpose_label': t.get_purpose_display(),
                'destination': t.destination,
                'checkout_at': t.checkout_at.isoformat() if t.checkout_at else None,
                'reasons': ', '.join(reasons) or 'flagged',
                'damage_notes': t.damage_notes,
            })

    out_now = sum(1 for v in vehicles if v.status == VehicleStatus.OUT)
    active = sum(1 for v in vehicles if v.status != VehicleStatus.RETIRED)
    utilisation = round(100 * out_now / active, 1) if active else 0.0

    return {
        'range': {'from': request.query_params.get('from'), 'to': request.query_params.get('to')},
        'totals': {
            'trips': len(trips),
            'km': sum((t.distance_km or 0) for t in trips),
            'flagged': len(flagged),
            'vehicles': len(vehicles),
            'utilisation_pct': utilisation,
        },
        'per_driver': sorted(
            [{'driver_name': k, **v} for k, v in per_driver.items()],
            key=lambda r: r['trips'], reverse=True),
        'per_purpose': sorted(
            [{'purpose': k, 'purpose_label': purpose_label.get(k, k), 'trips': n}
             for k, n in per_purpose.items()],
            key=lambda r: r['trips'], reverse=True),
        'per_vehicle': sorted(per_vehicle.values(), key=lambda r: r['km'], reverse=True),
        'flagged': flagged,
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def reports(request):
    return Response(_aggregate(request))


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def reports_export(request):
    """Alpha Direct-branded .xlsx of the register analytics."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = _aggregate(request)
    wb = openpyxl.Workbook()

    navy_fill = PatternFill('solid', fgColor=NAVY)
    orange_fill = PatternFill('solid', fgColor=ORANGE)
    hdr_font = Font(name=BRAND_FONT, bold=True, color='FFFFFF', size=11)
    title_font = Font(name=BRAND_FONT, bold=True, color=NAVY, size=15)
    body_font = Font(name=BRAND_FONT, size=10)

    def _sheet(ws, title, headers, rows):
        ws.sheet_view.showGridLines = False
        # Row 1 = the logo band; the title sits below it so the artwork never
        # overlaps the text (the logo is anchored at A1 and floats over cells).
        _add_logo(ws)
        ws.row_dimensions[1].height = 30
        ws['A2'] = 'Alpha Direct Insurance — Vehicle Register'
        ws['A2'].font = title_font
        rng = data['range']
        span = (f"{rng['from']} to {rng['to']}" if rng['from'] or rng['to'] else 'All dates')
        ws['A3'] = f'{title}  ·  {span}  ·  generated {timezone.localtime():%d %b %Y %H:%M}'
        ws['A3'].font = Font(name=BRAND_FONT, italic=True, color='6B7280', size=9)
        hr = 5
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=hr, column=c, value=h)
            cell.fill = navy_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal='left', vertical='center')
        for r, row in enumerate(rows, start=hr + 1):
            for c, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.font = body_font
        # Orange accent strip under the title.
        for c in range(1, len(headers) + 1):
            ws.cell(row=hr - 1, column=c).fill = orange_fill
        for c, h in enumerate(headers, start=1):
            width = max(len(str(h)), *([len(str(row[c - 1])) for row in rows] or [0])) + 3
            ws.column_dimensions[get_column_letter(c)].width = min(width, 48)

    ws1 = wb.active
    ws1.title = 'Summary'
    t = data['totals']
    _sheet(ws1, 'Summary', ['Metric', 'Value'], [
        ['Trips', t['trips']],
        ['Total km', t['km']],
        ['Flagged trips', t['flagged']],
        ['Vehicles', t['vehicles']],
        ['Fleet utilisation (out now)', f"{t['utilisation_pct']}%"],
    ])

    _sheet(wb.create_sheet('Trips per driver'),
           'Trips per driver', ['Driver', 'Trips', 'Km', 'Flagged'],
           [[r['driver_name'], r['trips'], r['km'], r['flagged']] for r in data['per_driver']])

    _sheet(wb.create_sheet('Purpose breakdown'),
           'Purpose breakdown', ['Purpose', 'Trips'],
           [[r['purpose_label'], r['trips']] for r in data['per_purpose']])

    _sheet(wb.create_sheet('Km per vehicle'),
           'Km per vehicle', ['Registration', 'Make / model', 'Status', 'Trips', 'Km', 'Odometer'],
           [[r['registration'], r['make_model'], r['status'], r['trips'], r['km'],
             r['odometer_km']] for r in data['per_vehicle']])

    _sheet(wb.create_sheet('Flagged'),
           'Flagged / review', ['Registration', 'Driver', 'Purpose', 'Destination',
                                 'Checked out', 'Reasons', 'Damage notes'],
           [[r['registration'], r['driver_name'], r['purpose_label'], r['destination'],
             (r['checkout_at'] or '')[:16].replace('T', ' '), r['reasons'], r['damage_notes']]
            for r in data['flagged']])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = timezone.localtime().strftime('%Y%m%d')
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="vehicle-register-{stamp}.xlsx"'
    return resp
