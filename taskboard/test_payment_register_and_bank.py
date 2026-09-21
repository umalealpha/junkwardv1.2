"""taskboard/test_payment_register_and_bank.py

Kago Tshutlhedi (Finance Manager) 2026-08-29, two asks:

  1. Authorisers cannot retrieve a paid/cleared request — "By company" showed 17
     but the table showed 2. A read-only PAYMENT REGISTER (every request, every
     status, every entity, paginated, exportable) fixes it without touching any
     approval control.
  2. Authorisers approve without seeing the beneficiary account. Surface the bank
     details + a changed-account acknowledgement on the authorisation path, and
     log every full-account view.

Every test here goes red if the corresponding feature is reverted.
Run in CI (needs a DB): manage.py test taskboard.test_payment_register_and_bank
"""
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework.test import APITestCase

from core.models import AuditLog, Company
from taskboard.models import PaymentRequest, PaymentRequestAttachment
from taskboard.test_helpers import window_always_open


@window_always_open
class PaymentRegisterTests(APITestCase):
    def setUp(self):
        self.cfo = User.objects.create_user('pganesharajah', email='pganesharajah@alphadirect.co.bw')
        self.kago = User.objects.create_user('ktshutlhedi', email='ktshutlhedi@alphadirect.co.bw')
        self.clerk = User.objects.create_user('lthebe', email='lthebe@alphadirect.co.bw')
        self.other = User.objects.create_user('mmpho', email='mmpho@alphadirect.co.bw')
        self.list_url = reverse('v1-payment-requests')

    def _mk(self, creator, payee, total, status, entity='Alpha Direct Insurance',
            acct='62011112222'):
        return PaymentRequest.objects.create(
            ref=f'PAY/T/{PaymentRequest.objects.count()+1:04d}',
            entity=entity, category=PaymentRequest.Category.SUPPLIER, currency='BWP',
            subject='Invoice', payee=payee, line_items=[], total=Decimal(total),
            account_name=payee, account_number=acct, bank_name='FNB',
            status=status, created_by=creator)

    # ── register scope ───────────────────────────────────────────────────────
    def test_regular_user_cannot_use_register_to_see_the_estate(self):
        """?register=1 from a non-approver stays scoped to their own rows."""
        self._mk(self.other, 'Supplier A', '1000', PaymentRequest.Status.PAID)
        self._mk(self.clerk, 'Supplier B', '2000', PaymentRequest.Status.PENDING_FINANCE)
        self.client.force_authenticate(self.clerk)
        data = self.client.get(self.list_url, {'register': '1'}).json()
        refs = {r['ref'] for r in data['requests']}
        # only their own request; the other user's PAID row must NOT leak
        self.assertEqual(len(refs), 1)
        self.assertNotIn('register', data)   # register meta only for authorised viewers

    def test_first_approver_register_sees_all_statuses_and_entities(self):
        self._mk(self.other, 'Supplier A', '1000', PaymentRequest.Status.PAID)
        self._mk(self.other, 'Supplier B', '2000', PaymentRequest.Status.CANCELLED)
        self._mk(self.clerk, 'Supplier C', '3000', PaymentRequest.Status.PENDING_CFO,
                 entity='Veritas Insurance')
        self.client.force_authenticate(self.kago)
        data = self.client.get(self.list_url, {'register': '1'}).json()
        self.assertEqual(len(data['requests']), 3)          # incl. paid + cancelled
        self.assertEqual(data['register']['total'], 3)
        # detective columns present
        self.assertIn('days_to_authorise', data['requests'][0])
        self.assertIn('first_approved_at', data['requests'][0])
        # account numbers never travel in list rows
        self.assertNotIn('account_number', data['requests'][0])

    def test_register_filters_narrow(self):
        self._mk(self.other, 'Alpha Payee', '1000', PaymentRequest.Status.PAID)
        self._mk(self.other, 'Beta Payee', '9000', PaymentRequest.Status.PAID,
                 entity='Veritas Insurance')
        self.client.force_authenticate(self.cfo)
        # payee contains
        d = self.client.get(self.list_url, {'register': '1', 'payee': 'Beta'}).json()
        self.assertEqual([r['payee'] for r in d['requests']], ['Beta Payee'])
        # amount min
        d = self.client.get(self.list_url, {'register': '1', 'min': '5000'}).json()
        self.assertEqual({r['payee'] for r in d['requests']}, {'Beta Payee'})
        # entity contains
        d = self.client.get(self.list_url, {'register': '1', 'entity': 'Veritas'}).json()
        self.assertEqual({r['entity'] for r in d['requests']}, {'Veritas Insurance'})

    def test_register_pagination(self):
        for i in range(5):
            self._mk(self.other, f'P{i}', '100', PaymentRequest.Status.PAID)
        self.client.force_authenticate(self.cfo)
        d = self.client.get(self.list_url, {'register': '1', 'limit': '2', 'offset': '0'}).json()
        self.assertEqual(len(d['requests']), 2)
        self.assertEqual(d['register']['total'], 5)
        self.assertTrue(d['register']['has_more'])

    def test_register_export_returns_xlsx(self):
        self._mk(self.other, 'Supplier A', '1000', PaymentRequest.Status.PAID)
        self.client.force_authenticate(self.cfo)
        res = self.client.get(self.list_url, {'register': '1', 'export': 'xlsx'})
        self.assertEqual(res.status_code, 200)
        self.assertIn('spreadsheetml', res['Content-Type'])
        self.assertIn('attachment', res['Content-Disposition'])
        self.assertGreater(len(res.content), 100)

    def test_regular_user_export_is_ignored(self):
        """A non-approver cannot pull the estate as a spreadsheet either."""
        self._mk(self.other, 'Supplier A', '1000', PaymentRequest.Status.PAID)
        self.client.force_authenticate(self.clerk)
        res = self.client.get(self.list_url, {'register': '1', 'export': 'xlsx'})
        # falls through to the normal JSON list scoped to the clerk (no rows)
        self.assertEqual(res.status_code, 200)
        self.assertNotIn('spreadsheetml', res.get('Content-Type', ''))


class MaskAccountTests(APITestCase):
    def test_mask(self):
        from taskboard.payment_views import _mask_account
        self.assertEqual(_mask_account('62011112222'), '••••2222')
        self.assertEqual(_mask_account('12-34 56'), '••••3456')
        self.assertEqual(_mask_account(''), '')
        self.assertEqual(_mask_account('99'), '••••')


@window_always_open
class BankAckAndLoggingTests(APITestCase):
    def setUp(self):
        self.cfo = User.objects.create_user('pganesharajah', email='pganesharajah@alphadirect.co.bw')
        self.kago = User.objects.create_user('ktshutlhedi', email='ktshutlhedi@alphadirect.co.bw')
        self.clerk = User.objects.create_user('lthebe', email='lthebe@alphadirect.co.bw')
        self.detail_url = lambda i: reverse('v1-payment-request-detail', args=[i])
        self.decide_url = lambda i: reverse('v1-payment-request-decide', args=[i])

    def _mk(self, creator, payee, acct, status, **over):
        kw = dict(
            ref=f'PAY/B/{PaymentRequest.objects.count()+1:04d}',
            entity='Alpha Direct Insurance', category=PaymentRequest.Category.SUPPLIER,
            currency='BWP', subject='Invoice', payee=payee, line_items=[],
            total=Decimal('1000.00'), account_name=payee, account_number=acct,
            bank_name='FNB', status=status, created_by=creator)
        kw.update(over)
        return PaymentRequest.objects.create(**kw)

    def test_changed_account_blocks_signoff_without_ack(self):
        # history: paid Kalahari into account A
        self._mk(self.clerk, 'Kalahari Motors', '11111111', PaymentRequest.Status.PAID)
        # new request: same payee, DIFFERENT account B, awaiting finance
        req = self._mk(self.clerk, 'Kalahari Motors', '99999999', PaymentRequest.Status.PENDING_FINANCE)
        self.client.force_authenticate(self.kago)
        # sign off WITHOUT bank_ack → 409 PAY-BANK-ACK
        r = self.client.post(self.decide_url(req.id), {'decision': 'approve'}, format='json')
        self.assertEqual(r.status_code, 409, r.content)
        self.assertEqual(r.json().get('control'), 'PAY-BANK-ACK')
        req.refresh_from_db()
        self.assertEqual(req.status, PaymentRequest.Status.PENDING_FINANCE)  # not advanced
        # Acknowledging the change is no longer the whole story: since the
        # bank cross-check became a hard gate (Finance spec 2026-09-08) a
        # changed account must ALSO be confirmed against the beneficiary
        # details on the supporting document, by somebody in Finance who is
        # not the preparer. So bank_ack alone is now held at PAY-BANK-04 —
        # nobody is named as verifier on this fixture.
        r2 = self.client.post(self.decide_url(req.id),
                              {'decision': 'approve', 'bank_ack': True}, format='json')
        self.assertEqual(r2.status_code, 409, r2.content)
        self.assertEqual(r2.json().get('control'), 'PAY-BANK-04')
        req.refresh_from_db()
        self.assertEqual(req.status, PaymentRequest.Status.PENDING_FINANCE)

        # Name the verifier and attach the document, then confirm the details
        # match → it proceeds.
        req.verifier = 'Kago Tshutlhedi'
        req.save(update_fields=['verifier'])
        PaymentRequestAttachment.objects.create(
            request=req, original_name='supplier-letterhead.pdf',
            file=SimpleUploadedFile('supplier-letterhead.pdf', b'%PDF-1.4 test'))
        r3 = self.client.post(self.decide_url(req.id),
                              {'decision': 'approve', 'bank_ack': True,
                               'bank_doc_ack': True}, format='json')
        self.assertEqual(r3.status_code, 200, r3.content)
        req.refresh_from_db()
        self.assertEqual(req.status, PaymentRequest.Status.PENDING_CFO)
        # ...and both confirmations are on the record.
        self.assertIn('PAY-BANK-ACK', req.decision_notes)
        self.assertIn('PAY-BANK-DOC', req.decision_notes)

    def test_unchanged_account_signs_off_without_ack(self):
        """No change history → no bank-ack friction (only a change gates)."""
        req = self._mk(self.clerk, 'Brand New Payee', '55555555', PaymentRequest.Status.PENDING_FINANCE)
        self.client.force_authenticate(self.kago)
        r = self.client.post(self.decide_url(req.id), {'decision': 'approve'}, format='json')
        self.assertEqual(r.status_code, 200, r.content)

    def test_detail_returns_bank_block_and_logs_the_view(self):
        req = self._mk(self.clerk, 'Brand New Payee', '55555555', PaymentRequest.Status.PENDING_CFO)
        before = AuditLog.objects.filter(action=AuditLog.Action.READ,
                                         table_name='PaymentRequest').count()
        self.client.force_authenticate(self.cfo)
        d = self.client.get(self.detail_url(req.id)).json()
        self.assertIn('bank', d)
        self.assertTrue(d['bank']['first_payment'])      # never paid before
        self.assertIn('branch_code', d)
        self.assertIn('account_type', d)
        after = AuditLog.objects.filter(action=AuditLog.Action.READ,
                                        table_name='PaymentRequest').count()
        self.assertEqual(after, before + 1)              # the view was logged


@window_always_open
class DigestEntityRollupTests(APITestCase):
    def setUp(self):
        self.cfo = User.objects.create_user('pganesharajah', email='pganesharajah@alphadirect.co.bw')
        Company.objects.get_or_create(code='ADIC', defaults={'name': 'Alpha Direct Insurance'})
        self.summary_url = reverse('v1-payment-summary')

    def _mk(self, entity):
        return PaymentRequest.objects.create(
            ref=f'PAY/E/{PaymentRequest.objects.count()+1:04d}',
            entity=entity, category=PaymentRequest.Category.SUPPLIER, currency='BWP',
            subject='Invoice', payee='X', line_items=[], total=Decimal('100.00'),
            account_name='X', account_number='62000000000', bank_name='FNB',
            status=PaymentRequest.Status.PENDING_CFO, created_by=self.cfo)

    def test_two_adic_spellings_collapse_to_one_row(self):
        self._mk('Alpha Direct Insurance')
        self._mk('Alpha Direct Insurance Company')
        self.client.force_authenticate(self.cfo)
        rows = self.client.get(self.summary_url).json()['by_entity']
        adic_rows = [r for r in rows if 'Alpha Direct' in r['name']]
        self.assertEqual(len(adic_rows), 1, adic_rows)   # collapsed
        self.assertEqual(adic_rows[0]['count'], 2)       # both counted


@window_always_open
class PaymentRegisterExportContentTests(APITestCase):
    """Kago 2026-08-31 — the export CONTENT defects. Each assertion goes red on
    the pre-fix exporter (blank payee, no claim column, 0-day bug, ADIC spellings,
    text dates, header in row 3, bullet account digits)."""

    def setUp(self):
        from datetime import datetime
        from django.utils import timezone
        self.cfo = User.objects.create_user('pganesharajah', email='pganesharajah@alphadirect.co.bw')
        self.kago = User.objects.create_user('ktshutlhedi', email='ktshutlhedi@alphadirect.co.bw')
        self.list_url = reverse('v1-payment-requests')
        p = PaymentRequest.objects.create(
            ref='PAY/ADIC/2026/08/28/0012',
            entity='ADIC',                      # a variant spelling
            category=PaymentRequest.Category.CLAIM, currency='BWP',
            subject='G2026005213 COLDLINE (PTY) LTD',
            payee='',                            # blank — must fall back to account_name
            account_name='Coldline (Pty) Ltd',
            account_number='62011112222',
            line_items=[{'gl_code': 'CLAIMS PAYABLE', 'ref': 'INV-77', 'amount': '189000'}],
            total=Decimal('189000.00'), bank_name='FNB',
            status=PaymentRequest.Status.PENDING_CFO, created_by=self.kago)
        tz = timezone.get_current_timezone()
        # Raised 28th 17:00, authorised 29th 09:00 (local) → 1 calendar day, but a
        # raw timedelta is < 24h and truncates to 0 (the reported bug).
        PaymentRequest.objects.filter(pk=p.pk).update(
            created_at=timezone.make_aware(datetime(2026, 8, 28, 17, 0), tz),
            first_approved_at=timezone.make_aware(datetime(2026, 8, 29, 9, 0), tz),
            first_approver=self.cfo)

    def _export_sheets(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.client.force_authenticate(self.cfo)
        res = self.client.get(self.list_url, {'register': '1', 'export': 'xlsx'})
        self.assertEqual(res.status_code, 200)
        return load_workbook(BytesIO(res.content))

    def test_full_account_number_never_leaves_in_the_workbook(self):
        """C5 regression guard — only the last 4 digits may appear, never the
        full account string, anywhere in the generated file."""
        self.client.force_authenticate(self.cfo)
        res = self.client.get(self.list_url, {'register': '1', 'export': 'xlsx'})
        self.assertEqual(res.status_code, 200)
        self.assertNotIn(b'62011112222', res.content)

    def test_export_content_fixes(self):
        from datetime import date, datetime
        wb = self._export_sheets()
        # Defect 6: title/criteria on a separate COVER sheet; data sheet is second.
        self.assertEqual(wb.sheetnames[0], 'Export info')
        self.assertIn('Payment register', wb.sheetnames)
        ws = wb['Payment register']
        # Defect 6: header in ROW 1.
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[0], 'Reference')
        self.assertEqual(headers, [
            'Reference', 'Date raised', 'Date authorised', 'Days to authorise',
            'Requester', 'Authoriser', 'Type / Category', 'Entity', 'Payee',
            'Subject', 'Claim #', 'Ref # / Invoice no.', 'GL code', 'Currency',
            'Amount', 'Status', 'Account (last 4)', 'Claim description'])
        row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(headers)}
        # Defect 1: blank payee falls back to the account holder name.
        self.assertEqual(row['Payee'], 'Coldline (Pty) Ltd')
        # Defect 2: claim number is exported, pulled from the subject.
        self.assertEqual(row['Claim #'], 'G2026005213')
        # Defect 3: 28th → 29th is 1 day, not 0.
        self.assertEqual(row['Days to authorise'], 1)
        # Defect 4: 'ADIC' normalised to the canonical entity label.
        self.assertEqual(row['Entity'], 'Alpha Direct Insurance Company')
        # Defect 5: dates are real Excel dates, not text.
        self.assertIsInstance(row['Date raised'], (date, datetime))
        self.assertEqual((row['Date raised'].year, row['Date raised'].month,
                          row['Date raised'].day), (2026, 8, 28))
        # Account (last 4): plain digits, no bullet characters.
        self.assertEqual(row['Account (last 4)'], '2222')
        self.assertNotIn('•', str(row['Account (last 4)']))
        # Currency column present and correct.
        self.assertEqual(row['Currency'], 'BWP')
