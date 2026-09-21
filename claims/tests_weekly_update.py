"""Tests for the weekly claims update (B1, Bokani Makosha).

What is worth pinning here is not "does it build a spreadsheet". It is the small
number of decisions that, if they quietly changed, would send Finance the wrong
picture of the claims book every Monday without anybody noticing for months.
Each test below exists because getting it wrong is SILENT — the report still
builds, the columns still foot, and the number is wrong.
"""
from __future__ import annotations

import datetime
from decimal import Decimal

from django.core import mail
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone

from claims import weekly_update, weekly_update_report
from integrations.models import GraphiteClaim
from reporting.models import ReportRecipient


def _claim(policy='COMG2026212367', registered='2026-08-14', reserve='1000.00',
           payment='250.00', gid=None, detail=True, created=None):
    """One mirrored claim, shaped the way pull_graphite_claims writes it."""
    _claim.n = getattr(_claim, 'n', 0) + 1
    return GraphiteClaim.objects.create(
        graphite_id=gid if gid is not None else _claim.n,
        claim_number=f'CLM{_claim.n:05d}',
        policy_number=policy,
        registered_date=(datetime.date.fromisoformat(registered)
                         if registered else None),
        graphite_created_at=created,
        total_reserve=Decimal(reserve),
        total_payment=Decimal(payment),
        status='Open',
        detail_synced_at=timezone.now() if detail else None,
    )


class GroupDerivationTests(TestCase):
    """The group comes off the policy prefix, and an unknown one is CONFESSED."""

    def test_the_four_known_prefixes_map_to_their_groups(self):
        self.assertEqual(weekly_update.group_for('COMG2026212367'),
                         ('Commercial', True))
        self.assertEqual(weekly_update.group_for('DOMG2024112241'),
                         ('Domestic', True))
        self.assertEqual(weekly_update.group_for('MIS2025153125'),
                         ('Miscellaneous', True))
        self.assertEqual(weekly_update.group_for('BONU2026000001'),
                         ('Bonus', True))

    def test_an_unknown_prefix_is_flagged_not_bucketed(self):
        # THE WHOLE POINT OF B1's prefix rule. 'COMD' really exists on
        # production (one claim, measured 2026-09-13) and is one letter away
        # from Commercial. Folded into an "Other" pile, a whole product line can
        # disappear out of the group totals for months while the sheet still
        # foots perfectly.
        group, known = weekly_update.group_for('COMD2024129965')
        self.assertFalse(known)
        self.assertEqual(group, weekly_update.UNKNOWN_GROUP)
        self.assertNotEqual(group.lower(), 'other')

    def test_a_blank_policy_number_is_flagged_not_bucketed(self):
        # One blank policy number exists on production today.
        self.assertEqual(weekly_update.group_for(''),
                         (weekly_update.UNKNOWN_GROUP, False))
        self.assertEqual(weekly_update.group_for(None),
                         (weekly_update.UNKNOWN_GROUP, False))

    def test_the_longest_prefix_wins(self):
        # If matching were shortest-first and a 'COM' entry were ever added,
        # every COMG claim would answer to it. Pinned so the map can be extended
        # safely.
        weekly_update.PREFIX_GROUPS['COM'] = 'Should never win'
        try:
            self.assertEqual(weekly_update.group_for('COMG2026212367'),
                             ('Commercial', True))
        finally:
            del weekly_update.PREFIX_GROUPS['COM']

    def test_the_case_of_the_policy_number_does_not_decide_the_group(self):
        self.assertEqual(weekly_update.group_for('comg2026212367'),
                         ('Commercial', True))


class MonthBucketingTests(TestCase):
    """Reported date decides the month. Never date of loss, never today."""

    def test_the_month_is_the_reported_date_not_the_date_of_loss(self):
        c = _claim(registered='2026-08-14')
        c.date_of_loss = datetime.date(2026, 5, 2)
        c.save()
        self.assertEqual(weekly_update.month_of(c), ('2026-08', True))

    def test_a_claim_with_no_reported_date_is_reported_not_dated_today(self):
        # 604 of 4,600 production claims have no registered_date. Putting them
        # in the current month would inflate this month every single week.
        c = _claim(registered=None, detail=True)
        month, dated = weekly_update.month_of(c)
        self.assertFalse(dated)
        self.assertEqual(month, weekly_update.UNKNOWN_MONTH)

    def test_the_timestamp_fallback_is_converted_to_botswana_time(self):
        # 2026-08-31 23:30 Gaborone is 2026-08-31 21:30 UTC. Taking the month
        # off the raw UTC timestamp happens to agree here, so the discriminating
        # case is the other side: 2026-09-01 00:30 Gaborone = 2026-08-31 22:30
        # UTC, which off the raw UTC value lands in AUGUST. It is September.
        utc = datetime.datetime(2026, 8, 31, 22, 30,
                                tzinfo=datetime.timezone.utc)
        c = _claim(registered=None, created=utc)
        self.assertEqual(weekly_update.month_of(c), ('2026-09', True))


class TableTests(TestCase):
    """Table 1: group x month, summing reserve and payment."""

    def test_claims_are_grouped_by_group_and_month_with_decimal_sums(self):
        _claim(policy='COMG1', registered='2026-08-01', reserve='100.10', payment='10.05')
        _claim(policy='COMG2', registered='2026-08-20', reserve='200.20', payment='20.05')
        _claim(policy='DOMG1', registered='2026-08-03', reserve='5.00', payment='1.00')
        _claim(policy='COMG3', registered='2026-09-02', reserve='7.00', payment='0.00')
        s = weekly_update.summarise(weekly_update.rows())

        cell = [r for r in s['table']
                if r['group'] == 'Commercial' and r['month'] == '2026-08'][0]
        self.assertEqual(cell['count'], 2)
        self.assertEqual(cell['reserve'], Decimal('300.30'))
        self.assertEqual(cell['payment'], Decimal('30.10'))
        self.assertIsInstance(cell['reserve'], Decimal)

    def test_the_totals_equal_the_sum_of_the_rows(self):
        # A report whose headline disagrees with its own rows is unusable.
        for i in range(6):
            _claim(policy=f'COMG{i}', reserve='11.11', payment='2.22')
        s = weekly_update.summarise(weekly_update.rows())
        self.assertEqual(sum((r['reserve'] for r in s['table']), Decimal('0')),
                         s['reserve'])
        self.assertEqual(sum((g['payment'] for g in s['by_group']), Decimal('0')),
                         s['payment'])

    def test_unrecognised_claims_are_counted_and_named(self):
        _claim(policy='COMD2024129965', reserve='40.00', payment='0.00')
        _claim(policy='COMG2026212367', reserve='60.00', payment='0.00')
        s = weekly_update.summarise(weekly_update.rows())
        self.assertEqual(s['unrecognised_count'], 1)
        self.assertEqual(s['unrecognised'][0]['prefix'], 'COMD')
        self.assertEqual(s['unrecognised'][0]['example'], 'COMD2024129965')
        self.assertEqual(s['unrecognised'][0]['reserve'], Decimal('40.00'))
        # And it is still inside the grand total — confessed, not discarded.
        self.assertEqual(s['reserve'], Decimal('100.00'))

    def test_the_undated_month_sorts_last_and_is_not_a_date(self):
        _claim(policy='COMG1', registered='2026-09-01')
        _claim(policy='COMG2', registered=None)
        s = weekly_update.summarise(weekly_update.rows())
        self.assertEqual(s['table'][-1]['month'], weekly_update.UNKNOWN_MONTH)
        self.assertEqual(s['undated'], 1)

    def test_claims_never_detail_synced_are_counted(self):
        # Their reserve and payment read as zero. Silently, if nobody counts.
        _claim(policy='COMG1', detail=False, reserve='0.00', payment='0.00')
        _claim(policy='COMG2', detail=True)
        s = weekly_update.summarise(weekly_update.rows())
        self.assertEqual(s['no_detail'], 1)

    def test_recently_reported_counts_off_botswana_today(self):
        today = timezone.localdate()
        _claim(policy='COMG1', registered=(today - datetime.timedelta(days=2)).isoformat())
        _claim(policy='COMG2', registered=(today - datetime.timedelta(days=30)).isoformat())
        s = weekly_update.summarise(weekly_update.rows(), weeks=1, today=today)
        self.assertEqual(s['recent'], 1)


class HealthTests(TestCase):
    """It must refuse to send when the source has nothing to say."""

    def test_an_empty_mirror_is_not_reporting(self):
        self.assertFalse(weekly_update.mirror_is_reporting()['reporting'])

    def test_a_mirror_with_no_detail_sync_at_all_is_not_reporting(self):
        # Every claim present, every money figure 0.00, every column footing.
        # It reads as a quiet month; it is a dead feed.
        _claim(detail=False, reserve='0.00', payment='0.00')
        _claim(policy='DOMG1', detail=False, reserve='0.00', payment='0.00')
        h = weekly_update.mirror_is_reporting()
        self.assertFalse(h['reporting'])
        self.assertEqual(h['claims'], 2)
        self.assertEqual(h['detailed'], 0)

    def test_a_healthy_mirror_is_reporting(self):
        _claim()
        self.assertTrue(weekly_update.mirror_is_reporting()['reporting'])


class SendTests(TestCase):
    """The command's refusals, which are the reason it can be trusted."""

    def setUp(self):
        ReportRecipient.objects.create(
            report_slug='weekly-claims-update', email='bokani@alphadirect.co.bw')
        mail.outbox = []

    def test_it_refuses_to_send_a_cheerful_zero_when_the_mirror_is_empty(self):
        with self.assertRaises(SystemExit) as e:
            call_command('send_weekly_claims_report')
        self.assertEqual(e.exception.code, 4)
        self.assertEqual(len(mail.outbox), 0)

    def test_it_refuses_when_no_claim_has_ever_been_detail_synced(self):
        _claim(detail=False, reserve='0.00', payment='0.00')
        with self.assertRaises(SystemExit) as e:
            call_command('send_weekly_claims_report')
        self.assertEqual(e.exception.code, 4)
        self.assertEqual(len(mail.outbox), 0)

    def test_it_refuses_to_send_to_nobody(self):
        ReportRecipient.objects.all().delete()
        _claim()
        with self.assertRaises(SystemExit) as e:
            call_command('send_weekly_claims_report')
        self.assertEqual(e.exception.code, 3)
        self.assertEqual(len(mail.outbox), 0)

    def test_dry_run_sends_nothing(self):
        _claim()
        call_command('send_weekly_claims_report', '--dry-run')
        self.assertEqual(len(mail.outbox), 0)

    def test_it_sends_to_the_list_with_the_spreadsheet_attached(self):
        _claim()
        call_command('send_weekly_claims_report')
        self.assertEqual(len(mail.outbox), 1)
        msg = mail.outbox[0]
        self.assertIn('bokani@alphadirect.co.bw', msg.to)
        self.assertEqual(len(msg.attachments), 1)
        self.assertTrue(msg.attachments[0][0].endswith('.xlsx'))

    def test_an_unrecognised_prefix_reaches_the_reader_in_the_email_body(self):
        # Not a log line, not a footnote in the attachment only. If the only
        # place a missing product line is mentioned is a file nobody opens, it
        # has not been reported.
        _claim(policy='COMD2024129965')
        call_command('send_weekly_claims_report')
        body = mail.outbox[0].body + ''.join(
            str(a[0]) for a in mail.outbox[0].alternatives or [])
        self.assertIn('COMD', body)


class WorkbookTests(TestCase):
    """The attachment says out loud what the totals could otherwise hide."""

    def test_the_workbook_names_an_unrecognised_prefix(self):
        _claim(policy='COMD2024129965')
        _claim(policy='COMG2026212367')
        s = weekly_update.summarise(weekly_update.rows())
        buf = weekly_update_report.build_xlsx(s)
        from openpyxl import load_workbook
        ws = load_workbook(buf).active
        text = '\n'.join(str(c.value) for row in ws.iter_rows()
                         for c in row if c.value is not None)
        self.assertIn('COMD', text)
        self.assertIn('COMD2024129965', text)

    def test_the_workbook_carries_the_month_and_group_totals(self):
        _claim(policy='COMG1', registered='2026-08-14', reserve='1000.00',
               payment='250.00')
        s = weekly_update.summarise(weekly_update.rows())
        buf = weekly_update_report.build_xlsx(s)
        from openpyxl import load_workbook
        ws = load_workbook(buf).active
        text = '\n'.join(str(c.value) for row in ws.iter_rows()
                         for c in row if c.value is not None)
        self.assertIn('Aug 2026', text)
        self.assertIn('Commercial', text)
        self.assertIn('TOTAL', text)


class RecipientScreenTests(TestCase):
    """Finance must be able to fill the list in without a deploy."""

    def test_the_report_is_offered_on_the_recipient_screen(self):
        # THE WHOLE REPORT DEPENDS ON THIS ONE LINE. The command refuses to send
        # to nobody and has no hardcoded fallback, so if the slug is not on the
        # screen's allowlist Finance cannot add anybody, and the job installs,
        # fires every Monday, and delivers nothing — which is precisely what
        # happened to the failed-debits report on 11 September 2026.
        from realpay.failed_debits_api import ALLOWED_SLUGS
        from claims.management.commands.send_weekly_claims_report import (
            REPORT_SLUG)
        self.assertIn(REPORT_SLUG, ALLOWED_SLUGS)
