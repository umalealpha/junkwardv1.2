"""procurement/test_vendor_bank_upload.py — loading a supplier list in one go.

CFO 2026-08-20: "Create a place where people can upload the supplier names. We
already have the list of suppliers and the bank accounts and they load it now
itself so they don't need to really do hard work." Confirmed 2026-08-21: build
the screen; the list comes from his file.

The rule these tests exist to defend is his other one, from the same day:
"if a person is changing the bank account details it rejects, saying 'Why are
you doing this because you paid this person with another bank account?'" A
spreadsheet is the easiest way to move a supplier's money to a new account
without anyone noticing, so the upload must never do it.
"""
import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from rest_framework.test import APITestCase

from billing.models import Contact
from core.models import Company, Currency
from procurement.models import VendorBankAccount
from procurement.vendor_bank_upload import (
    Verdict, apply_plan, build_plan, map_headers, normalise_name, read_table,
)


def _csv(rows) -> bytes:
    return '\n'.join(','.join(str(c) for c in r) for r in rows).encode('utf-8')


HEAD = ['Supplier Name', 'Bank', 'Account Number', 'Branch Code']


class ReadingWhateverSheetArrivesTests(SimpleTestCase):
    """The point of matching columns by meaning: nobody's supplier list has a
    heading called `account_holder_name`."""

    def test_it_understands_the_headings_people_actually_use(self):
        m = map_headers(['Supplier Name', 'Bankers', 'A/C No.', 'Sort Code',
                         'CCY', 'Remarks'])
        self.assertEqual(m['vendor_name'], 'Supplier Name')
        self.assertEqual(m['bank_name'], 'Bankers')
        self.assertEqual(m['account_number'], 'A/C No.')
        self.assertEqual(m['branch_code'], 'Sort Code')
        self.assertEqual(m['currency_code'], 'CCY')
        self.assertEqual(m['notes'], 'Remarks')

    def test_a_heading_with_extra_words_around_it_still_lands(self):
        m = map_headers(['Vendor', 'Supplier Bank Account Number (BWP)', 'Bank'])
        self.assertEqual(m['account_number'],
                         'Supplier Bank Account Number (BWP)')

    def test_an_exact_heading_beats_a_partial_one(self):
        """A sheet with both 'Name' and 'Account Holder Name' must not put the
        second one on the supplier name."""
        m = map_headers(['Name', 'Account Holder Name', 'Bank', 'Account No'])
        self.assertEqual(m['vendor_name'], 'Name')
        self.assertEqual(m['account_holder_name'], 'Account Holder Name')

    def test_semicolon_separated_exports_are_read_too(self):
        data = b'Supplier;Bank;Account No\nABC;FNB;10000000001'
        headers, rows, _hdr = read_table(data, 'list.csv')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['Supplier'], 'ABC')

    def test_a_real_excel_file_is_read(self):
        """He is sending a spreadsheet, not a CSV. Read through the same reader
        the rest of the codebase uses, so .xls and .xlsb work too."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Supplier Name', 'Bank', 'A/C No.'])
        ws.append(['A Panel Repairer', 'FNB', 10000000001])
        buf = io.BytesIO()
        wb.save(buf)
        headers, rows, _hdr = read_table(buf.getvalue(), 'suppliers.xlsx')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['Supplier Name'], 'A Panel Repairer')
        # A number typed into a spreadsheet must not arrive as 1e+10.
        self.assertEqual(rows[0]['A/C No.'], '10000000001')

    def test_a_spreadsheet_saved_with_a_csv_name_is_still_read(self):
        """People rename files. The bytes decide, not the extension."""
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(['Supplier', 'Bank', 'Account No'])
        wb.active.append(['ABC', 'FNB', '10000000001'])
        buf = io.BytesIO()
        wb.save(buf)
        headers, rows, _hdr = read_table(buf.getvalue(), 'list.csv')
        self.assertEqual(len(rows), 1)

    def test_a_title_row_above_the_table_is_skipped(self):
        """Real sheets open with a heading line before the real columns."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['SUPPLIER BANKING DETAILS 2026'])
        ws.append(['Supplier Name', 'Bank', 'Account No'])
        ws.append(['ABC', 'FNB', '10000000001'])
        buf = io.BytesIO()
        wb.save(buf)
        headers, rows, _hdr = read_table(buf.getvalue(), 'suppliers.xlsx')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['Supplier Name'], 'ABC')

    def test_names_match_the_way_a_person_would(self):
        self.assertEqual(normalise_name('A PANEL REPAIRER (PTY) LTD'),
                         normalise_name('a panel repairer'))
        self.assertNotEqual(normalise_name('A Panel Repairer'), normalise_name('A Different Repairer'))


class _Fixture(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.bwp, _ = Currency.objects.get_or_create(
            code='BWP', defaults={'name': 'Pula', 'symbol': 'P'})
        cls.company, _ = Company.objects.get_or_create(
            code='VBUP', defaults={'name': 'Upload Test Co',
                                   'base_currency': cls.bwp})
        cls.other, _ = Company.objects.get_or_create(
            code='VBUP2', defaults={'name': 'Another Co',
                                    'base_currency': cls.bwp})
        cls.user = User.objects.create_user('vb_maker', password='x')
        cls.known = Contact.objects.create(
            contact_type=Contact.ContactType.VENDOR,
            name='A Panel Repairer (Pty) Ltd', company=cls.company)

    def _plan(self, rows, head=None, company=None):
        return build_plan(_csv([head or HEAD] + rows), 'suppliers.csv',
                          str((company or self.company).id))

    def _apply(self, plan, company=None):
        return apply_plan(plan, self.user, str((company or self.company).id))


class WhatWillHappenToEveryRowTests(_Fixture):
    """The old loader counted skips. A row that quietly did nothing looked
    exactly like a row that worked, which is how nobody noticed the register was
    empty."""

    def test_a_new_supplier_is_named_as_new(self):
        plan = self._plan([['Brand New Traders', 'FNB', '10000000001', '293567']])
        self.assertEqual(plan.rows[0].verdict, Verdict.NEW_SUPPLIER)
        self.assertTrue(plan.rows[0].will_write)

    def test_a_known_supplier_with_a_new_account_loads(self):
        plan = self._plan([['A Panel Repairer', 'FNB', '10000000001', '293567']])
        self.assertEqual(plan.rows[0].verdict, Verdict.LOAD)
        self.assertEqual(plan.rows[0].contact_id, str(self.known.pk))

    def test_a_row_missing_the_bank_is_refused_and_says_which_field(self):
        plan = self._plan([['A Panel Repairer', '', '10000000001', '293567']])
        self.assertEqual(plan.rows[0].verdict, Verdict.MISSING)
        self.assertIn('bank name', plan.rows[0].reason)

    def test_a_phone_number_in_the_account_column_is_caught(self):
        plan = self._plan([['A Panel Repairer', 'FNB', '+267', '293567']])
        self.assertEqual(plan.rows[0].verdict, Verdict.BAD_ACCOUNT)

    def test_a_currency_we_do_not_hold_is_refused(self):
        plan = self._plan([['A Panel Repairer', 'FNB', '10000000001', '293567', 'XYZ']],
                          head=HEAD + ['Currency'])
        self.assertEqual(plan.rows[0].verdict, Verdict.BAD_CURRENCY)

    def test_the_account_holder_falls_back_to_the_supplier_name(self):
        """A sheet without a separate holder column is telling us they are the
        same, not that the field is missing."""
        plan = self._plan([['A Panel Repairer', 'FNB', '10000000001', '293567']])
        self.assertEqual(plan.rows[0].account_holder_name, 'A Panel Repairer')

    def test_nothing_is_written_by_a_preview(self):
        before = VendorBankAccount.objects.count()
        self._plan([['A Panel Repairer', 'FNB', '10000000001', '293567'],
                    ['Brand New Traders', 'Stanbic', '9012345678', '']])
        self.assertEqual(VendorBankAccount.objects.count(), before)

    def test_the_preview_never_shows_a_whole_account_number(self):
        """A full list of supplier account numbers on screen is the thing worth
        stealing, and whoever uploaded already has the file."""
        plan = self._plan([['A Panel Repairer', 'FNB', '10000000001', '293567']])
        shown = plan.public()
        self.assertNotIn('10000000001', str(shown))
        self.assertEqual(shown['rows'][0]['account_ends'], '0001')


class ASheetMustNotMoveASuppliersMoneyTests(_Fixture):
    """The CFO's rule, applied to a spreadsheet."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        VendorBankAccount.objects.create(
            contact=cls.known, bank_name='FNB',
            account_holder_name='A Panel Repairer (Pty) Ltd',
            account_number='10000000001', currency_code_id='BWP',
            status=VendorBankAccount.Status.ACTIVE, created_by=cls.user)

    def test_the_same_account_again_is_simply_already_on_file(self):
        plan = self._plan([['A Panel Repairer (Pty) Ltd', 'FNB', '10000000001', '293567']])
        self.assertEqual(plan.rows[0].verdict, Verdict.ALREADY_ON_FILE)
        self.assertFalse(plan.rows[0].will_write)

    def test_spacing_and_dashes_do_not_make_it_a_different_account(self):
        plan = self._plan([['A Panel Repairer', 'FNB', '1000-0000 001', '293567']])
        self.assertEqual(plan.rows[0].verdict, Verdict.ALREADY_ON_FILE)

    def test_a_different_account_is_held_and_says_which_one_we_use_now(self):
        plan = self._plan([['A Panel Repairer', 'FNB', '99999999999', '293567']])
        row = plan.rows[0]
        self.assertEqual(row.verdict, Verdict.BANK_CHANGED)
        self.assertFalse(row.will_write)
        self.assertIn('0001', row.reason)
        self.assertIn('say why', row.reason)

    def test_a_held_row_writes_nothing_even_when_the_load_runs(self):
        plan = self._plan([['A Panel Repairer', 'FNB', '99999999999', '293567']])
        before = VendorBankAccount.objects.count()
        out = self._apply(plan)
        self.assertEqual(VendorBankAccount.objects.count(), before)
        self.assertEqual(out['created'], 0)
        self.assertEqual(out['held'], 1)


class LoadingIsStillMakerCheckerTests(_Fixture):
    def test_every_loaded_account_lands_as_a_draft_awaiting_approval(self):
        plan = self._plan([['A Panel Repairer', 'FNB', '10000000001', '293567'],
                           ['Brand New Traders', 'Stanbic', '9012345678', '']])
        out = self._apply(plan)
        self.assertEqual(out['created'], 2)
        self.assertEqual(out['suppliers_created'], 1)
        for b in VendorBankAccount.objects.all():
            self.assertEqual(b.status, VendorBankAccount.Status.DRAFT)

    def test_the_new_supplier_is_created_as_a_vendor(self):
        plan = self._plan([['Brand New Traders', 'Stanbic', '9012345678', '']])
        self._apply(plan)
        c = Contact.objects.get(name='Brand New Traders')
        self.assertEqual(c.contact_type, Contact.ContactType.VENDOR)

    def test_one_bad_row_does_not_lose_the_rest(self):
        """4,000 good rows must not be thrown away by one that trips."""
        plan = self._plan([['A Panel Repairer', 'FNB', '10000000001', '293567'],
                           ['Brand New Traders', 'Stanbic', '9012345678', '']])
        plan.rows[0].currency_code = 'NOPE'      # will fail on the FK
        out = self._apply(plan)
        self.assertEqual(out['created'], 1)
        self.assertEqual(len(out['problems']), 1)
        self.assertIn('A Panel Repairer', out['problems'][0])

    def test_loading_the_same_sheet_twice_adds_nothing_the_second_time(self):
        rows = [['A Panel Repairer', 'FNB', '10000000001', '293567']]
        self._apply(self._plan(rows))
        n = VendorBankAccount.objects.count()
        self._apply(self._plan(rows))     # a fresh plan, as the screen does
        self.assertEqual(VendorBankAccount.objects.count(), n)


class TheScreenTests(APITestCase):
    """Driving the real endpoints, because a test that calls the functions
    directly cannot tell me the screen is wired to them."""

    def setUp(self):
        bwp, _ = Currency.objects.get_or_create(
            code='BWP', defaults={'name': 'Pula', 'symbol': 'P'})
        self.company, _ = Company.objects.get_or_create(
            code='VBUP3', defaults={'name': 'Screen Test Co',
                                    'base_currency': bwp})
        self.maker = User.objects.create_user('vb_maker2', password='x')
        self.outsider = User.objects.create_user('vb_outsider', password='x')
        # A supplier belongs to a company, so the request has to say which one.
        # NOTE: these tests append it by hand. That is NOT proof the browser
        # sends it — Fable's review found the screen was sending nothing while
        # these were green, so the browser side is pinned separately in
        # frontend/src/lib/__tests__/withChosenCompany.test.ts and by looking at
        # the real request on the live screen.
        self.preview_url = (reverse('vendor-bank-account-upload-preview')
                            + f'?company={self.company.code}')
        self.load_url = (reverse('vendor-bank-account-upload-load')
                         + f'?company={self.company.code}')

    def _file(self, rows=None, name='suppliers.csv'):
        rows = rows or [HEAD, ['Brand New Traders', 'FNB', '10000000001', '293567']]
        return SimpleUploadedFile(name, _csv(rows), content_type='text/csv')

    def _as_maker(self):
        """An Accountant is a maker on controlled transactions; a Finance
        Manager is deliberately not (core.UserProfile.SOD_MAKER_TITLES)."""
        from core.models import UserCompanyAccess, UserProfile
        UserProfile.objects.update_or_create(
            user=self.maker,
            defaults={'role': UserProfile.Role.ACCOUNTANT,
                      'title': UserProfile.Title.ACCOUNTANT,
                      'is_active': True})
        # The entity grant matters as much as the title: entity isolation
        # (SEC-02, CFO 2026-08-08) shows nothing to a person with no recorded
        # entity, so a maker who belongs nowhere cannot load a list for anyone.
        UserCompanyAccess.objects.update_or_create(
            user=self.maker, company=self.company,
            defaults={'can_view': True, 'can_write': True})
        self.client.force_authenticate(self.maker)

    def test_a_maker_cannot_load_a_list_into_another_entity(self):
        """Entity isolation: their own company only, unless granted more."""
        other, _ = Company.objects.get_or_create(
            code='VBUP4', defaults={'name': 'Not Theirs',
                                    'base_currency': self.company.base_currency})
        self._as_maker()          # granted VBUP3 only
        r = self.client.post(
            reverse('vendor-bank-account-upload-load') + f'?company={other.code}',
            {'file': self._file()}, format='multipart')
        self.assertEqual(r.status_code, 403, r.content)
        self.assertEqual(VendorBankAccount.objects.count(), 0)

    def test_a_finance_manager_may_approve_but_not_load(self):
        """The checker side of the same control. Their job is to approve what a
        maker loaded, so the upload is not theirs to run."""
        from core.models import UserProfile
        UserProfile.objects.update_or_create(
            user=self.outsider,
            defaults={'role': UserProfile.Role.ACCOUNTANT,
                      'title': UserProfile.Title.FINANCE_MANAGER,
                      'is_active': True})
        self.client.force_authenticate(self.outsider)
        r = self.client.post(self.load_url, {'file': self._file()},
                             format='multipart')
        self.assertEqual(r.status_code, 403, r.content)
        self.assertEqual(VendorBankAccount.objects.count(), 0)

    def test_a_maker_gets_a_row_by_row_answer_and_nothing_is_written(self):
        self._as_maker()
        r = self.client.post(self.preview_url, {'file': self._file()},
                             format='multipart')
        self.assertEqual(r.status_code, 200, r.content)
        body = r.json()
        self.assertEqual(body['total_rows'], 1)
        self.assertEqual(body['will_load'], 1)
        self.assertEqual(VendorBankAccount.objects.count(), 0)

    def test_the_load_creates_them_and_says_so_in_plain_words(self):
        self._as_maker()
        r = self.client.post(self.load_url, {'file': self._file()},
                             format='multipart')
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()['created'], 1)
        self.assertIn('waiting for approval', r.json()['message'])
        self.assertEqual(VendorBankAccount.objects.count(), 1)

    def test_a_partial_load_cannot_report_as_a_clean_one(self):
        """The screen shows one line. A failure that is not in it is a failure
        nobody hears about."""
        from unittest import mock
        self._as_maker()
        fake = {'created': 1, 'suppliers_created': 0, 'held': 0,
                'already_on_file': 0, 'not_loaded': 2,
                'problems': ['row 7 (ABC): something went wrong',
                             'row 9 (DEF): something else']}
        with mock.patch('procurement.vendor_bank_upload.apply_plan',
                        return_value=fake):
            r = self.client.post(self.load_url, {'file': self._file()},
                                 format='multipart')
        self.assertEqual(r.status_code, 200, r.content)
        msg = r.json()['message']
        self.assertIn('2 row(s) FAILED', msg)
        self.assertIn('row 7', msg)

    def test_someone_who_may_not_create_one_by_hand_may_not_upload_either(self):
        """An upload must not be a way around the maker rule."""
        self.client.force_authenticate(self.outsider)
        r = self.client.post(self.load_url, {'file': self._file()},
                             format='multipart')
        self.assertEqual(r.status_code, 403, r.content)
        self.assertEqual(VendorBankAccount.objects.count(), 0)

    def test_a_signed_out_visitor_gets_nowhere(self):
        r = self.client.post(self.preview_url, {'file': self._file()},
                             format='multipart')
        self.assertIn(r.status_code, (401, 403))

    def test_without_choosing_a_company_it_says_so(self):
        """Rather than creating suppliers nobody can find afterwards."""
        self._as_maker()
        r = self.client.post(reverse('vendor-bank-account-upload-load'),
                             {'file': self._file()}, format='multipart')
        self.assertEqual(r.status_code, 400, r.content)
        self.assertIn('company', r.json()['detail'])
        self.assertEqual(VendorBankAccount.objects.count(), 0)

    def test_no_file_says_what_to_do(self):
        self._as_maker()
        r = self.client.post(self.preview_url, {}, format='multipart')
        self.assertEqual(r.status_code, 400)
        self.assertIn('Attach', r.json()['detail'])

    def test_a_sheet_without_a_holder_column_is_fine(self):
        """Almost every real supplier list has one name column. Refusing the
        file over a column that would just repeat it is the busywork this screen
        exists to remove."""
        self._as_maker()
        r = self.client.post(self.preview_url, {'file': self._file()},
                             format='multipart')
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()['rows'][0]['account_holder'],
                         'Brand New Traders')

    def test_a_sheet_with_no_account_column_is_told_which_column_is_missing(self):
        self._as_maker()
        f = SimpleUploadedFile('bad.csv', _csv([['Supplier', 'Bank'],
                                                ['ABC', 'FNB']]),
                               content_type='text/csv')
        r = self.client.post(self.preview_url, {'file': f}, format='multipart')
        self.assertEqual(r.status_code, 400)
        self.assertIn('account number', r.json()['detail'])

    def test_a_file_that_is_not_a_sheet_is_answered_not_a_crash(self):
        self._as_maker()
        f = SimpleUploadedFile('photo.png', b'\x89PNG\r\n\x1a\n\x00broken',
                               content_type='image/png')
        r = self.client.post(self.preview_url, {'file': f}, format='multipart')
        self.assertIn(r.status_code, (400,))
        self.assertIn('detail', r.json())


class ASupplierBelongsToACompanyTests(_Fixture):
    """CFO 2026-05-18: ADIC's vendors must not appear when ADSA is selected. A
    Contact with no company is filtered out of every company-scoped screen, so
    loading thousands of them would report success and leave a register nobody
    can select from.
    """

    def test_a_created_supplier_is_stamped_with_the_company(self):
        plan = self._plan([['Brand New Traders', 'FNB', '10000000001', '293567']])
        self._apply(plan)
        c = Contact.objects.get(name='Brand New Traders')
        self.assertEqual(c.company_id, self.company.id)

    def test_with_no_company_a_new_supplier_is_refused_not_orphaned(self):
        plan = self._plan([['Brand New Traders', 'FNB', '10000000001', '293567']])
        out = apply_plan(plan, self.user, None)
        self.assertEqual(out['created'], 0)
        self.assertFalse(Contact.objects.filter(name='Brand New Traders').exists())
        self.assertIn('no company was chosen', out['problems'][0])

    def test_another_companys_supplier_is_not_treated_as_already_known(self):
        """Otherwise a list loaded for one entity attaches its accounts to
        another entity's supplier record."""
        Contact.objects.create(contact_type=Contact.ContactType.VENDOR,
                               name='Their Own Supplier', company=self.other)
        plan = self._plan([['Their Own Supplier', 'FNB', '10000000001', '293567']])
        self.assertEqual(plan.rows[0].verdict, Verdict.NEW_SUPPLIER)
        self._apply(plan)
        self.assertEqual(
            Contact.objects.filter(name='Their Own Supplier').count(), 2)
        self.assertTrue(Contact.objects.filter(
            name='Their Own Supplier', company=self.company).exists())


class DistinctCompaniesStayDistinctTests(SimpleTestCase):
    """Stripping the wrong words off a name attaches an account to the wrong
    legal entity. 'Pty' and 'Ltd' say HOW a company is incorporated; 'Holdings'
    and 'Group' say WHICH company it is."""

    def test_holdings_and_company_are_not_the_same_business(self):
        self.assertNotEqual(normalise_name('Botswana Insurance Company'),
                            normalise_name('Botswana Insurance Holdings'))

    def test_group_is_not_noise(self):
        self.assertNotEqual(normalise_name('Kgalagadi Group'),
                            normalise_name('Kgalagadi Holdings'))

    def test_how_it_is_incorporated_is_still_noise(self):
        for a, b in (('A Panel Repairer (Pty) Ltd', 'a panel repairer'),
                     ('Acme Limited', 'ACME'),
                     ('Acme Incorporated', 'acme inc')):
            self.assertEqual(normalise_name(a), normalise_name(b), f'{a} vs {b}')


class AColumnIsNotGuessedFromOneWordTests(SimpleTestCase):
    """'Account Opened' is a date. Matched as the account number, 2024-01-15
    strips to eight digits, passes every check, and loads as a real bank
    account."""

    def test_a_date_column_is_not_taken_for_the_account_number(self):
        m = map_headers(['Supplier', 'Bank', 'Account Opened'])
        self.assertIsNone(m.get('account_number'))

    def test_bank_charges_is_not_the_bank(self):
        m = map_headers(['Supplier', 'Bank Charges', 'Account No'])
        self.assertIsNone(m.get('bank_name'))

    def test_branch_manager_is_not_the_branch_code(self):
        m = map_headers(['Supplier', 'Bank', 'Account No', 'Branch Manager'])
        self.assertIsNone(m.get('branch_code'))

    def test_the_long_heading_it_was_written_for_still_matches(self):
        m = map_headers(['Vendor', 'Supplier Bank Account Number (BWP)', 'Bank'])
        self.assertEqual(m['account_number'],
                         'Supplier Bank Account Number (BWP)')


class AFailureMessageNeverCarriesAnAccountNumberTests(_Fixture):
    """The preview masks account numbers on purpose. A database error quoting
    the row it choked on must not be the hole that undoes it.

    The first version of this test forced a currency FK error, whose message
    never mentions the account number — so it passed with the masking removed
    and certified nothing. It now exercises a failure that really does carry the
    number, which is what a driver-level error looks like.
    """

    def test_the_masker_cuts_a_number_to_its_last_four(self):
        from procurement.vendor_bank_upload import _no_account_numbers
        out = _no_account_numbers(
            'DETAIL: Key (account_number)=(10000000001) already exists.')
        self.assertNotIn('10000000001', out)
        self.assertIn('0001', out)

    def test_short_numbers_are_left_alone(self):
        """A row number and a branch code are not secrets, and mangling them
        would make the message useless."""
        from procurement.vendor_bank_upload import _no_account_numbers
        self.assertEqual(_no_account_numbers('row 7 failed'), 'row 7 failed')

    def test_a_failed_row_reports_without_the_number(self):
        from unittest import mock
        plan = self._plan([['Brand New Traders', 'FNB', '10000000001', '293567']])
        boom = ValueError('duplicate key value: account_number 10000000001')
        with mock.patch(
                'procurement.models.VendorBankAccount.objects.create',
                side_effect=boom):
            out = self._apply(plan)
        self.assertEqual(out['created'], 0)
        self.assertTrue(out['problems'])
        joined = ' '.join(out['problems'])
        self.assertNotIn('10000000001', joined)
        self.assertIn('0001', joined)
