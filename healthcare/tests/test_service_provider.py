"""Tests for the ADH service-provider registry (import / derive / export / API).

Synthetic fixtures only — fake practice numbers, never real provider PII.
Mirrors the EXACT headers of ADH's working sheet so the importer is proven
against the real column names (incl. the trailing space on "Contract Status "
and the sheet's "DATE CONTANCTED" spelling).
"""
import io

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from rest_framework.test import APITestCase, APIClient

from healthcare import provider_registry as reg
from healthcare.models import ServiceProvider
from core.models import AuditLog

HEADERS = [
    "ADH Acceptance (Ready)", "PRACTICE", "Discipline", "Prac Name", "Email Addr",
    "Town", "Contact No", "Location", "Contract Status ", "Welcome Pack Provided",
    "ADH 'Accepted Here' Sticker Displayed", "Provider Orientation",
    "Provider Onboarding Link", "DATE CONTANCTED", "COMMENT", "VENDOR",
]


def _sheet(rows):
    """rows: list of dicts keyed by a short alias -> build a workbook blob."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([
            r.get("ready", ""), r.get("practice", ""), r.get("discipline", ""),
            r.get("name", ""), r.get("email", ""), r.get("town", ""),
            r.get("contact", ""), r.get("location", ""), r.get("contract", ""),
            r.get("welcome", ""), r.get("sticker", ""), r.get("orientation", ""),
            r.get("link", ""), r.get("date", ""), r.get("comment", ""), r.get("vendor", ""),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class DerivedReadinessTests(TestCase):
    def test_afa_registered_derives_from_contract_status(self):
        self.assertEqual(ServiceProvider(contract_status="Signed").afa_registered, "Yes")
        self.assertEqual(ServiceProvider(contract_status="AFA Signing").afa_registered, "Pending")
        self.assertEqual(ServiceProvider(contract_status="").afa_registered, "No")

    def test_adh_ready_needs_registered_and_qc(self):
        p = ServiceProvider(contract_status="Signed", qc_confirmed=False)
        self.assertFalse(p.adh_ready)          # registered but not QC'd
        p.qc_confirmed = True
        self.assertTrue(p.adh_ready)
        p.contract_status = "AFA Signing"
        self.assertFalse(p.adh_ready)          # QC'd but not registered

    def test_the_sheets_yes_alone_does_not_make_a_provider_ready(self):
        """The manual column must NOT be able to answer the readiness question.

        It used to: `adh_ready` returned True on adh_acceptance == YES before it
        ever looked at the contract or the QC tick, so the number was the sheet
        read back and could not disagree with it. Readiness is now the contract
        AND a person's QC, and the sheet is only reconciled against it.
        """
        p = ServiceProvider(contract_status="", qc_confirmed=False, adh_acceptance="YES")
        self.assertFalse(p.adh_ready)
        p2 = ServiceProvider(contract_status="AFA Signing", qc_confirmed=False, adh_acceptance="yes")
        self.assertFalse(p2.adh_ready)
        # Signed but nobody has QC'd them -> still not ready, and it shows up.
        p3 = ServiceProvider(contract_status="Signed", qc_confirmed=False, adh_acceptance="YES")
        self.assertFalse(p3.adh_ready)
        self.assertTrue(p3.ready_mismatch)
        # Signed AND QC'd by a person -> ready, and the two now agree.
        p4 = ServiceProvider(contract_status="Signed", qc_confirmed=True, adh_acceptance="YES")
        self.assertTrue(p4.adh_ready)
        self.assertFalse(p4.ready_mismatch)

    def test_ready_mismatch_flags_disagreement(self):
        # Manual says YES, derived says No (not QC'd) -> mismatch.
        p = ServiceProvider(contract_status="Signed", qc_confirmed=False, adh_acceptance="YES")
        self.assertTrue(p.ready_mismatch)
        # NA / blank never counts as a disagreement.
        p2 = ServiceProvider(contract_status="Signed", qc_confirmed=False, adh_acceptance="NA")
        self.assertFalse(p2.ready_mismatch)


class ParseAndImportTests(TestCase):
    def test_practice_number_kept_as_text_with_leading_zeros(self):
        rows, warnings = reg.parse_sheet(_sheet([
            {"practice": "0060178", "name": "Lint Pharmacy", "discipline": "PHARMACY"},
        ]))
        self.assertEqual(rows[0]["practice_number"], "0060178")

    def test_float_practice_number_becomes_plain_int(self):
        # Excel often stores a bare number as a float (60178.0).
        wb = openpyxl.Workbook(); ws = wb.active; ws.append(HEADERS)
        ws.append(["YES", 60178, "GP", "Kadiyala Surgery", "", "GABORONE",
                   "", "", "Signed", "YES", "YES", "", "", "", "", ""])
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        rows, _ = reg.parse_sheet(buf)
        self.assertEqual(rows[0]["practice_number"], "60178")

    def test_blank_practice_rows_skipped_and_warned(self):
        rows, warnings = reg.parse_sheet(_sheet([
            {"practice": "111", "name": "Has key"},
            {"practice": "", "name": "No key"},
        ]))
        self.assertEqual(len(rows), 1)
        self.assertTrue(any("no practice number" in w.lower() for w in warnings))

    def test_dropped_rows_are_named_with_sheet_row_and_reason(self):
        """15-Sep-2026: rows vanished silently on import (226 ready -> 218; a
        duplicate practice number collapsed 226 -> 225). Every row that will NOT
        be imported must now be reported with its real sheet row + reason so the
        user can see which providers to fix, not just a count."""
        preview = reg.preview_import(_sheet([
            {"practice": "500", "name": "Keeps its key"},       # sheet row 2 -> imported
            {"practice": "", "name": "Ready but unkeyed"},       # sheet row 3 -> dropped (no key)
            {"practice": "500", "name": "Same number again"},    # sheet row 4 -> dropped (duplicate)
        ]))
        self.assertEqual(preview["dropped_count"], 2)
        reasons = {d["row"]: d["reason"] for d in preview["dropped"]}
        self.assertIn("no practice number", reasons[3])
        self.assertIn("duplicate", reasons[4].lower())
        self.assertTrue(any(d["name"] == "Ready but unkeyed" for d in preview["dropped"]),
                        "the dropped provider is named so the user knows which one to fix")

    def test_file_row_count_is_reported_and_a_duplicate_is_not_double_counted(self):
        """The screen tells the user "you uploaded a file with N rows", so N has
        to be the file. `total_rows` counts only what could be matched, and a
        duplicate lands in BOTH `rows` and `dropped` — adding the two would
        report 4 rows for a 3-row file."""
        preview = reg.preview_import(_sheet([
            {"practice": "500", "name": "Keeps its key"},
            {"practice": "", "name": "No key at all"},      # in the file, not imported
            {"practice": "500", "name": "Same number again"},  # in the file AND imported
        ]))
        self.assertEqual(preview["file_rows"], 3, "three data rows were uploaded")
        self.assertEqual(preview["dropped_count"], 2, "both are worth the user's eye")
        kinds = sorted(d["kind"] for d in preview["dropped"])
        self.assertEqual(kinds, ["not_imported", "overwrites_earlier"],
                         "a duplicate overwrites, it is not a row lost from the file")
        # The box on screen tells the user the file adds up. Prove the identity
        # it prints — file = providers written + repeats + unmatchable — or the
        # reconciliation is just another number for them to distrust.
        providers = len(preview["new"]) + len(preview["changed"]) + preview["unchanged"]
        dupes = sum(1 for d in preview["dropped"] if d["kind"] == "overwrites_earlier")
        no_key = sum(1 for d in preview["dropped"] if d["kind"] == "not_imported")
        self.assertEqual(providers, 1, "the two rows sharing 500 are one provider")
        self.assertEqual(preview["file_rows"], providers + dupes + no_key)

    def test_preview_classifies_new_changed_unchanged(self):
        ServiceProvider.objects.create(
            practice_number="14065", name="Kadiyala Surgery",
            discipline="GP", contract_status="AFA Signing")
        preview = reg.preview_import(_sheet([
            {"practice": "14065", "name": "Kadiyala Surgery", "discipline": "GP",
             "contract": "Signed"},                       # changed: contract status
            {"practice": "99999", "name": "Brand New Clinic", "contract": "Signed"},  # new
        ]))
        self.assertEqual(len(preview["new"]), 1)
        self.assertEqual(len(preview["changed"]), 1)
        self.assertIn("contract_status", preview["changed"][0]["changes"])

    def test_the_sheet_never_grants_qc(self):
        """21-Sep-2026: readiness answered out of the column it was meant to
        check. A create-time seed copied "ADH Acceptance (Ready) = YES" into
        qc_confirmed, so the derived number could not disagree with the sheet —
        222 rows said YES and 222 counted ready. QC is a person's call in Omni,
        never the spreadsheet's, so an imported YES must leave QC untouched and
        show up as a mismatch to work through.
        """
        p = reg.preview_import(_sheet([
            {"practice": "1", "name": "Signed Ready", "contract": "Signed", "ready": "YES"},
            {"practice": "2", "name": "Signed NotYes", "contract": "Signed", "ready": "NO"},
        ]))
        reg.commit_import([r["values"] for r in p["new"]], source_file="s.xlsx")
        a = ServiceProvider.objects.get(practice_number="1")
        b = ServiceProvider.objects.get(practice_number="2")
        self.assertFalse(a.qc_confirmed, "the sheet must not grant QC")
        self.assertFalse(a.adh_ready, "readiness needs a real QC, not a YES on a sheet")
        self.assertTrue(a.ready_mismatch, "their YES vs our not-ready is the work queue")
        self.assertFalse(b.qc_confirmed)

        # And the two numbers must now be able to disagree at all — the whole
        # point. A person QCs provider 1; only then does it count as ready.
        a.qc_confirmed = True
        a.save()
        counts = reg.dashboard_counts()
        self.assertEqual(counts["afa_says_ready"], 1, "AFA's own claim, counted separately")
        self.assertEqual(counts["adh_ready"], 1, "our derived check, agreeing on its own evidence")

    def test_reimport_never_clobbers_an_in_system_qc_decision(self):
        p = reg.preview_import(_sheet([
            {"practice": "1", "name": "Signed Ready", "contract": "Signed", "ready": "YES"},
        ]))
        reg.commit_import([r["values"] for r in p["new"]], source_file="s.xlsx")
        a = ServiceProvider.objects.get(practice_number="1")
        a.qc_confirmed = True          # a person QC'd it in Omni
        a.save()
        p2 = reg.preview_import(_sheet([
            {"practice": "1", "name": "Signed Ready", "contract": "Signed", "ready": "NO"},
        ]))
        rows = [r["values"] for r in (p2["new"] + p2["changed"])]
        self.assertTrue(rows, "the row must actually reach the importer, or this "
                              "test passes without exercising anything")
        reg.commit_import(rows, source_file="s2.xlsx")
        a.refresh_from_db()
        self.assertTrue(a.qc_confirmed,
                        "a later sheet must not undo a QC a person made in Omni")
        self.assertEqual(a.adh_acceptance, "NO", "but the sheet's own column does update")
        self.assertTrue(a.ready_mismatch, "and the disagreement is now visible")

    def test_yes_no_columns_are_stored_in_one_spelling(self):
        """The sheet answered four ways where it meant two (NO/No/YES/Yes), so
        a count of "YES" saw a fraction of the real number. Canonicalise in, by
        allowlist — anything that is not a yes/no answer survives verbatim."""
        p = reg.preview_import(_sheet([
            {"practice": "1", "name": "Mixed case", "ready": "yes", "welcome": "Yes",
             "sticker": " no ", "orientation": "N/A"},
            {"practice": "2", "name": "Not an answer", "sticker": "Posted 3 Aug"},
        ]))
        reg.commit_import([r["values"] for r in p["new"]], source_file="s.xlsx")
        a = ServiceProvider.objects.get(practice_number="1")
        self.assertEqual(a.adh_acceptance, "YES")
        self.assertEqual(a.welcome_pack, "YES")
        self.assertEqual(a.sticker_displayed, "NO")
        self.assertEqual(a.provider_orientation, "NA")
        b = ServiceProvider.objects.get(practice_number="2")
        self.assertEqual(b.sticker_displayed, "Posted 3 Aug",
                         "an answer we do not recognise is preserved, never mangled")

    def test_recased_sheet_is_not_reported_as_a_change(self):
        """Re-importing the same answers in different capitals must read as
        unchanged, or every casing tidy-up looks like 291 edits to approve."""
        ServiceProvider.objects.create(practice_number="1", name="Same",
                                       sticker_displayed="YES", welcome_pack="NO")
        preview = reg.preview_import(_sheet([
            {"practice": "1", "name": "Same", "sticker": "Yes", "welcome": "no"},
        ]))
        self.assertEqual(len(preview["changed"]), 0)
        self.assertEqual(preview["unchanged"], 1)

    def test_normalise_command_fixes_stored_rows_and_is_dry_by_default(self):
        from io import StringIO
        from django.core.management import call_command
        ServiceProvider.objects.create(practice_number="1", name="Mixed",
                                       sticker_displayed="Yes", welcome_pack="No",
                                       provider_orientation="Posted 3 Aug")
        call_command("normalise_provider_answers", stdout=StringIO())
        p = ServiceProvider.objects.get(practice_number="1")
        self.assertEqual(p.sticker_displayed, "Yes", "dry run must not write")

        call_command("normalise_provider_answers", "--apply", stdout=StringIO())
        p.refresh_from_db()
        self.assertEqual(p.sticker_displayed, "YES")
        self.assertEqual(p.welcome_pack, "NO")
        self.assertEqual(p.provider_orientation, "Posted 3 Aug",
                         "a non-answer is left alone")

    def test_commit_creates_updates_and_audits(self):
        u = User.objects.create_user("medu", email="mtlagae@alphadirect.co.bw")
        preview = reg.preview_import(_sheet([
            {"practice": "14065", "name": "Kadiyala Surgery", "contract": "Signed",
             "ready": "YES"},
        ]))
        rows = [r["values"] for r in preview["new"]]
        result = reg.commit_import(rows, source_file="test.xlsx", user=u)
        self.assertEqual(result["created"], 1)
        p = ServiceProvider.objects.get(practice_number="14065")
        self.assertEqual(p.name, "Kadiyala Surgery")
        self.assertEqual(p.adh_acceptance, "YES")
        # An audit row exists for the create.
        self.assertTrue(AuditLog.objects.filter(
            table_name="ServiceProvider", record_id=str(p.id)).exists())

        # Re-commit with a change -> update, not duplicate.
        p2 = reg.preview_import(_sheet([
            {"practice": "14065", "name": "Kadiyala Surgery", "contract": "AFA Signing"},
        ]))
        reg.commit_import([r["values"] for r in p2["changed"]], source_file="t2.xlsx", user=u)
        self.assertEqual(ServiceProvider.objects.filter(practice_number="14065").count(), 1)
        self.assertEqual(
            ServiceProvider.objects.get(practice_number="14065").contract_status, "AFA Signing")


class ExportTests(TestCase):
    def setUp(self):
        ServiceProvider.objects.create(practice_number="1", name="Ready Co",
            contract_status="Signed", qc_confirmed=True)
        ServiceProvider.objects.create(practice_number="2", name="Reg NoQC",
            contract_status="Signed", qc_confirmed=False)
        ServiceProvider.objects.create(practice_number="3", name="Pending Co",
            contract_status="AFA Signing")

    def test_export_filters(self):
        self.assertEqual(reg.dashboard_counts()["total"], 3)
        self.assertEqual(reg.dashboard_counts()["afa_registered"], 2)
        self.assertEqual(reg.dashboard_counts()["adh_ready"], 1)
        self.assertEqual(reg.dashboard_counts()["registered_not_qc"], 1)

        for key in ("all", "afa_registered", "adh_ready", "registered_not_qc"):
            data, fname = reg.build_export(key)
            self.assertTrue(fname.endswith(".xlsx"))
            wb = openpyxl.load_workbook(io.BytesIO(data))
            self.assertGreaterEqual(wb.active.max_row, 1)  # header at least

        # adh_ready filter yields exactly the one ready provider (+ header).
        data, _ = reg.build_export("adh_ready")
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.active.max_row, 2)


class ApiTests(APITestCase):
    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user(
            "medu", email="mtlagae@alphadirect.co.bw", password="x")
        self.outsider = User.objects.create_user(
            "random", email="random@alphadirect.co.bw", password="x")

    def test_outsider_denied(self):
        self.client.force_authenticate(self.outsider)
        r = self.client.get("/api/v1/health/service-providers/")
        self.assertEqual(r.status_code, 403)

    def test_the_whole_adh_health_team_can_reach_the_registry(self):
        # Meduduetso Tlagae asked for her team on 2026-09-08 so they can keep
        # the provider network up to date: Keneilwe Jane, Amantle Thake and
        # Onkgolotse Sebetlela (Ritah Tonkope and Loapi Keotlhoboge already had
        # it). Named here so nobody quietly loses access in a later tidy-up.
        for username, email in [
            ("kjane", "kjane@alphadirect.co.bw"),
            ("athake", "athake@alphadirect.co.bw"),
            ("osebetlela", "osebetlela@alphadirect.co.bw"),
            ("ritah", "rtonkope@alphadirect.co.bw"),
            ("loapi", "lkeotlhoboge@alphadirect.co.bw"),
        ]:
            user = User.objects.create_user(username, email=email, password="x")
            self.client.force_authenticate(user)
            r = self.client.get("/api/v1/health/service-providers/")
            self.assertEqual(r.status_code, 200, f"{email} cannot reach the registry")

    def test_import_preview_commit_and_list(self):
        self.client.force_authenticate(self.staff)
        blob = _sheet([
            {"practice": "14065", "name": "Kadiyala Surgery", "contract": "Signed"},
        ])
        r = self.client.post("/api/v1/health/service-providers/import/preview/",
                             {"file": blob}, format="multipart")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(len(r.data["new"]), 1)

        rows = [x["values"] for x in r.data["new"]]
        r2 = self.client.post("/api/v1/health/service-providers/import/commit/",
                              {"rows": rows, "source_file": "s.xlsx"}, format="json")
        self.assertEqual(r2.status_code, 200, r2.content)
        self.assertEqual(r2.data["created"], 1)

        r3 = self.client.get("/api/v1/health/service-providers/")
        self.assertEqual(r3.status_code, 200)
        self.assertEqual(r3.data["counts"]["total"], 1)

    def test_patch_rejects_derived_field_and_edits_qc(self):
        self.client.force_authenticate(self.staff)
        p = ServiceProvider.objects.create(
            practice_number="14065", name="Kadiyala", contract_status="Signed")
        # adh_ready is derived — attempting to set it is rejected, not applied.
        r = self.client.patch(f"/api/v1/health/service-providers/{p.id}/",
                              {"adh_ready": True, "qc_confirmed": True}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertIn("adh_ready", r.data["rejected"])
        p.refresh_from_db()
        self.assertTrue(p.qc_confirmed)
        self.assertTrue(p.adh_ready)   # now derived True (registered + QC)

    def test_export_returns_xlsx(self):
        self.client.force_authenticate(self.staff)
        ServiceProvider.objects.create(practice_number="1", name="Co",
            contract_status="Signed", qc_confirmed=True)
        r = self.client.get("/api/v1/health/service-providers/export/?filter=adh_ready")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("attachment", r["Content-Disposition"])


class ProviderApplyTests(APITestCase):
    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user("medu", email="mtlagae@alphadirect.co.bw")

    def test_public_apply_creates_pending_no_login(self):
        # No auth -> AllowAny public endpoint.
        r = self.client.post("/api/v1/health/provider-apply/",
            {"name": "Sunrise Pharmacy", "discipline": "Pharmacy",
             "town": "Gaborone", "email": "s@example.bw"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        from healthcare.models import ServiceProviderApplication as A
        a = A.objects.get()
        self.assertEqual(a.status, "pending")
        self.assertEqual(a.name, "Sunrise Pharmacy")

    def test_apply_requires_name_and_a_contact(self):
        r1 = self.client.post("/api/v1/health/provider-apply/", {"email": "x@y.bw"}, format="json")
        self.assertEqual(r1.status_code, 400)  # no name
        r2 = self.client.post("/api/v1/health/provider-apply/", {"name": "X"}, format="json")
        self.assertEqual(r2.status_code, 400)  # no email or phone

    def test_staff_can_list_and_review_applications(self):
        from healthcare.models import ServiceProviderApplication as A
        app = A.objects.create(name="Dr K", discipline="GP", email="k@x.bw")
        self.client.force_authenticate(self.staff)
        r = self.client.get("/api/v1/health/service-providers/applications/?status=pending")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.data["results"]), 1)
        r2 = self.client.patch(f"/api/v1/health/service-providers/applications/{app.id}/",
                               {"status": "accepted"}, format="json")
        self.assertEqual(r2.status_code, 200, r2.content)
        app.refresh_from_db()
        self.assertEqual(app.status, "accepted")

    def test_applications_gated_to_staff(self):
        outsider = User.objects.create_user("rand", email="rand@alphadirect.co.bw")
        self.client.force_authenticate(outsider)
        r = self.client.get("/api/v1/health/service-providers/applications/")
        self.assertEqual(r.status_code, 403)


class DashboardCommandTests(TestCase):
    def test_dashboard_builds_and_counts_applications(self):
        from healthcare.models import ServiceProviderApplication as A
        ServiceProvider.objects.create(practice_number="1", name="Ready Co",
            contract_status="Signed", qc_confirmed=True)
        A.objects.create(name="New Applicant", status="pending")
        from healthcare.management.commands.send_provider_dashboard import build_email
        subject, html, text, wa = build_email()
        self.assertIn("Alpha Direct Health", html)
        self.assertIn("ready to accept ADH clients", html)
        self.assertEqual(wa[1], "1")   # ready count in the WhatsApp params
        self.assertEqual(reg.dashboard_counts()["pending_applications"], 1)

    def test_switch_off_by_default_skips_send(self):
        from unittest import mock
        from django.core.management import call_command
        from healthcare.models import ProviderDashboardConfig
        self.assertFalse(ProviderDashboardConfig.current().enabled)  # OFF by default
        with mock.patch("core.notifications.send_html_with_cfo_cc") as send:
            call_command("send_provider_dashboard")
        self.assertFalse(send.called)   # switch OFF -> nothing sent

    def test_dashboard_delivers_to_named_execs_when_switch_on(self):
        # The CEO/COO are on the _NEVER_CC list; the send MUST pass
        # allow_named_exec=True or they are silently stripped (Fable 2026-09-01).
        from unittest import mock
        from django.core.management import call_command
        from healthcare.models import ProviderDashboardConfig
        cfg = ProviderDashboardConfig.current(); cfg.enabled = True; cfg.save()
        with self.settings(PROVIDER_DASHBOARD_ENABLED=True,
                           PROVIDER_DASHBOARD_TO=["aiyer@alphadirect.co.bw",
                                                  "arjuniyer@alphadirect.co.bw"]):
            with mock.patch("core.notifications.send_html_with_cfo_cc") as send, \
                 mock.patch("healthcare.management.commands.send_provider_dashboard._send_whatsapp",
                            return_value="wa skipped"):
                call_command("send_provider_dashboard")
        self.assertTrue(send.called)
        _args, kwargs = send.call_args
        self.assertTrue(kwargs.get("allow_named_exec"),
                        "must pass allow_named_exec so CEO/COO are not stripped")
        recips = _args[2] if len(_args) > 2 else kwargs.get("to")
        self.assertIn("aiyer@alphadirect.co.bw", recips)
        self.assertIn("arjuniyer@alphadirect.co.bw", recips)


class DashboardSwitchApiTests(APITestCase):
    def setUp(self):
        self.client = APIClient()
        self.staff = User.objects.create_user("medu", email="mtlagae@alphadirect.co.bw")

    def test_staff_can_read_and_flip_the_switch(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get("/api/v1/health/service-providers/dashboard-switch/")
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.data["enabled"])            # OFF by default
        r2 = self.client.post("/api/v1/health/service-providers/dashboard-switch/",
                              {"enabled": True}, format="json")
        self.assertEqual(r2.status_code, 200)
        self.assertTrue(r2.data["enabled"])
        from healthcare.models import ProviderDashboardConfig
        self.assertTrue(ProviderDashboardConfig.current().enabled)

    def test_switch_gated_to_staff(self):
        outsider = User.objects.create_user("rand2", email="rand2@alphadirect.co.bw")
        self.client.force_authenticate(outsider)
        r = self.client.get("/api/v1/health/service-providers/dashboard-switch/")
        self.assertEqual(r.status_code, 403)


class SnapshotTests(TestCase):
    def test_snapshot_command_creates_row(self):
        from django.core.management import call_command
        ServiceProvider.objects.create(practice_number="1", name="Ready",
            contract_status="Signed", qc_confirmed=True)
        ServiceProvider.objects.create(practice_number="2", name="Pending",
            contract_status="AFA Signing")
        from healthcare.models import ProviderDailySnapshot
        call_command("snapshot_provider_counts")
        snap = ProviderDailySnapshot.objects.get()
        self.assertEqual(snap.total, 2)
        self.assertEqual(snap.adh_ready, 1)
        self.assertEqual(snap.afa_pending, 1)

    def test_snapshot_idempotent(self):
        from django.core.management import call_command
        ServiceProvider.objects.create(practice_number="1", name="A", contract_status="Signed", qc_confirmed=True)
        from healthcare.models import ProviderDailySnapshot
        call_command("snapshot_provider_counts")
        call_command("snapshot_provider_counts")
        self.assertEqual(ProviderDailySnapshot.objects.count(), 1)

    def test_the_two_readiness_numbers_are_counted_separately(self):
        """The tile used to fold the sheet's YES into the derived count, so the
        two could never be seen to differ. They are now two numbers."""
        ServiceProvider.objects.create(practice_number="1", name="Derived Ready",
            contract_status="Signed", qc_confirmed=True, adh_acceptance="YES")
        ServiceProvider.objects.create(practice_number="2", name="Accepted Only",
            contract_status="", qc_confirmed=False, adh_acceptance="YES")
        c = reg.dashboard_counts()
        self.assertEqual(c["adh_ready"], 1, "only the one we actually verified")
        self.assertEqual(c["afa_says_ready"], 2, "what their list claims")
        self.assertEqual(c["mismatches"], 1, "and the gap is visible, not hidden")


class TrendApiTests(APITestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user("medu", email="mtlagae@alphadirect.co.bw")
        self.outsider = User.objects.create_user(
            "random", email="random@alphadirect.co.bw")

    def test_trend_returns_snapshots(self):
        from datetime import date
        from healthcare.models import ProviderDailySnapshot
        ProviderDailySnapshot.objects.create(date=date(2026, 9, 8), total=200, adh_ready=110)
        ProviderDailySnapshot.objects.create(date=date(2026, 9, 9), total=205, adh_ready=114)
        self.client.force_authenticate(self.user)
        r = self.client.get("/api/v1/health/service-providers/trend/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.data["snapshots"]), 2)
        self.assertEqual(r.data["snapshots"][0]["adh_ready"], 110)

    def test_trend_requires_auth(self):
        r = self.client.get("/api/v1/health/service-providers/trend/")
        self.assertEqual(r.status_code, 401)

    def test_trend_denied_to_staff_outside_the_adh_team(self):
        """The counts are the registry's own figures — gate them like the registry.
        Shipped 9-Sep-2026 on IsAuthenticated, which let any signed-in staff member
        read the ADH network's position."""
        self.client.force_authenticate(self.outsider)
        r = self.client.get("/api/v1/health/service-providers/trend/")
        self.assertEqual(r.status_code, 403)
