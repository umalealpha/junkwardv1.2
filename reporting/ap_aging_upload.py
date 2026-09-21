"""
reporting/ap_aging_upload.py — upload the maintained AP age analysis (CFO 2026-07-13).

Omni's AP Aging is computed from posted vendor bills; until opening AP is loaded
it reads empty. Finance (Pako / Bontle / Kago) upload the age-analysis
spreadsheet (as at e.g. 30 June 2026) here and the AP Aging page shows it as an
as-at snapshot. Read-only against the GL — this does NOT post any journals.

  POST /api/v1/reports/ap-aging/upload/     multipart: file, as_of, [company]
  GET  /api/v1/reports/ap-aging/snapshot/   ?as_of=&company=  -> latest snapshot
"""
from __future__ import annotations

import csv
import datetime
import io
import re
from decimal import Decimal, InvalidOperation

from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .models import APAgingSnapshot, APAgingSnapshotLine


def _is_finance(user) -> bool:
    """Creditor-level AP data is finance-only: FM / Financial Controller / CFO
    (or superuser). Without this gate any authenticated employee could read
    every entity's vendor balances or upload a fabricated snapshot that becomes
    the official AP Aging figure (Fable review fix, 2026-07-13)."""
    if getattr(user, 'is_superuser', False):
        return True
    from core.models import UserProfile, get_user_profile
    p = get_user_profile(user)
    return bool(p and p.is_active and p.title in {
        UserProfile.Title.FINANCE_MANAGER,
        UserProfile.Title.FINANCIAL_CONTROLLER,
        UserProfile.Title.CFO,
    })


def _clean_company(raw) -> tuple[str | None, bool]:
    """Return (company_id or None, ok). A malformed uuid must 400, not 500."""
    import uuid as _uuid
    s = (str(raw or '')).strip()
    if not s:
        return None, True
    try:
        return str(_uuid.UUID(s)), True
    except ValueError:
        return None, False

MAX_MB = 10
VENDOR_KEYS = ('vendor', 'supplier', 'creditor', 'name', 'account')
BUCKETS = [
    ('b_0_30', ('0-30', '0 - 30', '0to30', 'current', '0–30', 'not yet due', '0 30')),
    ('b_31_60', ('31-60', '31 - 60', '31–60', '30-60', '31 60')),
    ('b_61_90', ('61-90', '61 - 90', '61–90', '60-90', '61 90')),
    ('b_90_plus', ('90+', '90 +', 'over 90', '90 plus', '90plus', '>90', '90 and', '120', '91')),
]
TOTAL_KEYS = ('total', 'balance', 'outstanding', 'grand')


def _num(v) -> Decimal:
    if v is None:
        return Decimal('0')
    s = re.sub(r'[^0-9.\-()]', '', str(v)).strip()
    if not s or s in ('-', '.'):
        return Decimal('0')
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()')
    try:
        d = Decimal(s)
    except InvalidOperation:
        return Decimal('0')
    return -d if neg else d


def _rows_from_upload(f) -> list[list]:
    """Return a list of cell-rows from an xlsx or csv upload."""
    name = (getattr(f, 'name', '') or '').lower()
    raw = f.read()
    if name.endswith('.csv') or (raw[:1] not in (b'P',) and not name.endswith(('.xlsx', '.xlsm'))):
        text = raw.decode('utf-8', 'replace')
        return [list(r) for r in csv.reader(io.StringIO(text))]
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _find_header(rows: list[list]) -> tuple[int, dict]:
    """Locate the header row + a column map. Returns (header_idx, colmap)."""
    for idx, row in enumerate(rows[:15]):
        cells = [str(c or '').strip().lower() for c in row]
        vendor_col = next((i for i, c in enumerate(cells) if any(k in c for k in VENDOR_KEYS)), None)
        if vendor_col is None:
            continue
        colmap = {'vendor': vendor_col}
        for field, keys in BUCKETS:
            col = next((i for i, c in enumerate(cells) if any(k in c for k in keys)), None)
            if col is not None:
                colmap[field] = col
        tot = next((i for i, c in enumerate(cells) if any(k in c for k in TOTAL_KEYS)), None)
        if tot is not None:
            colmap['total'] = tot
        # a valid header needs a vendor col plus at least one bucket or a total
        if len(colmap) >= 2:
            return idx, colmap
    return -1, {}


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def ap_aging_upload(request):
    if not _is_finance(request.user):
        return Response({'detail': 'Finance only.'}, status=403)
    f = request.FILES.get('file')
    if f is None:
        return Response({'detail': 'No file uploaded (field "file").'}, status=400)
    if f.size > MAX_MB * 1024 * 1024:
        return Response({'detail': f'File exceeds {MAX_MB} MB.'}, status=400)
    try:
        as_of = datetime.date.fromisoformat((request.data.get('as_of') or '').strip())
    except ValueError:
        return Response({'detail': 'as_of must be YYYY-MM-DD.'}, status=400)
    company, ok = _clean_company(request.data.get('company'))
    if not ok:
        return Response({'detail': 'company must be a valid id.'}, status=400)

    try:
        rows = _rows_from_upload(f)
    except Exception as e:   # noqa: BLE001
        return Response({'detail': f'Could not read the file: {e}'}, status=400)
    hidx, colmap = _find_header(rows)
    if hidx < 0:
        return Response({'detail': 'Could not find a vendor column. Expected a header '
                                   'row with "Vendor" (or Supplier) and the age buckets '
                                   '(0-30, 31-60, 61-90, 90+) or a Total column.'}, status=400)

    snap = APAgingSnapshot.objects.create(
        as_of_date=as_of, company_id=company,
        uploaded_by=request.user if request.user.is_authenticated else None,
        source_filename=(getattr(f, 'name', '') or '')[:200])
    lines, grand = [], Decimal('0')
    for row in rows[hidx + 1:]:
        if not row:
            continue
        vname = str(row[colmap['vendor']] if colmap['vendor'] < len(row) else '').strip()
        if not vname or vname.lower() in ('total', 'totals', 'grand total', 'nan'):
            continue
        def cell(field):
            c = colmap.get(field)
            return _num(row[c]) if c is not None and c < len(row) else Decimal('0')
        b = {f: cell(f) for f, _ in BUCKETS}
        total = cell('total') if 'total' in colmap else sum(b.values())
        if total == 0 and all(v == 0 for v in b.values()):
            continue
        lines.append(APAgingSnapshotLine(snapshot=snap, vendor_name=vname[:200],
                     b_0_30=b['b_0_30'], b_31_60=b['b_31_60'], b_61_90=b['b_61_90'],
                     b_90_plus=b['b_90_plus'], total=total))
        grand += total
    APAgingSnapshotLine.objects.bulk_create(lines, batch_size=500)
    snap.vendor_count = len(lines)
    snap.grand_total = grand
    snap.save(update_fields=['vendor_count', 'grand_total'])
    return Response({'snapshot_id': str(snap.id), 'as_of': as_of.isoformat(),
                     'vendor_count': len(lines), 'grand_total': str(grand)}, status=201)


def _snapshot_json(snap):
    return {
        'id': str(snap.id), 'as_of': snap.as_of_date.isoformat(),
        'vendor_count': snap.vendor_count, 'grand_total': str(snap.grand_total),
        'uploaded_by': (snap.uploaded_by.get_full_name() or snap.uploaded_by.username) if snap.uploaded_by else '',
        'source_filename': snap.source_filename, 'created_at': snap.created_at.isoformat(),
        'lines': [{
            'vendor': ln.vendor_name, 'b_0_30': str(ln.b_0_30), 'b_31_60': str(ln.b_31_60),
            'b_61_90': str(ln.b_61_90), 'b_90_plus': str(ln.b_90_plus), 'total': str(ln.total),
        } for ln in snap.lines.all()],
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ap_aging_snapshot(request):
    if not _is_finance(request.user):
        return Response({'detail': 'Finance only.'}, status=403)
    qs = APAgingSnapshot.objects.all()
    as_of = (request.query_params.get('as_of') or '').strip()
    if as_of:
        qs = qs.filter(as_of_date=as_of)
    company, ok = _clean_company(request.query_params.get('company'))
    if not ok:
        return Response({'detail': 'company must be a valid id.'}, status=400)
    if company:
        qs = qs.filter(company_id=company)
    snap = qs.first()   # ordering = -created_at
    if not snap:
        return Response({'found': False})
    return Response({'found': True, 'snapshot': _snapshot_json(snap)})
