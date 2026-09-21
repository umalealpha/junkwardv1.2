"""Reading the filed A.1 tab and comparing it to Omni's own A.1.

Finance asked for the filed figure beside Omni's figure per line (Oprah
Mogomotsi, 2026-08-18). Storing the workbook was already done; these tests
cover actually READING it.

The load-bearing test here is `test_property_maps_by_section_not_by_position`.
'Property' is an insurance class in section 2 AND an asset bucket in section 4,
and Omni's INSURANCE_CLASSES list is in a different order from the workbook's
rows. Mapping by row position instead of by (section, label) puts one class's
filed figure onto another class's line and still renders plausibly — a silently
wrong comparison, which is worse than no comparison.
"""
import datetime as dt
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APIClient, APITestCase

from core.models import Company, UserProfile
from nbfira.filed_parser import parse_a1
from nbfira.models import NBFIRAFiledDocument, NBFIRAReturn, NBFIRAReturnLine

URL = '/api/v1/nbfira/returns'

# Row layout and labels lifted from the four FY2026 filed workbooks. Values are
# the real Q4 (June 2026) figures, so the arithmetic assertions below are
# checked against a return that was actually filed.
A1_ROWS = [
    (8,  '1. MINIMUM CAPITAL REQUIREMENT', None),
    (10, None, 5000),
    (12, '2. INSURANCE RISK CAPITAL:', None),
    (15, 'Property', 1929.6999999999998),
    (16, 'Transportation', 483.27500000000003),
    (17, 'Motor', 2153.95),
    (18, 'Accident', 2371.875),
    (19, 'Health', 420.375),
    (20, 'Guarantee', 234.65),
    (21, 'Liability', 2643.075),
    (22, 'Engineering', 300.95),
    (23, 'Miscellaneous', 808.5),
    (24, 'Total', 11346.349999999999),
    (25, 'IRC adjusted', 8158.571059708153),
    (28, '3. MAXIMUM EVENT RETENTION', None),
    (29, 'Total MER', 300),
    (30, 'MER adjusted', 215.71442075314494),
    (31, 'g insurance', 0.825),
    (33, '4. MARKET RISK CAPITAL', None),
    (35, 'Cash or near cash (less current and other liabilites)', 0),
    (36, 'Fixed Interest (Outstanding Term = 1 year)', 0),
    (37, 'Fixed Interest (Outstanding Term = 2 years)', 0),
    (38, 'Fixed Interest (Outstanding Term = 5 years)', 0),
    (39, 'Fixed Interest (Outstanding Term = 7 years)', 0),
    (40, 'Fixed Interest (Outstanding Term = 10 years)', 0),
    (41, 'Property', 0),
    (42, 'Listed Equities', 0),
    (43, 'Other assets', 8017.334499999999),
    (44, 'Unlisted equities', 0),
    (45, 'Total MRCTC', 8017.334499999999),
    (46, 'MRCTR adjusted', 5764.848892172349),
    (47, 'g market', 0.65),
    (49, 'Asset allocation ', None),
    (65, '5. PRESCRIBED CAPITAL TARGET', 13479.421621650197),
]


def build_a1_workbook(rows=None, sheet_name='A.1', extra_sheets=('Information',)):
    """A workbook shaped like the filed return: labels in A, figures in E."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws['A1'] = 'Statement A.1'
    ws['A4'] = 'as at the end of financial period 2026/6/30'
    for row, label, value in (rows if rows is not None else A1_ROWS):
        if label is not None:
            ws.cell(row, 1).value = label
        if value is not None:
            ws.cell(row, 5).value = value
    for name in extra_sheets:
        wb.create_sheet(name)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def upload_of(buf, name='ADIC_2026Q4_Return.xlsx'):
    return SimpleUploadedFile(
        name, buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class ParseA1Test(APITestCase):
    """The parser on its own — no HTTP, no database."""

    def test_reads_the_prescribed_capital_target(self):
        got = parse_a1(build_a1_workbook())
        self.assertTrue(got['ok'], got.get('error'))
        self.assertEqual(got['values']['A1_PCT'], Decimal('13479.4216'))

    def test_reads_mcr_from_below_its_unlabelled_header(self):
        got = parse_a1(build_a1_workbook())
        self.assertEqual(got['values']['A1_MCR_01'], Decimal('5000.0000'))

    def test_property_maps_by_section_not_by_position(self):
        """'Property' is an IRC class (1,929.70) and an MRC bucket (0.00).

        Omni's INSURANCE_CLASSES order differs from the workbook's row order,
        so a positional mapping would put Accident's figure on Property's line.
        Both must land on their own code, with their own value.
        """
        vals = parse_a1(build_a1_workbook())['values']
        self.assertEqual(vals['A1_IRC_PROPERTY'], Decimal('1929.7000'))
        self.assertEqual(vals['A1_MRC_PROPERTY'], Decimal('0.0000'))
        # And every class keeps its OWN filed figure.
        self.assertEqual(vals['A1_IRC_ACCIDENT'], Decimal('2371.8750'))
        self.assertEqual(vals['A1_IRC_MOTOR'], Decimal('2153.9500'))
        self.assertEqual(vals['A1_IRC_LIABILITY'], Decimal('2643.0750'))

    def test_bare_total_is_the_irc_total_not_the_allocation_total(self):
        """'Total' appears in several sections; only the IRC one is claimed."""
        vals = parse_a1(build_a1_workbook())['values']
        self.assertEqual(vals['A1_IRC_TOTAL'], Decimal('11346.3500'))
        self.assertEqual(vals['A1_MRC_TOTAL'], Decimal('8017.3345'))

    def test_g_factors_come_back_as_ratios_not_amounts(self):
        got = parse_a1(build_a1_workbook())
        self.assertEqual(got['ratios']['G_INSURANCE'], Decimal('0.8250'))
        self.assertEqual(got['ratios']['G_MARKET'], Decimal('0.6500'))
        # A ratio must never be offered as a money line to compare.
        self.assertNotIn('G_INSURANCE', got['values'])

    def test_layout_survives_an_inserted_row(self):
        """A regulator template that gains a row must still parse."""
        shifted = [(r + 3 if r >= 12 else r, lbl, v) for r, lbl, v in A1_ROWS]
        got = parse_a1(build_a1_workbook(rows=shifted))
        self.assertTrue(got['ok'], got.get('error'))
        self.assertEqual(got['values']['A1_PCT'], Decimal('13479.4216'))
        self.assertEqual(got['values']['A1_IRC_PROPERTY'], Decimal('1929.7000'))

    def test_full_pula_file_is_refused_rather_than_compared(self):
        """Omni stores A.1 in P'000. A file in full pula must be refused, not
        compared and reported as a 1000x variance."""
        pula = [(r, lbl, (v * 1000 if v is not None else None))
                for r, lbl, v in A1_ROWS]
        got = parse_a1(build_a1_workbook(rows=pula))
        self.assertFalse(got['ok'])
        self.assertIn('1000', got['error'])

    def test_missing_a1_tab_says_so(self):
        got = parse_a1(build_a1_workbook(sheet_name='Summary',
                                         extra_sheets=('IS',)))
        self.assertFalse(got['ok'])
        self.assertIn('A.1', got['error'])

    def test_no_pct_line_refuses_rather_than_guessing(self):
        without_pct = [r for r in A1_ROWS if r[0] != 65]
        got = parse_a1(build_a1_workbook(rows=without_pct))
        self.assertFalse(got['ok'])
        self.assertIn('Prescribed Capital Target', got['error'])

    def test_a_pdf_or_junk_file_is_a_plain_message_not_a_crash(self):
        got = parse_a1(io.BytesIO(b'%PDF-1.4 not a workbook at all'))
        self.assertFalse(got['ok'])
        self.assertIn('Excel', got['error'])


@override_settings(MEDIA_ROOT='/tmp/nbfira-parse-test-media')
class ReconcileEndpointTest(APITestCase):

    @classmethod
    def setUpTestData(cls):
        cls.company = Company.objects.create(code='NBFP', name='NBFIRA Parse Co.')
        cls.user = User.objects.create_user('fc', 'fc@test.example', 'x')
        UserProfile.objects.update_or_create(
            user=cls.user,
            defaults={'title': 'financial_controller', 'is_active': True},
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.ret = NBFIRAReturn.objects.create(
            type='quarterly', period_label='2026Q4', company=self.company,
            period_start=dt.date(2026, 4, 1), period_end=dt.date(2026, 6, 30),
        )
        # Omni's own A.1, as prod actually has it: MCR only, risk parts zero.
        for i, (code, label, value) in enumerate([
            ('A1_MCR_01',       'Minimum Capital Requirement (BWP)', '5000.0000'),
            ('A1_IRC_PROPERTY', 'IRC - Property',                    '0.0000'),
            ('A1_IRC_TOTAL',    'Total IRC',                         '0.0000'),
            ('A1_MER_TOTAL',    'Total MER',                         '0.0000'),
            ('A1_MRC_TOTAL',    'Total MRC TC',                      '0.0000'),
            ('A1_PCT',          'Prescribed Capital Target',         '5000.0000'),
        ]):
            NBFIRAReturnLine.objects.create(
                return_obj=self.ret, schedule='A.1', line_code=code,
                label=label, value=Decimal(value), sort_order=i)
        self.doc = NBFIRAFiledDocument.objects.create(
            return_obj=self.ret,
            file=upload_of(build_a1_workbook()),
            original_name='ADIC_2026Q4_Return.xlsx',
            size_bytes=1, file_hash_sha256='a' * 64,
            uploaded_by=self.user,
        )

    def _reconcile(self):
        return self.client.get(
            f'{URL}/{self.ret.id}/filed/{self.doc.id}/reconcile/')

    def test_it_puts_the_filed_pct_beside_omnis_pct(self):
        r = self._reconcile()
        self.assertEqual(r.status_code, 200, r.data)
        self.assertTrue(r.data['ok'])
        self.assertEqual(r.data['summary']['filed_pct'], '13479.4216')
        self.assertEqual(r.data['summary']['omni_pct'], '5000.0000')
        self.assertEqual(r.data['summary']['difference'], '8479.4216')

    def test_every_omni_a1_line_appears_with_both_columns(self):
        rows = {x['line_code']: x for x in self._reconcile().data['rows']}
        self.assertEqual(rows['A1_MCR_01']['filed'], '5000.0000')
        self.assertEqual(rows['A1_MCR_01']['omni'], '5000.0000')
        self.assertTrue(rows['A1_MCR_01']['agrees'])
        # The risk components are where Omni reads zero and the return does not.
        self.assertEqual(rows['A1_IRC_TOTAL']['omni'], '0.0000')
        self.assertEqual(rows['A1_IRC_TOTAL']['filed'], '11346.3500')
        self.assertFalse(rows['A1_IRC_TOTAL']['agrees'])

    def test_filed_only_lines_are_not_silently_dropped(self):
        """The workbook carries class and bucket lines Omni's stub never made.
        They must still be listed, or the comparison hides the gap."""
        codes = {x['line_code'] for x in self._reconcile().data['rows']}
        self.assertIn('A1_IRC_MOTOR', codes)
        self.assertIn('A1_MRC_OTHER_ASSETS', codes)

    def test_why_the_columns_differ_is_stated_on_every_comparison(self):
        """A variance must not read as 'Omni misread the ledger'. Since the
        method was adopted (2026-08-18) the remaining gap is unentered inputs."""
        warning = self._reconcile().data['method_warning']
        self.assertIn('25%', warning)
        self.assertIn('square root of squares', warning)
        # The gap is missing inputs now, NOT a method dispute — saying otherwise
        # would send Finance hunting a ledger error that does not exist.
        self.assertIn('missing inputs', warning)

    def test_reconcile_changes_no_figure_in_the_return(self):
        before = list(NBFIRAReturnLine.objects
                      .filter(return_obj=self.ret)
                      .order_by('line_code').values_list('line_code', 'value'))
        self._reconcile()
        after = list(NBFIRAReturnLine.objects
                     .filter(return_obj=self.ret)
                     .order_by('line_code').values_list('line_code', 'value'))
        self.assertEqual(before, after)

    def test_an_unreadable_file_reports_instead_of_500ing(self):
        doc = NBFIRAFiledDocument.objects.create(
            return_obj=self.ret,
            file=SimpleUploadedFile('scan.pdf', b'%PDF-1.4 nope',
                                    content_type='application/pdf'),
            original_name='scan.pdf', size_bytes=1,
            file_hash_sha256='b' * 64, uploaded_by=self.user,
        )
        r = self.client.get(f'{URL}/{self.ret.id}/filed/{doc.id}/reconcile/')
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.data['ok'])
        self.assertIn('Excel', r.data['error'])

    def test_ordinary_staff_cannot_read_the_comparison(self):
        other = User.objects.create_user('clerk', 'clerk@test.example', 'x')
        UserProfile.objects.update_or_create(
            user=other, defaults={'title': 'accountant', 'is_active': True})
        c = APIClient()
        c.force_authenticate(other)
        r = c.get(f'{URL}/{self.ret.id}/filed/{self.doc.id}/reconcile/')
        self.assertIn(r.status_code, (403, 404))

    def test_a_document_on_another_return_is_not_reachable_through_this_one(self):
        other = NBFIRAReturn.objects.create(
            type='quarterly', period_label='2026Q3', company=self.company,
            period_start=dt.date(2026, 1, 1), period_end=dt.date(2026, 3, 31),
        )
        r = self.client.get(f'{URL}/{other.id}/filed/{self.doc.id}/reconcile/')
        self.assertEqual(r.status_code, 404)


class ScheduleRouteTest(APITestCase):
    """The A.1 schedule endpoint used to 404 because its URL pattern excluded
    dots, and 'A.1' is the only schedule code that has one."""

    @classmethod
    def setUpTestData(cls):
        cls.company = Company.objects.create(code='NBFR', name='NBFIRA Route Co.')
        cls.user = User.objects.create_user('fm', 'fm@test.example', 'x')
        UserProfile.objects.update_or_create(
            user=cls.user,
            defaults={'title': 'financial_controller', 'is_active': True},
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.ret = NBFIRAReturn.objects.create(
            type='quarterly', period_label='2026Q4', company=self.company,
            period_start=dt.date(2026, 4, 1), period_end=dt.date(2026, 6, 30),
        )
        NBFIRAReturnLine.objects.create(
            return_obj=self.ret, schedule='A.1', line_code='A1_PCT',
            label='Prescribed Capital Target', value=Decimal('5000.0000'))

    def test_the_a1_schedule_can_actually_be_fetched(self):
        r = self.client.get(f'{URL}/{self.ret.id}/schedule/A.1/')
        self.assertEqual(r.status_code, 200, 'A.1 must not 404 - it has a dot')
        codes = [l['line_code'] for l in r.data['lines']]
        self.assertIn('A1_PCT', codes)
