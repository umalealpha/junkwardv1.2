"""
agent_portal/review_xlsx.py — the per-agent commission REVIEW workbook.

A one-click Excel of every agent's commission for a cycle (by stream + gross /
tax / net, net of tax on the UniCoin agent scale), with the rows that need
attention highlighted (no email, no bank on file, or has not-paid items). Built
for finance to check agent-by-agent and advise before payout (CFO 2026-07-27).

Money-safe: it carries bank/email PRESENCE (Y/N) only — never the account number.
Figures come from the already-built AgentPayslip snapshots, so the totals tie to
the approved pay-run exactly.
"""
from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import AgentBankAccount, AgentPayslip

NAVY = "1B2A6B"; TINT = "F3F5FC"; WARN = "FDE7D8"
HEAD = ["Payslip", "Agent", "Agent ID", "New Sales", "RealPay Conversion", "Collection",
        "Motor Comprehensive", "KYC & Banking", "Backlog", "Incentives",
        "Gross", "Tax", "Net", "Email?", "Bank?", "Not-paid items"]
_STREAM_COLS = HEAD[3:10]           # the agent-facing stream labels, in order
_MONEY_COLS = set(range(4, 13 + 1))  # New Sales .. Net (1-based)


def _f(x):
    try:
        return float(Decimal(str(x)))
    except (InvalidOperation, ValueError, TypeError):
        return None


def agent_review_workbook(cycle) -> bytes:
    """Return the per-agent review workbook for `cycle` as .xlsx bytes."""
    payslips = list(
        AgentPayslip.objects.filter(cycle=cycle).select_related("agent").order_by("-net")
    )
    banked = set(
        AgentBankAccount.objects
        .exclude(account_number="").exclude(account_number="0")
        .values_list("agent_id", flat=True)
    )

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Agents"
    ws["A1"] = f"UniCoin · Instant Insurance — {cycle.label} · Agent commission review"
    ws["A1"].font = Font(name="Book Antiqua", size=14, bold=True, color=NAVY)
    period = ""
    if cycle.start_date and cycle.end_date:
        period = f"{cycle.start_date:%d %b} – {cycle.end_date:%d %b %Y}"
    ws["A2"] = (f"Cycle {period} · {len(payslips)} agents · net of tax (UniCoin agent scale). "
                "Highlighted rows need attention: missing email / bank, or has not-paid items.")
    ws["A2"].font = Font(name="Book Antiqua", size=9, italic=True, color="555555")

    hdr = 4
    for c, h in enumerate(HEAD, 1):
        cell = ws.cell(row=hdr, column=c, value=h)
        cell.font = Font(name="Book Antiqua", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    gross_t = tax_t = net_t = 0.0
    r = hdr + 1
    for p in payslips:
        bd = {lbl: amt for lbl, amt in (p.breakdown or [])}
        email = "Y" if (p.agent.email or "").strip() else "N"
        bank = "Y" if p.agent_id in banked else "N"
        np_n = len(p.not_paid or [])
        row = ([p.number, p.agent.name, p.agent.ref_id]
               + [bd.get(s, "") for s in _STREAM_COLS]
               + [str(p.gross), str(p.tax), str(p.net), email, bank, str(np_n)])
        flag = (email == "N") or (bank == "N") or (np_n > 0)
        for c, val in enumerate(row, 1):
            v = _f(val) if c in _MONEY_COLS else val
            cell = ws.cell(row=r, column=c, value=v if v is not None else (val or ""))
            cell.font = Font(name="Book Antiqua", size=10); cell.border = border
            if c in _MONEY_COLS and isinstance(v, float):
                cell.number_format = "#,##0.00"
            if flag:
                cell.fill = PatternFill("solid", fgColor=WARN)
        gross_t += _f(p.gross) or 0; tax_t += _f(p.tax) or 0; net_t += _f(p.net) or 0
        r += 1

    ws.cell(row=r, column=2, value="TOTAL").font = Font(name="Book Antiqua", bold=True, color=NAVY)
    for c, tot in ((11, gross_t), (12, tax_t), (13, net_t)):
        cell = ws.cell(row=r, column=c, value=round(tot, 2))
        cell.font = Font(name="Book Antiqua", bold=True, color=NAVY)
        cell.number_format = "#,##0.00"; cell.fill = PatternFill("solid", fgColor=TINT)

    for i, w in enumerate([14, 26, 9, 11, 13, 11, 15, 13, 10, 11, 11, 9, 12, 7, 6, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
