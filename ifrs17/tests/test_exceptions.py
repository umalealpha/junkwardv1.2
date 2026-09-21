"""
Tests for the IFRS 17 Exception Register.

The two things that must be true, because they ARE the CFO's instruction:
  1. Only material exceptions demand an explanation. Small variances are seeded
     closed and cannot be "answered".
  2. A material exception cannot be closed with a tick-box answer.
"""
from __future__ import annotations

from decimal import Decimal as D

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase

from ifrs17 import constants as K
from ifrs17 import exceptions_service as XS
from ifrs17 import materiality as M
from ifrs17.models import MIN_EXPLANATION_CHARS, IFRS17Exception

GOOD = ('The commission is booked at the provisional 25% rate because the final '
        'underwriting-year loss ratio is not yet determinable. Empirica has '
        'confirmed the sliding-scale uplift is a sensitivity, not a bookable '
        'adjustment, and we will recognise it when the loss ratio is final.')


class MaterialityTests(TestCase):
    def test_threshold_is_derived_from_insurance_revenue(self):
        revenue = K.REPORTED['FY2026']['insurance_revenue']
        self.assertEqual(M.overall_materiality(),
                         (revenue * D('0.005')).quantize(D('0.01')))
        # Performance materiality rounds down to the register threshold.
        self.assertLessEqual(M.REGISTER_THRESHOLD, M.performance_materiality())

    def test_pbt_cross_check_is_looser_so_revenue_basis_is_the_binding_one(self):
        self.assertGreater(M.pbt_cross_check(), M.REGISTER_THRESHOLD)

    def test_big_amount_is_material_small_amount_is_not(self):
        big = M.classify({'ref': 'X', 'amount': D('8712000')})
        self.assertEqual(big['band'], M.BAND_MATERIAL)
        self.assertTrue(big['explanation_required'])

        small = M.classify({'ref': 'Y', 'amount': D('21490')})
        self.assertEqual(small['band'], M.BAND_IMMATERIAL)
        self.assertFalse(small['explanation_required'])

    def test_small_amount_can_still_be_material_by_nature(self):
        # DQ-03 is 70,353 — far below the threshold — but an unmodelled
        # product line is a completeness issue, not a rounding.
        c = M.classify({'ref': 'DQ-03', 'amount': D('70353')})
        self.assertEqual(c['band'], M.BAND_MATERIAL)
        self.assertEqual(c['basis'], 'qualitative')
        self.assertTrue(c['explanation_required'])

    def test_the_report_ten_flags_land_in_the_material_band(self):
        bands = {v['ref']: v['materiality_band'] for v in M.classified_variances()}
        for ref in ('DQ-01', 'DQ-02', 'DQ-03', 'DQ-04', 'DQ-05', 'DQ-06', 'DQ-08'):
            self.assertEqual(bands[ref], M.BAND_MATERIAL, f'{ref} must be material')
        # The pure rounding differences must NOT create work.
        for ref in ('DQ-09', 'DQ-11', 'DQ-12', 'DQ-13', 'DQ-14'):
            self.assertNotEqual(bands[ref], M.BAND_MATERIAL,
                                f'{ref} is a rounding difference, not an exception to answer')

    def test_material_set_is_a_minority_of_the_register(self):
        """"Don't put everything — major ones only." Guard against the filter
        silently degrading into "everything is material"."""
        rows = M.classified_variances()
        material = [r for r in rows if r['materiality_band'] == M.BAND_MATERIAL]
        self.assertLess(len(material), len(rows))
        self.assertGreaterEqual(len(material), 5)


class RegisterTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.fin = User.objects.create_user(username='fin', password='x')
        self.cfo = User.objects.create_user(username='cfo', password='x')
        XS.seed_register('FY2026')

    def test_seed_is_idempotent(self):
        before = IFRS17Exception.objects.count()
        out = XS.seed_register('FY2026')
        self.assertEqual(out['created'], 0)
        self.assertEqual(IFRS17Exception.objects.count(), before)

    def test_immaterial_rows_are_seeded_needing_no_response(self):
        e = IFRS17Exception.objects.get(ref='DQ-09', financial_year='FY2026')
        self.assertEqual(e.status, IFRS17Exception.Status.NO_RESPONSE_REQUIRED)
        self.assertFalse(e.explanation_required)
        self.assertFalse(e.is_outstanding)

    def test_material_rows_are_seeded_open(self):
        e = IFRS17Exception.objects.get(ref='DQ-01', financial_year='FY2026')
        self.assertEqual(e.status, IFRS17Exception.Status.OPEN)
        self.assertTrue(e.is_outstanding)

    def test_cannot_answer_an_immaterial_exception(self):
        e = IFRS17Exception.objects.get(ref='DQ-13', financial_year='FY2026')
        with self.assertRaises(ValidationError):
            XS.answer(e, explanation=GOOD, user=self.fin)

    def test_tickbox_answer_is_rejected(self):
        e = IFRS17Exception.objects.get(ref='DQ-01', financial_year='FY2026')
        for junk in ('', '   ', 'noted', 'agreed', 'N/A', 'ok will fix'):
            with self.assertRaises(ValidationError, msg=f'accepted junk: {junk!r}'):
                XS.answer(e, explanation=junk, user=self.fin)
        e.refresh_from_db()
        self.assertEqual(e.status, IFRS17Exception.Status.OPEN)

    def test_real_answer_then_accept(self):
        e = IFRS17Exception.objects.get(ref='DQ-01', financial_year='FY2026')
        XS.answer(e, explanation=GOOD, action='Recognise once the UWY loss ratio is final.',
                  user=self.fin)
        e.refresh_from_db()
        self.assertEqual(e.status, IFRS17Exception.Status.ANSWERED)
        self.assertEqual(e.answered_by, self.fin)
        self.assertIsNotNone(e.answered_at)

        XS.review(e, accept=True, note='Agreed, consistent with the treaty.', user=self.cfo)
        e.refresh_from_db()
        self.assertEqual(e.status, IFRS17Exception.Status.ACCEPTED)
        self.assertEqual(e.reviewed_by, self.cfo)
        self.assertFalse(e.is_outstanding)

    def test_return_needs_a_reason_and_reopens_the_item(self):
        e = IFRS17Exception.objects.get(ref='DQ-06', financial_year='FY2026')
        XS.answer(e, explanation=GOOD, user=self.fin)
        with self.assertRaises(ValidationError):
            XS.review(e, accept=False, note='', user=self.cfo)
        XS.review(e, accept=False, note='Post the difference or say why not.', user=self.cfo)
        e.refresh_from_db()
        self.assertEqual(e.status, IFRS17Exception.Status.RETURNED)
        self.assertTrue(e.is_outstanding)

    def test_answering_clears_a_previous_review(self):
        e = IFRS17Exception.objects.get(ref='DQ-06', financial_year='FY2026')
        XS.answer(e, explanation=GOOD, user=self.fin)
        XS.review(e, accept=False, note='More detail please.', user=self.cfo)
        XS.answer(e, explanation=GOOD + ' Further detail added.', user=self.fin)
        e.refresh_from_db()
        self.assertEqual(e.status, IFRS17Exception.Status.ANSWERED)
        self.assertIsNone(e.reviewed_by)
        self.assertEqual(e.review_note, '')

    def test_cannot_review_an_unanswered_exception(self):
        e = IFRS17Exception.objects.get(ref='DQ-01', financial_year='FY2026')
        with self.assertRaises(ValidationError):
            XS.review(e, accept=True, note='fine', user=self.cfo)

    def test_seed_refresh_never_overwrites_a_human_answer(self):
        e = IFRS17Exception.objects.get(ref='DQ-01', financial_year='FY2026')
        XS.answer(e, explanation=GOOD, user=self.fin)
        XS.seed_register('FY2026')
        e.refresh_from_db()
        self.assertEqual(e.explanation, GOOD)
        self.assertEqual(e.status, IFRS17Exception.Status.ANSWERED)

    def test_summary_counts_what_the_cfo_asks_for(self):
        s = XS.register_summary('FY2026')
        self.assertEqual(s['total'], len(K.KNOWN_VARIANCES))
        self.assertEqual(s['outstanding'], s['material'])
        e = IFRS17Exception.objects.get(ref='DQ-01', financial_year='FY2026')
        XS.answer(e, explanation=GOOD, user=self.fin)
        s2 = XS.register_summary('FY2026')
        self.assertEqual(s2['awaiting_review'], 1)
        self.assertEqual(s2['outstanding'], s['material'] - 1)

    def test_min_length_boundary(self):
        e = IFRS17Exception.objects.get(ref='DQ-01', financial_year='FY2026')
        just_short = 'a' * (MIN_EXPLANATION_CHARS - 1)
        with self.assertRaises(ValidationError):
            XS.answer(e, explanation=just_short, user=self.fin)
        XS.answer(e, explanation='a' * MIN_EXPLANATION_CHARS, user=self.fin)
        e.refresh_from_db()
        self.assertEqual(e.status, IFRS17Exception.Status.ANSWERED)


class ExportTests(TestCase):
    def test_workbook_builds_and_separates_the_bands(self):
        XS.seed_register('FY2026')
        from ifrs17.exceptions_export import register_workbook_bytes
        data = register_workbook_bytes('FY2026')
        self.assertTrue(data.startswith(b'PK'))
        self.assertGreater(len(data), 5000)

        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ['Materiality', 'Exceptions', 'Below threshold'])
        # Every material exception, and only those, on the Exceptions sheet.
        material = IFRS17Exception.objects.filter(band='material').count()
        self.assertEqual(wb['Exceptions'].max_row - 1, material)
