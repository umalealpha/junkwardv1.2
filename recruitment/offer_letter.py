from decimal import Decimal
from io import BytesIO

from docx import Document
from django.utils import timezone

from core.models import Company


def _format_money(amount, currency: str) -> str:
    value = Decimal(str(amount or 0))
    prefix = "P" if currency == "BWP" else currency
    return f"{prefix}{value:,.2f}"


def _start_date_text(authority, offer) -> str:
    start = None
    if offer is not None and getattr(offer, "start_date", None):
        start = offer.start_date
    elif getattr(authority, "effective_date", None):
        start = authority.effective_date
    if start is None:
        return "a date to be agreed"
    return start.strftime("%d %B %Y")


def build_docx(authority, offer) -> bytes:
    """Return the offer letter as a plain, professional .docx byte string."""
    doc = Document()

    header = doc.add_paragraph()
    run = header.add_run("Alpha Direct")
    run.bold = True

    doc.add_paragraph(timezone.localdate().strftime("%d %B %Y"))
    doc.add_paragraph(f"Dear {authority.person_name},")

    title = doc.add_paragraph()
    title_run = title.add_run(f"Offer of employment — {authority.position}")
    title_run.bold = True

    entity = authority.entity or "Alpha Direct Insurance Company"
    department = authority.department or "the company"
    start = _start_date_text(authority, offer)

    doc.add_paragraph(
        f"We are pleased to offer you the position of {authority.position} "
        f"in {department} at {entity}, starting {start}."
    )

    employment_type = authority.get_employment_type_display() or authority.employment_type
    doc.add_paragraph(f"Employment type: {employment_type}")

    salary_lines = authority.salary_lines or []
    if salary_lines:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        header_cells = table.rows[0].cells
        header_cells[0].text = "Component"
        header_cells[1].text = "Monthly amount"

        for row in salary_lines:
            if not isinstance(row, dict):
                continue
            label = row.get("label") or row.get("item") or ""
            amount = row.get("amount")
            if amount is None:
                amount = row.get("monthly", 0)
            cells = table.add_row().cells
            cells[0].text = str(label)
            cells[1].text = _format_money(amount, authority.currency)
    else:
        doc.add_paragraph(
            f"Total monthly cost to company: "
            f"{_format_money(authority.quoted_ctc_monthly, authority.currency)}"
        )

    doc.add_paragraph(
        "This offer is subject to satisfactory references, police clearance and "
        "a probation period of up to six months "
        "(Employment and Labour Relations Act, 2025, s.155)."
    )
    doc.add_paragraph(
        "Please confirm your acceptance by signing and returning a copy of this letter."
    )
    doc.add_paragraph("Yours sincerely,")
    doc.add_paragraph("Human Capital Department")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def company_for(authority):
    """Return the Company matching the authority entity, or None."""
    entity = (authority.entity or "").strip()
    if not entity:
        return None

    company = Company.objects.filter(name__iexact=entity).first()
    if company is not None:
        return company

    words = entity.split()[:2]
    if words:
        phrase = " ".join(words)
        return Company.objects.filter(name__icontains=phrase).first()
    return None
