"""
core/smart_upload/extractors.py

Pull tabular rows out of an uploaded file. Supports:
  - .xlsx  (openpyxl)            — first non-empty sheet by default
  - .xls   (xlrd if installed, else fall back to openpyxl best-effort)
  - .csv   (csv.reader)
  - .pdf   (pdfplumber tables)
  - .docx  (python-docx tables)
  - .txt   (TSV / pipe / semicolon heuristic)

Output shape is the same regardless of source:
  ExtractedTable(headers=[...], rows=[[...], ...], source_hint=str)

Header detection: take the first row that has >=2 non-empty cells AND
where >50% of cells are non-empty. That tolerates spreadsheets with a
big "title" cell in row 1 and a blank row 2 (CFO's xlsx packs do this).
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


# Row cap for extraction. This MUST be high enough that real GL/TB files commit
# in full: the commit path reuses these extracted rows, so a low cap silently
# dropped everything past it (e.g. a 3,390-row GL committed only 1,000 -> the
# rest vanished, and the upload rejected as "unbalanced"). Bug fix 2026-06-08
# (Legakwa, RSA/Unicoin truncation). The balance guard in committers.commit_tb/
# commit_gl is the backstop: anything that still truncated would fail the
# debits==credits check rather than post a partial set. 50k covers every real
# entity GL; only a genuinely abnormal file would trip `truncated`.
MAX_PREVIEW_ROWS = 50000   # was 1000 — silently truncated commits > 1k rows


@dataclass
class ExtractedTable:
    headers: list[str]
    rows: list[list[str]] = field(default_factory=list)
    source_hint: str = ''      # e.g. 'xlsx:Sheet1', 'pdf:page-3'
    truncated: bool = False

    def as_dicts(self) -> list[dict]:
        out = []
        for r in self.rows:
            d = {}
            for i, h in enumerate(self.headers):
                d[h] = r[i] if i < len(r) else ''
            out.append(d)
        return out


class ExtractError(Exception):
    pass


# ---------------------------------------------------------------------------
# Top-level dispatcher
# ---------------------------------------------------------------------------

def extract(file_bytes: bytes, filename: str,
            sheet: Optional[str] = None) -> ExtractedTable:
    """
    Look at the file extension, dispatch to the right reader.
    """
    suffix = Path(filename or '').suffix.lower()

    if suffix in ('.xlsx', '.xlsm'):
        return _from_xlsx(file_bytes, sheet)
    if suffix == '.xls':
        try:
            return _from_xls(file_bytes, sheet)
        except Exception:
            # xlrd may not be installed; fall back
            return _from_xlsx(file_bytes, sheet)
    if suffix == '.csv':
        return _from_csv(file_bytes)
    if suffix == '.tsv' or suffix == '.txt':
        return _from_delimited(file_bytes)
    if suffix == '.pdf':
        return _from_pdf(file_bytes)
    if suffix == '.docx':
        return _from_docx(file_bytes)

    # Unknown extension — try CSV then XLSX
    try:
        return _from_csv(file_bytes)
    except Exception:
        return _from_xlsx(file_bytes, sheet)


# ---------------------------------------------------------------------------
# Header heuristic
# ---------------------------------------------------------------------------

def _find_header_row(rows: list[list[str]]) -> int:
    """
    Return the 0-based index of the row that looks like a header.
    Header row = >=2 non-empty cells, >50% of cells non-empty.
    """
    for idx, r in enumerate(rows[:20]):   # only inspect first 20 rows
        non_empty = sum(1 for c in r if str(c).strip())
        if non_empty >= 2 and non_empty / max(len(r), 1) > 0.5:
            return idx
    return 0


def _clean(v) -> str:
    if v is None:
        return ''
    s = str(v).strip()
    return s


def _slice(rows: list[list[str]]) -> ExtractedTable:
    if not rows:
        return ExtractedTable(headers=[], rows=[])
    idx = _find_header_row(rows)
    headers = [_clean(c) for c in rows[idx]]
    # de-dup blank headers as column1..N
    headers = [h or f'column_{i+1}' for i, h in enumerate(headers)]
    data = []
    for r in rows[idx + 1:]:
        cells = [_clean(c) for c in r]
        if not any(cells):
            continue
        data.append(cells)

    truncated = False
    if len(data) > MAX_PREVIEW_ROWS:
        data = data[:MAX_PREVIEW_ROWS]
        truncated = True
    return ExtractedTable(headers=headers, rows=data, truncated=truncated)


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def _from_xlsx(b: bytes, sheet: Optional[str]) -> ExtractedTable:
    try:
        import openpyxl
    except ImportError as e:
        raise ExtractError(f'openpyxl not installed: {e}')

    try:
        wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True, read_only=True)
    except Exception as e:
        raise ExtractError(f'Could not open xlsx: {e}')

    target = None
    if sheet:
        for s in wb.sheetnames:
            if s.lower() == sheet.lower():
                target = s
                break
    if not target:
        # Pick first non-empty sheet
        for s in wb.sheetnames:
            ws = wb[s]
            if ws.max_row and ws.max_row > 0:
                target = s
                break
    if not target:
        return ExtractedTable(headers=[], rows=[], source_hint='xlsx:empty')

    ws = wb[target]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    out = _slice(rows)
    out.source_hint = f'xlsx:{target}'
    return out


def _from_xls(b: bytes, sheet: Optional[str]) -> ExtractedTable:
    import xlrd  # type: ignore
    book = xlrd.open_workbook(file_contents=b)
    name = sheet or book.sheet_names()[0]
    sh = book.sheet_by_name(name)
    rows = []
    for r in range(sh.nrows):
        rows.append([sh.cell_value(r, c) for c in range(sh.ncols)])
    out = _slice(rows)
    out.source_hint = f'xls:{name}'
    return out


# ---------------------------------------------------------------------------
# CSV / TSV / delimited
# ---------------------------------------------------------------------------

def _from_csv(b: bytes) -> ExtractedTable:
    text = _decode(b)
    reader = csv.reader(io.StringIO(text))
    rows = [list(r) for r in reader]
    out = _slice(rows)
    out.source_hint = 'csv'
    return out


def _from_delimited(b: bytes) -> ExtractedTable:
    text = _decode(b)
    first = (text.splitlines() or [''])[0]
    # Pick the most common single-char separator
    candidates = ['\t', '|', ';', ',']
    sep = max(candidates, key=lambda s: first.count(s))
    rows = []
    for line in text.splitlines():
        rows.append([c.strip() for c in line.split(sep)])
    out = _slice(rows)
    out.source_hint = f'delim:{sep!r}'
    return out


def _decode(b: bytes) -> str:
    for enc in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue
    return b.decode('utf-8', errors='replace')


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _from_pdf(b: bytes) -> ExtractedTable:
    try:
        import pdfplumber  # type: ignore
    except ImportError as e:
        raise ExtractError(f'pdfplumber not installed: {e}')

    rows: list[list[str]] = []
    page_hint = 'pdf'
    with pdfplumber.open(io.BytesIO(b)) as pdf:
        for i, page in enumerate(pdf.pages):
            tables = page.extract_tables() or []
            for t in tables:
                for r in t:
                    rows.append([_clean(c) for c in r])
            if rows:
                page_hint = f'pdf:page-{i+1}'
                break   # first page with a recognisable table wins
        if not rows:
            # Fall back to "lines of text split by whitespace"
            text = '\n'.join((p.extract_text() or '') for p in pdf.pages)
            for line in text.splitlines():
                parts = re.split(r'\s{2,}', line.strip())
                if len(parts) >= 2:
                    rows.append(parts)

    out = _slice(rows)
    out.source_hint = page_hint
    return out


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _from_docx(b: bytes) -> ExtractedTable:
    try:
        import docx   # type: ignore  # python-docx
    except ImportError as e:
        raise ExtractError(f'python-docx not installed: {e}')

    d = docx.Document(io.BytesIO(b))
    rows: list[list[str]] = []
    for tbl in d.tables:
        for row in tbl.rows:
            rows.append([_clean(cell.text) for cell in row.cells])
        if rows:
            break   # first table wins

    if not rows:
        # Fallback: paragraph text split by tabs
        for para in d.paragraphs:
            if '\t' in para.text:
                rows.append([c.strip() for c in para.text.split('\t')])

    out = _slice(rows)
    out.source_hint = 'docx'
    return out
