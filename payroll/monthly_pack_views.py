# -*- coding: utf-8 -*-
"""HTTP endpoints for the monthly payroll pack + control check.

  GET /api/v1/payroll/monthly-pack/?period=2026-08&company=<uuid>   -> JSON
  GET /api/v1/payroll/monthly-pack/export/?period=&company=         -> xlsx

Read-only. Gated by user_can_view_payroll and entity-grant scoped inside
build_monthly_pack (via _scoped_payslips)."""
from decimal import Decimal

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied, NotFound, ValidationError

from .amendment_views import user_can_view_payroll, _xlsx_response
from .monthly_pack import build_monthly_pack


def _jsonable(v):
    if isinstance(v, Decimal):
        return str(v)
    if isinstance(v, dict):
        return {k: _jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_jsonable(x) for x in v]
    return v


def _resolve(request):
    if not user_can_view_payroll(request.user):
        raise PermissionDenied("Payroll data is restricted to HR / Finance / payroll administrators.")
    period = (request.GET.get("period") or "").strip()
    company_id = (request.GET.get("company") or "").strip()
    if not period or not company_id:
        raise ValidationError("period and company are required.")
    from core.models import Company
    try:
        company = Company.objects.get(id=company_id)
    except (Company.DoesNotExist, ValueError, Exception):
        raise NotFound("Company not found.")
    try:
        pack = build_monthly_pack(period, company, request.user)
    except LookupError as e:
        raise NotFound(str(e))
    return pack, company, period


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def monthly_pack(request):
    pack, _company, _period = _resolve(request)
    return Response(_jsonable(pack))


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def monthly_pack_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    pack, company, period = _resolve(request)
    NAVY = "0D1B2A"
    hdr = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor=NAVY)
    bold = Font(bold=True)

    def money(v):
        return float(Decimal(str(v))) if v not in (None, "") else 0.0

    wb = Workbook()

    # --- Sheet 1: Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"{company.name} — Payroll {period} vs {pack['meta']['prior_period'] or 'n/a'}"])
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.append([])
    cur = pack["totals"]["current"]
    pri = pack["totals"]["prior"] or {}
    ws.append(["Measure", period, pack["meta"]["prior_period"] or "-", "Change"])
    for c in ws[ws.max_row]:
        c.font = hdr; c.fill = hdr_fill
    def row(label, key):
        a = money(cur.get(key, 0)); b = money(pri.get(key, 0))
        ws.append([label, a, b, round(a - b, 2)])
    ws.append(["Employees paid", cur.get("headcount", 0), pri.get("headcount", 0),
               cur.get("headcount", 0) - (pri.get("headcount", 0) or 0)])
    for lbl, k in [("Basic salaries", "basic"), ("Commission", "commission"),
                   ("Incentives", "incentive"), ("Bonus", "bonus"),
                   ("Allowances & benefits", "allowances"), ("Gross pay", "gross"),
                   ("PAYE", "paye"), ("Net pay", "net"), ("Cost to company", "ctc")]:
        row(lbl, k)
    for col in "ABCD":
        ws.column_dimensions[col].width = 24

    # --- Sheet 2: Incentive reconciliation ---
    ws2 = wb.create_sheet("Incentive check")
    ir = pack["incentive_recon"]
    ws2.append(["Incentive paid in payroll", money(ir["pay_total"])])
    ws2.append(["Module approved total", money(ir["mod_approved_total"])])
    ws2.append(["Module PROCESSED total", money(ir["mod_processed_total"])])
    ws2.append(["Module pending", money(ir["mod_pending_total"])])
    ws2.append(["Module rejected", money(ir["mod_rejected_total"])])
    ws2.append([])
    ws2.append(["Employee", "Department", "Paid", "Flag", "Note"])
    for c in ws2[ws2.max_row]:
        c.font = hdr; c.fill = hdr_fill
    for r in ir["rows"]:
        ws2.append([r["name"], r["dept"], money(r["amount"]), r["flag"], r["note"]])
    if ir["not_paid"]:
        ws2.append([])
        ws2.append(["Approved in module but NOT paid in payroll"])
        ws2[ws2.max_row][0].font = bold
        ws2.append(["Employee", "Approved", "Request"])
        for c in ws2[ws2.max_row]:
            c.font = hdr; c.fill = hdr_fill
        for r in ir["not_paid"]:
            ws2.append([r["name"], money(r["amount"]), r["req"]])
    for col in "ABCDE":
        ws2.column_dimensions[col].width = 26

    # --- Sheet 3: Authority to Recruit / Regrade ---
    ws3 = wb.create_sheet("Authority check")
    ws3.append(["Person", "Change", "Authority", "Kind", "Flag", "Note"])
    for c in ws3[ws3.max_row]:
        c.font = hdr; c.fill = hdr_fill
    for r in pack["atr_check"]:
        ws3.append([r["person"], r["change"], r["authority"], r["kind"], r["flag"], r["note"]])
    for col in "ABCDEF":
        ws3.column_dimensions[col].width = 26

    # --- Sheet 4: Major changes ---
    ws4 = wb.create_sheet("Major changes")
    mc = pack["major_changes"]
    ws4.append(["Joiners"]); ws4[ws4.max_row][0].font = bold
    ws4.append(["Name", "Department", "Basic"])
    for c in ws4[ws4.max_row]:
        c.font = hdr; c.fill = hdr_fill
    for j in mc["joiners"]:
        ws4.append([j["name"], j["dept"], money(j["basic"])])
    ws4.append([])
    ws4.append(["Leavers"]); ws4[ws4.max_row][0].font = bold
    ws4.append(["Name", "Department", "Basic (prior)"])
    for c in ws4[ws4.max_row]:
        c.font = hdr; c.fill = hdr_fill
    for l in mc["leavers"]:
        ws4.append([l["name"], l["dept"], money(l["basic"])])
    ws4.append([])
    ws4.append(["Salary / gross changes"]); ws4[ws4.max_row][0].font = bold
    ws4.append(["Name", "Department", "Basic prior", "Basic now", "Basic Δ", "Gross prior", "Gross now", "Gross Δ"])
    for c in ws4[ws4.max_row]:
        c.font = hdr; c.fill = hdr_fill
    for m in mc["moves"]:
        ws4.append([m["name"], m["dept"], money(m["basic_prior"]), money(m["basic_cur"]),
                    money(m["basic_change"]), money(m["gross_prior"]), money(m["gross_cur"]),
                    money(m["gross_change"])])
    for col in "ABCDEFGH":
        ws4.column_dimensions[col].width = 18

    fname = f"Payroll Pack {company.name} {period}.xlsx".replace("/", "-")
    return _xlsx_response(wb, fname)
