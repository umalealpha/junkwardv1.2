from datetime import date, datetime
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.datetime import from_excel
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from core.hris_access import user_can_access_hris
from core.models import AuditLog
from hris.contract_module import active_contracts, contract_row, months_before, DEFAULT_MONTHS
from hris.models import ContractRenewalDecision, Grade, HRISProfile
from payroll.models import Employee, EmploymentContract


def _has_hr_access(request):
    return user_can_access_hris(request.user)


def _profile_for(employee):
    try:
        return employee.hris_profile
    except AttributeError:
        return None


def _date_text(value):
    if value is None:
        return ''
    return value.isoformat()


def _cell_value(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _string_value(row, idx):
    value = _cell_value(row, idx)
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_date(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value).date()
    value_text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(value_text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f'Invalid date: {value_text}')


def _parse_bool(value):
    if value is None or str(value).strip() == '':
        return False
    return str(value).strip().lower() in ('y', 'yes', 'true', '1')


def _find_col(headers, *terms):
    for index, header in enumerate(headers):
        if any(term in header for term in terms):
            return index
    return None


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def contract_register(request):
    if not _has_hr_access(request):
        return Response({'detail': 'HR access required'}, status=403)

    today = timezone.localdate()
    contracts = active_contracts()

    counts = {
        'all': active_contracts().count(),
        'probation': active_contracts().filter(probation_end_date__gte=today).count(),
        'fixed_term': active_contracts().filter(contract_type='fixed_term').count(),
        'due': sum(1 for c in active_contracts() if contract_row(c, today)['due_for_renewal']),
    }

    q = request.query_params.get('q', '').strip()
    company_id = request.query_params.get('company', '').strip()
    filter_param = request.query_params.get('filter', 'all')

    if q:
        contracts = contracts.filter(employee__full_name__icontains=q)
    if company_id:
        contracts = contracts.filter(employee__company_id=company_id)

    if filter_param == 'probation':
        contracts = contracts.filter(probation_end_date__gte=today)
        rows = [contract_row(c, today) for c in contracts]
    elif filter_param == 'fixed_term':
        contracts = contracts.filter(contract_type='fixed_term')
        rows = [contract_row(c, today) for c in contracts]
    elif filter_param == 'due':
        rows = [contract_row(c, today) for c in contracts if contract_row(c, today)['due_for_renewal']]
    else:
        rows = [contract_row(c, today) for c in contracts]

    no_contract_qs = Employee.objects.filter(status__in=['active', 'on_leave']).exclude(contracts__status='active').distinct()
    if q:
        no_contract_qs = no_contract_qs.filter(full_name__icontains=q)
    if company_id:
        no_contract_qs = no_contract_qs.filter(company_id=company_id)

    no_contract = [
        {
            'employee_id': e.id,
            'name': e.full_name,
            'company': e.company.name if e.company else '',
            'department': e.department or '',
        }
        for e in no_contract_qs
    ]

    rules = {category: months_before(category) for category in DEFAULT_MONTHS}

    return Response({'rows': rows, 'counts': counts, 'rules': rules, 'no_contract': no_contract})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def contract_decision(request, contract_id):
    if not _has_hr_access(request):
        return Response({'detail': 'HR access required'}, status=403)

    try:
        contract = EmploymentContract.objects.get(id=contract_id)
    except EmploymentContract.DoesNotExist:
        return Response({'detail': 'Contract not found'}, status=404)

    decision = request.data.get('decision')
    valid_decisions = set(ContractRenewalDecision.Decision.values)
    if decision not in valid_decisions:
        return Response({'detail': 'Invalid decision'}, status=400)

    note = request.data.get('note', '')

    ContractRenewalDecision.objects.create(
        contract=contract,
        decision=decision,
        note=note,
        decided_by=request.user,
    )

    AuditLog.objects.create(
        table_name='payroll_employmentcontract',
        record_id=str(contract.id),
        action='update',
        new_values={'decision': decision, 'note': note},
        user=request.user,
        description='Contract renewal decision',
    )

    return Response(contract_row(contract, timezone.localdate()))


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def contract_template(request):
    if not _has_hr_access(request):
        return Response({'detail': 'HR access required'}, status=403)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Contract sheet'

    headers = [
        'Employee number',
        'Full name',
        'Company',
        'Department',
        'Job title',
        'Grade code (current)',
        'Contract type (permanent/fixed_term/probation/internship)',
        'Start date',
        'End date',
        'Probation end date',
        'Pension or provident (pension/provident/none)',
        'Controller (Y/N)',
        'Expatriate (Y/N)',
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.freeze_panes = 'A2'
    column_widths = {
        'A': 20,
        'B': 32,
        'C': 24,
        'D': 22,
        'E': 26,
        'F': 18,
        'G': 38,
        'H': 16,
        'I': 16,
        'J': 20,
        'K': 34,
        'L': 16,
        'M': 16,
    }
    for column, width in column_widths.items():
        ws.column_dimensions[column].width = width

    employees = Employee.objects.filter(status__in=['active', 'on_leave']).order_by('employee_number', 'full_name')

    for employee in employees:
        profile = _profile_for(employee)
        grade_code = ''
        if profile is not None and profile.grade is not None:
            grade_code = profile.grade.code or ''

        active_contract = employee.contracts.filter(status='active').first()

        ws.append([
            employee.employee_number,
            employee.full_name,
            employee.company.name if employee.company else '',
            employee.department or '',
            employee.job_title or '',
            grade_code,
            active_contract.contract_type if active_contract else '',
            active_contract.start_date if active_contract and active_contract.start_date else '',
            active_contract.end_date if active_contract and active_contract.end_date else '',
            active_contract.probation_end_date if active_contract and active_contract.probation_end_date else '',
            active_contract.retirement_fund if active_contract else '',
            'Y' if profile and profile.is_controller else 'N',
            'Y' if profile and profile.is_expatriate else 'N',
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="contract-sheet-{timezone.localdate().isoformat()}.xlsx"'
    return response


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def contract_upload(request):
    if not _has_hr_access(request):
        return Response({'detail': 'HR access required'}, status=403)

    commit = request.query_params.get('commit') == '1'
    uploaded = request.FILES.get('file')
    if not uploaded:
        return Response({'detail': 'No file uploaded'}, status=400)

    if uploaded.size > 5 * 1024 * 1024:
        return Response({'detail': 'File too large (max 5 MB)'}, status=400)

    try:
        wb = openpyxl.load_workbook(uploaded, data_only=True)
    except Exception as exc:
        return Response({'detail': f'Invalid workbook: {exc}'}, status=400)

    ws = wb.active
    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        return Response({'detail': 'Empty worksheet'}, status=400)

    headers = [str(h).strip().lower() if h is not None else '' for h in header_row]

    col_empno = _find_col(headers, 'employee number', 'employee_number', 'employee no')
    col_name = _find_col(headers, 'full name', 'name')
    col_contract_type = _find_col(headers, 'contract type')
    col_start = _find_col(headers, 'start date')
    col_end = _find_col(headers, 'end date')
    col_prob_end = _find_col(headers, 'probation end date')
    col_fund = _find_col(headers, 'pension or provident', 'fund')
    col_grade = _find_col(headers, 'grade code')
    col_controller = _find_col(headers, 'controller')
    col_expat = _find_col(headers, 'expatriate')

    prepared = []
    summary = {'ok': 0, 'error': 0, 'unmatched': 0}

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or str(value).strip() == '' for value in row):
            continue

        employee_number = _string_value(row, col_empno)
        full_name = _string_value(row, col_name)

        employee = None
        if employee_number:
            employee = Employee.objects.filter(status__in=['active', 'on_leave'], employee_number=employee_number).first()
        if employee is None and full_name:
            employee = Employee.objects.filter(status__in=['active', 'on_leave'], full_name__iexact=full_name).first()

        if employee is None:
            prepared.append({
                'row': row_number,
                'name': full_name,
                'status': 'unmatched',
                'messages': ['No active employee matched'],
                'changes': {},
            })
            summary['unmatched'] += 1
            continue

        messages = []
        changes = {}

        try:
            start_date = _parse_date(_cell_value(row, col_start))
        except ValueError as exc:
            start_date = None
            messages.append(str(exc))

        try:
            end_date = _parse_date(_cell_value(row, col_end))
        except ValueError as exc:
            end_date = None
            messages.append(str(exc))

        try:
            probation_end_date = _parse_date(_cell_value(row, col_prob_end))
        except ValueError as exc:
            probation_end_date = None
            messages.append(str(exc))

        contract_type = _string_value(row, col_contract_type).lower()
        valid_contract_types = {'permanent', 'fixed_term', 'probation', 'internship'}
        if contract_type and contract_type not in valid_contract_types:
            messages.append(f'Invalid contract type: {contract_type}')
            contract_type = ''

        retirement_fund = _string_value(row, col_fund).lower()
        valid_funds = {'', 'pension', 'provident', 'none'}
        if retirement_fund not in valid_funds:
            messages.append(f'Invalid pension/provident value: {retirement_fund}')
            retirement_fund = ''

        grade_code = _string_value(row, col_grade)
        grade = None
        if grade_code:
            try:
                grade = Grade.objects.get(code=grade_code)
            except Grade.DoesNotExist:
                messages.append(f'Unknown grade code: {grade_code}')
                grade_code = ''

        controller_flag = _parse_bool(_cell_value(row, col_controller))
        expatriate_flag = _parse_bool(_cell_value(row, col_expat))

        existing_profile = _profile_for(employee)
        old_controller = bool(existing_profile and existing_profile.is_controller)
        old_expat = bool(existing_profile and existing_profile.is_expatriate)

        active_contract = employee.contracts.filter(status='active').first()
        is_new = active_contract is None

        if active_contract is None:
            if not start_date:
                messages.append('Start date is required for a new contract')
            else:
                active_contract = EmploymentContract(employee=employee, status='active', start_date=start_date)

        if is_new and not contract_type:
            messages.append('Contract type is required for a new contract')

        if not messages and active_contract is not None:
            old_contract = {
                'contract_type': active_contract.contract_type if not is_new else '',
                'start_date': active_contract.start_date if not is_new else start_date,
                'end_date': active_contract.end_date if not is_new else end_date,
                'probation_end_date': active_contract.probation_end_date if not is_new else probation_end_date,
                'retirement_fund': active_contract.retirement_fund if not is_new else retirement_fund,
                'grade_code': active_contract.grade.code if active_contract.grade_id and active_contract.grade else '',
                'is_controller': old_controller,
                'is_expatriate': old_expat,
            }

            if contract_type:
                active_contract.contract_type = contract_type
            if start_date:
                active_contract.start_date = start_date
            active_contract.end_date = end_date
            active_contract.probation_end_date = probation_end_date
            active_contract.retirement_fund = retirement_fund
            if grade is not None:
                active_contract.grade = grade

            try:
                active_contract.full_clean()
            except ValidationError as exc:
                for field, errors in exc.message_dict.items():
                    for error in errors:
                        messages.append(str(error))

            new_contract = {
                'contract_type': active_contract.contract_type,
                'start_date': active_contract.start_date,
                'end_date': active_contract.end_date,
                'probation_end_date': active_contract.probation_end_date,
                'retirement_fund': active_contract.retirement_fund,
                'grade_code': active_contract.grade.code if active_contract.grade_id and active_contract.grade else '',
                'is_controller': controller_flag,
                'is_expatriate': expatriate_flag,
            }

            for field in new_contract:
                old_value = old_contract.get(field)
                new_value = new_contract[field]
                if old_value != new_value:
                    changes[field] = {
                        'old': _date_text(old_value) if isinstance(old_value, date) else old_value,
                        'new': _date_text(new_value) if isinstance(new_value, date) else new_value,
                    }

        if messages:
            status = 'error'
            summary['error'] += 1
        else:
            status = 'ok'
            summary['ok'] += 1

        prepared.append({
            'row': row_number,
            'name': employee.full_name,
            'status': status,
            'messages': messages,
            'changes': changes,
            '_employee': employee,
            '_contract': active_contract,
            '_is_new': is_new,
            '_controller_flag': controller_flag,
            '_expat_flag': expatriate_flag,
        })

    written = 0
    if commit:
        try:
            with transaction.atomic():
                for item in prepared:
                    if item['status'] != 'ok':
                        continue

                    employee = item['_employee']
                    contract = item['_contract']
                    is_new = item['_is_new']

                    if contract is None:
                        continue

                    if is_new:
                        contract.save()
                    else:
                        contract.save()

                    profile, _ = HRISProfile.objects.get_or_create(employee=employee)
                    profile.is_controller = item['_controller_flag']
                    profile.is_expatriate = item['_expat_flag']
                    profile.save()

                    AuditLog.objects.create(
                        table_name='payroll_employmentcontract',
                        record_id=str(contract.id),
                        action='create' if is_new else 'update',
                        old_values={k: v['old'] for k, v in item['changes'].items()},
                        new_values={k: v['new'] for k, v in item['changes'].items()},
                        user=request.user,
                        description='Contract upload' if is_new else 'Contract upload update',
                    )
                    written += 1
        except Exception as exc:
            return Response({'detail': f'Commit failed: {exc}'}, status=400)

    response_rows = []
    for item in prepared:
        response_rows.append({
            'row': item['row'],
            'name': item['name'],
            'status': item['status'],
            'messages': item['messages'],
            'changes': item['changes'],
        })

    payload = {'rows': response_rows, 'summary': summary}
    if commit:
        payload['written'] = written

    return Response(payload)
