"""FY2025 Development Dialogue workbook import + the permanent lock.

Dorothy's `Development Dialogues 2025.zip` put 37 workbooks in the HR vault; only
the 7 cockpit rows reached the 9-box, so 30 reviewed people were invisible on it.

The trap these tests exist for: the two axes are weighted 80/20 in the workbook,
not 50/50. Bharath's real file reads performance 0.7275 and potential 0.1970, and
'Overall Score (All Sections)' 0.9245 — the two summed. Store the raw potential on
a 0..1 axis and almost every employee bands as low-potential, which would be a
false statement about 35 people's careers.
"""
from __future__ import annotations

import io

from django.db import transaction
from django.test import TestCase

from hris.dd_workbook_import import (POTENTIAL_MAX, PERFORMANCE_MAX, band_normalise,
                                     parse_workbook)
from hris.talent_cockpit_models import DevelopmentDialogue


def _workbook(perf_rows, pot_rows, name='Bharath Balasubramanian',
              dept='Finance & Planning', position='Financial Controller',
              sheet_title='Nine Box Grid '):
    """A minimal stand-in for the real PMS workbook's Nine Box sheet."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws['A1'] = 'Alpha Direct 9-Box Potential vs Performance '
    ws['A2'], ws['B2'], ws['C2'] = 'Name ', 'Department ', 'Position '
    ws['A3'], ws['B3'], ws['C3'] = name, dept, position
    ws['A4'], ws['B4'] = 'Performance Competency', 'Final Score'
    r = 5
    for label, val in perf_rows:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=val)
        r += 1
    ws.cell(row=r, column=1, value='Total')
    ws.cell(row=r, column=2, value=round(sum(v for _, v in perf_rows), 4))
    r += 2
    ws.cell(row=r, column=1, value='Potential Indicator')
    ws.cell(row=r, column=2, value='Final Score')
    r += 1
    for label, val in pot_rows:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=val)
        r += 1
    ws.cell(row=r, column=1, value='Total')
    ws.cell(row=r, column=2, value=round(sum(v for _, v in pot_rows), 4))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ParseTests(TestCase):
    def test_bharaths_real_numbers(self):
        """0.7275 of 0.80 = 90.9% of section = Above target; 0.1970 of 0.20 = 98.5%."""
        f = _workbook([('Leadership', 0.2), ('Business', 0.1875),
                       ('Building Relationship', 0.19), ('Self', 0.15)],
                      [('Skills', 0.045), ('Values', 0.152)])
        got = parse_workbook(f)
        self.assertEqual(got['error'], '')
        self.assertAlmostEqual(got['raw_performance'], 0.7275, places=4)
        self.assertAlmostEqual(got['raw_potential'], 0.197, places=4)
        self.assertGreaterEqual(got['performance'], 2 / 3)   # HIGH band
        self.assertGreaterEqual(got['potential'], 2 / 3)     # HIGH band
        self.assertAlmostEqual(got['overall'], 92.45, places=1)


class BandNormaliseTests(TestCase):
    """The workbook's Point Scale (10-20/30-40/50-60/70-80 of 80) and its Box Grid
    Guide (Above / On / Below target) — mapped onto the grid's thirds.

    A straight raw/max put 45% of the company in 'Star' and nobody below Core
    Player on the first live run. These pin the boundaries."""

    def test_seventy_points_is_exactly_the_high_boundary(self):
        # 70 of the section's 80 points = 87.5%
        self.assertAlmostEqual(band_normalise(0.70, 0.80), 2 / 3, places=3)

    def test_fifty_points_is_exactly_the_medium_boundary(self):
        self.assertAlmostEqual(band_normalise(0.50, 0.80), 1 / 3, places=3)

    def test_a_mid_range_score_lands_medium_not_high(self):
        # 0.61 of 0.80 = 76% -> On target. Straight division gave 0.7625 = HIGH.
        v = band_normalise(0.61, 0.80)
        self.assertGreaterEqual(v, 1 / 3)
        self.assertLess(v, 2 / 3)

    def test_a_weak_score_lands_low(self):
        self.assertLess(band_normalise(0.32, 0.80), 1 / 3)     # 32 points -> Below target

    def test_full_and_zero_marks(self):
        self.assertEqual(band_normalise(0.80, 0.80), 1.0)
        self.assertEqual(band_normalise(0.0, 0.80), 0.0)

    def test_potential_uses_the_same_rule_on_its_own_max(self):
        # 0.152 of 0.20 = 76% of the section -> On target (below the 87.5% high mark).
        v = band_normalise(0.152, 0.20)
        self.assertGreaterEqual(v, 1 / 3)
        self.assertLess(v, 2 / 3)

    def test_a_zero_max_never_divides_by_zero(self):
        self.assertEqual(band_normalise(0.5, 0), 0.0)


class ParseTests2(TestCase):

    def test_reads_name_department_and_position(self):
        got = parse_workbook(_workbook([('L', 0.4)], [('V', 0.1)]))
        self.assertEqual(got['name'], 'Bharath Balasubramanian')
        self.assertEqual(got['department'], 'Finance & Planning')
        self.assertEqual(got['position'], 'Financial Controller')

    def test_scores_are_capped_at_full_marks(self):
        f = _workbook([('L', 0.95)], [('V', 0.4)])       # above both maxima
        got = parse_workbook(f)
        self.assertEqual(got['performance'], 1.0)
        self.assertEqual(got['potential'], 1.0)

    def test_a_workbook_without_the_sheet_is_reported_not_guessed(self):
        got = parse_workbook(_workbook([('L', 0.4)], [('V', 0.1)], sheet_title='Summary'))
        self.assertIn('Nine Box', got['error'])
        self.assertIsNone(got['performance'])

    def test_a_corrupt_file_is_reported_not_raised(self):
        got = parse_workbook(io.BytesIO(b'this is not a workbook'))
        self.assertTrue(got['error'])
        self.assertIsNone(got['performance'])

    def test_maxima_are_the_documented_80_20_split(self):
        self.assertEqual(PERFORMANCE_MAX + POTENTIAL_MAX, 1.0)
        self.assertEqual((PERFORMANCE_MAX, POTENTIAL_MAX), (0.80, 0.20))


class PermanentLockTests(TestCase):
    """CFO 2026-07-31: "Nobody should be able to delete it, including me or you"."""

    def test_a_locked_dialogue_cannot_be_deleted(self):
        row = DevelopmentDialogue.objects.create(ref='x::FY2025', name='Someone',
                                                 period='FY2025', locked=True)
        with self.assertRaises(PermissionError):
            row.delete()
        self.assertTrue(DevelopmentDialogue.objects.filter(pk=row.pk).exists())

    def test_the_error_names_the_person_and_period(self):
        row = DevelopmentDialogue.objects.create(ref='y::FY2025', name='Bonno Ben',
                                                 period='FY2025', locked=True)
        with self.assertRaises(PermissionError) as ctx:
            row.delete()
        self.assertIn('Bonno Ben', str(ctx.exception))
        self.assertIn('FY2025', str(ctx.exception))

    def test_bulk_queryset_delete_cannot_bypass_the_lock(self):
        """Django's QuerySet.delete() never calls Model.delete(). Without the
        pre_delete signal, one ORM one-liner wipes signed-off reviews — the same
        footgun class that destroyed 6 real records once (Fable, 2026-07-31)."""
        DevelopmentDialogue.objects.create(ref='bulk::FY2025', name='Locked One',
                                           period='FY2025', locked=True)
        # The raise aborts the delete's own transaction, so scope it and check after.
        with self.assertRaises(PermissionError):
            with transaction.atomic():
                DevelopmentDialogue.objects.filter(period='FY2025').delete()
        self.assertTrue(DevelopmentDialogue.objects.filter(ref='bulk::FY2025').exists())

    def test_an_open_dialogue_can_still_be_removed(self):
        row = DevelopmentDialogue.objects.create(ref='z::FY2026', name='Draft',
                                                 period='FY2026', locked=False)
        row.delete()
        self.assertFalse(DevelopmentDialogue.objects.filter(pk=row.pk).exists())
