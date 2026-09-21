"""
claims/services.py

CSV / XLSX parser + commit logic for the Subrogation / Salvage importers.

Append-only: rows whose claim_reference already exists are SKIPPED, not
overwritten. The CFO's rule: imports must never erase existing data.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from typing import Optional

from django.contrib.auth.models import User
from django.db import transaction

from core.models import Company

from .models import RecoveryImportBatch, Salvage, Subrogation


ZERO = Decimal('0.00')


def _looks_like_xlsx(raw: bytes) -> bool:
    return isinstance(raw, (bytes, bytearray)) and len(raw) >= 4 and bytes(raw[:4]) == b'PK\x03\x04'


def _to_decimal(value) -> Optional[Decimal]:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    for ch in [',', ' ', 'P', 'BWP', 'R', '$']:
        s = s.replace(ch, '')
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _to_date(value) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y', '%Y/%m/%d', '%d.%m.%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


SUBROGATION_ALIASES = {
    'claim_reference':     ['claim ref', 'claim reference', 'claim_ref', 'claim no', 'claim number', 'reference'],
    'incident_date':       ['incident date', 'date of loss', 'incident_date', 'loss date'],
    'third_party_name':    ['third party', 'third-party name', 'third_party', 'tp name', 'recovery from', 'name'],
    'third_party_insurer': ['tp insurer', 'third party insurer', 'insurer', 'third-party insurer'],
    'claim_paid_amount':   ['paid', 'claim paid', 'amount paid', 'claim_paid', 'gross paid', 'paid amount'],
    'expected_recovery':   ['expected', 'expected recovery', 'estimated recovery', 'recoverable'],
    'actual_recovery':     ['actual', 'recovered', 'actual recovery', 'amount recovered'],
    'last_recovery_date':  ['recovery date', 'last recovery', 'recovered on'],
    'status':              ['status', 'state'],
    'notes':               ['notes', 'comments', 'remarks'],
}

SALVAGE_ALIASES = {
    'claim_reference':   ['claim ref', 'claim reference', 'claim_ref', 'claim no', 'claim number', 'reference'],
    'incident_date':     ['incident date', 'date of loss', 'loss date'],
    'asset_description': ['asset', 'asset description', 'description', 'item', 'salvage item'],
    'estimated_value':   ['estimated value', 'estimate', 'pre-sale value', 'estimated_value'],
    'sale_proceeds':     ['sale proceeds', 'proceeds', 'sold for', 'sale amount', 'realised'],
    'sale_date':         ['sale date', 'sold on', 'date sold'],
    'buyer_name':        ['buyer', 'purchaser', 'sold to', 'buyer name'],
    'status':            ['status', 'state'],
    'notes':             ['notes', 'comments', 'remarks'],
}


def _normalise(h):
    return (h or '').strip().lower().replace('  ', ' ')


def _build_header_map(headers, aliases):
    norm = [_normalise(h) for h in headers]
    out = {}
    for canonical, alts in aliases.items():
        for alt in alts:
            a = _normalise(alt)
            for i, h in enumerate(norm):
                if h == a and canonical not in out:
                    out[canonical] = i
                    break
            if canonical in out:
                break
    return out


def _read_rows(raw: bytes, file_name: str = '') -> tuple[list[str], list[list]]:
    """Returns (headers, data_rows) regardless of CSV or XLSX format."""
    if _looks_like_xlsx(raw) or file_name.lower().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        from io import BytesIO
        wb = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return [], []
        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration:
            return [], []
        headers = ['' if c is None else str(c) for c in header_row]
        data = []
        for r in it:
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in r):
                continue
            data.append(['' if c is None else c for c in r])
        return headers, data

    # CSV
    text = ''
    for enc in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect=dialect)
    try:
        headers = next(reader)
    except StopIteration:
        return [], []
    return headers, list(reader)


def parse_subrogation_file(file_obj, file_name='') -> tuple[list[dict], list[dict]]:
    return _parse(file_obj, file_name, SUBROGATION_ALIASES, kind='subrogation')


def parse_salvage_file(file_obj, file_name='') -> tuple[list[dict], list[dict]]:
    return _parse(file_obj, file_name, SALVAGE_ALIASES, kind='salvage')


def _parse(file_obj, file_name, aliases, *, kind) -> tuple[list[dict], list[dict]]:
    raw = file_obj.read() if hasattr(file_obj, 'read') else file_obj
    raw_bytes = raw if isinstance(raw, (bytes, bytearray)) else raw.encode('utf-8')
    errors: list[dict] = []

    headers, data_rows = _read_rows(bytes(raw_bytes), file_name)
    if not headers:
        errors.append({'row_index': -1, 'field': 'file', 'message': 'File is empty.'})
        return [], errors

    hmap = _build_header_map(headers, aliases)
    if 'claim_reference' not in hmap:
        errors.append({
            'row_index': -1, 'field': 'headers',
            'message': f'Missing required claim_reference column. Headers seen: {headers}',
        })
        return [], errors

    out: list[dict] = []
    for idx, row in enumerate(data_rows, start=1):
        def get(field_name):
            i = hmap.get(field_name)
            if i is None or i >= len(row):
                return None
            v = row[i]
            return v.strip() if isinstance(v, str) else v

        def numeric(field_name):
            v = _to_decimal(get(field_name))
            return str(v) if v is not None else '0'

        def isodate(field_name):
            d = _to_date(get(field_name))
            return d.isoformat() if d else None

        if kind == 'subrogation':
            out.append({
                'row_index':           idx,
                'claim_reference':     str(get('claim_reference') or '').strip(),
                'incident_date':       isodate('incident_date'),
                'third_party_name':    str(get('third_party_name') or '').strip(),
                'third_party_insurer': str(get('third_party_insurer') or '').strip(),
                'claim_paid_amount':   numeric('claim_paid_amount'),
                'expected_recovery':   numeric('expected_recovery'),
                'actual_recovery':     numeric('actual_recovery'),
                'last_recovery_date':  isodate('last_recovery_date'),
                'status':              (str(get('status') or 'pending')
                                        .strip().lower().replace(' ', '_').replace('-', '_'))[:20] or 'pending',
                'notes':               str(get('notes') or '').strip(),
            })
        else:
            out.append({
                'row_index':         idx,
                'claim_reference':   str(get('claim_reference') or '').strip(),
                'incident_date':     isodate('incident_date'),
                'asset_description': str(get('asset_description') or '').strip(),
                'estimated_value':   numeric('estimated_value'),
                'sale_proceeds':     numeric('sale_proceeds'),
                'sale_date':         isodate('sale_date'),
                'buyer_name':        str(get('buyer_name') or '').strip(),
                'status':            (str(get('status') or 'pending')
                                      .strip().lower().replace(' ', '_').replace('-', '_'))[:20] or 'pending',
                'notes':             str(get('notes') or '').strip(),
            })
    return out, errors


def validate_subrogation_rows(rows):
    errors: list[dict] = []
    for r in rows:
        if not r.get('claim_reference'):
            errors.append({'row_index': r['row_index'], 'field': 'claim_reference', 'message': 'Missing.'})
        if not r.get('third_party_name'):
            errors.append({'row_index': r['row_index'], 'field': 'third_party_name', 'message': 'Missing.'})
        if (Decimal(r.get('claim_paid_amount') or '0')) <= ZERO:
            errors.append({'row_index': r['row_index'], 'field': 'claim_paid_amount', 'message': 'Must be > 0.'})
    return errors


def validate_salvage_rows(rows):
    errors: list[dict] = []
    for r in rows:
        if not r.get('claim_reference'):
            errors.append({'row_index': r['row_index'], 'field': 'claim_reference', 'message': 'Missing.'})
        if not r.get('asset_description'):
            errors.append({'row_index': r['row_index'], 'field': 'asset_description', 'message': 'Missing.'})
    return errors


@transaction.atomic
def commit_subrogation_import(batch: RecoveryImportBatch, user: User):
    if batch.kind != RecoveryImportBatch.Kind.SUBROGATION:
        raise ValueError('Wrong batch kind for subrogation commit.')
    return _commit(batch, user, model=Subrogation, kind='subrogation')


@transaction.atomic
def commit_salvage_import(batch: RecoveryImportBatch, user: User):
    if batch.kind != RecoveryImportBatch.Kind.SALVAGE:
        raise ValueError('Wrong batch kind for salvage commit.')
    return _commit(batch, user, model=Salvage, kind='salvage')


def _commit(batch, user, *, model, kind):
    """Append-only — duplicates by claim_reference are skipped."""
    existing = set(model.objects.values_list('claim_reference', flat=True))
    created = 0
    skipped = 0
    errors: list[dict] = []

    for r in batch.parsed_rows:
        ref = (r.get('claim_reference') or '').strip()
        if not ref:
            errors.append({'row_index': r['row_index'], 'field': 'claim_reference',
                           'message': 'Missing — row not imported.'})
            continue
        if ref in existing:
            skipped += 1
            continue

        try:
            if kind == 'subrogation':
                obj = Subrogation(
                    claim_reference=ref,
                    incident_date=date.fromisoformat(r['incident_date']) if r.get('incident_date') else None,
                    third_party_name=r.get('third_party_name') or '',
                    third_party_insurer=r.get('third_party_insurer') or '',
                    claim_paid_amount=Decimal(r.get('claim_paid_amount') or '0'),
                    expected_recovery=Decimal(r.get('expected_recovery') or '0'),
                    actual_recovery=Decimal(r.get('actual_recovery') or '0'),
                    last_recovery_date=date.fromisoformat(r['last_recovery_date']) if r.get('last_recovery_date') else None,
                    status=r.get('status') or 'pending',
                    notes=r.get('notes') or '',
                    company=batch.company,
                    created_by=user,
                )
            else:
                obj = Salvage(
                    claim_reference=ref,
                    incident_date=date.fromisoformat(r['incident_date']) if r.get('incident_date') else None,
                    asset_description=r.get('asset_description') or '',
                    estimated_value=Decimal(r.get('estimated_value') or '0'),
                    sale_proceeds=Decimal(r.get('sale_proceeds') or '0'),
                    sale_date=date.fromisoformat(r['sale_date']) if r.get('sale_date') else None,
                    buyer_name=r.get('buyer_name') or '',
                    status=r.get('status') or 'pending',
                    notes=r.get('notes') or '',
                    company=batch.company,
                    created_by=user,
                )
            obj.save(audit_user=user, audit_description=f'Imported via batch {batch.id}')
            created += 1
            existing.add(ref)
        except Exception as exc:  # noqa: BLE001
            errors.append({'row_index': r['row_index'], 'field': '__row__', 'message': str(exc)})

    batch.rows_imported = created
    batch.rows_skipped_dup = skipped
    batch.status = RecoveryImportBatch.Status.COMMITTED
    from django.utils import timezone
    batch.committed_at = timezone.now()
    batch.save(audit_user=user, audit_description=f'Committed: {created} new, {skipped} skipped')

    return created, skipped, errors
