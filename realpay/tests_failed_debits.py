"""Tests for the weekly failed-debits chase list (CFO 2026-09-11).

What is worth pinning here is not "does it build a spreadsheet". It is the small
number of decisions that, if they quietly changed, would send the wrong thing to
Underwriting and the agents every Monday without anybody noticing for weeks.
Each test below exists because getting it wrong is SILENT.
"""
from __future__ import annotations

import datetime
from decimal import Decimal
from unittest import mock

from django.contrib.auth.models import User
from django.core import mail
from django.core.management import call_command
from django.test import TestCase
from rest_framework.test import APIClient

from realpay import failed_debits, failed_debits_report
from reporting.models import ReportRecipient


def _built(policy='DOMG2026212170', amount=594.0, retries=0, pstatus='Active'):
    """A row shaped the way reporting.finance_monitoring.build_failed_debits emits."""
    return {
        'policy_number': policy, 'result': 'Failed', 'amount': amount,
        'action_date': '2026-09-08T08:52:21Z', 'retry_count': retries,
        'policy_status': pstatus, 'reason_plain': 'No money in the account',
        'reason_mapped': True,
        'list_bucket': 'call_centre' if retries >= 2 else 'finance',
        'severity': 'warning',
    }


def _patch_source(rows, agents=None):
    """Patch the daily report builder and the agent lookup this sits on."""
    return (
        mock.patch('reporting.finance_monitoring.build_failed_debits',
                   return_value={'rows': rows, 'summary': {}, 'meta': {}}),
        mock.patch.object(failed_debits, '_agents_for',
                          return_value=agents if agents is not None else {}),
    )


class ScopeTests(TestCase):
    """It covers the book a person can actually chase, and nothing else."""

    def test_the_instant_book_is_excluded(self):
        # MIS fails in the thousands weekly and nobody phones those clients one
        # by one. Leaked in, they would bury the rows somebody can act on and
        # the email would stop being read inside a month.
        self.assertTrue(failed_debits.is_chaseable('DOMG2026212170'))
        self.assertTrue(failed_debits.is_chaseable('COMG2026212072'))
        self.assertTrue(failed_debits.is_chaseable('DOM2020000222'))
        self.assertFalse(failed_debits.is_chaseable('MIS2024087054'))
        self.assertFalse(failed_debits.is_chaseable(''))

    def test_mis_rows_are_dropped_from_the_result(self):
        rows = [_built('MIS2024087054'), _built('DOMG2026212170')]
        p1, p2 = _patch_source(rows)
        with p1, p2:
            out = failed_debits.failed_rows()
        self.assertEqual([r['policy_number'] for r in out], ['DOMG2026212170'])

    def test_the_window_is_inclusive_and_consecutive_runs_do_not_skip_a_day(self):
        start, end = failed_debits.window(7, datetime.date(2026, 9, 14))
        self.assertEqual(start, datetime.date(2026, 9, 8))
        self.assertEqual(end, datetime.date(2026, 9, 14))
        prev_start, prev_end = failed_debits.window(7, datetime.date(2026, 9, 7))
        # The previous run ends the day before this one starts: no gap, no overlap.
        self.assertEqual(prev_end + datetime.timedelta(days=1), start)


class UnavailableTests(TestCase):
    """"We could not look" must never be reported as "there were none"."""

    def test_a_dead_replica_raises_rather_than_returning_empty(self):
        from reporting.finance_monitoring import ReplicaUnavailable
        with mock.patch('reporting.finance_monitoring.build_failed_debits',
                        side_effect=ReplicaUnavailable('replica down')):
            with self.assertRaises(failed_debits.GraphiteUnavailable):
                failed_debits.failed_rows()

    def test_the_command_sends_nothing_when_the_replica_is_down(self):
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='debtors@alphadirect.co.bw')
        with mock.patch.object(failed_debits, 'failed_rows',
                               side_effect=failed_debits.GraphiteUnavailable('x')):
            with self.assertRaises(SystemExit) as cm:
                call_command('send_failed_debits_report')
        self.assertEqual(cm.exception.code, 2)
        self.assertEqual(len(mail.outbox), 0)


class ShapeTests(TestCase):
    def test_a_row_carries_what_the_chaser_needs(self):
        agents = {'DOMG2026212170': {'agent': 'Saba Sultan-Chalira',
                                     'agency': 'Alpha Direct', 'premium': '594.00'}}
        p1, p2 = _patch_source([_built()], agents)
        with p1, p2:
            out = failed_debits.failed_rows()[0]
        self.assertEqual(out['policy_number'], 'DOMG2026212170')
        self.assertEqual(out['action_date'], '2026-09-08')   # not the Z timestamp
        self.assertEqual(out['amount'], Decimal('594.00'))
        self.assertEqual(out['agent'], 'Saba Sultan-Chalira')
        # Plain English from the shared bank-code dictionary, not a raw code.
        self.assertEqual(out['reason'], 'No money in the account')

    def test_a_policy_with_no_agent_is_still_reported(self):
        p1, p2 = _patch_source([_built()], {})
        with p1, p2:
            out = failed_debits.failed_rows()
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0]['agent'], '')
        self.assertEqual(failed_debits.summarise(out)['by_agent'][0]['agent'],
                         'Unassigned')

    def test_rows_are_ordered_biggest_first(self):
        rows = [_built('DOMG1', 100.0), _built('COMG2', 900.0), _built('DOM3', 500.0)]
        p1, p2 = _patch_source(rows)
        with p1, p2:
            out = failed_debits.failed_rows()
        self.assertEqual([r['policy_number'] for r in out],
                         ['COMG2', 'DOM3', 'DOMG1'])


class SummaryTests(TestCase):
    def _rows(self, specs):
        p1, p2 = _patch_source([_built(*s) for s in specs])
        with p1, p2:
            return failed_debits.failed_rows()

    def test_the_total_equals_the_sum_of_the_rows(self):
        out = self._rows([('DOMG1', 206.84), ('DOMG2', 551.61), ('COMG3', 1314.0)])
        s = failed_debits.summarise(out)
        self.assertEqual(s['count'], 3)
        self.assertEqual(s['amount'], Decimal('2072.45'))

    def test_non_active_policies_are_counted_apart(self):
        out = self._rows([('DOMG1', 10.0, 0, 'Active'),
                          ('DOMG2', 10.0, 0, 'Cancelled'),
                          ('DOMG3', 10.0, 0, 'Lapsed')])
        # Chasing a cancelled policy for money is the wrong job; the mandate
        # needs stopping instead, so the two are never merged into one figure.
        self.assertEqual(failed_debits.summarise(out)['non_active'], 2)

    def test_repeat_failures_are_counted(self):
        out = self._rows([('DOMG1', 10.0, 0), ('DOMG2', 10.0, 2), ('DOMG3', 10.0, 3)])
        self.assertEqual(failed_debits.summarise(out)['repeat'], 2)


class BrokerSplitTests(TestCase):
    """The 2026-09-12 change: the report is grouped by BROKER (the agency on
    the policy), with a per-broker subtotal, so a broker's own list can be
    lifted straight out of the sheet — Bokani Makosha's request, met by
    changing the existing report rather than building a second one."""

    def _rows(self, specs):
        """specs: list of (policy, amount, agency, agent)."""
        built = [_built(p, a) for p, a, _, _ in specs]
        agents = {p: {'agent': agent, 'agency': agency, 'premium': '0'}
                 for p, _, agency, agent in specs}
        p1, p2 = _patch_source(built, agents)
        with p1, p2:
            return failed_debits.failed_rows()

    def test_summarise_groups_by_broker_not_by_agent(self):
        rows = self._rows([
            ('DOMG1', 100.0, 'Aware Brokers', 'Kalvin Kimani'),
            ('DOMG2', 200.0, 'Aware Brokers', 'Modiri Katai'),
            ('COMG3', 50.0, 'Centric Sure', 'Kalvin Kimani'),
        ])
        s = failed_debits.summarise(rows)
        by_broker = {b['broker']: b for b in s['by_broker']}
        # Two different agents at the same broker land in ONE broker bucket.
        self.assertEqual(by_broker['Aware Brokers']['count'], 2)
        self.assertEqual(by_broker['Aware Brokers']['amount'], Decimal('300.00'))
        self.assertEqual(by_broker['Centric Sure']['amount'], Decimal('50.00'))
        # Biggest broker first.
        self.assertEqual(s['by_broker'][0]['broker'], 'Aware Brokers')

    def test_a_policy_with_no_broker_is_grouped_as_unassigned(self):
        rows = self._rows([('DOMG1', 10.0, '', 'Someone')])
        s = failed_debits.summarise(rows)
        self.assertEqual(s['by_broker'][0]['broker'], 'Unassigned')

    def test_the_workbook_groups_each_brokers_rows_together_with_a_subtotal(self):
        rows = self._rows([
            ('DOMG1', 100.0, 'Aware Brokers', 'Kalvin Kimani'),
            ('COMG2', 900.0, 'Centric Sure', 'Modiri Katai'),
            ('DOM3', 50.0, 'Aware Brokers', 'Kalvin Kimani'),
        ])
        summary = failed_debits.summarise(rows)
        buf = failed_debits_report.build_xlsx(
            rows, summary, datetime.date(2026, 9, 5), datetime.date(2026, 9, 11))
        from openpyxl import load_workbook
        ws = load_workbook(buf).active
        col_a = [c.value for c in ws['A'] if c.value not in (None, '')]
        # Centric Sure is the bigger broker (900) so its block comes first;
        # within a broker the existing biggest-first row order still holds;
        # each broker's block ends with its own subtotal row, before the next
        # broker's rows start — i.e. a broker's own list is contiguous.
        centric_idx = col_a.index('COMG2')
        centric_subtotal_idx = next(
            i for i, v in enumerate(col_a) if v == 'Subtotal — Centric Sure')
        aware_first_idx = col_a.index('DOMG1')
        self.assertLess(centric_idx, centric_subtotal_idx)
        self.assertLess(centric_subtotal_idx, aware_first_idx)
        # And the subtotal is the broker's own total, not the grand total.
        subtotal_row = next(row for row in ws.iter_rows()
                            if row[0].value == 'Subtotal — Aware Brokers')
        self.assertEqual(subtotal_row[2].value, 150.0)


class RecipientListTests(TestCase):
    def test_the_same_address_cannot_be_listed_twice_for_one_report(self):
        # Two rows means two copies, which reads to the recipient as the system
        # double-sending and to Finance as a bug in the report.
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='Rose@alphadirect.co.bw')
        with self.assertRaises(BaseException):
            ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                           email='rose@alphadirect.co.bw')

    def test_the_same_address_may_be_on_two_different_reports(self):
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='rose@alphadirect.co.bw')
        ReportRecipient.objects.create(report_slug='failed-debits',
                                       email='rose@alphadirect.co.bw')
        self.assertEqual(ReportRecipient.objects.count(), 2)

    def test_switched_off_recipients_do_not_receive(self):
        ReportRecipient.objects.create(report_slug='r', email='on@alphadirect.co.bw')
        ReportRecipient.objects.create(report_slug='r', email='off@alphadirect.co.bw',
                                       active=False)
        self.assertEqual(ReportRecipient.route_for('r')['to'],
                         ['on@alphadirect.co.bw'])

    def test_a_list_with_nobody_addressed_is_treated_as_not_set_up(self):
        # Otherwise a report goes out every Monday with an empty To line, copied
        # to two people, because somebody picked the wrong option once.
        ReportRecipient.objects.create(report_slug='r', email='cc@alphadirect.co.bw',
                                       kind=ReportRecipient.Kind.CC)
        self.assertIsNone(ReportRecipient.route_for('r'))

    def test_the_existing_daily_reports_keep_working_with_no_rows(self):
        # Turning database-held lists on must not silently stop a report that is
        # delivering correctly today.
        from reporting.management.commands.send_finance_monitoring import (
            REPORT_RECIPIENTS, _route_for,
        )
        self.assertEqual(_route_for('failed-debits'),
                         REPORT_RECIPIENTS['failed-debits'])

    def test_a_database_list_overrides_the_hardcoded_one(self):
        from reporting.management.commands.send_finance_monitoring import _route_for
        ReportRecipient.objects.create(report_slug='failed-debits',
                                       email='new@alphadirect.co.bw')
        self.assertEqual(_route_for('failed-debits')['to'],
                         ['new@alphadirect.co.bw'])


class SendTests(TestCase):
    def setUp(self):
        # The health gate is exercised by its own tests below; these ones are
        # about delivery, so the book is assumed to be reporting normally.
        patcher = mock.patch.object(failed_debits, 'book_is_reporting',
                                    return_value={'reporting': True,
                                                  'outcomes': 5, 'no_outcome': 0})
        patcher.start()
        self.addCleanup(patcher.stop)

    def _rows(self):
        p1, p2 = _patch_source([_built()],
                               {'DOMG2026212170': {'agent': 'Saba Sultan-Chalira',
                                                   'agency': 'Alpha Direct',
                                                   'premium': '594.00'}})
        with p1, p2:
            return failed_debits.failed_rows()

    def test_it_refuses_to_send_to_nobody(self):
        # An empty list is a configuration mistake, not a quiet success. Exiting
        # 0 here would let the job look healthy for weeks while nobody was told.
        with mock.patch.object(failed_debits, 'failed_rows', return_value=self._rows()):
            with self.assertRaises(SystemExit) as cm:
                call_command('send_failed_debits_report')
        self.assertEqual(cm.exception.code, 3)
        self.assertEqual(len(mail.outbox), 0)

    def test_it_sends_to_the_list_with_the_spreadsheet_attached(self):
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='uw@alphadirect.co.bw')
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='debtors@alphadirect.co.bw',
                                       kind=ReportRecipient.Kind.CC)
        with mock.patch.object(failed_debits, 'failed_rows', return_value=self._rows()):
            call_command('send_failed_debits_report')
        self.assertEqual(len(mail.outbox), 1)
        msg = mail.outbox[0]
        self.assertIn('uw@alphadirect.co.bw', msg.to)
        self.assertIn('debtors@alphadirect.co.bw', msg.cc)
        self.assertEqual(len(msg.attachments), 1)
        self.assertTrue(msg.attachments[0][0].endswith('.xlsx'))

    def test_dry_run_sends_nothing(self):
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='uw@alphadirect.co.bw')
        with mock.patch.object(failed_debits, 'failed_rows', return_value=self._rows()):
            call_command('send_failed_debits_report', '--dry-run')
        self.assertEqual(len(mail.outbox), 0)

    def test_a_clean_week_still_reports(self):
        # Silence on a good week is indistinguishable from the job having died.
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='uw@alphadirect.co.bw')
        with mock.patch.object(failed_debits, 'failed_rows', return_value=[]):
            call_command('send_failed_debits_report')
        self.assertEqual(len(mail.outbox), 1)

    def test_no_customer_contact_details_leave_the_building(self):
        # The spreadsheet this replaces carried the client's name and email.
        self.assertNotIn('Client email', failed_debits_report.COLUMNS)
        self.assertNotIn('Client name', failed_debits_report.COLUMNS)
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='uw@alphadirect.co.bw')
        with mock.patch.object(failed_debits, 'failed_rows', return_value=self._rows()):
            call_command('send_failed_debits_report')
        body = (mail.outbox[0].body + str(mail.outbox[0].alternatives)).lower()
        # Deliberately invented strings. A REAL customer address was used here
        # first and the /fabe PII tripwire caught it — a live client's email
        # does not belong in a test file, even in an assertion that it must be
        # absent. The guard is the same; the example is not a real person.
        for banned in ('omang', 'client.example@invalid.test', 'account number'):
            self.assertNotIn(banned, body)


class ApiTests(TestCase):
    URL = '/api/v1/report-recipients/'

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user('acct', password='x',
                                             is_superuser=True, is_staff=True)
        self.client.force_authenticate(self.user)

    def test_an_accountant_can_add_and_switch_off_a_recipient(self):
        r = self.client.post(self.URL, {'email': 'rose@alphadirect.co.bw',
                                        'name': 'Rose'}, format='json')
        self.assertEqual(r.status_code, 201)
        rid = r.json()['id']
        d = self.client.delete(f'{self.URL}{rid}/')
        self.assertEqual(d.status_code, 200)
        self.assertFalse(d.json()['active'])
        # Switched off, not erased — who received a financial report survives.
        self.assertTrue(ReportRecipient.objects.filter(id=rid).exists())

    def test_a_bad_address_is_refused_at_the_screen(self):
        r = self.client.post(self.URL, {'email': 'rose at alphadirect'},
                             format='json')
        self.assertEqual(r.status_code, 400)

    def test_re_adding_a_switched_off_address_turns_it_back_on(self):
        row = ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                             email='rose@alphadirect.co.bw',
                                             active=False)
        r = self.client.post(self.URL, {'email': 'ROSE@alphadirect.co.bw'},
                             format='json')
        self.assertEqual(r.status_code, 200)
        row.refresh_from_db()
        self.assertTrue(row.active)
        self.assertEqual(ReportRecipient.objects.count(), 1)

    def test_an_unknown_report_is_refused(self):
        # A typo would otherwise create a list for a report that does not exist:
        # it looks saved on screen and silently never sends.
        r = self.client.post(self.URL, {'email': 'a@b.com',
                                        'report_slug': 'typo'}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_signing_out_blocks_the_list(self):
        self.client.force_authenticate(None)
        self.assertIn(self.client.get(self.URL).status_code, (401, 403))


class RouteResilienceTests(TestCase):
    """Reading the address list must never be able to stop a report going out."""

    def test_a_broken_recipient_lookup_falls_back_to_the_built_in_list(self):
        # Before this, resolving recipients was the one step in the send that
        # could not fail. Putting it in the database added a way for a
        # monitoring report to stop arriving silently — which is the exact
        # failure the monitoring pack exists to prevent.
        from reporting.management.commands.send_finance_monitoring import (
            REPORT_RECIPIENTS, _route_for,
        )
        with mock.patch('reporting.models.ReportRecipient.route_for',
                        side_effect=RuntimeError('database gone')):
            self.assertEqual(_route_for('failed-debits'),
                             REPORT_RECIPIENTS['failed-debits'])


class FalseAllClearTests(TestCase):
    """The failure this nearly shipped with, and the guard that stops it.

    Measured on live data 2026-09-11: Graphite holds 312 September instalments
    for the commercial/domestic book, every one status 'A' (raised, no result),
    and no success or failure recorded since May 2026 — while the Instant book
    logged 3,336 successes and 2,666 failures over the same eleven days. So an
    unguarded weekly report would have emailed "0 failed debits" for a week that
    really had eighteen worth P31,016, and it would have read as good news.
    """

    def _health(self, outcomes, no_outcome):
        rows = [{'outcomes': outcomes, 'no_outcome': no_outcome}]
        return mock.patch.object(failed_debits.graphite_ro, 'query',
                                 return_value=rows)

    def test_scheduled_debits_with_no_result_is_not_reporting(self):
        with mock.patch.object(failed_debits.graphite_ro, 'is_configured',
                               return_value=True), self._health(0, 312):
            self.assertFalse(failed_debits.book_is_reporting()['reporting'])

    def test_a_book_with_real_outcomes_is_reporting(self):
        with mock.patch.object(failed_debits.graphite_ro, 'is_configured',
                               return_value=True), self._health(183, 12):
            self.assertTrue(failed_debits.book_is_reporting()['reporting'])

    def test_a_genuinely_quiet_week_with_no_debits_due_still_counts_as_reporting(self):
        # Nothing scheduled and nothing reported is a real quiet week, not an
        # outage — it must not be mistaken for the broken case above.
        with mock.patch.object(failed_debits.graphite_ro, 'is_configured',
                               return_value=True), self._health(0, 0):
            self.assertTrue(failed_debits.book_is_reporting()['reporting'])

    def test_the_command_refuses_to_send_a_false_all_clear(self):
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='uw@alphadirect.co.bw')
        with mock.patch.object(failed_debits, 'failed_rows', return_value=[]),              mock.patch.object(failed_debits, 'book_is_reporting',
                               return_value={'reporting': False, 'outcomes': 0,
                                             'no_outcome': 312}):
            with self.assertRaises(SystemExit) as cm:
                call_command('send_failed_debits_report')
        self.assertEqual(cm.exception.code, 4)
        self.assertEqual(len(mail.outbox), 0)


class FableRoundOneTests(TestCase):
    """The four defects Fable found on the first review pass (2026-09-11).

    Each is here because it would have shipped otherwise, and each is the kind
    that nobody notices from the outside for weeks.
    """

    def _rows(self):
        p1, p2 = _patch_source([_built()], {})
        with p1, p2:
            return failed_debits.failed_rows()

    def _healthy(self):
        return mock.patch.object(failed_debits, 'book_is_reporting',
                                 return_value={'reporting': True, 'outcomes': 5,
                                               'no_outcome': 0})

    def test_the_email_never_carries_a_do_not_reply_banner(self):
        # The body asks in as many words for a reply — "reply to Debtors with
        # the proof of payment". A banner telling the reader not to reply would
        # contradict the sentence above it.
        ReportRecipient.objects.create(report_slug='failed-debits-weekly',
                                       email='uw@alphadirect.co.bw')
        with mock.patch.object(failed_debits, 'failed_rows',
                               return_value=self._rows()), self._healthy():
            call_command('send_failed_debits_report')
        html = str(mail.outbox[0].alternatives).lower()
        self.assertNotIn('do not reply', html)
        self.assertIn('reply to debtors', html)

    def test_the_spreadsheet_is_not_forced_to_whole_pula(self):
        # aware.reporting._wb formats a DECLARED money column as '#,##0'. Feeding
        # thebe into it would show 594.50 as 595 while the KPI above read
        # P 2,072.45 — rows that visibly do not sum to their own total.
        #
        # This opens the real workbook rather than reading the source, because
        # an assertion on source text would still pass if whole-pula rounding
        # came back by another route (a number_format set in a loop, say).
        # Fable made exactly that point on review, and it was right.
        import datetime as _dt
        from openpyxl import load_workbook
        rows = [{'policy_number': 'DOMG1', 'action_date': '2026-09-08',
                 'amount': Decimal('594.50'), 'retry_count': 0,
                 'reason': 'No money in the account', 'policy_status': 'Active',
                 'premium': Decimal('1234.56'), 'agency': 'A', 'agent': 'B'}]
        summary = failed_debits.summarise(rows)
        buf = failed_debits_report.build_xlsx(
            rows, summary, _dt.date(2026, 9, 5), _dt.date(2026, 9, 11))
        ws = load_workbook(buf).active
        amounts = [c for row in ws.iter_rows() for c in row
                   if c.value == 594.5]
        self.assertTrue(amounts, 'the 594.50 row is not in the workbook')
        for c in amounts:
            self.assertNotEqual(c.number_format, '#,##0')
        # And the thebe survived rather than being rounded to 595.
        self.assertEqual(amounts[0].value, 594.5)

    def test_an_address_omni_will_never_send_to_is_refused(self):
        # Otherwise the row saves, the screen shows "Receiving", and the sender
        # strips it silently — the 2026-08-10 "the email said Arun was copied
        # and he was not" incident, this time reachable from a form.
        client = APIClient()
        user = User.objects.create_user('acct2', password='x',
                                        is_superuser=True, is_staff=True)
        client.force_authenticate(user)
        r = client.post('/api/v1/report-recipients/',
                        {'email': 'admin@alphadirect.co.bw'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(ReportRecipient.objects.count(), 0)

    def test_the_never_send_check_is_case_insensitive(self):
        client = APIClient()
        user = User.objects.create_user('acct3', password='x',
                                        is_superuser=True, is_staff=True)
        client.force_authenticate(user)
        r = client.post('/api/v1/report-recipients/',
                        {'email': 'Admin@AlphaDirect.co.bw'}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_an_ordinary_address_is_still_accepted(self):
        # The guard must not become a wall.
        client = APIClient()
        user = User.objects.create_user('acct4', password='x',
                                        is_superuser=True, is_staff=True)
        client.force_authenticate(user)
        r = client.post('/api/v1/report-recipients/',
                        {'email': 'rose@alphadirect.co.bw'}, format='json')
        self.assertEqual(r.status_code, 201)


class PreviewNeverShowsAFalseAllClearTests(TestCase):
    """The screen must not say "no failed debits" when the feed is dead.

    Found by LOOKING at the post-deploy QC screenshot, not by a test — every
    test in this file mocks the data source, so all of them passed while the
    live page rendered a green tick and "No failed debits in the last 7 days"
    at the same moment the Monday job was refusing to send for exactly the
    opposite reason. A false all-clear on money owed is no better on a screen
    than in an email.
    """

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user('prev', password='x',
                                             is_superuser=True, is_staff=True)
        self.client.force_authenticate(self.user)

    def test_a_dead_feed_is_reported_as_unknown_not_as_none(self):
        with mock.patch.object(failed_debits, 'failed_rows', return_value=[]), \
             mock.patch.object(failed_debits, 'book_is_reporting',
                               return_value={'reporting': False, 'outcomes': 0,
                                             'no_outcome': 214}):
            r = self.client.get('/api/v1/realpay/failed-debits/preview/')
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertFalse(body['available'])
        self.assertIn('214', body['detail'])
        self.assertIn('not a clean week', body['detail'])
        # And it must NOT report a count of zero, which the page renders as
        # "No failed debits in the last 7 days".
        self.assertNotIn('count', body)

    def test_a_genuinely_clean_week_still_reports_zero(self):
        # The guard must not turn every quiet week into an alarm.
        with mock.patch.object(failed_debits, 'failed_rows', return_value=[]), \
             mock.patch.object(failed_debits, 'book_is_reporting',
                               return_value={'reporting': True, 'outcomes': 40,
                                             'no_outcome': 2}):
            r = self.client.get('/api/v1/realpay/failed-debits/preview/')
        body = r.json()
        self.assertTrue(body['available'])
        self.assertEqual(body['count'], 0)

    def test_a_broken_health_check_is_unknown_not_all_clear(self):
        # The first version of the preview guard fell back to reporting=True
        # when the health check itself raised — reproducing, one level down,
        # the false all-clear it was written to prevent. The review panel
        # caught it. "We could not check" is the only honest answer here.
        with mock.patch.object(failed_debits, 'failed_rows', return_value=[]), \
             mock.patch.object(failed_debits, 'book_is_reporting',
                               side_effect=failed_debits.GraphiteUnavailable('x')):
            r = self.client.get('/api/v1/realpay/failed-debits/preview/')
        body = r.json()
        self.assertFalse(body['available'])
        self.assertIn('could not check', body['detail'].lower())
        self.assertNotIn('count', body)
