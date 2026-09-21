from datetime import timedelta
from io import BytesIO

from django.conf import settings
from django.core.exceptions import ValidationError
from django.utils import timezone
from django.utils.html import escape

from core.models import AuditLog
from core.notifications import send_html_with_cfo_cc
from payroll.models import EmploymentContract, Employee
from hris.hr_settings import get_setting
from hris.models import (
    ContractRenewalDecision,
    LoginClassification,
    OffboardingCase,
    ProbationDecision,
)

BASE_URL = getattr(settings, "SITE_URL", "https://omni.alphadirect.co.bw")


def probation_due(today=None):
    """Return active contracts whose probation ends in the next 30 days and have no decision."""
    if today is None:
        today = timezone.localdate()
    start = today
    end = today + timedelta(days=30)
    return (
        EmploymentContract.objects.filter(
            status="active",
            probation_end_date__isnull=False,
            probation_end_date__gte=start,
            probation_end_date__lte=end,
        )
        .exclude(probation_decisions__isnull=False)
        .distinct()
    )


def _recently_reminded(contract):
    """Return True if a probation_reminder audit row was written in the last 7 days."""
    since = timezone.now() - timedelta(days=7)
    return AuditLog.objects.filter(
        table_name="probation_reminder",
        record_id=str(contract.id),
        created_at__gte=since,
    ).exists()


def _probation_email_html(contract):
    name = escape(contract.employee.full_name)
    date_str = escape(str(contract.probation_end_date))
    link = escape(f"{BASE_URL}/hris/contracts?probation={contract.id}")
    choices = [
        "Confirm the appointment",
        "Extend probation",
        "End employment",
    ]
    choice_items = "".join(f"<li>{escape(choice)}</li>" for choice in choices)
    return f"""<!DOCTYPE html>
<html><body style="margin:0;padding:0;background:#f5f5f5;font-family:Arial,sans-serif;">
  <div style="width:640px;margin:0 auto;background:#ffffff;">
    <div style="background:#0D1B2A;padding:20px;">
      <div style="color:#ffffff;font-size:12px;letter-spacing:2px;font-variant:small-caps;">ALPHA DIRECT · OMNI</div>
      <div style="color:#F4A623;font-size:24px;font-weight:bold;margin-top:8px;">Probation Review</div>
    </div>
    <div style="padding:20px;">
      <p>Probation for <strong>{name}</strong> ends on <strong>{date_str}</strong>.</p>
      <p>Please choose one of the following before or on the end date:</p>
      <ol>{choice_items}</ol>
      <p><a href="{link}">Review this contract in Omni</a></p>
      <p>Thank you,<br>Human Capital Department</p>
    </div>
  </div>
</body></html>"""


def run_probation_reminders(today=None, commit=False):
    """Email HR and line managers for probation contracts ending within 30 days.

    De-duplication is based on core.AuditLog rows with table_name='probation_reminder'
    within the last 7 days.
    """
    if today is None:
        today = timezone.localdate()
    recipients = get_setting("contract_reminder_recipients", ["dikgopoleng@alphadirect.co.bw", "ubutale@alphadirect.co.bw", "hc@alphadirect.co.bw"])
    due = probation_due(today)
    sent = 0
    skipped = 0

    for contract in due:
        if _recently_reminded(contract):
            skipped += 1
            continue

        to = list(recipients)
        try:
            profile = contract.employee.hris_profile
            manager = profile.manager
            if manager and manager.email:
                to.append(manager.email)
        except Exception:
            pass

        # Remove duplicates while preserving order.
        to = list(dict.fromkeys(to))

        subject = f"Probation ends {contract.probation_end_date} — {contract.employee.full_name}"
        html = _probation_email_html(contract)

        if commit:
            send_html_with_cfo_cc(subject=subject, html=html, to=to, cc_cfo=False,
                                  allow_named_exec=True)
            AuditLog.objects.create(
                table_name="probation_reminder",
                record_id=str(contract.id),
                action="create",
                old_values={},
                new_values={"recipients": to, "contract_id": str(contract.id)},
                user=None,
                description=f"Probation reminder sent for {contract.employee.full_name}",
            )
            sent += 1

    return {"due": due.count(), "sent": sent, "skipped": skipped}


def record_probation(contract, decision, user, new_end=None, note=""):
    """Record the outcome of probation. `decision` is one of confirm/extend/end."""
    today = timezone.localdate()

    if decision == ProbationDecision.Decision.EXTEND:
        if not new_end:
            raise ValidationError("new_end is required when decision is extend")
        if new_end <= contract.probation_end_date:
            raise ValidationError("new_end must be after the current probation_end_date")
        contract.probation_end_date = new_end
        contract.full_clean()
        contract.save()
    elif decision == ProbationDecision.Decision.END:
        from hris.offboarding_service import open_case

        open_case(
            contract.employee,
            reason="other",
            last_working_day=contract.probation_end_date or today,
            user=user,
        )
    elif decision == ProbationDecision.Decision.CONFIRM:
        if contract.contract_type == "probation":
            contract.contract_type = "permanent"
            contract.save()
    else:
        raise ValidationError("Invalid decision")

    probation_decision = ProbationDecision.objects.create(
        contract=contract,
        decision=decision,
        new_probation_end=new_end if decision == ProbationDecision.Decision.EXTEND else None,
        note=note,
        decided_by=user,
    )

    AuditLog.objects.create(
        table_name="probation_decision",
        record_id=str(contract.id),
        action="create",
        old_values={},
        new_values={"decision": decision, "note": note},
        user=user,
        description=f"Probation decision recorded for {contract.employee.full_name}",
    )

    return probation_decision


def follow_through_renewal(decision_obj, user):
    """Execute the follow-through action for a ContractRenewalDecision.

    Never raises; returns a dictionary, possibly containing 'error'.
    """
    old = decision_obj.contract
    try:
        decision = decision_obj.decision

        if decision == ContractRenewalDecision.Decision.RENEW_SAME:
            from hris.contract_module import add_months

            if not old.end_date:
                return {"error": "Contract has no end_date"}

            start = old.end_date + timedelta(days=1)
            duration = old.end_date - old.start_date
            end = start + duration

            if old.contract_type == "fixed_term":
                max_end = add_months(start, 12) - timedelta(days=1)
                if end > max_end:
                    end = max_end

            new_contract = EmploymentContract(
                employee=old.employee,
                start_date=start,
                end_date=end,
                basic=old.basic,
                currency_code=old.currency_code,
                frequency=old.frequency,
                grade=old.grade,
                contract_type=old.contract_type,
                retirement_fund=old.retirement_fund,
                allowance_template=old.allowance_template,
                status="pending",
            )
            new_contract.full_clean()
            new_contract.save()

            return {
                "new_contract_id": new_contract.id,
                "letter_url": f"/api/v1/hris/contracts/{new_contract.id}/letter/",
            }

        if decision == ContractRenewalDecision.Decision.CHANGE:
            return {
                "next": "/hris/amendments",
                "message": "Record the new terms as a payroll amendment.",
            }

        if decision == ContractRenewalDecision.Decision.END:
            from hris.offboarding_service import open_case

            case = open_case(
                old.employee,
                reason="end_of_contract",
                last_working_day=old.end_date,
                user=user,
            )
            return {"offboarding_case_id": case.id}

        return {"error": f"Unknown decision: {decision}"}

    except Exception as exc:
        return {"error": str(exc)}


def renewal_letter_docx(contract):
    """Return a DOCX renewal letter as bytes."""
    from docx import Document

    employee = contract.employee
    type_label = contract.contract_type.replace("_", " ").title() if contract.contract_type else "contract"
    today = timezone.localdate()

    doc = Document()
    doc.add_heading("Contract renewal", 0)
    doc.add_paragraph(f"Date: {today.strftime('%d %B %Y')}")
    doc.add_paragraph(f"Dear {employee.full_name},")
    doc.add_paragraph(
        f"We are pleased to renew your {type_label} contract as {employee.job_title} "
        f"from {contract.start_date} to {contract.end_date} on the same terms; "
        "probation not applicable; please sign and return."
    )
    doc.add_paragraph("Human Capital Department")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def orphan_logins():
    """Return active non-superusers with no active payroll Employee link (by FK or email)."""
    from django.contrib.auth import get_user_model

    User = get_user_model()

    classifications = {
        lc.user_id: lc.kind
        for lc in LoginClassification.objects.all()
    }

    rows = []
    for user in User.objects.filter(is_active=True, is_superuser=False).order_by("username"):
        if Employee.objects.filter(user=user, status__in=["active", "on_leave"]).exists():
            continue
        if Employee.objects.filter(email__iexact=user.email, status__in=["active", "on_leave"]).exists():
            continue

        rows.append(
            {
                "user_id": user.pk,
                "username": user.username,
                "email": user.email or "",
                "name": user.get_full_name() or "",
                "last_login": user.last_login.isoformat() if user.last_login else "",
                "classification": classifications.get(user.pk, ""),
            }
        )

    return rows


def classify_login(user_obj, kind, actor, note=""):
    """Record HR classification for a login. If kind=='close', deactivate secure non-protected users."""
    from core.models import UserProfile
    from core.session_teardown import clear_browser_tokens, revoke_device_sessions

    existing = LoginClassification.objects.filter(user=user_obj).first()
    old_kind = existing.kind if existing else ""

    login_classification, created = LoginClassification.objects.update_or_create(
        user=user_obj,
        defaults={"kind": kind, "note": note, "classified_by": actor},
    )

    action = "create" if created else "update"

    AuditLog.objects.create(
        table_name="login_classification",
        record_id=str(user_obj.id),
        action=action,
        old_values={"kind": old_kind} if old_kind else {},
        new_values={"kind": kind, "note": note},
        user=actor,
        description=f"Login classification for {user_obj.username}",
    )

    if (
        kind == LoginClassification.Kind.CLOSE
        and not user_obj.is_superuser
        and user_obj.email.lower() != "pganesharajah@alphadirect.co.bw"
    ):
        revoke_device_sessions(user_obj, timezone.now())
        clear_browser_tokens(user_obj)

        user_obj.is_active = False
        user_obj.save(update_fields=["is_active"])

        profile, _ = UserProfile.objects.get_or_create(user=user_obj)
        profile.is_active = False
        profile.save(update_fields=["is_active"])

    return login_classification


def _digest_data(today=None):
    """Build HR Monday digest data and return (subject, html, counts)."""
    if today is None:
        today = timezone.localdate()

    # Contracts due for renewal without any renewal decision (next 30 days)
    due_renewals = []
    for contract in EmploymentContract.objects.filter(
        status="active",
        end_date__isnull=False,
        end_date__lte=today + timedelta(days=30),
    ):
        if not ContractRenewalDecision.objects.filter(contract=contract).exists():
            due_renewals.append(contract)

    due_renewal_strs = [
        f"{escape(c.employee.full_name)} — ends {escape(str(c.end_date))}"
        for c in due_renewals
    ]

    probations = list(probation_due(today))
    probation_strs = [
        f"{escape(c.employee.full_name)} — ends {escape(str(c.probation_end_date))}"
        for c in probations
    ]

    offboarding_qs = OffboardingCase.objects.filter(status="open").select_related("employee")
    open_offboarding = offboarding_qs.count()
    offboarding_rows = []
    for case in offboarding_qs:
        outstanding = case.steps.filter(done_at__isnull=True).count()
        offboarding_rows.append(
            f"{escape(case.employee.full_name)} — {outstanding} outstanding steps"
        )

    # Joiners in the last 60 days
    joiners = Employee.objects.filter(
        hire_date__gte=today - timedelta(days=60),
        hire_date__lte=today,
    )
    joiner_rows = []
    try:
        from hris.joiner_pack import documents_status
    except ImportError:
        documents_status = None

    for employee in joiners:
        incomplete = False
        unsigned_acks = employee.acknowledgements.filter(employee_signed_at__isnull=True).exists()
        if unsigned_acks:
            incomplete = True

        if documents_status is not None:
            try:
                result = documents_status(employee)
                if isinstance(result, dict):
                    if not result.get("complete", False):
                        incomplete = True
                elif result is False:
                    incomplete = True
            except Exception:
                incomplete = True

        if incomplete:
            joiner_rows.append(escape(employee.full_name))

    orphan_rows = [row for row in orphan_logins() if row["classification"] == ""]

    m365_rows = []
    try:
        from licensing.models import M365AccountSeen

        # The first run seeded every existing account silently; those are not new.
        first = M365AccountSeen.objects.order_by('first_seen_at').values_list(
            'first_seen_at', flat=True).first()
        m365_qs = M365AccountSeen.objects.filter(
            first_seen_at__date__gte=today - timedelta(days=7),
            is_agent=False,
        )
        if first:
            m365_qs = m365_qs.filter(first_seen_at__gt=first + timedelta(hours=1))
        m365_qs = m365_qs[:20]
        for account in m365_qs:
            user = getattr(account, "user", None)
            if user:
                m365_rows.append(escape(getattr(user, "email", str(user))))
            else:
                m365_rows.append(escape(str(account)))
    except ImportError:
        m365_rows = []

    counts = {
        "due_renewals": len(due_renewals),
        "probations_ending": len(probations),
        "open_offboarding": open_offboarding,
        "joiners_needing_action": len(joiner_rows),
        "orphan_logins": len(orphan_rows),
        "new_m365_accounts": len(m365_rows),
    }

    subject = f"Monday HR Digest — {today.isoformat()}"

    html = _render_digest_html(
        today,
        due_renewal_strs,
        probation_strs,
        offboarding_rows,
        joiner_rows,
        orphan_rows,
        m365_rows,
    )

    return subject, html, counts


def _render_digest_html(
    today,
    due_renewal_strs,
    probation_strs,
    offboarding_rows,
    joiner_rows,
    orphan_rows,
    m365_rows,
):
    def section(title, rows):
        safe_title = escape(title)
        if rows:
            items = "".join(
                f'<tr><td style="padding:4px 8px;border-bottom:1px solid #e0e0e0;">{row}</td></tr>'
                for row in rows
            )
        else:
            items = '<tr><td style="padding:4px 8px;border-bottom:1px solid #e0e0e0;">None</td></tr>'
        return (
            f'<h3 style="color:#0D1B2A;margin-top:18px;margin-bottom:6px;">{safe_title}</h3>'
            f'<table style="border-collapse:collapse;width:100%;">{items}</table>'
        )

    sections = []
    sections.append(section("Contracts due for renewal without decision", due_renewal_strs))
    sections.append(section("Probations ending in 30 days", probation_strs))
    sections.append(section("Open offboarding cases", offboarding_rows))
    sections.append(section("Joiners needing action (last 60 days)", joiner_rows))
    sections.append(section("Unclassified Omni logins without payroll", orphan_rows))
    sections.append(section("New Microsoft accounts in last 7 days", m365_rows))

    body_sections = "".join(sections)

    return f"""<!DOCTYPE html>
<html><body style="margin:0;padding:0;background:#f5f5f5;font-family:Arial,sans-serif;">
  <div style="width:640px;margin:0 auto;background:#ffffff;">
    <div style="background:#0D1B2A;padding:20px;">
      <div style="color:#ffffff;font-size:12px;letter-spacing:2px;font-variant:small-caps;">ALPHA DIRECT · OMNI</div>
      <div style="color:#F4A623;font-size:24px;font-weight:bold;margin-top:8px;">Monday HR Digest</div>
    </div>
    <div style="padding:20px;">
      <p style="color:#333333;">Date: {escape(today.isoformat())}</p>
      {body_sections}
    </div>
  </div>
</body></html>"""


def digest_html(today=None):
    """Return (subject, html) for the Monday HR digest."""
    subject, html, _ = _digest_data(today)
    return subject, html


def _digest_counts(today=None):
    """Return only the counts for the Monday HR digest."""
    _, _, counts = _digest_data(today)
    return counts
