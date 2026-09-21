"""
nbfira/filed_parser.py — read the A.1 tab of a FILED NBFIRA workbook.

Why this exists: NBFIRAFiledDocument stored the filed return but never read
it, so the only figures on screen were Omni's own GL-generated ones. Finance
asked for the filed figure beside Omni's figure, per line, per period
(Oprah Mogomotsi, 2026-08-18).

Scope is deliberately A.1 only — the Prescribed Capital Target schedule.
That is the schedule the question was about, and it is the one whose layout
has been verified against all four FY2026 filed workbooks.

Two rules drive the design:

1. **Anchor on row LABELS, never on row numbers.** The four FY2026 workbooks
   happen to agree on row positions, but a regulator template gains a row and
   every fixed address silently reads its neighbour.

2. **Anchor on (section, label), not label alone.** 'Property' is BOTH an
   insurance class in section 2 and an asset bucket in section 4, and 'Total'
   appears four times. Worse, Omni's INSURANCE_CLASSES is in a different order
   from the workbook's rows, so mapping by position would put Accident's filed
   value on Property's line and still look plausible. Section-scoped label
   matching is the only safe key.

Units: the workbook states P'000 and Omni's lines are stored P'000 (see
builders._p000), so values compare directly. `_units_ok` refuses the file
rather than comparing figures 1000x apart if that ever stops being true.
"""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

from .constants import CLASS_LABEL, DEFAULT_MRC

SHEET_NAME = 'A.1'
VALUE_COL = 5          # column E — every A.1 figure of interest sits here
LABEL_COL = 1          # column A
MAX_SCAN_ROWS = 200

# The MCR figure sits a couple of rows BELOW its section header, with no label
# of its own in column A. Scan a small window rather than hard-coding E10.
MCR_LOOKAHEAD = 4


def _norm(v: Any) -> str:
    """Lowercase, collapse whitespace, drop punctuation — so 'Total  MER' and
    'Total MER' and 'TOTAL MER:' all match."""
    if v is None:
        return ''
    s = str(v).replace('\xa0', ' ')
    s = re.sub(r'[^a-z0-9 ]+', ' ', s.lower())
    return re.sub(r'\s+', ' ', s).strip()


def _dec(v: Any) -> Optional[Decimal]:
    """Excel cell → Decimal, or None when the cell is not a number.
    Text, dates and blanks all return None rather than raising."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        try:
            return Decimal(str(v)).quantize(Decimal('0.0001'))
        except (InvalidOperation, ValueError):
            return None
    s = str(v).strip().replace(',', '').replace('P', '').strip()
    if not s or s in {'-', '—'}:
        return None
    neg = s.startswith('(') and s.endswith(')')
    if neg:
        s = s[1:-1]
    try:
        d = Decimal(s)
    except (InvalidOperation, ValueError):
        return None
    return (-d if neg else d).quantize(Decimal('0.0001'))


# ── (section, normalised label) → Omni A.1 line code ─────────────────────
# Built from the SAME constants the builder uses, so the two can never drift
# apart into a comparison that silently comes out wrong.
_IRC_BY_LABEL = {_norm(lbl): f'A1_IRC_{key.upper()}'
                 for key, lbl in CLASS_LABEL.items()}
_MRC_BY_LABEL = {_norm(lbl): f'A1_MRC_{key.upper()}'
                 for key, lbl, _fac in DEFAULT_MRC}

_SECTION_LABELS: dict[str, dict[str, str]] = {
    'irc': {**_IRC_BY_LABEL,
            'total': 'A1_IRC_TOTAL',
            'irc adjusted': 'A1_IRC_ADJ'},
    'mer': {'total mer': 'A1_MER_TOTAL',
            'mer adjusted': 'A1_MER_ADJ',
            'g insurance': 'G_INSURANCE'},
    'mrc': {**_MRC_BY_LABEL,
            'total mrctc': 'A1_MRC_TOTAL',
            'mrctr adjusted': 'A1_MRC_ADJ',
            'g market': 'G_MARKET'},
}

# Numbered section headers in column A. Matched on the leading number so a
# reworded heading still lands in the right section.
_SECTION_HEADERS = [
    (re.compile(r'^1\b.*minimum capital'), 'mcr'),
    (re.compile(r'^2\b.*insurance risk'), 'irc'),
    (re.compile(r'^3\b.*maximum event'), 'mer'),
    (re.compile(r'^4\b.*market risk'), 'mrc'),
    (re.compile(r'^5\b.*prescribed capital'), 'pct'),
    (re.compile(r'^asset allocation'), 'alloc'),   # ends the mrc section
]

# g-factors are ratios, not money — they must not be compared as amounts.
RATIO_CODES = {'G_INSURANCE', 'G_MARKET'}


def _section_for(label_norm: str) -> Optional[str]:
    for pattern, name in _SECTION_HEADERS:
        if pattern.match(label_norm):
            return name
    return None


def _units_ok(values: dict[str, Decimal]) -> tuple[bool, str]:
    """The MCR is a known constant (P5,000,000 = 5,000 in P'000). Use it as the
    units witness: if it reads like full pula, refuse rather than compare
    against Omni's P'000 lines and report a 1000x variance as a finding."""
    mcr = values.get('A1_MCR_01')
    if mcr is None:
        return True, ''                     # nothing to check against
    if mcr >= Decimal('1000000'):
        return False, (
            f"A.1 appears to be in full pula (MCR reads {mcr:,.2f}), but Omni "
            f"stores this schedule in P'000. Refusing to compare — the figures "
            f"would be out by a factor of 1000."
        )
    return True, ''


def parse_a1(file_obj) -> dict:
    """Parse the A.1 tab of a filed NBFIRA workbook.

    Returns a dict always carrying 'ok'. On success:
        ok, sheet, as_at, values {line_code: Decimal}, ratios {...}, note
    On failure 'ok' is False and 'error' says why, in words Finance can read.
    """
    try:
        import openpyxl
    except ImportError:                          # pragma: no cover
        return {'ok': False, 'error': 'openpyxl is not installed on the server.'}

    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    except Exception as exc:                     # noqa: BLE001 — any bad file
        return {'ok': False,
                'error': f'Not a readable Excel workbook ({type(exc).__name__}). '
                         f'A .xlsx file is needed to compare figures; a PDF or '
                         f'CSV can be stored but not read.'}

    try:
        if SHEET_NAME not in wb.sheetnames:
            return {'ok': False,
                    'error': f'No "{SHEET_NAME}" tab in this workbook '
                             f'(found: {", ".join(wb.sheetnames)}).'}
        ws = wb[SHEET_NAME]

        # Take one full pass into memory: read_only sheets stream forward only,
        # and the MCR probe needs to look a few rows ahead of its header.
        grid: list[tuple] = list(
            ws.iter_rows(min_row=1, max_row=MAX_SCAN_ROWS,
                         max_col=VALUE_COL, values_only=True))

        values: dict[str, Decimal] = {}
        ratios: dict[str, Decimal] = {}
        as_at = ''
        section: Optional[str] = None
        rows_matched = 0

        for idx, row in enumerate(grid):
            label_raw = row[LABEL_COL - 1] if len(row) >= LABEL_COL else None
            val_raw = row[VALUE_COL - 1] if len(row) >= VALUE_COL else None
            label = _norm(label_raw)

            if not as_at and label.startswith('as at'):
                as_at = str(label_raw).strip()

            new_section = _section_for(label)
            if new_section:
                section = new_section
                if section == 'pct':
                    # PCT carries its value on the header row itself.
                    d = _dec(val_raw)
                    if d is not None:
                        values['A1_PCT'] = d
                        rows_matched += 1
                elif section == 'mcr':
                    # MCR sits a row or two below, unlabelled.
                    for probe in grid[idx + 1: idx + 1 + MCR_LOOKAHEAD]:
                        d = _dec(probe[VALUE_COL - 1]
                                 if len(probe) >= VALUE_COL else None)
                        if d is not None:
                            values['A1_MCR_01'] = d
                            rows_matched += 1
                            break
                continue

            if not section or not label:
                continue

            code = _SECTION_LABELS.get(section, {}).get(label)
            if not code:
                continue
            d = _dec(val_raw)
            if d is None:
                continue
            if code in RATIO_CODES:
                ratios[code] = d
            else:
                values[code] = d
            rows_matched += 1
    finally:
        wb.close()

    if 'A1_PCT' not in values:
        return {'ok': False,
                'error': 'Found the A.1 tab but no Prescribed Capital Target '
                         'line (section 5). The layout is not the one Omni '
                         'knows — nothing compared, rather than compared wrongly.'}

    ok_units, why = _units_ok(values)
    if not ok_units:
        return {'ok': False, 'error': why}

    return {
        'ok': True,
        'sheet': SHEET_NAME,
        'as_at': as_at,
        'values': values,
        'ratios': ratios,
        'note': f'Parsed {rows_matched} A.1 lines from the filed workbook.',
    }


# ── reading the INPUTS, not just the answers ──────────────────────────────
# The A.1 answers are worth comparing, but Finance also needs the figures the
# workbook was DRIVEN by, so Omni can be given the same starting point instead
# of anyone retyping 30 numbers off a screen. Same (section, label) anchoring —
# see the module docstring for why position is never safe.
IRC_FACTOR_COL = 3     # column C — the class loading
IRC_ANWP_COL   = 4     # column D — assumed annual NWP, next 12 months
MRC_NET_COL    = 3     # column C — net total assets in that bucket
MRC_ALLOC_COL  = 4     # column D — how much is allocated to market-risk cover

_IRC_KEY = {_norm(lbl): key for key, lbl in CLASS_LABEL.items()}
_MRC_KEY = {_norm(lbl): key for key, lbl, _f in DEFAULT_MRC}


def parse_a1_inputs(file_obj) -> dict:
    """Pull the ENTERED inputs out of a filed A.1 tab.

    Returns {'ok', 'anwp', 'mer_total', 'net_assets', 'alloc_mrctr'} with plain
    strings (JSON-safe, Decimal-exact — never floats, these reach a regulatory
    return). On failure 'ok' is False with a readable 'error'.
    """
    try:
        import openpyxl
    except ImportError:                          # pragma: no cover
        return {'ok': False, 'error': 'openpyxl is not installed on the server.'}
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    except Exception as exc:                     # noqa: BLE001
        return {'ok': False,
                'error': f'Not a readable Excel workbook ({type(exc).__name__}).'}
    try:
        if SHEET_NAME not in wb.sheetnames:
            return {'ok': False,
                    'error': f'No "{SHEET_NAME}" tab (found: '
                             f'{", ".join(wb.sheetnames)}).'}
        ws = wb[SHEET_NAME]
        grid = list(ws.iter_rows(min_row=1, max_row=MAX_SCAN_ROWS,
                                 max_col=VALUE_COL, values_only=True))

        anwp: dict[str, str] = {}
        net_assets: dict[str, str] = {}
        alloc: dict[str, str] = {}
        mer = Decimal('0')
        section: Optional[str] = None

        def cell(row, col):
            return row[col - 1] if len(row) >= col else None

        for row in grid:
            label = _norm(cell(row, LABEL_COL))
            new_section = _section_for(label)
            if new_section:
                section = new_section
                continue
            if not section or not label:
                continue
            if section == 'irc' and label in _IRC_KEY:
                v = _dec(cell(row, IRC_ANWP_COL))
                if v is not None:
                    anwp[_IRC_KEY[label]] = str(v)
            elif section == 'mer' and label == 'total mer':
                v = _dec(cell(row, VALUE_COL))
                if v is not None:
                    mer = v
            elif section == 'mrc' and label in _MRC_KEY:
                key = _MRC_KEY[label]
                n = _dec(cell(row, MRC_NET_COL))
                a = _dec(cell(row, MRC_ALLOC_COL))
                if n is not None:
                    net_assets[key] = str(n)
                if a is not None:
                    alloc[key] = str(a)
    finally:
        wb.close()

    if not anwp and not net_assets:
        return {'ok': False,
                'error': 'Found the A.1 tab but no assumed-premium or asset rows '
                         '— nothing loaded, rather than something wrong loaded.'}

    return {
        'ok': True,
        'anwp': anwp,
        'mer_total': str(mer),
        'net_assets': net_assets,
        'alloc_mrctr': alloc,
        'note': (f'Read {len(anwp)} class premium rows and '
                 f'{len(net_assets)} asset rows from the filed workbook.'),
    }
