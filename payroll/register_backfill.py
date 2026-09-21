"""
payroll/register_backfill.py — shared logic to backfill payslip COMPONENT LINES
from a payroll register (Odoo export or the Omni layout), used by both the
management command and the in-app upload endpoint (CFO 2026-07-14).

Why: payslips imported "totals-only" have header Gross/PAYE/Net/CTC but no
PayslipLine components, so the payroll report shows blanks. This loads the
per-component detail for ANY period/company.

SAFETY (never corrupts approved pay): an employee's lines are written ONLY if
the register Gross ties the payslip's approved gross_amount within `tolerance`.
Non-tying / unmatched employees are skipped and reported. Earnings +, employee
deductions −, employer (company) contributions forced + (CTC = gross + employer).
"""
from __future__ import annotations

import io
from decimal import Decimal

from payroll.models import Payslip, PayslipComponent, PayslipLine, PayrollPeriod

# Register column header (lower-cased, trimmed) -> Omni PayslipComponent.code.
COL_TO_CODE = {
    'basic salary': 'BASIC', 'basic': 'BASIC', 'commission': 'COMMISSION',
    'incentive': 'INCENTIVE', 'bonus': 'BONUS',
    'principal officer allowance': 'PO_ALLOWANCE', 'allowance': 'ALLOWANCE',
    'housing allowance': 'HOUSING_ALLOWANCE', 'loans deduction': 'LOANS_DEDUCTION',
    'housing tax deduction': 'HOUSING_TAX', 'housing deduction': 'HOUSING_DEDUCTION',
    'medical aid allowance': 'MEDICAL_AID_ALLOWANCE',
    'vehicle allowance': 'VEHICLE_ALLOWANCE', 'health insurance allowance': 'HEALTH_INS_ALLOWANCE',
    'fuel allowance': 'FUEL_ALLOWANCE', 'internet allowance': 'INTERNET_ALLOWANCE',
    'mobile allowance': 'MOBILE_ALLOWANCE', 'sales allowance': 'SALES_ALLOWANCE',
    # 'severence' = the Odoo export's misspelling of Severance (real register header).
    'leave pay': 'LEAVE_PAY', 'severance': 'SEVERANCE', 'severence': 'SEVERANCE',
    'housing benefit': 'HOUSING_BENEFIT',
    'furniture benefit': 'FURNITURE_BENEFIT', 'non-cash benefits': 'NON_CASH_BENEFIT',
    'medical aid employee contribution': 'MEDICAL_AID_EE',
    'medical aid company contribution': 'MEDICAL_AID_ER',
    'pension employee contribution': 'PENSION_EE', 'pension company contribution': 'PENSION_ER',
    'provident fund employee contribution': 'PROVIDENT_EE',
    'provident fund company contribution': 'PROVIDENT_ER',
}
GROSS_HDR = 'gross'


def _toks(n):
    return [t for t in (n or '').replace('.', ' ').split() if len(t) > 1]


def _norm(n):
    return ' '.join(_toks(n)).lower()


def _firstlast(n):
    t = _toks(n)
    return f'{t[0]} {t[-1]}'.lower() if len(t) >= 2 else (t[0].lower() if t else '')


def _dec(v):
    try:
        return Decimal(str(round(float(v), 2)))
    except (TypeError, ValueError):
        return Decimal('0')


def _rows(file_bytes: bytes, filename: str = '') -> list[list]:
    name = (filename or '').lower()
    if name.endswith('.csv') or file_bytes[:1] != b'P':
        import csv
        return [list(r) for r in csv.reader(io.StringIO(file_bytes.decode('utf-8', 'replace')))]
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def backfill(period_name: str, company_code: str, file_bytes: bytes, *,
             filename: str = '', commit: bool = False, tolerance: float = 1.0) -> dict:
    """Backfill payslip lines for (period, company) from a register. Returns a
    summary dict — never raises on a bad row, only on unusable period/company/file."""
    period = PayrollPeriod.objects.filter(period_name=period_name).first()
    if not period:
        return {'ok': False, 'error': f'Period {period_name!r} not found.'}
    code = (company_code or '').strip().upper()
    payslips = list(Payslip.objects.filter(period=period, company__code=code)
                    .select_related('employee', 'company').prefetch_related('lines'))
    if not payslips:
        return {'ok': False, 'error': f'No payslips for {period_name} / {code}.'}

    by_norm, by_fl = {}, {}
    for ps in payslips:
        nm = ps.employee.full_name if ps.employee else ''
        by_norm.setdefault(_norm(nm), []).append(ps)
        by_fl.setdefault(_firstlast(nm), []).append(ps)

    rows = _rows(file_bytes, filename)
    hidx = next((i for i, r in enumerate(rows[:20])
                 if any(str(c or '').strip().lower() == 'employee' for c in r)), -1)
    if hidx < 0:
        return {'ok': False, 'error': 'Could not find an "Employee" header row in the file.'}
    hdr = [str(c or '').strip().lower() for c in rows[hidx]]
    colmap = {h: i for i, h in enumerate(hdr)}
    comp_col = {code_: colmap[h] for h, code_ in COL_TO_CODE.items() if h in colmap}
    gross_col = colmap.get(GROSS_HDR)
    comps = {c.code: c for c in PayslipComponent.objects.filter(code__in=set(comp_col.keys()))}

    matched = tied = written = recon_fail = unmatched = 0
    fails, unmatched_names = [], []
    for r in rows[hidx + 1:]:
        if not r or not str(r[0] or '').strip():
            continue
        name = str(r[0]).strip()
        cand = by_norm.get(_norm(name)) or by_fl.get(_firstlast(name))
        if not cand or len(cand) != 1:
            unmatched += 1
            if name.lower() not in ('total', 'totals', 'grand total'):
                unmatched_names.append(name)
            continue
        ps = cand[0]; matched += 1
        odoo_gross = _dec(r[gross_col]) if gross_col is not None and gross_col < len(r) else None
        stored = ps.gross_amount or Decimal('0')
        if odoo_gross is None or abs(odoo_gross - stored) > Decimal(str(tolerance)):
            recon_fail += 1
            fails.append({'employee': name, 'register_gross': str(odoo_gross),
                          'system_gross': str(stored)})
            continue
        tied += 1
        lines = []
        for code_, ci in comp_col.items():
            if code_ not in comps or ci >= len(r):
                continue
            comp = comps[code_]
            amt = _dec(r[ci])
            if amt == 0:
                continue
            if comp.kind == 'company_contribution':
                amt = abs(amt)
            lines.append(PayslipLine(payslip=ps, component=comp, amount=amt))
        if commit:
            from django.db import transaction
            with transaction.atomic():
                ps.lines.all().delete()
                PayslipLine.objects.bulk_create(lines)
            written += 1

    return {
        'ok': True, 'period': period_name, 'company': code, 'commit': commit,
        'payslips': len(payslips), 'matched': matched, 'reconciled': tied,
        'written': written, 'recon_fail': recon_fail, 'unmatched': unmatched,
        'fails': fails[:60], 'unmatched_names': unmatched_names[:60],
    }
