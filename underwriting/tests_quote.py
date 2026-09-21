"""Quotations — the controls that make one standard template worth having.

Ten real quotes were reviewed before this was built. Two clients shared a file,
one was a competitor's schedule reused as our base, and VAT appeared three
different ways. Each test below is one of those failures, made impossible.
"""
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.test import TestCase
from rest_framework.test import APITestCase

from core.models import Company
from underwriting.models import (Quote, QuoteRateFloor, quote_rates_below_floor,
                                 quote_below_min_premium, GraphiteMapping)
from underwriting.quote_parse import (
    _fallback, premium_appears_in, price, price_from_rate, total_sum_insured,
    premium_from_section_rates, unrated_covered_rows, section_premium_rows,
    instalment_plan, PAYMENT_PLAN_PCT, prorate,
)
from underwriting.quote_render import (_display_rows, _rows_with_premiums,
                                       qr_svg_data_uri, verify_url, quote_html,
                                       _png_data_uri, _LOGO_COLOUR, _STAMP,
                                       _woff2_data_uri, _FONT_DISPLAY, _FONT_TEXT)


class QuoteMoneyTests(TestCase):
    """VAT is computed in one place. It is the reason this module exists."""

    def test_vat_is_always_14_percent_of_premium(self):
        m = price('418600')
        self.assertEqual(m['premium'], Decimal('418600.00'))
        self.assertEqual(m['vat'], Decimal('58604.00'))
        self.assertEqual(m['total'], Decimal('477204.00'))

    def test_it_reads_the_shorthand_underwriters_actually_type(self):
        self.assertEqual(price('8.4m')['premium'], Decimal('8400000.00'))
        self.assertEqual(price('15k')['premium'], Decimal('15000.00'))
        self.assertEqual(price('P 96,400')['premium'], Decimal('96400.00'))

    def test_rubbish_becomes_zero_rather_than_an_exception(self):
        self.assertEqual(price('abc')['total'], Decimal('0.00'))
        self.assertEqual(price(None)['total'], Decimal('0.00'))

    def test_vat_and_total_cannot_be_set_by_hand(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('100000'),
                                 vat=Decimal('1'), total=Decimal('1'))
        self.assertEqual(q.vat, Decimal('14000.00'))
        self.assertEqual(q.total, Decimal('114000.00'))

    def test_changing_the_premium_recomputes_the_total(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('100000'))
        q.premium = Decimal('200000')
        q.save()
        self.assertEqual(q.total, Decimal('228000.00'))


class QuoteRateTests(TestCase):
    """A premium can be a RATE on the sum insured — worked out, never typed.

    The exact case the CFO reported: '3% is the premium rate incl VAT' on a
    million pula must give 30,000 total on its own, not be left for someone to
    multiply by hand.
    """

    def test_total_sum_insured_adds_numeric_rows_only(self):
        sections = [
            {'name': 'Motor', 'sum_insured': '1,000,000', 'excess': '5000'},
            {'name': 'Computers', 'sum_insured': '200000', 'excess': '2000'},
            {'name': 'Liability', 'sum_insured': 'Included', 'excess': ''},
        ]
        self.assertEqual(total_sum_insured(sections), Decimal('1200000.00'))

    def test_rate_incl_vat_backs_the_vat_out_of_the_figure(self):
        # 3% of 1,000,000 = 30,000, and that 30,000 already carries VAT.
        m = price_from_rate('1000000', '3', incl_vat=True)
        self.assertEqual(m['total'], Decimal('30000.00'))
        self.assertEqual(m['premium'], Decimal('26315.79'))
        self.assertEqual(m['vat'], Decimal('3684.21'))
        # premium + VAT must equal the total shown, to the thebe.
        self.assertEqual(m['premium'] + m['vat'], m['total'])

    def test_rate_excl_vat_adds_vat_on_top(self):
        m = price_from_rate('1000000', '3', incl_vat=False)
        self.assertEqual(m['premium'], Decimal('30000.00'))
        self.assertEqual(m['vat'], Decimal('4200.00'))
        self.assertEqual(m['total'], Decimal('34200.00'))

    def test_a_rated_quote_computes_its_own_premium_on_save(self):
        q = Quote.objects.create(
            client_name='Rated Risk Ltd',
            sections=[{'name': 'Contents', 'sum_insured': '1,000,000', 'excess': '5000'}],
            rate_pct=Decimal('3'), rate_incl_vat=True,
        )
        self.assertEqual(q.total, Decimal('30000.00'))
        self.assertEqual(q.premium, Decimal('26315.79'))
        self.assertEqual(q.vat, Decimal('3684.21'))

    def test_no_rate_keeps_the_typed_premium(self):
        q = Quote.objects.create(client_name='X', premium=Decimal('50000'))
        self.assertEqual(q.premium, Decimal('50000.00'))
        self.assertEqual(q.total, Decimal('57000.00'))


class QuotePerProductRateTests(TestCase):
    """Different products, different rates, one quote — motor rated apart from
    contents. The premium is the sum of each row's sum insured x its own rate.
    """

    def test_section_rates_sum_and_back_the_vat_out(self):
        sections = [
            {'name': 'Motor', 'sum_insured': '500,000', 'rate': '2', 'excess': '5000'},
            {'name': 'Contents', 'sum_insured': '200,000', 'rate': '1', 'excess': '2000'},
        ]
        m = premium_from_section_rates(sections, incl_vat=True)
        # 500,000 x 2% + 200,000 x 1% = 10,000 + 2,000 = 12,000 gross (incl VAT)
        self.assertEqual(m['total'], Decimal('12000.00'))
        self.assertEqual(m['premium'], Decimal('10526.32'))
        self.assertEqual(m['vat'], Decimal('1473.68'))
        self.assertEqual(m['premium'] + m['vat'], m['total'])

    def test_no_row_rate_returns_none_so_caller_falls_back(self):
        sections = [{'name': 'Contents', 'sum_insured': '200000', 'excess': '2000'}]
        self.assertIsNone(premium_from_section_rates(sections, incl_vat=True))

    def test_section_rates_win_over_a_quote_level_rate_on_save(self):
        q = Quote.objects.create(
            client_name='Mixed Risk Ltd',
            sections=[
                {'name': 'Motor', 'sum_insured': '500,000', 'rate': '2'},
                {'name': 'Contents', 'sum_insured': '200,000', 'rate': '1'},
            ],
            rate_pct=Decimal('9'), rate_incl_vat=True,   # should be ignored
        )
        self.assertEqual(q.total, Decimal('12000.00'))
        self.assertEqual(q.premium, Decimal('10526.32'))

    def test_a_covered_row_left_unrated_is_flagged(self):
        # Motor is rated; Contents has a value but no rate -> it would be free.
        s = [{'name': 'Motor', 'sum_insured': '500000', 'rate': '2'},
             {'name': 'Contents', 'sum_insured': '200000'}]
        self.assertEqual(unrated_covered_rows(s), ['Contents'])

    def test_no_flag_when_the_quote_is_not_section_rated(self):
        self.assertEqual(unrated_covered_rows([{'name': 'Contents', 'sum_insured': '200000'}]), [])

    def test_included_row_needs_no_rate(self):
        s = [{'name': 'Motor', 'sum_insured': '500000', 'rate': '2'},
             {'name': 'Liability', 'sum_insured': 'Included'}]
        self.assertEqual(unrated_covered_rows(s), [])

    def test_a_sub_limit_of_a_priced_section_does_not_block(self):
        # Q-2026-00009: public liability is rated once, then states its own
        # sub-limits. Those limits are what the section's premium buys — asking
        # for a rate on each of them stopped a correctly-priced quote issuing.
        s = [{'group': 'Public liability', 'name': 'Limit of liability',
              'sum_insured': '5,000,000', 'rate': '0.30'},
             {'name': 'Products liability', 'sum_insured': '5,000,000'},
             {'name': 'Defective workmanship', 'sum_insured': '10,000,000'},
             {'name': 'Legal defence', 'sum_insured': '300,000'}]
        self.assertEqual(unrated_covered_rows(s), [])

    def test_an_unpriced_section_still_blocks_however_it_is_headed(self):
        # A whole section with a value and no rate is a product being given away,
        # heading or no heading. The exemption is for sub-limits of a section that
        # IS priced — never for a section nobody priced.
        s = [{'group': 'Motor', 'name': 'Fleet', 'sum_insured': '500,000', 'rate': '2'},
             {'group': 'Contents', 'name': 'Office contents', 'sum_insured': '200,000'}]
        self.assertEqual(unrated_covered_rows(s), ['Office contents'])

    def test_the_exemption_cannot_be_had_without_a_section_heading(self):
        # The guard must not be defeated by a flat ungrouped list: one rated row
        # at the top would otherwise excuse every unrated row beneath it.
        s = [{'name': 'Motor', 'sum_insured': '500,000', 'rate': '2'},
             {'name': 'Contents', 'sum_insured': '200,000'},
             {'name': 'Plant', 'sum_insured': '900,000'}]
        self.assertEqual(unrated_covered_rows(s), ['Contents', 'Plant'])


class PerProductPremiumBreakdownTests(TestCase):
    """The client sees what EACH product costs, and those figures add up to the
    total exactly. A breakdown that does not reconcile is the whole disease this
    module exists to end.
    """

    SECTIONS = [
        {'group': 'Fire & Allied Perils', 'name': 'Buildings', 'sum_insured': '12,500,000', 'rate': '0.15'},
        {'name': 'Plant & machinery', 'sum_insured': '3,200,000', 'rate': '0.22'},
        {'group': 'Stock', 'name': 'Stock in trade', 'sum_insured': '4,750,000', 'rate': '0.35'},
        {'group': 'Money', 'name': 'Money in transit', 'sum_insured': '150,000', 'premium': '1,250'},
        {'group': 'Liability', 'name': 'Public liability', 'sum_insured': 'Included'},
    ]

    def test_each_row_is_priced_by_its_own_rate_or_flat_premium(self):
        b = section_premium_rows(self.SECTIONS, incl_vat=False)
        nets = [p['net'] for p in b['rows']]
        self.assertEqual(nets, [Decimal('18750.00'), Decimal('7040.00'),
                               Decimal('16625.00'), Decimal('1250.00')])

    def test_the_printed_rows_sum_to_the_net_exactly(self):
        b = section_premium_rows(self.SECTIONS, incl_vat=False)
        self.assertEqual(sum(p['net'] for p in b['rows']), b['net'])
        self.assertEqual(sum(g['net'] for g in b['groups']), b['net'])

    def test_group_subtotals_follow_the_printed_grouping(self):
        b = section_premium_rows(self.SECTIONS, incl_vat=False)
        self.assertEqual(
            [(g['group'], g['net']) for g in b['groups']],
            [('Fire & Allied Perils', Decimal('25790.00')),
             ('Stock', Decimal('16625.00')),
             ('Money', Decimal('1250.00'))])

    def test_a_typed_row_premium_beats_that_rows_rate(self):
        s = [{'name': 'X', 'sum_insured': '1,000,000', 'rate': '5', 'premium': '900'}]
        self.assertEqual(section_premium_rows(s, incl_vat=False)['net'], Decimal('900.00'))

    def test_incl_vat_rows_are_rounded_individually_then_summed(self):
        # Two rows whose VAT back-out does not divide cleanly: each row is rounded
        # to the thebe, and the net is the sum OF THOSE ROUNDED figures.
        s = [{'name': 'A', 'sum_insured': '1,000', 'rate': '1'},
             {'name': 'B', 'sum_insured': '1,000', 'rate': '1'}]
        b = section_premium_rows(s, incl_vat=True)
        self.assertEqual([p['net'] for p in b['rows']], [Decimal('8.77'), Decimal('8.77')])
        self.assertEqual(b['net'], Decimal('17.54'))

    def test_incl_vat_total_equals_the_figure_the_client_sees(self):
        # Gomolemo, ADIC 2026-08-19: a motor quote at 3.2% on a SI of 260,000 is
        # 8,320.00 gross — SI x rate, exact. Backing out net (7,298.25) and then
        # forward-adding VAT gave 8,320.01, so the system tile and the PDF
        # disagreed by a thebe. The gross total the client pays must equal the
        # figure the rate produces, to the thebe, for every incl_vat quote.
        s = [{'name': 'Motor', 'sum_insured': '260000', 'rate': '3.2'}]
        m = premium_from_section_rates(s, incl_vat=True)
        self.assertEqual(m['total'], Decimal('8320.00'))
        self.assertEqual(m['premium'] + m['vat'], m['total'])

    def test_saved_quote_stored_total_matches_the_screen(self):
        # The full save path (recalc) must return the same 8,320.00 the tile
        # shows — forward-VATing the net drifted the PDF to 8,320.01.
        q = Quote.objects.create(
            client_name='Gomolemo test 8320',
            sections=[{'name': 'Motor', 'sum_insured': '260000', 'rate': '3.2'}],
            rate_incl_vat=True,
        )
        self.assertEqual(q.total, Decimal('8320.00'))
        self.assertEqual(q.premium + q.vat, q.total)

    def test_the_stored_premium_equals_the_printed_net(self):
        q = Quote.objects.create(client_name='Recon Ltd', sections=self.SECTIONS,
                                 rate_incl_vat=False)
        b = section_premium_rows(self.SECTIONS, incl_vat=False)
        self.assertEqual(q.premium, b['net'])
        self.assertEqual(q.premium + q.vat, q.total)

    def test_no_priced_row_means_no_breakdown(self):
        self.assertIsNone(section_premium_rows(
            [{'name': 'X', 'sum_insured': '1,000,000'}], incl_vat=False))

    def test_every_priced_row_carries_its_premium_onto_the_document(self):
        # The bug this guards: premiums were matched on comma-formatted text, so
        # every row printed "—" while the totals were right (QC 2026-08-10).
        q = Quote.objects.create(client_name='Print Ltd', sections=self.SECTIONS,
                                 rate_incl_vat=False)
        rows, breakdown = _rows_with_premiums(q)
        priced = [r for r in rows if r.get('premium_display')]
        self.assertEqual(len(priced), 4)
        self.assertEqual(priced[0]['premium_display'], '18,750.00')
        self.assertEqual(breakdown['net'], '43,665.00')

    def test_a_priced_reinsurance_row_suppresses_the_breakdown(self):
        # The RI row is hidden from the client but still priced, so a per-row
        # breakdown would not add up on the page. Fall back to the single figure;
        # the stored premium must not change (Fable review 2026-08-10).
        s = list(self.SECTIONS) + [
            {'group': 'Facultative Reinsurance', 'name': 'Fac placement',
             'sum_insured': '5,000,000', 'rate': '0.10'}]
        q = Quote.objects.create(client_name='RI Ltd', sections=s, rate_incl_vat=False)
        rows, breakdown = _rows_with_premiums(q)
        self.assertIsNone(breakdown)
        self.assertEqual(q.premium, premium_from_section_rates(s, False)['premium'])

    def test_a_flat_premium_row_counts_as_priced_for_the_issue_guard(self):
        s = [{'name': 'A', 'sum_insured': '100,000', 'rate': '1'},
             {'name': 'B', 'sum_insured': '50,000', 'premium': '500'}]
        self.assertEqual(unrated_covered_rows(s), [])


class InstalmentPlanTests(TestCase):
    """The monthly-payment option. Paying by instalment carries the payment plan
    charge, and every figure printed is what the client is actually debited.
    """

    def test_monthly_is_the_total_over_twelve_plus_the_charge(self):
        p = instalment_plan('834394.50')
        self.assertEqual(p['monthly_premium'], Decimal('69532.88'))
        self.assertEqual(p['plan_charge'], Decimal('5562.63'))          # 8% of 69,532.88
        self.assertEqual(p['monthly_instalment'], Decimal('75095.51'))
        self.assertEqual(p['monthly_premium'] + p['plan_charge'], p['monthly_instalment'])

    def test_the_twelve_month_total_is_the_instalment_times_twelve(self):
        p = instalment_plan('834394.50')
        self.assertEqual(p['total_over_term'], p['monthly_instalment'] * 12)

    def test_the_charge_is_the_agreed_rate(self):
        self.assertEqual(PAYMENT_PLAN_PCT, Decimal('8'))
        p = instalment_plan('1200.00')
        self.assertEqual(p['monthly_premium'], Decimal('100.00'))
        self.assertEqual(p['plan_charge'], Decimal('8.00'))
        self.assertEqual(p['monthly_instalment'], Decimal('108.00'))

    def test_no_plan_on_a_zero_total(self):
        self.assertIsNone(instalment_plan('0'))
        self.assertIsNone(instalment_plan(None))


class QuoteRateDisplayTests(TestCase):
    """Rates print with TWO decimals, always: "4%" beside "4.50%" reads as two
    different kinds of figure on a client document (CFO 2026-08-10).
    """

    def test_every_printed_rate_has_two_decimals(self):
        q = Quote.objects.create(
            client_name='Uniform Ltd', rate_incl_vat=False,
            sections=[{'name': 'A', 'sum_insured': '100,000', 'rate': '4'},
                      {'name': 'B', 'sum_insured': '100,000', 'rate': '4.5'},
                      {'name': 'C', 'sum_insured': '100,000', 'rate': '0.15'}])
        rows, _ = _rows_with_premiums(q)
        self.assertEqual([r.get('rate_display') for r in rows],
                         ['4.00%', '4.50%', '0.15%'])

    def test_a_flat_premium_row_shows_no_rate(self):
        q = Quote.objects.create(
            client_name='Flat Ltd', rate_incl_vat=False,
            sections=[{'name': 'A', 'sum_insured': '100,000', 'rate': '2'},
                      {'name': 'MIT', 'sum_insured': '50,000', 'premium': '500'}])
        rows, _ = _rows_with_premiums(q)
        self.assertIsNone(rows[1].get('rate_display'))


class QuoteVerifyQrTests(TestCase):
    """The QR code on the quotation opens a page that confirms the document is
    genuinely ours. The link is the quote's UUID — a capability, not an id to
    enumerate.
    """

    def test_the_link_carries_the_quote_uuid(self):
        q = Quote.objects.create(client_name='QR Ltd', premium=Decimal('1000'))
        self.assertIn(f'/api/verify/quote/{q.pk}/', verify_url(q))

    def test_the_qr_is_an_inline_svg_data_uri(self):
        uri = qr_svg_data_uri('https://omni.alphadirect.co.bw/api/verify/quote/abc/')
        self.assertTrue(uri.startswith('data:image/svg+xml;base64,'))
        import base64
        svg = base64.b64decode(uri.split(',', 1)[1]).decode('utf-8')
        self.assertIn('<svg', svg)
        self.assertIn('<path', svg)          # one path, not a rect per module
        self.assertLess(len(svg), 40000)     # renderSVG's rect-per-module was 93 KB

    def test_the_verify_page_answers_for_a_real_quote(self):
        from django.test import Client
        q = Quote.objects.create(client_name='QR Ltd', premium=Decimal('1000'))
        html = Client().get(f'/api/verify/quote/{q.pk}/').content.decode()
        self.assertIn('Genuine quotation', html)
        self.assertIn(q.quote_number, html)

    def test_an_unknown_code_is_not_recognised(self):
        import uuid
        from django.test import Client
        r = Client().get(f'/api/verify/quote/{uuid.uuid4()}/')
        self.assertEqual(r.status_code, 404)
        self.assertIn('Not recognised', r.content.decode())


class SectionedRowsTests(TestCase):
    """How the printed schedule is assembled from the stored rows — the CFO's
    fix list of 11 Aug 2026 (one excess per section, no dash-only rows, one
    sub-line per vehicle)."""

    def test_vehicle_sub_limits_fold_into_one_sub_line(self):
        # Graphite stores each vehicle as four rows; a 15-vehicle fleet printed
        # 45 sub-limit rows and the document ran to six pages (fix list #5).
        rows = _display_rows([
            {'group': 'Motor', 'name': 'B170ALK Mercedes', 'sum_insured': '131,800',
             'rate': '3', 'excess': 'Basic 10%'},
            {'name': 'Third party liability', 'sum_insured': '2,500,000'},
            {'name': 'Windscreen', 'sum_insured': '4,000'},
            {'name': 'Loss of keys', 'sum_insured': '1,500'},
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['sub'],
                         'Third party liability 2,500,000 · Windscreen 4,000 · Loss of keys 1,500')

    def test_valueless_sub_limit_rows_vanish(self):
        # The unvalued Prado carried three all-dash children (fix list #4).
        rows = _display_rows([
            {'group': 'Motor', 'name': 'B793BMG Prado', 'sum_insured': '', 'rate': '',
             'note': 'To be confirmed'},
            {'name': 'Third party liability', 'sum_insured': ''},
            {'name': 'Windscreen', 'sum_insured': ''},
            {'name': 'Loss of keys', 'sum_insured': ''},
        ])
        names = [r['name'] for r in rows]
        self.assertEqual(names, ['B793BMG Prado'])

    def test_an_unvalued_vehicle_with_no_note_stays_a_visible_row(self):
        # Prod's real Prado carries NO note. Because its own third-party /
        # windscreen / loss-of-keys rows follow it, it is a vehicle awaiting
        # figures, not a remark — it must stay a visible row (with a supplied
        # "to be confirmed" note), never be swallowed onto the section bar
        # (the fault caught on the prod render, 11 Aug 2026).
        rows = _display_rows([
            {'group': 'Motor', 'name': 'B793BMG Prado 2.7', 'sum_insured': '', 'rate': ''},
            {'name': 'Third party liability', 'sum_insured': ''},
            {'name': 'Windscreen', 'sum_insured': ''},
            {'name': 'Loss of keys', 'sum_insured': ''},
        ])
        self.assertEqual([r['name'] for r in rows], ['B793BMG Prado 2.7'])
        self.assertIn('to be confirmed', rows[0]['sub'].lower())
        self.assertNotIn('Prado', rows[0]['group_info'])     # not on the bar

    def test_a_bare_remark_row_still_folds_to_the_bar(self):
        # A no-value row NOT followed by vehicle sub-limits is a genuine remark
        # ("Occupation surcharge: nil") and still moves onto the section bar.
        rows = _display_rows([
            {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '500,000', 'rate': '0.1'},
            {'name': 'Occupation surcharge: nil', 'sum_insured': ''},
        ])
        self.assertEqual([r['name'] for r in rows], ['Buildings'])
        self.assertIn('Occupation surcharge: nil', rows[0]['group_info'])

    def test_a_priced_row_is_never_folded_whatever_its_name(self):
        # Folding is layout, not pricing: a windscreen the underwriter RATED is
        # cover with money on it and must stay a visible data row.
        rows = _display_rows([
            {'group': 'Motor', 'name': 'B170ALK Mercedes', 'sum_insured': '131,800', 'rate': '3'},
            {'name': 'Windscreen', 'sum_insured': '4,000', 'rate': '1'},
        ])
        self.assertEqual([r['name'] for r in rows], ['B170ALK Mercedes', 'Windscreen'])

    def test_a_remark_row_moves_onto_the_section_bar(self):
        # "Occupation surcharge: nil" printed as a line of dashes (fix list #4).
        rows = _display_rows([
            {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '500,000', 'rate': '0.1',
             'excess': '10% of claim'},
            {'name': 'Occupation surcharge: nil', 'sum_insured': ''},
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['group_info'], 'Excess: 10% of claim · Occupation surcharge: nil')

    def test_one_shared_excess_prints_on_the_bar_not_the_rows(self):
        rows = _display_rows([
            {'group': 'Motor', 'name': 'A', 'sum_insured': '100,000', 'rate': '3', 'excess': 'Basic 10%'},
            {'group': '', 'name': 'B', 'sum_insured': '200,000', 'rate': '3', 'excess': 'Basic 10%'},
        ])
        self.assertEqual(rows[0]['group_info'], 'Excess: Basic 10%')
        self.assertEqual([r['row_excess'] for r in rows], ['', ''])

    def test_a_row_with_its_own_different_excess_keeps_it(self):
        rows = _display_rows([
            {'group': 'Motor', 'name': 'A', 'sum_insured': '100,000', 'rate': '3', 'excess': 'Basic 10%'},
            {'group': '', 'name': 'B', 'sum_insured': '200,000', 'rate': '3', 'excess': '25%, min P450'},
        ])
        self.assertEqual(rows[0]['group_info'], '')          # no single shared excess
        self.assertEqual([r['row_excess'] for r in rows], ['Basic 10%', '25%, min P450'])


class DocumentGateTests(TestCase):
    """The quotation measures its own printed pages before issue. Built after two
    print faults reached the CFO on 11 Aug 2026 that the on-screen render hid.

    Fixtures are built with reportlab — already a dependency, and it needs no
    browser, so these run everywhere including CI (the real renderer needs
    Chromium, which the CI box does not carry). The measuring is what is under
    test, and it reads a reportlab PDF exactly as it reads a Chromium one:
    pypdfium2 sees text runs at points, whatever drew them. A4 = 595 x 842 pt;
    the footer strip is the bottom 34pt (quote_doc_gate._FOOTER_BAND_PT)."""

    A4 = (595.27, 841.89)

    def _pdf(self, draw, pages=1):
        # draw(canvas, page_index) places text; returns the PDF bytes.
        import io
        from reportlab.pdfgen import canvas
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=self.A4)
        for i in range(pages):
            draw(c, i)
            c.showPage()
        c.save()
        return buf.getvalue()

    def _clean_body(self, c, _i):
        # A dozen content lines well clear of the footer strip.
        c.setFont('Helvetica', 10)
        for row in range(12):
            c.drawString(60, 700 - row * 20, f'Cover line {row} — 100,000.00')

    def test_a_clean_page_has_no_geometry_findings(self):
        from underwriting.quote_doc_gate import geometry_findings
        self.assertEqual(geometry_findings(self._pdf(self._clean_body)), [])

    def test_a_schedule_row_in_the_footer_band_is_caught(self):
        # A content line down at y=12pt is inside the reserved footer strip —
        # the @page{margin:0} regression that put rows over the licence line on
        # pages 3-4 of the 11 Aug print.
        from underwriting.quote_doc_gate import geometry_findings

        def draw(c, _i):
            self._clean_body(c, _i)
            c.setFont('Helvetica', 10)
            c.drawString(44, 12, 'B290BUN Nissan Truck 1,552,000.00 schedule row over the footer')
        findings = geometry_findings(self._pdf(draw))
        self.assertTrue(any('footer band' in f for f in findings), findings)

    def test_the_footer_page_numbers_are_not_mistaken_for_intrusion(self):
        # The footer's OWN words live in the strip on purpose: the licence line
        # (footer words) and the short "Page 1 of 3" marker (too short to be a
        # content line). Neither may read as an intrusion.
        from underwriting.quote_doc_gate import geometry_findings

        def draw(c, _i):
            self._clean_body(c, _i)
            c.setFont('Helvetica', 8)
            c.drawString(44, 20, 'Alpha Direct Insurance Company (Pty) Ltd · Licensed by NBFIRA')
            c.drawString(500, 20, 'Page 1 of 3')
        self.assertEqual(geometry_findings(self._pdf(draw)), [])

    def test_a_stranded_near_blank_page_is_caught(self):
        # Payment options alone on a last page was the stranded-block fault. A
        # near-empty page BEFORE the last one is the signal.
        from underwriting.quote_doc_gate import geometry_findings

        def draw(c, i):
            if i == 0:
                c.setFont('Helvetica', 10)
                c.drawString(60, 400, 'Lonely line')     # < 5 runs on page 1
            else:
                self._clean_body(c, i)
        findings = geometry_findings(self._pdf(draw, pages=2))
        self.assertTrue(any('nearly blank' in f for f in findings), findings)

    def test_a_broken_pdf_is_reported_not_swallowed(self):
        from underwriting.quote_doc_gate import geometry_findings
        findings = geometry_findings(b'%PDF-1.4 not really a pdf')
        self.assertTrue(findings)          # names the problem rather than pass silently

    def test_an_exact_duplicate_line_hard_blocks(self):
        # The Win Win Samsung: same phone, same sum insured, same premium, entered
        # twice = a double-charge. Deterministic, so it blocks.
        from underwriting.quote_doc_gate import duplicate_findings
        rows = [
            {'name': 'Samsung S25 256GB', 'sum_insured': '15,650', 'premium': '1,095.50'},
            {'name': 'Samsung S25 256GB', 'sum_insured': '15,650', 'premium': '1,095.50'},
        ]
        findings = duplicate_findings(rows)
        self.assertEqual(len(findings), 1)
        self.assertIn('Samsung S25', findings[0])

    def test_two_labelled_units_are_not_a_duplicate(self):
        # Two genuinely-separate phones, labelled, are not identical and pass.
        from underwriting.quote_doc_gate import duplicate_findings
        rows = [
            {'name': 'Samsung S25 256GB — unit 1 of 2', 'sum_insured': '15,650', 'premium': '1,095.50'},
            {'name': 'Samsung S25 256GB — unit 2 of 2', 'sum_insured': '15,650', 'premium': '1,095.50'},
        ]
        self.assertEqual(duplicate_findings(rows), [])

    def test_aria_is_advisory_not_a_hard_block(self):
        # A hallucinating reviewer must never veto a clean document: the blocking
        # set holds ONLY the deterministic faults; Aria rides on advisory_findings.
        from underwriting import quote_doc_gate as g
        import inspect
        self.assertNotIn('aria_review', inspect.getsource(g.blocking_findings))
        self.assertIn('aria_review', inspect.getsource(g.advisory_findings))


class QuoteDocumentDesignTests(TestCase):
    """The quotation is a CLIENT document: it must carry the real mark, the total
    cover, and a stamp only once issued (CFO design review 2026-08-10).
    """

    SECTIONS = [
        {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '8,500,000', 'rate': '0.15'},
        {'group': 'Stock', 'name': 'Stock in trade', 'sum_insured': '9,400,000', 'rate': '0.45'},
        {'name': 'Liability', 'sum_insured': 'Included'},
    ]

    def _quote(self, **kw):
        return Quote.objects.create(client_name='Design Ltd', rate_incl_vat=False,
                                    sections=self.SECTIONS, **kw)

    def test_the_brand_assets_exist_on_disk(self):
        # If the asset moves, the letterhead silently falls back to typed text.
        self.assertTrue(_png_data_uri(str(_LOGO_COLOUR)).startswith('data:image/png;base64,'))
        self.assertTrue(_png_data_uri(str(_STAMP)).startswith('data:image/png;base64,'))

    def test_the_brand_typefaces_ship_with_the_code(self):
        # The render box has only the basic free fonts, so a missing face means
        # every quotation silently prints plain — the exact fault the CFO handed
        # back on 11 Aug 2026. Both faces must be on disk AND in the document.
        for face in (_FONT_DISPLAY, _FONT_TEXT):
            self.assertTrue(_woff2_data_uri(str(face)).startswith('data:font/woff2;base64,'),
                            f'{face.name} is missing')
        html = quote_html(self._quote())
        self.assertEqual(html.count('@font-face'), 2)
        self.assertIn('data:font/woff2;base64,', html)

    def test_the_logo_and_total_cover_print(self):
        html = quote_html(self._quote())
        self.assertIn('class="brand-logo"', html)
        self.assertIn('Total sums insured', html)
        self.assertIn('17,900,000.00', html)          # 8.5m + 9.4m; 'Included' skipped

    def test_a_draft_is_not_stamped_but_an_issued_quote_is(self):
        draft = self._quote()
        self.assertNotIn('class="stamp"', quote_html(draft))
        comp = Company.objects.create(name='ADIC-D', code='ADICD')
        u = User.objects.create(username='uw_d', email='uw_d@alphadirect.co.bw')
        issued = self._quote(company=comp, underwriter=u)
        issued.issue(u)
        self.assertIn('class="stamp"', quote_html(issued))

    def test_the_underwriter_email_is_printed_so_the_client_can_call(self):
        u = User.objects.create(username='uw_e', email='uw_e@alphadirect.co.bw')
        self.assertIn('uw_e@alphadirect.co.bw', quote_html(self._quote(underwriter=u)))

    def test_exclusions_print_when_given_and_the_pointer_always_stands(self):
        html = quote_html(self._quote(exclusions=['Wear and tear.', 'Deliberate damage.']))
        self.assertIn('What is not covered', html)
        self.assertIn('Wear and tear.', html)
        self.assertIn('policy wording', html)
        # No list typed: the block and its pointer sentence still print.
        bare = quote_html(self._quote())
        self.assertIn('What is not covered', bare)
        self.assertIn('policy wording', bare)

    def test_the_document_never_invents_an_exclusion(self):
        # Nothing but the underwriter's own list may appear as an exclusion.
        html = quote_html(self._quote())
        self.assertNotIn('<li>', html.split('What is not covered')[1].split('</div>')[0])


class ProRataTests(TestCase):
    """A part-year policy is charged months/12 of the annual premium."""

    SECTIONS = [{'name': 'Buildings', 'sum_insured': '8,500,000', 'rate': '0.15'},
                {'name': 'Stock', 'sum_insured': '9,400,000', 'rate': '0.45'}]

    def test_six_months_is_exactly_half_the_annual(self):
        self.assertEqual(prorate('731925.00', 6)['charged'], Decimal('365962.50'))

    def test_a_full_year_is_not_pro_rated(self):
        self.assertIsNone(prorate('731925.00', 12))
        self.assertIsNone(prorate('731925.00', None))

    def test_the_record_charges_the_shorter_period_and_vat_follows(self):
        annual = Quote.objects.create(client_name='Full Yr', sections=self.SECTIONS,
                                      rate_incl_vat=False)
        half = Quote.objects.create(client_name='Half Yr', sections=self.SECTIONS,
                                    rate_incl_vat=False, period_months=6)
        self.assertEqual(half.premium * 2, annual.premium)
        self.assertEqual(half.premium + half.vat, half.total)
        self.assertEqual(half.vat, (half.premium * Decimal('0.14')).quantize(Decimal('0.01')))

    def test_an_out_of_range_period_refuses_to_issue(self):
        # It used to fall through to "no pro-rata" and quietly charge a full
        # year, hiding the bad data (panel review 2026-08-11).
        comp = Company.objects.create(name='ADIC-PR', code='ADICPR')
        u = User.objects.create(username='uw_pr')
        q = Quote.objects.create(client_name='Bad Period Ltd', sections=self.SECTIONS,
                                 rate_incl_vat=False, company=comp, underwriter=u)
        q.issue(u)                       # 12 months issues fine
        q2 = Quote.objects.create(client_name='Bad Period 2', sections=self.SECTIONS,
                                  rate_incl_vat=False, company=comp, underwriter=u,
                                  period_months=99)
        with self.assertRaises(ValidationError):
            q2.issue(u)

    def test_a_typed_part_year_premium_is_not_pro_rated_twice(self):
        # recalc() runs on EVERY save, and issue() saves. Before annual_premium
        # was persisted, a typed 6-month quote went 60,000 -> 30,000 -> 15,000
        # and the issued PDF charged a quarter of the year (Fable 2026-08-11).
        q = Quote.objects.create(client_name='Typed 6mo', sections=[],
                                 premium=Decimal('120000'), period_months=6)
        self.assertEqual(q.annual_premium, Decimal('120000.00'))
        self.assertEqual(q.premium, Decimal('60000.00'))
        q.save(); q.save()
        q.refresh_from_db()
        self.assertEqual(q.premium, Decimal('60000.00'))

    def test_the_document_shows_the_real_annual_without_a_row_breakdown(self):
        q = Quote.objects.create(client_name='Typed 6mo doc', sections=[],
                                 premium=Decimal('120000'), period_months=6)
        html = quote_html(q)
        self.assertIn('120,000.00', html)          # the annual
        self.assertIn('60,000.00', html)           # what is charged
        self.assertIn('Charged for 6 months', html)

    def test_instalments_run_over_the_months_on_cover(self):
        half = Quote.objects.create(client_name='Half Yr 2', sections=self.SECTIONS,
                                    rate_incl_vat=False, period_months=6)
        self.assertIn('Monthly &mdash; 6 instalments', quote_html(half))


class QuoteExportTests(TestCase):
    """Excel must be LIVE (formulas, not pasted values) and must agree with the
    record; Word must build; neither carries the stamp or the QR (CFO 2026-08-11).
    """

    SECTIONS = [
        {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '8,500,000', 'rate': '0.15'},
        {'group': 'Money', 'name': 'Money in transit', 'sum_insured': '250,000', 'premium': '2,400'},
    ]

    def setUp(self):
        self.q = Quote.objects.create(client_name='Export Ltd', sections=self.SECTIONS,
                                      rate_incl_vat=False, exclusions=['Wear and tear.'])

    def _sheet(self):
        import io
        from openpyxl import load_workbook
        from underwriting.quote_export import build_xlsx
        return load_workbook(io.BytesIO(build_xlsx(self.q)))['Quotation']

    def test_the_premium_cells_are_formulas_not_values(self):
        ws = self._sheet()
        formulas = [c.value for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith('=')]
        self.assertTrue(any('*C' in f and '/100' in f for f in formulas))   # SI x rate
        self.assertTrue(any(f.startswith('=SUM(') for f in formulas))       # totals
        self.assertTrue(any('0.14' in f for f in formulas))                 # VAT

    def test_a_vat_inclusive_rate_does_not_inflate_the_workbook(self):
        # rate_incl_vat defaults True: the engine backs the net out of the gross,
        # so the sheet must divide by 1.14 too or the total came out 14% high
        # (Fable 2026-08-11).
        import io, re
        from openpyxl import load_workbook
        from underwriting.quote_export import build_xlsx
        q = Quote.objects.create(client_name='Incl VAT Ltd', rate_incl_vat=True,
                                 sections=[{'name': 'B', 'sum_insured': '1,000,000', 'rate': '3'}])
        ws = load_workbook(io.BytesIO(build_xlsx(q)))['Quotation']
        row_formula = next(str(c.value) for row in ws.iter_rows() for c in row
                           if isinstance(c.value, str) and c.value.startswith('=IF(N('))
        self.assertIn('/1.14', row_formula)

    def test_a_typed_row_premium_beats_the_rate_in_the_sheet(self):
        import io
        from openpyxl import load_workbook
        from underwriting.quote_export import build_xlsx
        q = Quote.objects.create(client_name='Flat Wins Ltd', rate_incl_vat=False,
                                 sections=[{'name': 'X', 'sum_insured': '1,000,000',
                                            'rate': '5', 'premium': '900'}])
        ws = load_workbook(io.BytesIO(build_xlsx(q)))['Quotation']
        vals = [c.value for row in ws.iter_rows() for c in row]
        self.assertIn(900.0, vals)                 # the typed figure, not 50,000

    def test_the_workbook_totals_agree_with_the_record(self):
        ws = self._sheet()
        labels = {str(c.value): c.row for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str)}
        # The VAT formula must reference the charged row, and the total must add
        # the two — checked structurally; the arithmetic itself is price()'s.
        vat_row = next(r for t, r in labels.items() if t.startswith('Value Added Tax'))
        tot_row = next(r for t, r in labels.items() if t == 'Total payable')
        self.assertIn('0.14', str(ws.cell(row=vat_row, column=6).value))
        self.assertTrue(str(ws.cell(row=tot_row, column=6).value).startswith('=F'))

    def test_word_builds_and_neither_export_carries_stamp_or_qr(self):
        from underwriting.quote_export import build_docx, build_xlsx
        self.assertGreater(len(build_docx(self.q)), 5000)
        import underwriting.quote_export as ex
        src = open(ex.__file__, encoding='utf-8').read()
        self.assertNotIn('verify_qr', src)
        self.assertNotIn('_STAMP', src)


class QuoteStyleTests(TestCase):
    """Simplified keeps the money and drops the detail."""

    SECTIONS = [
        {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '8,500,000', 'rate': '0.15'},
        {'group': 'Stock', 'name': 'Stock in trade', 'sum_insured': '9,400,000', 'rate': '0.45'},
    ]

    def setUp(self):
        self.q = Quote.objects.create(client_name='Style Ltd', sections=self.SECTIONS,
                                      rate_incl_vat=False,
                                      exclusions=['Wear and tear.'])

    def test_simplified_drops_the_schedule_and_the_exclusions(self):
        sim = quote_html(self.q, style='simple')
        self.assertNotIn('Cover Offered', sim)
        self.assertNotIn('What is not covered', sim)
        self.assertNotIn('Stock in trade', sim)

    def test_simplified_keeps_the_money_and_the_legal_notice(self):
        sim = quote_html(self.q, style='simple')
        self.assertIn('Total payable', sim)
        self.assertIn(f'{self.q.total:,.2f}', sim)
        self.assertIn('Payment Options', sim)
        self.assertIn('Duty of Disclosure', sim)

    def test_detailed_is_the_default_and_keeps_everything(self):
        det = quote_html(self.q)
        self.assertIn('Cover Offered', det)
        self.assertIn('What is not covered', det)
        self.assertIn(f'{self.q.total:,.2f}', det)


class GraphiteMappingTests(TestCase):
    """Our words -> Graphite's identifiers, for the policy-on-won handover."""

    def test_a_class_maps_to_a_product_and_plan_case_insensitively(self):
        GraphiteMapping.objects.create(kind=GraphiteMapping.Kind.CLASS_OF_BUSINESS,
                                      omni_value='Commercial Combined',
                                      product_id=7, plan_id=14)
        m = GraphiteMapping.resolve(GraphiteMapping.Kind.CLASS_OF_BUSINESS,
                                    'commercial combined')
        self.assertEqual((m.product_id, m.plan_id), (7, 14))

    def test_a_broker_maps_to_an_agent_code(self):
        GraphiteMapping.objects.create(kind=GraphiteMapping.Kind.BROKER,
                                      omni_value='Marsh Botswana (Pty) Ltd',
                                      agent_code='AG-0058')
        self.assertEqual(
            GraphiteMapping.resolve(GraphiteMapping.Kind.BROKER,
                                    'Marsh Botswana (Pty) Ltd').agent_code, 'AG-0058')

    def test_an_unmapped_value_returns_none_rather_than_a_guess(self):
        self.assertIsNone(GraphiteMapping.resolve(
            GraphiteMapping.Kind.CLASS_OF_BUSINESS, 'Something We Never Mapped'))

    def test_an_inactive_row_does_not_resolve(self):
        GraphiteMapping.objects.create(kind=GraphiteMapping.Kind.BROKER,
                                      omni_value='Old Broker', agent_code='AG-1',
                                      active=False)
        self.assertIsNone(GraphiteMapping.resolve(GraphiteMapping.Kind.BROKER, 'Old Broker'))


class MinimumPremiumTests(TestCase):
    """Straight pro-rata makes a short period cheap: six months is half, one month
    a twelfth. A correct RATE can still produce a premium too small to be worth
    writing, so the money has its own floor (CFO 2026-08-11).
    """

    SECTIONS = [{'name': 'Buildings', 'sum_insured': '1,000,000', 'rate': '1'}]

    def _quote(self, months=12, cls='Fire'):
        comp = Company.objects.create(name=f'ADIC-{months}-{cls}', code=f'AD{months}{cls[:2]}')
        u = User.objects.create(username=f'uw_min_{months}_{cls}')
        return Quote.objects.create(client_name='Min Ltd', class_of_business=cls,
                                    sections=self.SECTIONS, rate_incl_vat=False,
                                    period_months=months, company=comp, underwriter=u), u

    def test_no_minimum_set_means_no_check(self):
        q, u = self._quote()
        self.assertEqual(quote_below_min_premium(q)[1], None)
        q.issue(u)                                    # issues fine

    def test_a_full_year_above_the_minimum_issues(self):
        QuoteRateFloor.objects.create(class_of_business='', min_premium=Decimal('5000'))
        q, u = self._quote()                          # 1% of 1,000,000 = 10,000
        q.issue(u)
        self.assertEqual(q.status, Quote.Status.ISSUED)

    def test_a_short_period_that_falls_under_the_minimum_is_refused(self):
        # 1 month of a 10,000 annual = 833.33, under a 5,000 minimum.
        QuoteRateFloor.objects.create(class_of_business='', min_premium=Decimal('5000'))
        q, u = self._quote(months=1)
        minimum, under = quote_below_min_premium(q)
        self.assertEqual(minimum, Decimal('5000'))
        self.assertEqual(under, q.premium)
        with self.assertRaises(ValidationError):
            q.issue(u)

    def test_the_minimum_is_per_class(self):
        QuoteRateFloor.objects.create(class_of_business='', min_premium=Decimal('500'))
        QuoteRateFloor.objects.create(class_of_business='Fire', min_premium=Decimal('9000'))
        q, u = self._quote(months=6)                  # 5,000 charged
        self.assertEqual(quote_below_min_premium(q)[0], Decimal('9000'))
        with self.assertRaises(ValidationError):
            q.issue(u)


class QuoteRateFloorTests(TestCase):
    """No quotation may be rated below the floor set for its class."""

    def test_default_floor_applies_to_any_class(self):
        QuoteRateFloor.objects.create(class_of_business='', min_rate_pct=Decimal('0.5'))
        self.assertEqual(QuoteRateFloor.resolve('Anything'), Decimal('0.5'))

    def test_class_override_wins_over_default(self):
        QuoteRateFloor.objects.create(class_of_business='', min_rate_pct=Decimal('0.5'))
        QuoteRateFloor.objects.create(class_of_business='Motor Fleet', min_rate_pct=Decimal('2'))
        self.assertEqual(QuoteRateFloor.resolve('motor fleet'), Decimal('2'))  # case-insensitive

    def test_no_floor_set_means_no_check(self):
        self.assertEqual(QuoteRateFloor.resolve('Fire'), Decimal('0'))

    def test_a_rate_below_the_floor_is_flagged(self):
        QuoteRateFloor.objects.create(class_of_business='', min_rate_pct=Decimal('1'))
        q = Quote(client_name='X', class_of_business='Fire', rate_pct=Decimal('0.5'), rate_incl_vat=True,
                  sections=[{'name': 'Buildings', 'sum_insured': '1000000', 'rate': '0.4'}])
        floor, below = quote_rates_below_floor(q)
        self.assertEqual(floor, Decimal('1'))
        self.assertEqual(len(below), 2)  # whole-quote 0.5% AND the 0.4% row

    def test_a_floored_quote_refuses_to_issue(self):
        QuoteRateFloor.objects.create(class_of_business='', min_rate_pct=Decimal('1'))
        comp = Company.objects.create(name='ADIC', code='ADIC')
        u = User.objects.create(username='uw1')
        q = Quote.objects.create(client_name='Low Ball Ltd', class_of_business='Fire', company=comp,
                                 sections=[{'name': 'Buildings', 'sum_insured': '1000000', 'rate': '0.4'}],
                                 rate_incl_vat=True)
        with self.assertRaises(ValidationError):
            q.issue(u)


class RenewalLookupTests(TestCase):
    """Renewal pre-fill reads a Graphite policy by number and returns only
    non-personal fields; the Graphite call is mocked here."""

    def test_a_bad_policy_number_is_rejected(self):
        from underwriting.renewal_lookup import lookup_policy
        self.assertFalse(lookup_policy('not a policy!!')['found'])

    def test_it_returns_company_class_premium_and_claims(self):
        from unittest.mock import patch
        from underwriting.renewal_lookup import lookup_policy
        with patch('aware.engine.run_select_params') as q:
            q.side_effect = [
                [{'policy_number': 'COMG2025189299', 'business_name': 'Bonanza Equipment',
                  'annual_premium': '50000', 'sum_assured': '1000000', 'customer_id': 7,
                  'product_name': 'Fire & Perils', 'line_of_business': 'Commercial'}],
                [{'n': 3, 'reserve': '20000', 'paid': '5000'}],
            ]
            r = lookup_policy('COMG2025189299')
        self.assertTrue(r['found'])
        self.assertEqual(r['client_name'], 'Bonanza Equipment')
        self.assertEqual(r['class_of_business'], 'Fire & Perils')
        self.assertEqual(r['sum_insured'], '1000000')
        self.assertEqual(r['claims']['count'], 3)

    def test_an_unknown_policy_reports_not_found(self):
        from unittest.mock import patch
        from underwriting.renewal_lookup import lookup_policy
        with patch('aware.engine.run_select_params') as q:
            q.return_value = []
            self.assertFalse(lookup_policy('COMG9999999999')['found'])


class ExpiringQuoteReminderTests(TestCase):
    def test_it_reminds_for_a_quote_about_to_lapse(self):
        from datetime import timedelta
        from unittest.mock import patch
        from django.core.management import call_command
        from django.utils import timezone
        u = User.objects.create(username='uw_exp', email='uw_exp@alphadirect.co.bw')
        comp = Company.objects.create(name='ADIC-EXP', code='ADEXP')
        Quote.objects.create(client_name='Soon Ltd', company=comp, underwriter=u,
                             premium=Decimal('1000'), status=Quote.Status.ISSUED,
                             valid_until=timezone.localdate() + timedelta(days=3))
        with patch('core.notifications.send_html_with_cfo_cc', return_value=1) as send:
            call_command('remind_expiring_quotes', '--days', '5')
        self.assertTrue(send.called)


class QuoteNumberTests(TestCase):

    def test_every_quote_gets_its_own_number(self):
        a = Quote.objects.create(client_name='Client A', premium=Decimal('1000'))
        b = Quote.objects.create(client_name='Client B', premium=Decimal('1000'))
        self.assertTrue(a.quote_number.startswith('Q-'))
        self.assertNotEqual(a.quote_number, b.quote_number)

    def test_numbers_run_in_sequence(self):
        first = Quote.objects.create(client_name='Client A', premium=Decimal('1'))
        second = Quote.objects.create(client_name='Client B', premium=Decimal('1'))
        self.assertEqual(int(second.quote_number.rsplit('-', 1)[1]),
                         int(first.quote_number.rsplit('-', 1)[1]) + 1)


class QuoteIssueTests(TestCase):

    def setUp(self):
        self.user = User.objects.create_user('uw', password='x')
        # A quotation must belong to an entity before it can be issued.
        self.company = Company.objects.create(code='QT1', name='Quote Test Co')

    def test_a_normal_quote_issues(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('50000'),
                                 company=self.company)
        q.issue(self.user)
        self.assertEqual(q.status, Quote.Status.ISSUED)
        self.assertIsNotNone(q.issued_at)
        self.assertEqual(q.issued_by, self.user)

    def test_a_guessed_premium_cannot_be_issued(self):
        # The CFO allowed Aria to suggest a premium. This is the guard that stops
        # a guess reaching a broker: somebody has to look at it first.
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('50000'),
                                 premium_is_suggested=True)
        with self.assertRaises(ValidationError):
            q.issue(self.user)
        q.refresh_from_db()
        self.assertEqual(q.status, Quote.Status.DRAFT)

    def test_confirming_the_premium_releases_it(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('50000'),
                                 premium_is_suggested=True, company=self.company)
        q.premium_is_suggested = False
        q.save()
        q.issue(self.user)
        self.assertEqual(q.status, Quote.Status.ISSUED)

    def test_a_quote_with_no_client_cannot_be_issued(self):
        # Two clients sharing one file is how this started.
        q = Quote.objects.create(client_name='', premium=Decimal('50000'))
        with self.assertRaises(ValidationError):
            q.issue(self.user)

    def test_a_quote_with_no_premium_cannot_be_issued(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('0'))
        with self.assertRaises(ValidationError):
            q.issue(self.user)

    def test_it_cannot_be_issued_twice(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('50000'),
                                 company=self.company)
        q.issue(self.user)
        with self.assertRaises(ValidationError):
            q.issue(self.user)

    def test_validity_defaults_to_thirty_days(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('1'))
        self.assertIsNotNone(q.valid_until)


class QuoteBypassTests(TestCase):
    """The guards must not be walk-throughable by an ordinary API call.

    DeepSeek review 2026-08-08 found three ways round them: PATCH status to
    'issued' and skip issue() entirely; PATCH premium_is_suggested to false and
    release an unconfirmed premium; and change the premium after issue while the
    client holds a PDF with the old figures. Each one is a test here.
    """

    def setUp(self):
        self.user = User.objects.create_user('uw2', password='x')
        self.company = Company.objects.create(code='QT2', name='Quote Test Co 2')

    def test_status_cannot_be_set_by_the_serializer(self):
        from underwriting.serializers import QuoteSerializer
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('1000'))
        ser = QuoteSerializer(q, data={'status': 'issued'}, partial=True)
        self.assertTrue(ser.is_valid(), ser.errors)
        ser.save()
        q.refresh_from_db()
        self.assertEqual(q.status, Quote.Status.DRAFT,
                         'status must only move through issue()/outcome()')
        self.assertIsNone(q.issued_at)

    def test_the_suggested_flag_cannot_be_cleared_by_the_serializer(self):
        from underwriting.serializers import QuoteSerializer
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('1000'),
                                 premium_is_suggested=True)
        ser = QuoteSerializer(q, data={'premium_is_suggested': False}, partial=True)
        self.assertTrue(ser.is_valid(), ser.errors)
        ser.save()
        q.refresh_from_db()
        self.assertTrue(q.premium_is_suggested,
                        'only confirm-premium may clear the guess flag')
        with self.assertRaises(ValidationError):
            q.issue(self.user)

    def test_typing_a_premium_is_a_genuine_confirmation(self):
        # The one legitimate way to clear it: the underwriter types the figure.
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('1000'),
                                 premium_is_suggested=True, company=self.company)
        q.premium = Decimal('2000')
        q.premium_is_suggested = False        # what perform_update does
        q.save()
        q.issue(self.user)
        self.assertEqual(q.status, Quote.Status.ISSUED)

    def test_resending_the_same_suggested_figure_is_not_a_confirmation(self):
        # Echoing Aria's own number back is not somebody looking at it and
        # agreeing — that would clear the gate by accident.
        from underwriting.serializers import QuoteSerializer
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('1000'),
                                 premium_is_suggested=True)
        ser = QuoteSerializer(q, data={'premium': '1000'}, partial=True)
        self.assertTrue(ser.is_valid(), ser.errors)
        new_premium = ser.validated_data.get('premium')
        cleared = new_premium is not None and new_premium != q.premium
        self.assertFalse(cleared, 'the same figure must not clear the guess flag')

    def test_a_quote_with_no_company_cannot_be_issued(self):
        # An entity-less quote is invisible to every per-company view.
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('1000'))
        self.assertIsNone(q.company_id)
        with self.assertRaises(ValidationError):
            q.issue(self.user)

    def test_company_is_not_taken_from_the_payload(self):
        from underwriting.serializers import QuoteSerializer
        self.assertIn('company', QuoteSerializer.Meta.read_only_fields)

    def test_the_typed_note_is_kept_for_qc(self):
        # It was documented as the QC trail but silently dropped by the
        # serializer, so nothing was ever recorded.
        from underwriting.serializers import QuoteSerializer
        ser = QuoteSerializer(data={'client_name': 'Test Client', 'premium': '1000',
                                    'source_text': 'fleet of 14, premium 1000'})
        self.assertTrue(ser.is_valid(), ser.errors)
        q = ser.save()
        self.assertEqual(q.source_text, 'fleet of 14, premium 1000')


class SuggestedFlagIsDerivedTests(APITestCase):
    """The gate between a guessed premium and a broker — tested through the real
    request path.

    The previous version of this class carried its OWN private copy of the rule
    and never called the server. It passed no matter what shipped, and the copy it
    was testing was the substring algorithm that had already been proven wrong.
    A test that reimplements the thing it is testing tests nothing.
    """

    def setUp(self):
        self.company = Company.objects.create(code='QSG', name='Quote Gate Co')
        self.user = User.objects.create_user('uw_gate', password='x',
                                             is_superuser=True, is_staff=True)
        self.client.force_authenticate(self.user)

    def _create(self, **payload):
        body = {'client_name': 'Test Client', 'premium': '418600', 'sections': []}
        body.update(payload)
        return self.client.post(
            f'/api/v1/underwriting/quotes/?company={self.company.id}',
            body, format='json')

    def test_a_premium_not_in_the_note_is_flagged_as_a_guess(self):
        r = self._create(source_text='fleet of 14 vehicles, contents sum 186,000')
        self.assertEqual(r.status_code, 201, r.content[:200])
        self.assertTrue(r.data['premium_is_suggested'])

    def test_a_premium_the_underwriter_typed_is_not_flagged(self):
        r = self._create(source_text='fleet of 14, premium 418,600')
        self.assertEqual(r.status_code, 201, r.content[:200])
        self.assertFalse(r.data['premium_is_suggested'])

    def test_clearing_the_note_cannot_lower_the_flag(self):
        """Draft from a note, let Aria suggest, then empty the text box before
        saving. An empty note used to derive "not suggested" and the guess issued."""
        r = self._create(source_text='', premium_is_suggested=True)
        self.assertEqual(r.status_code, 201, r.content[:200])
        self.assertTrue(r.data['premium_is_suggested'])

    def test_a_flagged_quote_refuses_to_issue(self):
        r = self._create(source_text='', premium_is_suggested=True)
        issued = self.client.post(f"/api/v1/underwriting/quotes/{r.data['id']}/issue/")
        self.assertEqual(issued.status_code, 400, issued.content[:200])


class QuoteUrlLayerTests(APITestCase):
    """Real requests through the URL layer.

    Every test above passed while the feature was dead on arrival: the viewset
    named a throttle scope that does not exist, so DRF raised inside
    check_throttles and EVERY request 500'd — list included. Model-level tests
    cannot see that, because they never touch a URL. Fable found it by making
    actual calls. These are those calls.
    """

    def setUp(self):
        self.company = Company.objects.create(code='QTU', name='Quote URL Co')
        self.user = User.objects.create_user('uw_url', password='x', is_superuser=True,
                                             is_staff=True)
        self.client.force_authenticate(self.user)

    def test_the_register_actually_answers(self):
        r = self.client.get('/api/v1/underwriting/quotes/')
        self.assertEqual(r.status_code, 200, r.content[:300])

    def test_a_quote_can_be_created_with_the_entity_on_the_post(self):
        r = self.client.post(
            f'/api/v1/underwriting/quotes/?company={self.company.id}',
            {'client_name': 'Test Client', 'premium': '100000', 'sections': []},
            format='json')
        self.assertEqual(r.status_code, 201, r.content[:300])
        self.assertTrue(r.data['quote_number'].startswith('Q-'))
        self.assertEqual(r.data['vat'], '14000.00')
        self.assertEqual(r.data['total'], '114000.00')

    def test_creating_without_an_entity_is_refused_not_crashed(self):
        r = self.client.post('/api/v1/underwriting/quotes/',
                             {'client_name': 'Test Client', 'premium': '1000'},
                             format='json')
        self.assertEqual(r.status_code, 400, r.content[:300])

    def test_the_ask_box_endpoint_answers(self):
        r = self.client.post('/api/v1/underwriting/quotes/draft-from-text/',
                             {'text': 'Test Client, fire cover 1m, premium 5,000'},
                             format='json')
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertIn('draft', r.data)

    def test_an_issued_quote_cannot_be_deleted(self):
        q = Quote.objects.create(client_name='Test Client', premium=Decimal('1000'),
                                 company=self.company)
        q.issue(self.user)
        r = self.client.delete(f'/api/v1/underwriting/quotes/{q.id}/')
        self.assertIn(r.status_code, (400, 403, 405), r.content[:200])
        self.assertTrue(Quote.objects.filter(pk=q.pk).exists())

    def test_export_uses_fmt_not_the_drf_reserved_format_param(self):
        """Excel/Word download was dead on arrival (Motlatsi item 2, 2026-08-12).

        The chooser was `?format=`, but `format` is DRF's reserved content-
        negotiation query param — DRF looked for an 'xlsx' renderer, found none,
        and 404'd before the action ran. Every Excel/Word download failed; only
        the PDF (its own `/pdf/` action, no query param) survived. The endpoint
        now reads `fmt`. The `?fmt=docx` assertion FAILS without the fix — the old
        code, reading `format`, sees no chooser and hands back xlsx instead.
        """
        XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        q = Quote.objects.create(client_name='Export Co', premium=Decimal('100000'),
                                 company=self.company)
        base = f'/api/v1/underwriting/quotes/{q.id}/export/'
        r_xlsx = self.client.get(base + '?fmt=xlsx')
        self.assertEqual(r_xlsx.status_code, 200, r_xlsx.content[:200])
        self.assertEqual(r_xlsx['Content-Type'], XLSX)
        r_docx = self.client.get(base + '?fmt=docx')
        self.assertEqual(r_docx.status_code, 200, r_docx.content[:200])
        self.assertEqual(r_docx['Content-Type'], DOCX)      # xlsx here without the fix
        # The reserved name is still swallowed by DRF — documents WHY we renamed it.
        self.assertEqual(self.client.get(base + '?format=xlsx').status_code, 404)


class PremiumCameFromTheNoteTests(TestCase):
    """The single guard between a guessed premium and a broker.

    Fable defeated the first version: it stripped every non-digit and asked
    whether the premium was a substring, so "fleet of 14 vehicles, contents sum
    186,000" became "14186000" — which contains "418600". A premium Aria invented
    read as one the underwriter had typed, and issued.
    """

    def test_the_concatenation_case_that_defeated_the_first_version(self):
        note = 'fleet of 14 vehicles, contents sum 186,000, no price agreed'
        self.assertFalse(premium_appears_in('418600', note))

    def test_a_premium_the_underwriter_typed_is_recognised(self):
        self.assertTrue(premium_appears_in('418600', 'fleet of 14, premium 418,600'))

    def test_the_shorthand_forms_match_the_same_number(self):
        self.assertTrue(premium_appears_in('418600', 'premium 418.6k'))
        self.assertTrue(premium_appears_in('418600', 'premium 0.4186m'))

    def test_a_different_figure_does_not_match(self):
        self.assertFalse(premium_appears_in('418600', 'premium 418,700'))

    def test_no_note_means_it_did_not_come_from_one(self):
        self.assertFalse(premium_appears_in('418600', ''))


class FallbackReaderTests(TestCase):
    """The reader that runs when Aria is down — i.e. when the underwriter is
    already under pressure. Both cases below were invisible to every other test
    and only showed up when the screen was rendered and the fields read.
    """

    NOTE = ('Tsholo Wholesalers, fire and perils on the warehouse at Plot 5512 '
            'Gaborone. Buildings 12m reinstatement, stock 4.5m, excess 25k. '
            'Broker Minet. Premium 178,250.')

    def test_a_plot_number_never_becomes_a_sum_insured(self):
        fire = [s for s in _fallback(self.NOTE)['sections']
                if s['name'] == 'Fire and allied perils'][0]
        self.assertNotEqual(fire['sum_insured'], '5512')
        self.assertEqual(fire['sum_insured'], '12m')

    def test_the_broker_is_read_however_it_is_capitalised(self):
        for note in ('Broker Minet. Premium 1000', 'broker Minet. Premium 1000',
                     'BROKER Minet. Premium 1000'):
            self.assertEqual(_fallback(note)['broker'], 'Minet', note)

    def test_the_premium_still_reads(self):
        self.assertEqual(_fallback(self.NOTE)['premium'], '178,250')


class PrintedFiguresTests(TestCase):
    """What the broker actually reads."""

    def test_shorthand_is_printed_as_a_figure(self):
        rows = _display_rows([{'sum_insured': '8m', 'excess': '20k'}])
        self.assertEqual(rows[0]['sum_insured'], '8,000,000.00')
        self.assertEqual(rows[0]['excess'], '20,000.00')

    def test_wording_is_left_exactly_as_typed(self):
        rows = _display_rows([{'sum_insured': 'Included', 'excess': 'Nil'}])
        self.assertEqual(rows[0]['sum_insured'], 'Included')
        self.assertEqual(rows[0]['excess'], 'Nil')

    def test_a_blank_stays_blank_on_a_kept_row(self):
        # A blank sum must never be formatted into money. The row needs a price
        # to stay on the schedule at all — a fully empty row no longer prints
        # (CFO fix list 2026-08-11 #4: no dash-only rows).
        rows = _display_rows([{'name': 'Fees', 'sum_insured': '', 'excess': '', 'premium': '100'}])
        self.assertEqual(rows[0]['sum_insured'], '')

    def test_a_fully_empty_row_does_not_print(self):
        self.assertEqual(_display_rows([{'sum_insured': '', 'excess': ''}]), [])


class ClientNeverSeesReinsuranceTests(TestCase):
    """CFO 2026-08-08: the client does not need to know how we lay the risk off.

    The underwriter's note carries it as a matter of course, and everything in the
    note flows through Aria into the cover rows and onto the client's PDF.
    """

    def test_a_reinsurance_row_never_prints(self):
        rows = _display_rows([
            {'name': 'Fire and allied perils', 'sum_insured': '8m'},
            {'name': 'Facultative placement — 60% ceded', 'sum_insured': '4.8m'},
            {'group': 'Treaty', 'name': 'Surplus', 'sum_insured': '2m'},
        ])
        self.assertEqual([r['name'] for r in rows], ['Fire and allied perils'])

    def test_cover_survives_but_the_arrangement_behind_it_does_not(self):
        rows = _display_rows([
            {'name': 'Third-party liability', 'sum_insured': '5m',
             'note': 'we retain 1m, balance to the treaty'},
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['sum_insured'], '5,000,000.00')
        self.assertEqual(rows[0]['note'], '')

    def test_ordinary_cover_wording_is_untouched(self):
        rows = _display_rows([
            {'name': 'Riot, strike and civil commotion', 'sum_insured': 'Included',
             'note': 'as per policy wording'},
        ])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['note'], 'as per policy wording')


class OutcomeIsCorrectableTests(APITestCase):
    """A mis-click used to be permanent — status is read-only on the serializer,
    so not even a PATCH could undo it, and the conversion rate the register exists
    to report stayed wrong for good."""

    def setUp(self):
        self.company = Company.objects.create(code='QOC', name='Quote Outcome Co')
        self.user = User.objects.create_user('uw_outcome', password='x',
                                             is_superuser=True, is_staff=True)
        self.client.force_authenticate(self.user)
        self.quote = Quote.objects.create(client_name='Test Client',
                                          premium=Decimal('1000'), company=self.company)
        self.quote.issue(self.user)

    def _mark(self, result):
        return self.client.post(f'/api/v1/underwriting/quotes/{self.quote.id}/outcome/',
                                {'result': result}, format='json')

    def test_a_wrong_outcome_can_be_put_right(self):
        self.assertEqual(self._mark('won').status_code, 200)
        r = self._mark('lost')
        self.assertEqual(r.status_code, 200, r.content[:200])
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.status, Quote.Status.LOST)

    def test_a_draft_still_has_no_outcome(self):
        draft = Quote.objects.create(client_name='Test Client',
                                     premium=Decimal('500'), company=self.company)
        r = self.client.post(f'/api/v1/underwriting/quotes/{draft.id}/outcome/',
                             {'result': 'won'}, format='json')
        self.assertEqual(r.status_code, 400, r.content[:200])


class QuoteEntityWriteWallTests(APITestCase):
    """A quotation must not be saveable onto an entity the caller cannot see.

    Read-side scoping hides other companies' quotes, which makes this easy to
    miss: the caller cannot READ entity B, so nobody checks whether they can
    WRITE to it. They could. `resolve_company_id_param` returns the company
    named in the query string, not one the user is allowed to use, so a
    restricted user could put a quotation on another company's books by
    changing a URL. Found in the 8 Aug 2026 sweep.
    """

    def setUp(self):
        from core.models import UserCompanyAccess
        self.mine = Company.objects.create(code='EWA', name='My Entity')
        self.theirs = Company.objects.create(code='EWB', name='Someone Else')
        self.user = User.objects.create_user('uw_walled', password='x', is_staff=True)
        # An explicit grant to ONE entity is what makes scoping bite.
        UserCompanyAccess.objects.create(user=self.user, company=self.mine)
        self.client.force_authenticate(self.user)

    def test_cannot_save_a_quote_onto_another_entity(self):
        r = self.client.post(
            f'/api/v1/underwriting/quotes/?company={self.theirs.id}',
            {'client_name': 'Test Client', 'premium': '1000', 'sections': []},
            format='json')
        self.assertIn(r.status_code, (400, 403), r.content[:300])
        self.assertFalse(
            Quote.objects.filter(company=self.theirs).exists(),
            'a quotation was written onto an entity the user has no access to')

    def test_own_entity_still_works(self):
        r = self.client.post(
            f'/api/v1/underwriting/quotes/?company={self.mine.id}',
            {'client_name': 'Test Client', 'premium': '1000', 'sections': []},
            format='json')
        self.assertEqual(r.status_code, 201, r.content[:300])


class VatRoundsHalfUpTests(TestCase):
    """VAT rounds half UP, not to even (CFO 2026-08-09).

    Decimal's default is ROUND_HALF_EVEN — a programming default, not a tax
    decision. On a premium of 5,000.75 that put VAT at 700.10 while a broker
    working it out by hand gets 700.11, so the quotation disagreed with the
    person reading it by a thebe.
    """

    def test_a_half_thebe_rounds_up_not_to_even(self):
        m = price(Decimal('5000.75'))          # 14% = 700.105 exactly
        self.assertEqual(m['vat'], Decimal('700.11'))

    def test_the_total_is_the_two_figures_shown_added_up(self):
        # Whatever rounding is used, premium + VAT must equal the printed total,
        # or the quotation fails its own arithmetic in front of a broker.
        for prem in ('5000.75', '1000.25', '12345.68', '418600', '0'):
            m = price(Decimal(prem))
            self.assertEqual(m['premium'] + m['vat'], m['total'], f'premium {prem}')

    def test_a_case_that_was_already_right_is_unchanged(self):
        m = price(Decimal('418600'))
        self.assertEqual(m['vat'], Decimal('58604.00'))
        self.assertEqual(m['total'], Decimal('477204.00'))


class WCACommonLawLiabilityTests(TestCase):
    """Common Law Liability always prints on a Workmen's Compensation cover.

    Underwriting, 19 August 2026, feature request from /underwriting/quotes:
    she listed the benefits a captured WCA cover should show and marked one of them
    *"(should always be there)"* — Common Law Liability P1,000,000. Her screenshot
    proved the other four already print (Death, Permanent Total Disablement,
    Temporary Total Disablement, Medical Expenses) because the underwriter types
    them; that one row is the standard benefit most easily forgotten, and a quote
    that leaves it off understates the cover being sold.
    """

    def _rows(self, sections):
        from underwriting.quote_render import _display_rows
        return _display_rows(sections)

    def _find(self, rows, needle):
        return [r for r in rows
                if needle.lower() in str(r.get('name') or '').lower()]

    def test_it_is_added_to_a_workmens_section_that_omits_it(self):
        rows = self._rows([
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500', 'rate': '1.00'},
            {'group': "Workmen's Compensation", 'name': 'Death',
             'sum_insured': '200000'},
            {'group': "Workmen's Compensation", 'name': 'Medical Expenses',
             'sum_insured': '75000'},
        ])
        cll = self._find(rows, 'Common Law Liability')
        self.assertEqual(len(cll), 1, 'exactly one Common Law Liability row')
        self.assertEqual(cll[0]['sum_insured'], '1,000,000.00',
                         'formatted like every other figure in the column')

    def test_it_is_not_duplicated_when_the_underwriter_typed_it(self):
        rows = self._rows([
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500'},
            {'group': "Workmen's Compensation", 'name': 'Common Law Liability',
             'sum_insured': '2000000'},
        ])
        cll = self._find(rows, 'Common Law Liability')
        self.assertEqual(len(cll), 1, 'their own row stands, no second one added')
        self.assertEqual(cll[0]['sum_insured'], '2,000,000.00',
                         "the underwriter's own figure must win")

    def test_every_spelling_of_workmens_counts(self):
        for label in ("Workmen's Compensation", "Workman's Compensation",
                      'WORKMENS COMPENSATION', "WORKMAN'S COMPENSATION"):
            rows = self._rows([{'group': label, 'name': label,
                                'sum_insured': '1049500'}])
            self.assertEqual(len(self._find(rows, 'Common Law Liability')), 1,
                             f'not added for {label!r}')

    def test_a_quote_with_no_workmens_cover_is_untouched(self):
        rows = self._rows([
            {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '8000000'},
            {'group': 'Motor', 'name': 'Fleet', 'sum_insured': '1200000'},
        ])
        self.assertEqual(self._find(rows, 'Common Law Liability'), [])

    def test_it_carries_no_rate_and_no_premium_of_its_own(self):
        """A benefit limit, like Death and PTD — never a priced row."""
        rows = self._rows([
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500', 'rate': '1.00', 'premium': '9206.14'},
        ])
        cll = self._find(rows, 'Common Law Liability')[0]
        self.assertIn(str(cll.get('rate') or ''), ('', 'None'))
        self.assertIn(str(cll.get('premium') or ''), ('', 'None'))

    def test_it_cannot_change_the_premium(self):
        """The premium comes from quote.sections, which this never touches."""
        from decimal import Decimal
        from underwriting.quote_parse import section_premium_rows
        sections = [
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500', 'rate': '1.00'},
        ]
        before = section_premium_rows(list(sections), Decimal('1.00'))
        self._rows(sections)
        after = section_premium_rows(list(sections), Decimal('1.00'))
        self.assertEqual(before, after,
                         'rendering must not feed a new row into the pricing')
        self.assertEqual(len(sections), 1,
                         'the caller\'s own list must not gain a row')

    def test_two_workmens_sites_print_one_benefit_between_them(self):
        """Same-family sections MERGE into one printed section, so one row is right.

        The renderer folds "Workmen's Compensation — site A" and "— site B" into a
        single section (the same behaviour that merged three "Fire" sections on a
        real schedule). Printing the benefit twice inside one merged section would
        read as a duplicate to a broker.
        """
        rows = self._rows([
            {'group': "Workmen's Compensation — site A",
             'name': "Workmen's Compensation", 'sum_insured': '500000'},
            {'group': "Workmen's Compensation — site B",
             'name': "Workmen's Compensation", 'sum_insured': '600000'},
        ])
        self.assertEqual(len(self._find(rows, 'Common Law Liability')), 1)

    def test_it_does_not_mutate_the_list_it_is_given(self):
        """Raised on review: it built its answer by inserting into the caller's list.

        Safe where it is called from today, and that is exactly the kind of property
        that survives until somebody calls it from somewhere else.
        """
        from underwriting.quote_render import _ensure_common_law_liability
        given = [{'group': "Workmen's Compensation",
                  'name': "Workmen's Compensation", 'sum_insured': '1049500'}]
        returned = _ensure_common_law_liability(given)
        self.assertEqual(len(given), 1, "the caller's list must be untouched")
        self.assertEqual(len(returned), 2, 'the answer carries the added row')

    # ── the shape the real UI emits, not the one the developer imagined ──────
    def test_a_typed_benefit_on_a_blank_group_row_is_not_duplicated(self):
        """The frontend's add-row default is group:'' — rows INHERIT their section.

        Keying on the row's own fields found the heading row and missed every
        benefit line beneath it, so a Common Law Liability the underwriter had
        typed went undetected and a second was added: two contradictory liability
        figures on one client document. Same failure class as the A.1 suite being
        green against a payload no user sends.
        """
        rows = self._rows([
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500', 'rate': '1.00'},
            {'group': '', 'name': 'Death', 'sum_insured': '200000'},
            {'group': '', 'name': 'Common Law Liability', 'sum_insured': '2000000'},
        ])
        cll = self._find(rows, 'Common Law Liability')
        self.assertEqual(len(cll), 1, 'their typed row must be the only one')
        self.assertEqual(cll[0]['sum_insured'], '2,000,000.00',
                         "the underwriter's own figure wins")

    def test_it_is_still_added_when_the_benefits_are_blank_group_rows(self):
        rows = self._rows([
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500'},
            {'group': '', 'name': 'Death', 'sum_insured': '200000'},
            {'group': '', 'name': 'Medical Expenses', 'sum_insured': '75000'},
        ])
        self.assertEqual(len(self._find(rows, 'Common Law Liability')), 1)

    def test_the_added_row_lands_under_its_own_section(self):
        rows = self._rows([
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500'},
            {'group': '', 'name': 'Death', 'sum_insured': '200000'},
            {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '8000000'},
        ])
        names = [str(r.get('name') or '') for r in rows]
        self.assertLess(names.index('Common Law Liability'), names.index('Buildings'),
                        'it must not drift under the next cover')

    def test_the_excel_and_word_copies_carry_it_too(self):
        """quote_export promises all three documents agree. They must."""
        from underwriting.quote_render import sections_with_guaranteed_rows
        pairs = sections_with_guaranteed_rows([
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500', 'rate': '1.00'},
            {'group': '', 'name': 'Death', 'sum_insured': '200000'},
        ])
        added = [(i, r) for i, r in pairs if r.get('_cll_added')]
        self.assertEqual(len(added), 1)
        self.assertIsNone(added[0][0],
                          'a guaranteed row takes no premium, so it carries no index')

    def test_the_export_view_keeps_every_premium_on_its_own_row(self):
        """Inserting into the record would slide every later premium onto the
        wrong line — far worse than the missing line."""
        from underwriting.quote_render import sections_with_guaranteed_rows
        src = [
            {'group': "Workmen's Compensation", 'name': "Workmen's Compensation",
             'sum_insured': '1049500'},
            {'group': '', 'name': 'Death', 'sum_insured': '200000'},
            {'group': 'Fire', 'name': 'Buildings', 'sum_insured': '8000000'},
        ]
        pairs = sections_with_guaranteed_rows(src)
        for idx, row in pairs:
            if idx is None:
                continue
            self.assertEqual(row['name'], src[idx]['name'],
                             'each original row keeps its own index')
