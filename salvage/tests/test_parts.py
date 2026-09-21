"""Tests for Veritas Parts & Savings (CFO 2026-09-10).

Fixtures are built in code, in the exact shape of the two workbooks the Parts
& Assessments team sends — two-level header, stated-totals block in column M,
a SUMMARY sheet that is rounded and derived, TOTAL rows, and the 'GLASS'
spacer row inside DEALERSHIP. No real registration numbers or vehicles go
into the repo.

Each test covers something that would actually hurt if it broke:
  * the month must come from the Req Auth Date, never from the A1 title (the
    July and August sheets of the real June workbook both say "June 2026");
  * savings are recomputed as quote − assessment, and the file's own savings
    column is kept so the variance shows instead of being overwritten;
  * SUMMARY is never a row-level source — only its stated totals, its
    contract-pricing line and its five-FY history;
  * TOTAL rows and the GLASS spacer are never counted;
  * re-uploading a month replaces it and never doubles it.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.urls import reverse
from rest_framework.test import APITestCase

from salvage.parts_models import (
    AssessmentSaving, PartsContractPricing, PartsHistory, PartsSpend, PartsUpload,
)
from salvage.parts_parser import (
    detect_workbook_kind, fiscal_year_label, parse_assessment_savings, parse_parts_summary,
)


# ── fixture builders ───────────────────────────────────────────────────────

def assessment_workbook(rows_by_sheet: dict[str, list[list]],
                        stated: dict[str, list] | None = None) -> bytes:
    """Build an Assessment Savings workbook. `rows_by_sheet` maps sheet name to
    data rows of 17 cells. Every sheet's A1 deliberately says June 2026."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in rows_by_sheet.items():
        worksheet = workbook.create_sheet(sheet_name[:31])
        worksheet['A1'] = 'ASSESSMENT SAVINGS REPORT  (June 2026)'
        worksheet.append([])                                   # row 2 spacer
        worksheet.append(['Assessment ID', 'Reg No.', 'Vehicle', 'Repairer',
                          'Req Auth Date', 'REPAIRER LOWER QUOTATION (P)', None, None, None,
                          'ASSESSMENT REPORT (P)', None, None, None, 'SAVINGS', None, None, None])
        worksheet.append([None, None, None, None, None, 'Parts', 'Labour', 'Paint', 'Total',
                          'Parts', 'Labour', 'Paint', 'Total', 'Parts', 'Labour', 'Paint', 'TOTAL'])
        for row in rows:
            worksheet.append(row)
        for label, values in (stated or {}).get(sheet_name, []):
            worksheet.append([None] * 12 + [label] + values)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def parts_workbook() -> bytes:
    """Build a Parts Summary workbook shaped exactly like the real one."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    summary = workbook.create_sheet('SUMMARY')
    summary.append(['MONTHS ', 'JULY', 'AUGUST'])
    summary.append(['SUPPLIERS ', None, None])
    summary.append(['CFAO', 100000, 40000])           # rounded restatement
    summary.append(['ACE AUTO', 10000, 0])
    summary.append(['KIA MOTORS', 5000, 0])           # only on SUMMARY → a real gap
    summary.append(['TOTAL ', 115000, 40000])
    summary.append([])
    summary.append(['GLASS', None, None])
    summary.append(['PG GLASS', 2000, 3000])
    summary.append(['TOTAL ', 2000, 3000])
    summary.append([])
    summary.append(['CONTRACT PRICING', 50000, 20000])
    summary.append(['PARTS TOTAL', 117000, 43000])
    summary.append(['GRAND TOTAL', 167000, 63000])
    summary.append([])
    summary.append(['PARTS SUMMARY ', None, None])
    summary.append(['MONTHS ', 'FY-26', 'FY-27'])
    summary.append(['July', 800000, 117000])
    summary.append(['August', 900000, 43000])
    summary.append(['TOTAL', 1700000, 160000])

    dealership = workbook.create_sheet('DEALERSHIP')
    dealership.append(['MONTHS ', 'July', 'August'])
    dealership.append(['SUPPLIERS ', 2026, 2026])
    dealership.append(['CFAO', 100000.40, 40000.10])
    dealership.append(['GLASS', 0, None])             # spacer row — must be skipped
    dealership.append(['TOTAL ', 100000.40, 40000.10])

    aftermarket = workbook.create_sheet('AFTERMARKET')
    aftermarket.append(['MONTHS ', 'July', 'August'])
    aftermarket.append(['SUPPLIERS ', 2026, 2026])
    aftermarket.append(['ACE AUTO', 10000.25, 0])
    aftermarket.append(['TOTAL ', 10000.25, 0])

    windscreen = workbook.create_sheet('WINDSCREEN')
    windscreen.append(['GLASS', None, None])          # no month header of its own
    windscreen.append(['PG GLASS', 2000, 3000])
    windscreen.append(['TOTAL ', 2000, 3000])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


JULY_ROWS = [
    # id, reg, vehicle, repairer, date, quote p/l/pt/total, report p/l/pt/total, saving p/l/pt/total
    ['ALPHA-0000000001', 'TEST-1', 'SEDAN', 'Panel Beaters A', date(2026, 7, 2),
     10000, 2000, 1000, 13000,  8000, 1500, 800, 10300,  2000, 500, 200, 2700],
    # savings column dragged wrong: states 9000, actual quote − assessment is 6000
    ['ALPHA-0000000002', 'TEST-2', 'BAKKIE', 'Panel Beaters B', date(2026, 7, 9),
     20000, 3000, 2000, 25000,  15000, 2000, 2000, 19000,  5000, 4000, 0, 9000],
    # a job with no quote yet — must still be counted as a job, contributing zero
    ['ALPHA-0000000003', 'TEST-3', 'SUV', '', date(2026, 7, 15),
     None, None, None, 0,  None, None, None, 0,  0, 0, 0, 0],
]


class ParseAssessmentTests(APITestCase):
    def test_month_comes_from_auth_date_not_the_a1_title(self):
        data = assessment_workbook({'Assessment_Savings_Report_July': JULY_ROWS})
        sheet = parse_assessment_savings(data)['sheets'][0]
        # A1 says June 2026 on every sheet of the real workbook.
        self.assertEqual(sheet['period'], date(2026, 7, 1))

    def test_savings_are_recomputed_and_the_file_variance_is_kept(self):
        data = assessment_workbook({'Assessment_Savings_Report_July': JULY_ROWS})
        sheet = parse_assessment_savings(data)['sheets'][0]
        second = sheet['rows'][1]
        self.assertEqual(second['saving_total'], Decimal('6000.00'))    # 25000 − 19000
        self.assertEqual(second['file_saving_total'], Decimal('9000.00'))
        self.assertEqual(second['savings_variance'], Decimal('3000.00'))
        self.assertEqual(sheet['variance_rows'], 1)

    def test_sheet_total_block_is_read_and_ties_to_the_rows(self):
        rows = JULY_ROWS[:1]
        data = assessment_workbook(
            {'Assessment_Savings_Report_June': rows},
            stated={'Assessment_Savings_Report_June': [
                ('Parts Savings:',  [2000, None, None, None]),
                ('Labour Savings:', [None, 500, None, None]),
                ('Paint Savings:',  [None, None, 200, None]),
                ('TOTAL SAVINGS:',  [None, None, None, 2700]),
            ]},
        )
        sheet = parse_assessment_savings(data)['sheets'][0]
        self.assertEqual(sheet['stated']['total'], Decimal('2700.00'))
        self.assertEqual(sheet['computed']['total'], Decimal('2700.00'))
        self.assertEqual(len(sheet['rows']), 1)          # the totals block is not a row

    def test_a_job_with_no_quote_still_counts_as_a_job(self):
        data = assessment_workbook({'Assessment_Savings_Report_July': JULY_ROWS})
        sheet = parse_assessment_savings(data)['sheets'][0]
        self.assertEqual(len(sheet['rows']), 3)
        self.assertEqual(sheet['rows'][2]['saving_total'], Decimal('0.00'))

    def test_every_month_sheet_is_read(self):
        data = assessment_workbook({
            'Assessment_Savings_Report_July': JULY_ROWS,
            'Assessment_Savings_Report_Augus': [
                ['ALPHA-0000000009', 'TEST-9', 'TRUCK', 'Panel Beaters C', date(2026, 8, 4),
                 1000, 0, 0, 1000,  600, 0, 0, 600,  400, 0, 0, 400],
            ],
        })
        periods = [s['period'] for s in parse_assessment_savings(data)['sheets']]
        self.assertEqual(periods, [date(2026, 7, 1), date(2026, 8, 1)])


class ParsePartsTests(APITestCase):
    def setUp(self):
        self.parsed = parse_parts_summary(parts_workbook())

    def test_spend_comes_from_the_category_sheets_at_full_precision(self):
        july = {(r['category'], r['supplier']): r['amount']
                for r in self.parsed['spend'] if r['month'] == date(2026, 7, 1)}
        self.assertEqual(july[('DEALERSHIP', 'CFAO')], Decimal('100000.40'))
        self.assertEqual(july[('AFTERMARKET', 'ACE AUTO')], Decimal('10000.25'))
        self.assertEqual(july[('WINDSCREEN', 'PG GLASS')], Decimal('2000.00'))

    def test_total_rows_and_the_glass_spacer_are_never_counted(self):
        names = {r['supplier'].upper() for r in self.parsed['spend']}
        self.assertNotIn('TOTAL', names)
        self.assertNotIn('GLASS', names)
        self.assertNotIn('SUPPLIERS', names)

    def test_summary_is_not_used_as_a_row_source(self):
        # KIA MOTORS exists only on SUMMARY, so it must NOT appear as spend.
        self.assertNotIn('KIA MOTORS', {r['supplier'] for r in self.parsed['spend']})

    def test_the_summary_gap_is_reported_not_hidden(self):
        line = next(l for l in self.parsed['recon']
                    if l['month'] == '2026-07-01' and l['block'].startswith('Dealership'))
        self.assertEqual(line['stated'], Decimal('115000.00'))
        self.assertEqual(line['computed'], Decimal('110000.65'))
        self.assertEqual(line['difference'], Decimal('4999.35'))

    def test_windscreen_inherits_the_dealership_month_header(self):
        months = {r['month'] for r in self.parsed['spend'] if r['category'] == 'WINDSCREEN'}
        self.assertEqual(months, {date(2026, 7, 1), date(2026, 8, 1)})

    def test_contract_pricing_and_fy_history_are_read(self):
        self.assertEqual(
            {c['month']: c['amount'] for c in self.parsed['contract_pricing']},
            {date(2026, 7, 1): Decimal('50000.00'), date(2026, 8, 1): Decimal('20000.00')},
        )
        history = {(h['fy_label'], h['month_number']): h['amount'] for h in self.parsed['history']}
        self.assertEqual(history[('FY-27', 7)], Decimal('117000.00'))
        self.assertEqual(history[('FY-26', 8)], Decimal('900000.00'))

    def test_fiscal_year_runs_july_to_june(self):
        self.assertEqual(fiscal_year_label(date(2026, 7, 1)), 'FY-27')
        self.assertEqual(fiscal_year_label(date(2026, 6, 30)), 'FY-26')

    def test_workbook_kind_is_detected_from_the_sheets(self):
        self.assertEqual(detect_workbook_kind(parts_workbook()), 'PARTS')
        self.assertEqual(
            detect_workbook_kind(assessment_workbook({'Assessment_Savings_Report_July': JULY_ROWS})),
            'ASSESSMENT',
        )


class PartsApiTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser('parts-cfo', 'parts@example.com', 'x')
        self.client.force_authenticate(self.user)
        self.upload_url  = reverse('v1-salvage-parts-upload')
        self.summary_url = reverse('v1-salvage-parts-summary')

    def _post(self, name: str, data: bytes):
        return self.client.post(
            self.upload_url, {'file': _named(name, data)}, format='multipart',
        )

    def test_upload_then_summary_returns_the_figures(self):
        assessment = assessment_workbook({'Assessment_Savings_Report_July': JULY_ROWS})
        response = self._post('Assessment_Savings_Report July2026.xlsx', assessment)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['rows_created'], 3)

        response = self._post('Parts Summary July 2026.xlsx', parts_workbook())
        self.assertEqual(response.status_code, 201, response.data)

        summary = self.client.get(self.summary_url)
        self.assertEqual(summary.status_code, 200)
        july = next(t for t in summary.data['timeline'] if t['month'] == '2026-07-01')
        self.assertEqual(july['dealership'],       '100000.40')
        self.assertEqual(july['aftermarket'],      '10000.25')
        self.assertEqual(july['windscreen'],       '2000.00')
        self.assertEqual(july['parts_total'],      '112000.65')
        self.assertEqual(july['contract_pricing'], '50000.00')
        # 2700 + 6000 + 0 — recomputed, not the file's 2700 + 9000.
        self.assertEqual(july['assessment_saving'], '8700.00')
        self.assertEqual(july['fy'], 'FY-27')
        self.assertEqual(summary.data['totals']['jobs'], 3)
        self.assertEqual(len(summary.data['variance_rows']), 1)

    def test_uploading_the_same_month_twice_replaces_it(self):
        data = assessment_workbook({'Assessment_Savings_Report_July': JULY_ROWS})
        self._post('a.xlsx', data)
        self._post('a.xlsx', data)
        self.assertEqual(AssessmentSaving.objects.count(), 3)
        self.assertEqual(PartsUpload.objects.count(), 2)

        parts = parts_workbook()
        self._post('p.xlsx', parts)
        self._post('p.xlsx', parts)
        self.assertEqual(PartsSpend.objects.count(), 5)
        self.assertEqual(PartsContractPricing.objects.count(), 2)
        self.assertEqual(PartsHistory.objects.count(), 4)

    def test_an_unrecognised_workbook_is_rejected_cleanly(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = 'Cashbook'
        workbook.active['A1'] = 'nothing to see'
        buffer = io.BytesIO()
        workbook.save(buffer)
        response = self._post('random.xlsx', buffer.getvalue())
        self.assertEqual(response.status_code, 400)
        self.assertIn('Unrecognised workbook', response.data['detail'])

    def test_repairer_names_differing_only_by_case_are_one_entry(self):
        """The team types the same shop several ways — left ungrouped, the
        savings league shows it twice at half size."""
        data = assessment_workbook({'Assessment_Savings_Report_July': [
            ['ALPHA-0000000101', 'T1', 'SEDAN', 'Specialised Panel Beaters', date(2026, 7, 2),
             10000, 0, 0, 10000,  6000, 0, 0, 6000,  4000, 0, 0, 4000],
            ['ALPHA-0000000102', 'T2', 'SEDAN', 'SPECIALISED PANEL BEATERS', date(2026, 7, 3),
             20000, 0, 0, 20000,  15000, 0, 0, 15000,  5000, 0, 0, 5000],
            ['ALPHA-0000000103', 'T3', 'SEDAN', '  specialised   panel beaters ', date(2026, 7, 4),
             5000, 0, 0, 5000,  4000, 0, 0, 4000,  1000, 0, 0, 1000],
        ]})
        self._post('a.xlsx', data)
        summary = self.client.get(self.summary_url)
        league = summary.data['top_repairers']
        self.assertEqual(len(league), 1, league)
        self.assertEqual(league[0]['jobs'], 3)
        self.assertEqual(league[0]['quoted'], '35000.00')
        self.assertEqual(league[0]['saving'], '10000.00')
        self.assertEqual(league[0]['saving_pct'], '28.6')

    def test_supplier_names_differing_only_by_case_are_one_entry(self):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        dealership = workbook.create_sheet('DEALERSHIP')
        dealership.append(['MONTHS ', 'July'])
        dealership.append(['SUPPLIERS ', 2026])
        dealership.append(['CFAO', 100])
        dealership.append(['cfao ', 50])
        dealership.append(['TOTAL ', 150])
        buffer = io.BytesIO()
        workbook.save(buffer)
        self._post('p.xlsx', buffer.getvalue())
        summary = self.client.get(self.summary_url)
        suppliers = summary.data['top_suppliers']
        self.assertEqual(len(suppliers), 1, suppliers)
        self.assertEqual(suppliers[0]['supplier'], 'CFAO')     # the heavier spelling
        self.assertEqual(suppliers[0]['amount'], '150.00')

    def test_the_headline_month_skips_a_contract_pricing_only_stub(self):
        """The August workbook already carries a September contract-pricing
        figure with no purchases behind it — the page must not open on two
        zeros."""
        self._post('a.xlsx', assessment_workbook({'Assessment_Savings_Report_July': JULY_ROWS}))
        self._post('p.xlsx', parts_workbook())
        PartsContractPricing.objects.create(
            upload=PartsUpload.objects.filter(kind='PARTS').first(),
            month=date(2026, 9, 1), amount=Decimal('12141.25'),
        )
        summary = self.client.get(self.summary_url)
        self.assertEqual(summary.data['timeline'][-1]['month'], '2026-09-01')
        self.assertEqual(summary.data['latest_month']['month'], '2026-08-01')

    def test_upload_needs_a_file(self):
        response = self.client.post(self.upload_url, {}, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_summary_is_closed_to_anonymous_callers(self):
        self.client.force_authenticate(None)
        self.assertIn(self.client.get(self.summary_url).status_code, (401, 403))


def _named(name: str, data: bytes):
    """A file-like object DRF's multipart parser will accept, with a name."""
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile(
        name, data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


class RegisterWriteTests(APITestCase):
    """The register half (CFO 2026-09-10): Veritas staff type into this, Bharath
    manages it, and an uploaded workbook must never wipe what they typed."""

    def setUp(self):
        self.user = User.objects.create_superuser('register-cfo', 'reg@example.com', 'x')
        self.client.force_authenticate(self.user)
        self.savings_url  = reverse('salvage-parts-saving-list')
        self.spend_url    = reverse('salvage-parts-spend-list')
        self.contract_url = reverse('salvage-parts-contract-list')
        self.summary_url  = reverse('v1-salvage-parts-summary')
        self.upload_url   = reverse('v1-salvage-parts-upload')

    def _add_saving(self, **over):
        body = {
            'assessment_id': 'REG-0001', 'reg_no': 'T-1', 'vehicle': 'SEDAN',
            'repairer': 'Panel Beaters A', 'req_auth_date': '2026-07-08',
            'quote_parts': '10000', 'quote_labour': '2000', 'quote_paint': '1000',
            'report_parts': '8000', 'report_labour': '1500', 'report_paint': '800',
        }
        body.update(over)
        return self.client.post(self.savings_url, body, format='json')

    def test_a_typed_saving_derives_its_own_totals(self):
        response = self._add_saving()
        self.assertEqual(response.status_code, 201, response.data)
        row = AssessmentSaving.objects.get(assessment_id='REG-0001')
        self.assertEqual(row.quote_total,   Decimal('13000.00'))
        self.assertEqual(row.report_total,  Decimal('10300.00'))
        self.assertEqual(row.saving_total,  Decimal('2700.00'))
        self.assertEqual(row.savings_variance, Decimal('0.00'))
        self.assertEqual(row.period, date(2026, 7, 1))
        self.assertEqual(row.source, 'MANUAL')
        self.assertEqual(row.entered_by, self.user)

    def test_a_typed_total_cannot_override_the_derived_one(self):
        response = self._add_saving(saving_total='999999', quote_total='1')
        self.assertEqual(response.status_code, 201, response.data)
        row = AssessmentSaving.objects.get(assessment_id='REG-0001')
        self.assertEqual(row.saving_total, Decimal('2700.00'))
        self.assertEqual(row.quote_total,  Decimal('13000.00'))

    def test_editing_a_line_recomputes_the_saving(self):
        row_id = self._add_saving().data['id']
        response = self.client.patch(f'{self.savings_url}{row_id}/',
                                     {'report_parts': '6000'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        row = AssessmentSaving.objects.get(id=row_id)
        self.assertEqual(row.saving_total, Decimal('4700.00'))

    def test_a_line_can_be_deleted(self):
        row_id = self._add_saving().data['id']
        self.assertEqual(self.client.delete(f'{self.savings_url}{row_id}/').status_code, 204)
        self.assertFalse(AssessmentSaving.objects.filter(id=row_id).exists())

    def test_a_new_supplier_and_a_new_month_can_be_added(self):
        response = self.client.post(self.spend_url, {
            'supplier': '  Kagiso   Spares ', 'category': 'AFTERMARKET',
            'month': '2026-07-19', 'amount': '4500.50',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        row = PartsSpend.objects.get(supplier='Kagiso Spares')   # whitespace tidied
        self.assertEqual(row.month, date(2026, 7, 1))            # snapped to the month
        self.assertEqual(row.amount, Decimal('4500.50'))
        self.assertEqual(row.source, 'MANUAL')

        suppliers = self.client.get(f'{self.spend_url}suppliers/').data['suppliers']
        self.assertIn('Kagiso Spares', suppliers)

    def test_contract_pricing_can_be_typed_for_a_month(self):
        response = self.client.post(self.contract_url,
                                    {'month': '2026-09-01', 'amount': '12141.25'},
                                    format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(PartsContractPricing.objects.get().amount, Decimal('12141.25'))

    def test_typed_lines_survive_a_workbook_upload_of_the_same_month(self):
        """The whole point of the source marker."""
        self._add_saving()
        self.client.post(self.spend_url, {
            'supplier': 'Kagiso Spares', 'category': 'AFTERMARKET',
            'month': '2026-07-01', 'amount': '4500',
        }, format='json')

        self.client.post(self.upload_url,
                         {'file': _named('a.xlsx', assessment_workbook(
                             {'Assessment_Savings_Report_July': JULY_ROWS}))},
                         format='multipart')
        self.client.post(self.upload_url,
                         {'file': _named('p.xlsx', parts_workbook())},
                         format='multipart')

        self.assertTrue(AssessmentSaving.objects.filter(assessment_id='REG-0001').exists())
        self.assertTrue(PartsSpend.objects.filter(supplier='Kagiso Spares').exists())
        # and the uploaded rows landed alongside, marked as coming from a file
        self.assertEqual(AssessmentSaving.objects.filter(source='UPLOAD').count(), 3)
        self.assertEqual(
            AssessmentSaving.objects.filter(source='MANUAL').count(), 1)

    def test_a_typed_line_shows_up_in_the_dashboard_totals(self):
        self._add_saving()
        summary = self.client.get(self.summary_url)
        july = next(t for t in summary.data['timeline'] if t['month'] == '2026-07-01')
        self.assertEqual(july['assessment_saving'], '2700.00')
        self.assertEqual(summary.data['totals']['jobs'], 1)

    def test_a_line_needs_a_date_or_a_month(self):
        response = self._add_saving(req_auth_date='', period=None)
        self.assertEqual(response.status_code, 400)

    def test_an_outsider_cannot_write_to_the_register(self):
        outsider = User.objects.create_user('outsider', 'out@example.com', 'x')
        self.client.force_authenticate(outsider)
        self.assertEqual(self._add_saving().status_code, 403)
        self.client.force_authenticate(None)
        self.assertIn(self._add_saving().status_code, (401, 403))


class RegisterAccessTests(APITestCase):
    """Bharath manages the module, so his rights must not depend on which
    company row his employee record happens to carry (CFO 2026-09-10)."""

    def test_the_named_module_manager_is_allowed_in(self):
        from salvage.permissions import user_manages_parts_register
        bharath = User.objects.create_user(
            'bbalasubramanian@alphadirect.co.bw', 'bbalasubramanian@alphadirect.co.bw', 'x')
        self.assertTrue(user_manages_parts_register(bharath))

        by_email = User.objects.create_user(
            'bharath.b', 'BBalasubramanian@alphadirect.co.bw', 'x')
        self.assertTrue(user_manages_parts_register(by_email))

        someone_else = User.objects.create_user('nobody', 'nobody@alphadirect.co.bw', 'x')
        self.assertFalse(user_manages_parts_register(someone_else))

    def test_the_module_manager_can_write_without_a_salvage_company(self):
        bharath = User.objects.create_user(
            'bbalasubramanian@alphadirect.co.bw', 'bbalasubramanian@alphadirect.co.bw', 'x')
        self.client.force_authenticate(bharath)
        response = self.client.post(reverse('salvage-parts-spend-list'), {
            'supplier': 'CFAO', 'category': 'DEALERSHIP',
            'month': '2026-08-01', 'amount': '100',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)


class RegisterPagingTests(APITestCase):
    """The register must never show fewer rows than it has without saying so.

    Found by /qctest on 10-Sep-2026: the screen asked for 200 rows, the server
    caps a page at 100, and the tab badge counted the rows on screen — so 20 of
    120 savings lines were invisible with nothing to hint at them.
    """

    def setUp(self):
        self.user = User.objects.create_superuser('paging-cfo', 'p@example.com', 'x')
        self.client.force_authenticate(self.user)
        self.url = reverse('salvage-parts-saving-list')
        AssessmentSaving.objects.bulk_create([
            AssessmentSaving(
                period=date(2026, 7, 1) if i % 2 else date(2026, 8, 1),
                assessment_id=f'PAGE-{i:04d}', source='UPLOAD',
                quote_parts=Decimal('100'), report_parts=Decimal('60'),
                quote_total=Decimal('100'), report_total=Decimal('60'),
                saving_parts=Decimal('40'), saving_total=Decimal('40'),
            ) for i in range(120)
        ])

    def test_the_count_is_the_whole_register_not_the_page(self):
        first = self.client.get(f'{self.url}?page_size=100')
        self.assertEqual(first.data['count'], 120)
        self.assertEqual(len(first.data['results']), 100)
        self.assertIsNotNone(first.data['next'], 'a short page must advertise the next one')

    def test_asking_for_more_than_the_cap_does_not_hide_the_rest(self):
        capped = self.client.get(f'{self.url}?page_size=200')
        self.assertEqual(capped.data['count'], 120)
        self.assertIsNotNone(capped.data['next'])

    def test_the_second_page_holds_the_remainder_with_no_overlap(self):
        first  = self.client.get(f'{self.url}?page_size=100&page=1')
        second = self.client.get(f'{self.url}?page_size=100&page=2')
        ids = {r['id'] for r in first.data['results']} | {r['id'] for r in second.data['results']}
        self.assertEqual(len(ids), 120, 'every line appears exactly once across the pages')
        self.assertEqual(len(second.data['results']), 20)

    def test_a_month_filter_narrows_the_count_too(self):
        july = self.client.get(f'{self.url}?month=2026-07-01&page_size=100')
        self.assertEqual(july.data['count'], 60)
        self.assertTrue(all(r['period'] == '2026-07-01' for r in july.data['results']))
