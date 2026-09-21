"""
ifrs17/export.py — the disclosures out the door.

Auditors receive documents, not screens. Two exports:
  · workbook_bytes()  — one sheet per disclosure table (xlsx).
  · disclosure_docx() — the disclosure notes as they go into the financial
                        statements, house template.

Both read exactly what the /ifrs17 screen reads (disclosures.build_all over an
engine.compute), so a lever the CFO drags moves the document too. Nothing here
posts to the ledger.
"""
from __future__ import annotations

import io
from decimal import Decimal as D

from .disclosures import build_all
from .statistics import build_all as build_stats
from .engine import Computed

NAVY = '0D1B2A'
ORANGE = 'F4A623'


def _docx_safe(text: str) -> str:
    """python-docx serialises to XML, which rejects control chars and chokes on
    some emoji. Drop the flag markers and anything below the printable range —
    the words carry the meaning, the icon is decoration."""
    text = str(text or '').replace('\U0001f534', '').replace('🔴', '')
    return ''.join(ch for ch in text if ch >= ' ' or ch in '\n\t')


def _num(v):
    if isinstance(v, (int, float, D)):
        v = D(v)
        return f'({abs(v):,.0f})' if v < 0 else f'{v:,.0f}'
    return '' if v is None else str(v)


def workbook_bytes(c: Computed) -> bytes:
    """One sheet per disclosure table."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    tables = {**build_all(c), **build_stats(c)}
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill('solid', fgColor=NAVY)
    header_font = Font(color='FFFFFF', bold=True)
    title_font = Font(color=NAVY, bold=True, size=13)

    # A cover sheet so the pack is self-describing.
    cover = wb.create_sheet('Cover')
    cover['A1'] = 'Alpha Direct Insurance Company'
    cover['A1'].font = title_font
    cover['A2'] = f'IFRS 17 disclosures — {c.year}'
    cover['A3'] = 'Premium Allocation Approach · Empirica Actuaries · 30 June 2026'
    cover['A5'] = ('Figures at the current cockpit lever positions. At the signed '
                   'basis these reproduce the valuation report verbatim.')
    if c.notes:
        cover['A7'] = 'Notes:'
        for i, n in enumerate(c.notes):
            cover.cell(row=8 + i, column=1, value=f'· {n}')
    cover.column_dimensions['A'].width = 110

    used = {'Cover'}
    for key, t in tables.items():
        # Excel sheet names: <=31 chars, unique, no []:*?/\
        name = ''.join(ch for ch in t['title'][:28] if ch not in r'[]:*?/\\') or key[:28]
        n, base = name, name
        i = 2
        while n in used:
            n = f'{base[:26]} {i}'
            i += 1
        used.add(n)
        ws = wb.create_sheet(n)

        ws['A1'] = t['title']
        ws['A1'].font = title_font
        if t.get('ref'):
            ws['A2'] = t['ref']
            ws['A2'].font = Font(color='888888', italic=True)

        r0 = 4
        for ci, col in enumerate(t['columns'], start=1):
            cell = ws.cell(row=r0, column=ci, value=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='left' if ci == 1 else 'right')
        for ri, row in enumerate(t['rows'], start=r0 + 1):
            for ci, val in enumerate(row, start=1):
                cell = ws.cell(row=ri, column=ci,
                               value=float(val) if isinstance(val, (int, float, D))
                               and ci > 1 else _num(val) if ci == 1 else
                               (float(val) if isinstance(val, (int, float, D)) else val))
                if ci > 1 and isinstance(val, (int, float, D)):
                    cell.number_format = '#,##0;(#,##0)'
        if t.get('note'):
            note_row = r0 + 1 + len(t['rows']) + 1
            ws.cell(row=note_row, column=1, value=t['note'].replace('🔴 ', ''))
            ws.cell(row=note_row, column=1).font = Font(italic=True, color='B04E00')
        ws.column_dimensions['A'].width = 52
        for ci in range(2, len(t['columns']) + 1):
            ws.column_dimensions[chr(64 + ci)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def disclosure_docx(c: Computed) -> bytes:
    """The disclosure notes, in the house style, for the financial statements."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    tables = build_all(c)
    doc = Document()

    h = doc.add_heading('Alpha Direct Insurance Company', level=0)
    doc.add_heading(f'IFRS 17 disclosures — {c.year}', level=1)
    p = doc.add_paragraph(
        'Premium Allocation Approach. Prepared by Empirica Actuaries, valuation '
        'date 30 June 2026. The tables below are reproduced from the valuation '
        'model and present the disclosures required by IFRS 17 paragraphs 100 to '
        '105.')
    p.runs[0].font.size = Pt(10)

    for key, t in tables.items():
        doc.add_heading(t['title'] + (f'  ({t["ref"]})' if t.get('ref') else ''), level=2)
        tbl = doc.add_table(rows=1, cols=len(t['columns']))
        tbl.style = 'Light Grid Accent 1'
        for ci, col in enumerate(t['columns']):
            cell = tbl.rows[0].cells[ci]
            cell.text = str(col)
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
        for row in t['rows']:
            cells = tbl.add_row().cells
            for ci, val in enumerate(row):
                cells[ci].text = _docx_safe(_num(val) if ci > 0
                                            else (val if val is not None else ''))
        if t.get('note'):
            note = doc.add_paragraph(_docx_safe(t['note']))
            note.runs[0].italic = True
            note.runs[0].font.size = Pt(9)
            note.runs[0].font.color.rgb = RGBColor(0xB0, 0x4E, 0x00)

    if c.notes:
        doc.add_heading('Notes on the current basis', level=2)
        for n in c.notes:
            doc.add_paragraph(_docx_safe(n), style='List Bullet')

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
