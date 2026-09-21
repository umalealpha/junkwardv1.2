"""genric/exports.py — the pack as Excel, and the invoice as PDF.

Excel: one workbook, one sheet per report, same Alpha Navy / Direct Orange
house look as the Aware reports — the colours and the single-report builder are
imported from ``aware.reporting`` rather than copied, so a brand change happens
once.

The cession sheet carries LIVE FORMULAS, not pasted values. The build prompt
asks for "Excel export with live formulas, so every number can be traced", and
it is also the honest thing to ship: Finance can change the confirmed GWP cell
and watch the invoice total move, which is how they will check our arithmetic
against theirs.

PDF: the GENRIC invoice, reportlab, following healthcare/quote_pdf.py.
"""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from aware.reporting import GREY, NAVY, ORANGE, WHITE, _wb  # brand + single-report builder

from . import constants as K
from . import reports as R
from .money import money, q2


# ── Excel ───────────────────────────────────────────────────────────────────
def _safe_sheet_title(title: str) -> str:
    """Excel sheet names: 31 chars, and none of : \\ / ? * [ ]"""
    out = ''.join(ch for ch in title if ch not in ':\\/?*[]')
    return out[:31] or 'Sheet'


def _write_report(ws, report: R.Report) -> None:
    ncol = max(len(report.columns), 4)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, report.title + (f'  ({report.nmi})' if report.nmi else ''))
    c.font = Font(bold=True, size=15, color=NAVY)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c = ws.cell(2, 1, R.STATUS_LABELS[report.status])
    c.font = Font(size=10, bold=True,
                  color=('B00020' if report.status in (R.BLOCKED, R.NIL_NO_SOURCE)
                         else '555555'))

    r = 4
    for label, value in report.kpis:
        ws.cell(r, 1, label).font = Font(bold=True, color='555555')
        ws.cell(r, 2, value).font = Font(bold=True, color=NAVY)
        r += 1
    r += 1

    if report.columns:
        for j, name in enumerate(report.columns, 1):
            cell = ws.cell(r, j, name)
            cell.font = Font(bold=True, color=WHITE)
            cell.fill = PatternFill('solid', fgColor=NAVY)
            cell.alignment = Alignment(horizontal='center', vertical='center',
                                       wrap_text=True)
        r += 1
        for i, row in enumerate(report.rows):
            for j, val in enumerate(row, 1):
                cell = ws.cell(r, j, float(val) if isinstance(val, Decimal) else val)
                if isinstance(val, Decimal):
                    cell.number_format = '#,##0.00'
                if i % 2:
                    cell.fill = PatternFill('solid', fgColor=GREY)
            r += 1

    r += 1
    if report.reconciled_against:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        cell = ws.cell(r, 1, 'Reconciled against: ' + report.reconciled_against)
        cell.font = Font(size=9, bold=True, color=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    for n in report.notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        cell = ws.cell(r, 1, f'• {n}')
        cell.font = Font(size=9, italic=True, color='666666')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1

    for j in range(1, ncol + 1):
        ws.column_dimensions[get_column_letter(j)].width = 34 if j == 1 else 20


def _write_cession(ws, ctx) -> None:
    """The six steps, as LIVE FORMULAS off one input cell."""
    ces = ctx.cession
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    c = ws.cell(1, 1, 'Cession & GENRIC Invoice')
    c.font = Font(bold=True, size=15, color=NAVY)
    ws.cell(2, 1, f'{K.ENTITY_NAME} ({K.ENTITY_BOOK}) — {ctx.period_label}').font = \
        Font(size=10, color='555555')

    # Inputs block — the cells the formulas point at.
    ws.cell(4, 1, 'INPUTS').font = Font(bold=True, color=WHITE)
    ws.cell(4, 1).fill = PatternFill('solid', fgColor=ORANGE)
    ws.cell(5, 1, 'Confirmed GWP incl VAT (from the bank)')
    ws.cell(5, 2, float(ces.confirmed_gwp_incl_vat)).number_format = '#,##0.00'
    ws.cell(6, 1, 'SA VAT rate')
    ws.cell(6, 2, float(K.SA_VAT_RATE)).number_format = '0.0%'
    ws.cell(7, 1, 'Collection charges')
    ws.cell(7, 2, float(ces.collection_charges)).number_format = '#,##0.00'
    ws.cell(8, 1, 'Quota share ceded')
    ws.cell(8, 2, float(K.QUOTA_SHARE_CEDED)).number_format = '0.0%'
    ws.cell(9, 1, 'Ceding commission (signed treaty — the older MOU said 10%)')
    ws.cell(9, 2, float(K.CEDING_COMMISSION_RATE)).number_format = '0.0%'

    hdr = 11
    for j, name in enumerate(['#', 'Step', f'Amount ({K.CURRENCY})', 'Basis'], 1):
        cell = ws.cell(hdr, j, name)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=NAVY)

    # ROUND(...,2) in Excel is half-up for positive numbers, which is the same
    # decision as Decimal ROUND_HALF_UP on the Python side. The two must agree
    # or Finance's spreadsheet and Omni will differ by a cent.
    formulas = [
        (1, 'GWP excl VAT', '=ROUND(B5/(1+B6),2)'),
        (2, 'Less collection charges', '=-B7'),
        (3, 'Net premium base', '=ROUND(C12-B7,2)'),
        (4, f'Ceded at {K.QUOTA_SHARE_CEDED:.0%}', '=ROUND(C14*B8,2)'),
        (5, f'Less ceding commission {K.CEDING_COMMISSION_RATE:.1%}', '=-ROUND(C15*B9,2)'),
        (6, 'NET REINSURANCE PREMIUM DUE', '=ROUND(C15+C16,2)'),
    ]
    r = hdr + 1
    for (num, label, formula), step in zip(formulas, ces.steps):
        ws.cell(r, 1, num)
        ws.cell(r, 2, label)
        cell = ws.cell(r, 3, formula)
        cell.number_format = '#,##0.00'
        if num == 6:
            cell.font = Font(bold=True, color=NAVY, size=12)
            ws.cell(r, 2).font = Font(bold=True, color=NAVY, size=12)
        ws.cell(r, 4, step.basis).font = Font(size=9, italic=True, color='666666')
        r += 1

    r += 1
    ws.cell(r, 1, 'Invoice number').font = Font(bold=True)
    ws.cell(r, 2, ctx.invoice_number or '(unconfirmed)')
    r += 1
    ws.cell(r, 1, 'Invoice VAT').font = Font(bold=True)
    ws.cell(r, 2, 0.00).number_format = '#,##0.00'
    ws.cell(r, 3, 'Zero-rated — cross-border reinsurance').font = \
        Font(size=9, italic=True, color='666666')
    r += 2
    ws.cell(r, 1, 'Python and this sheet must agree to the cent. Omni computed '
                  f'{money(ces.net_reinsurance_premium_due, K.CURRENCY)}.').font = \
        Font(size=9, italic=True, color='666666')

    ws.column_dimensions['A'].width = 46
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 60


def pack_workbook(result) -> io.BytesIO:
    """The whole pack: Master first, then each report, then the cession."""
    wb = Workbook()
    wb.remove(wb.active)

    ordered = ([r for r in result.reports if r.key == 'master']
               + [r for r in result.reports if r.key != 'master'])
    for rep in ordered:
        ws = wb.create_sheet(_safe_sheet_title(rep.title))
        _write_report(ws, rep)

    _write_cession(wb.create_sheet('Cession & Invoice'), result.context)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def report_workbook(report: R.Report, ctx) -> io.BytesIO:
    """One report on its own — straight through aware.reporting._wb."""
    return _wb(
        title=report.title + (f' ({report.nmi})' if report.nmi else ''),
        subtitle=(f'{K.ENTITY_NAME} ({K.ENTITY_BOOK}) — {ctx.period_label} · '
                  f'{R.STATUS_LABELS[report.status]}'),
        kpis=list(report.kpis),
        columns=list(report.columns),
        rows=[[float(v) if isinstance(v, Decimal) else v for v in row]
              for row in report.rows],
        notes=([report.reconciled_against] if report.reconciled_against else [])
              + list(report.notes),
    )


# ── Who pays whom, and into what account ────────────────────────────────────
def remit_to_block() -> tuple:
    """(direction, label, value) for the invoice's payment block.

    CFO ruling, 13 Sep 2026: Alpha Direct PAYS GENRIC. This is a quota-share
    cession and the cedant remits the net reinsurance premium to the reinsurer.
    The first build printed "PAYABLE TO" over Alpha Direct's own FNB account
    62403392335 — i.e. it read as GENRIC paying us. That is inverted.

    GENRIC's account is never guessed. With the setting unset the invoice says
    so, loudly, and names the setting; it does not fall back to any account.
    """
    from .config import (
        GenricConfigurationError, REMIT_TO_SETTING, reinsurer_remit_to,
    )
    direction = f'PAYMENT DIRECTION: {K.PAYMENT_DIRECTION}'
    try:
        return (direction,
                f'{K.ENTITY_NAME} PAYS GENRIC — REMIT TO',
                reinsurer_remit_to())
    except GenricConfigurationError:
        return (direction,
                'REMIT TO — 🔴 NOT ON FILE, DO NOT PAY AGAINST THIS DOCUMENT',
                f"GENRIC's bank details have not been supplied, so this invoice "
                f"cannot say where to pay. Nothing was assumed and no account "
                f"is printed. Set {REMIT_TO_SETTING} once GENRIC has given the "
                f"bank, account number, branch and SWIFT.")


# ── PDF invoice ─────────────────────────────────────────────────────────────
def invoice_pdf(result) -> bytes:
    """The GENRIC reinsurance premium invoice. VAT R0.00, and it says why."""
    from reportlab.lib import colors
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    PDF_NAVY = HexColor('#0D1B2A')
    PDF_ORANGE = HexColor('#F4A623')
    PDF_GREY = HexColor('#6B7280')
    PDF_LIGHT = HexColor('#F3F4F6')

    ctx = result.context
    ces = ctx.cession

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
        topMargin=14 * mm, bottomMargin=16 * mm,
        title=f'{ctx.invoice_number} GENRIC {ctx.period_label}',
    )
    ss = getSampleStyleSheet()
    h_title = ParagraphStyle('t', parent=ss['Title'], textColor=PDF_ORANGE,
                             fontSize=18, spaceAfter=2, alignment=0)
    lbl = ParagraphStyle('l', parent=ss['Normal'], textColor=PDF_GREY, fontSize=8)
    val = ParagraphStyle('v', parent=ss['Normal'], textColor=PDF_NAVY, fontSize=9)
    foot = ParagraphStyle('f', parent=ss['Normal'], textColor=PDF_GREY,
                          fontSize=7, leading=9)

    el = [
        Paragraph('Reinsurance Premium Invoice', h_title),
        Paragraph(f'{K.ENTITY_NAME} — {K.ENTITY_BOOK}', val),
        Spacer(1, 6),
        HRFlowable(width='100%', thickness=2, color=PDF_ORANGE),
        Spacer(1, 8),
    ]

    head = Table([
        [Paragraph('INVOICE NUMBER', lbl), Paragraph('PERIOD', lbl),
         Paragraph('VAT', lbl)],
        [Paragraph(ctx.invoice_number or '(unconfirmed)', val),
         Paragraph(ctx.period_label, val),
         Paragraph(f'{K.CURRENCY}0.00 — zero-rated (cross-border)', val)],
    ], colWidths=[60 * mm, 55 * mm, 63 * mm])
    head.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), PDF_LIGHT),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
    ]))
    el += [head, Spacer(1, 10)]

    body = [['#', 'Step', f'Amount ({K.CURRENCY})']]
    for s in ces.steps:
        body.append([str(s.number), s.label, f'{q2(s.amount):,.2f}'])
    t = Table(body, colWidths=[10 * mm, 108 * mm, 60 * mm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), PDF_NAVY),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, PDF_LIGHT]),
        ('BACKGROUND', (0, -1), (-1, -1), PDF_LIGHT),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('TEXTCOLOR', (0, -1), (-1, -1), PDF_NAVY),
        ('GRID', (0, 0), (-1, -1), 0.4, HexColor('#D1D5DB')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    el += [t, Spacer(1, 10)]

    direction_label, remit_label, remit_value = remit_to_block()
    el += [
        Paragraph(direction_label, val),
        Spacer(1, 6),
        Paragraph(remit_label, lbl),
        Paragraph(remit_value, val),
        Spacer(1, 8),
        HRFlowable(width='100%', thickness=0.6, color=PDF_GREY),
        Spacer(1, 6),
        Paragraph(
            f'Ceding commission is {K.CEDING_COMMISSION_RATE:.1%} per the signed '
            f'treaty. The older MOU said 10% — the signed treaty wins.', foot),
        Paragraph(
            f'Confirmed GWP is bank-anchored: {ctx.summary.net_collection_count} '
            f'collection(s) into FNB {K.FNB_COLLECTION_ACCOUNT} for '
            f'{ctx.period_label}, net of {ctx.summary.reversal_count} reversal(s).',
            foot),
        Paragraph(
            'VAT is zero because cross-border reinsurance is zero-rated, not '
            'because VAT was omitted.', foot),
    ]
    if ctx.invoice_needs_confirmation:
        el.append(Paragraph(
            'DRAFT — the invoice number was derived from the previous pack and '
            'has not been confirmed by Finance.', foot))

    doc.build(el)
    return buf.getvalue()
