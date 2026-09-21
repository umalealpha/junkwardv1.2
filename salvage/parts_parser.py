"""salvage/parts_parser.py — readers for the two monthly Parts & Assessments
workbooks that Lemogang Machola (Coordinator, Parts & Assessments) sends to
Bharath, and that Bharath forwards to the CFO + EXCO.

CFO directive 2026-09-10 ("a new feature under Veritas"): the numbers in those
two spreadsheets must live in Omni, under the Veritas surface, instead of in a
mailbox.

Workbook 1 — "Assessment_Savings_Report <month>.xlsx"
    One sheet per month. Header on rows 3-4 (two-level: block / sub-column),
    data from row 5, and an optional stated-totals block in column M.
    Columns: Assessment ID, Reg No., Vehicle, Repairer, Req Auth Date,
             repairer-quote  Parts/Labour/Paint/Total,
             assessment-report Parts/Labour/Paint/Total,
             savings          Parts/Labour/Paint/Total.

Workbook 2 — "Parts Summary <months>.xlsx"
    SUMMARY (derived + rounded), then DEALERSHIP / WINDSCREEN / AFTERMARKET
    with one row per supplier and one column per month. SUMMARY also carries
    two things that exist nowhere else: the monthly CONTRACT PRICING figure
    and a five-financial-year month-by-month history.

Traps this module handles deliberately (do not "simplify" them away):
  * The A1 title on every sheet of the June workbook reads "(June 2026)" —
    including the July and August sheets. The month is therefore taken from
    the Req Auth Date values, with the sheet name as fallback. Never A1.
  * The savings columns in the file are hand-dragged formulas and do not
    always equal quote − assessment (e.g. ALPHA-0000002763 in July states
    30,983.28 against a computed 25,873.28). Both figures are returned so the
    caller can report the variance instead of silently picking one.
  * SUMMARY's supplier block is DEALERSHIP + AFTERMARKET rounded to the pula,
    and its July total (880,121) is 14,058 above the two category sheets
    (866,063.11). SUMMARY is never used as the row-level source — only its
    stated totals are kept, for reconciliation.
  * TOTAL rows, the 'GLASS' spacer row inside DEALERSHIP, and blank rows are
    skipped so nothing is double-counted.
  * SUMMARY's month header runs JULY..MAY (June is missing); the category
    sheets run Jul..Jun. Each sheet's months are read from its own header,
    and WINDSCREEN (which has no month header of its own) inherits the
    DEALERSHIP header.
"""
from __future__ import annotations

import calendar
import io
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import openpyxl
from django.utils.timezone import localdate

# Fiscal year runs July → June. FY-27 therefore starts 1 July 2026.
FY_START_MONTH = 7

MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
MONTHS.update({m.lower(): i for i, m in enumerate(calendar.month_abbr) if m})

CATEGORY_SHEETS = {
    'DEALERSHIP':  'DEALERSHIP',
    'AFTERMARKET': 'AFTERMARKET',
    'WINDSCREEN':  'WINDSCREEN',
}

TWOPLACES = Decimal('0.01')


def money(value) -> Decimal:
    """Coerce a cell to a 2dp Decimal. Blank / text / '-' become 0.00.

    VAT and every other rounding in this company is HALF UP (CFO, 9-Aug-2026).
    """
    if value is None or value == '':
        return Decimal('0.00')
    if isinstance(value, Decimal):
        raw = value
    elif isinstance(value, (int, float)):
        raw = Decimal(str(value))
    else:
        cleaned = re.sub(r'[^0-9.\-]', '', str(value))
        if cleaned in ('', '-', '.', '-.'):
            return Decimal('0.00')
        try:
            raw = Decimal(cleaned)
        except InvalidOperation:
            return Decimal('0.00')
    return raw.quantize(TWOPLACES, rounding=ROUND_HALF_UP)


def fiscal_year_label(d: date) -> str:
    """1 Jul 2026 → 'FY-27'; 1 Jun 2026 → 'FY-26'."""
    year = d.year + 1 if d.month >= FY_START_MONTH else d.year
    return f'FY-{year % 100:02d}'


def _month_start(value) -> date | None:
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    return None


def _month_from_name(text: str, default_year: int) -> date | None:
    """'July' / 'Jul 2026' / 'Assessment_Savings_Report_Augus' → month start."""
    lowered = (text or '').lower()
    year = default_year
    year_hit = re.search(r'(20\d{2})', lowered)
    if year_hit:
        year = int(year_hit.group(1))
    for name, number in sorted(MONTHS.items(), key=lambda kv: -len(kv[0])):
        if name in lowered:
            return date(year, number, 1)
    # Truncated sheet names ('Augus', 'Septemb') — match on prefix.
    for name, number in sorted(MONTHS.items(), key=lambda kv: -len(kv[0])):
        token = re.search(r'[a-z]{3,}$', lowered.strip())
        if token and name.startswith(token.group(0)):
            return date(year, number, 1)
    return None


# ---------------------------------------------------------------------------
# Workbook 1 — assessment savings
# ---------------------------------------------------------------------------

ASSESSMENT_HEADER_ROW = 4          # sub-column row; data starts on the next row
STATED_LABEL_COL = 13              # column M carries 'Parts Savings:' etc.
STATED_LABELS = {
    'parts savings':  'parts',
    'labour savings': 'labour',
    'paint savings':  'paint',
    'total savings':  'total',
}


def parse_assessment_savings(data: bytes, default_year: int | None = None) -> dict:
    """Read every month sheet of an Assessment Savings workbook.

    Returns
        {'sheets': [{'sheet','period','rows','stated','computed','variance_rows'}, ...]}

    Each row carries the file's own savings figures *and* the recomputed
    quote − assessment figures. `variance_rows` counts the rows where those
    two disagree by more than one thebe.
    """
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets: list[dict] = []

    for worksheet in workbook.worksheets:
        grid = [list(r) for r in worksheet.iter_rows(min_row=1, max_col=17, values_only=True)]
        rows: list[dict] = []
        stated: dict[str, Decimal] = {}
        months: Counter = Counter()

        for raw in grid[ASSESSMENT_HEADER_ROW:]:
            raw = list(raw) + [None] * (17 - len(raw))
            label = str(raw[STATED_LABEL_COL - 1] or '').strip().lower().rstrip(':')
            if label in STATED_LABELS:
                value = next((c for c in raw[STATED_LABEL_COL:] if c not in (None, '')), None)
                stated[STATED_LABELS[label]] = money(value)
                continue

            assessment_id = str(raw[0] or '').strip()
            if not assessment_id or assessment_id.lower().startswith('total'):
                continue

            auth_date = _month_start(raw[4])
            if auth_date:
                months[auth_date] += 1

            quote = {k: money(raw[i]) for k, i in (('parts', 5), ('labour', 6), ('paint', 7))}
            report = {k: money(raw[i]) for k, i in (('parts', 9), ('labour', 10), ('paint', 11))}
            file_saving = {k: money(raw[i]) for k, i in (('parts', 13), ('labour', 14), ('paint', 15))}

            row = {
                'assessment_id': assessment_id,
                'reg_no':        str(raw[1] or '').strip(),
                'vehicle':       str(raw[2] or '').strip(),
                'repairer':      str(raw[3] or '').strip(),
                'req_auth_date': raw[4].date() if isinstance(raw[4], datetime) else raw[4],
                'quote_parts':   quote['parts'],
                'quote_labour':  quote['labour'],
                'quote_paint':   quote['paint'],
                'quote_total':   money(raw[8]) or (quote['parts'] + quote['labour'] + quote['paint']),
                'report_parts':  report['parts'],
                'report_labour': report['labour'],
                'report_paint':  report['paint'],
                'report_total':  money(raw[12]) or (report['parts'] + report['labour'] + report['paint']),
                'file_saving_parts':  file_saving['parts'],
                'file_saving_labour': file_saving['labour'],
                'file_saving_paint':  file_saving['paint'],
                'file_saving_total':  money(raw[16]),
            }
            for part in ('parts', 'labour', 'paint'):
                row[f'saving_{part}'] = quote[part] - report[part]
            row['saving_total'] = (
                row['saving_parts'] + row['saving_labour'] + row['saving_paint']
            )
            row['savings_variance'] = row['file_saving_total'] - row['saving_total']
            rows.append(row)

        if not rows:
            continue

        period = months.most_common(1)[0][0] if months else _month_from_name(
            # localdate(), never date.today() — the server clock is UTC and
            # reads as yesterday between 00:00 and 02:00 in Gaborone.
            worksheet.title, default_year or localdate().year
        )
        computed = {
            part: sum((r[f'saving_{part}'] for r in rows), Decimal('0.00'))
            for part in ('parts', 'labour', 'paint', 'total')
        }
        sheets.append({
            'sheet':         worksheet.title,
            'period':        period,
            'rows':          rows,
            'stated':        stated or None,
            'computed':      computed,
            'variance_rows': sum(1 for r in rows if abs(r['savings_variance']) > Decimal('0.01')),
        })

    workbook.close()
    return {'sheets': sheets}


# ---------------------------------------------------------------------------
# Workbook 2 — parts purchase summary
# ---------------------------------------------------------------------------

SKIP_SUPPLIER_ROWS = {'total', 'suppliers', 'months', 'glass', 'grand total',
                      'parts total', 'contract pricing'}


def _month_header(worksheet, default_year: int) -> dict[int, date]:
    """Column index (1-based) → month start, read from rows 1-2."""
    names = [c for c in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), [])]
    years = [c for c in next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True), [])]
    header: dict[int, date] = {}
    rolled_year = None
    for idx, name in enumerate(names[1:], start=2):
        if not name:
            continue
        year_cell = years[idx - 1] if len(years) >= idx else None
        year = int(year_cell) if isinstance(year_cell, (int, float)) and year_cell > 1900 else default_year
        parsed = _month_from_name(str(name), year)
        if not parsed:
            continue
        # Jul→Jun sheets roll into the next calendar year at January.
        if rolled_year is not None and parsed.month < FY_START_MONTH and not (
            isinstance(year_cell, (int, float)) and year_cell > 1900
        ):
            parsed = date(rolled_year + 1, parsed.month, 1)
        else:
            rolled_year = parsed.year
        header[idx] = parsed
    return header


def parse_parts_summary(data: bytes, default_year: int | None = None) -> dict:
    """Read a Parts Summary workbook.

    Returns
        {'spend': [...], 'contract_pricing': [...], 'history': [...],
         'stated': {...}, 'recon': [...]}

    `spend` comes only from the three category sheets — the row-level truth.
    SUMMARY contributes the contract-pricing line, the five-FY history, and
    its stated totals, which are reconciled against the category sheets.
    """
    default_year = default_year or localdate().year
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)

    dealership_header: dict[int, date] = {}
    if 'DEALERSHIP' in workbook.sheetnames:
        dealership_header = _month_header(workbook['DEALERSHIP'], default_year)

    spend: list[dict] = []
    for sheet_name, category in CATEGORY_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        header = _month_header(worksheet, default_year) or dealership_header
        if not header:
            continue
        for raw in worksheet.iter_rows(min_row=1, values_only=True):
            supplier = str(raw[0] or '').strip()
            if not supplier or supplier.lower().strip() in SKIP_SUPPLIER_ROWS:
                continue
            for col, month in header.items():
                if len(raw) < col:
                    continue
                amount = money(raw[col - 1])
                if amount == 0:
                    continue
                spend.append({
                    'category': category,
                    'supplier': supplier,
                    'month':    month,
                    'amount':   amount,
                })

    contract_pricing: list[dict] = []
    history: list[dict] = []
    stated: dict[str, dict[str, Decimal]] = {}

    if 'SUMMARY' in workbook.sheetnames:
        summary = workbook['SUMMARY']
        summary_header = _month_header(summary, default_year) or dealership_header
        grid = [list(r) for r in summary.iter_rows(values_only=True)]

        stated_labels = {
            'parts total': 'parts_total',
            'grand total': 'grand_total',
        }
        seen_totals = 0
        fy_columns: dict[int, str] = {}

        for raw in grid:
            first = str(raw[0] or '').strip()
            lowered = first.lower()

            if lowered == 'contract pricing':
                for col, month in summary_header.items():
                    if len(raw) >= col:
                        amount = money(raw[col - 1])
                        if amount:
                            contract_pricing.append({'month': month, 'amount': amount})
                continue

            if lowered in stated_labels:
                stated[stated_labels[lowered]] = {
                    month.isoformat(): money(raw[col - 1])
                    for col, month in summary_header.items() if len(raw) >= col
                }
                continue

            if lowered == 'total':
                # First TOTAL = dealership + aftermarket block, second = glass.
                key = 'summary_supplier_total' if seen_totals == 0 else 'summary_glass_total'
                seen_totals += 1
                if seen_totals <= 2:
                    stated[key] = {
                        month.isoformat(): money(raw[col - 1])
                        for col, month in summary_header.items() if len(raw) >= col
                    }
                continue

            # Five-FY history block: a month name in column A and FY headers.
            if re.match(r'^fy[-\s]?\d{2}$', lowered.replace('_', '-')):
                continue
            if lowered == 'months' or lowered.startswith('months '):
                fy_columns = {
                    idx: str(cell).strip().upper().replace(' ', '-')
                    for idx, cell in enumerate(raw[1:], start=2)
                    if cell and re.match(r'^fy[-\s]?\d{2}$', str(cell).strip().lower().replace('_', '-'))
                }
                continue
            if fy_columns and lowered in MONTHS:
                month_number = MONTHS[lowered]
                for col, fy in fy_columns.items():
                    if len(raw) < col:
                        continue
                    amount = money(raw[col - 1])
                    if amount == 0:
                        continue
                    history.append({
                        'fy_label':     fy,
                        'month_number': month_number,
                        'amount':       amount,
                    })

    workbook.close()

    # Reconcile SUMMARY's stated supplier total against the category sheets.
    recon: list[dict] = []
    stated_supplier = stated.get('summary_supplier_total') or {}
    stated_glass = stated.get('summary_glass_total') or {}
    for month_iso, stated_amount in sorted(stated_supplier.items()):
        computed = sum(
            (r['amount'] for r in spend
             if r['month'].isoformat() == month_iso and r['category'] != 'WINDSCREEN'),
            Decimal('0.00'),
        )
        recon.append({
            'month':      month_iso,
            'block':      'Dealership + Aftermarket',
            'stated':     stated_amount,
            'computed':   computed,
            'difference': stated_amount - computed,
        })
    for month_iso, stated_amount in sorted(stated_glass.items()):
        computed = sum(
            (r['amount'] for r in spend
             if r['month'].isoformat() == month_iso and r['category'] == 'WINDSCREEN'),
            Decimal('0.00'),
        )
        recon.append({
            'month':      month_iso,
            'block':      'Windscreen / glass',
            'stated':     stated_amount,
            'computed':   computed,
            'difference': stated_amount - computed,
        })

    return {
        'spend':            spend,
        'contract_pricing': contract_pricing,
        'history':          history,
        'stated':           stated,
        'recon':            [r for r in recon if r['stated'] or r['computed']],
    }


def detect_workbook_kind(data: bytes) -> str:
    """'ASSESSMENT' | 'PARTS' — decided on sheet names, not the file name."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    names = {n.upper() for n in workbook.sheetnames}
    workbook.close()
    if names & set(CATEGORY_SHEETS) or 'SUMMARY' in names:
        return 'PARTS'
    if any('ASSESSMENT' in n or 'SAVING' in n for n in names):
        return 'ASSESSMENT'
    raise ValueError(
        'Unrecognised workbook. Expected the Assessment Savings report '
        '(a sheet per month) or the Parts Summary '
        '(SUMMARY / DEALERSHIP / WINDSCREEN / AFTERMARKET).'
    )
