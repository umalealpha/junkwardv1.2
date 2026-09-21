"""Excel builders for the Aware canned reports.

The weekly job (management command `generate_aware_reports`, Sunday cron) and
the on-demand path both call these. They reuse the SAME report functions that
feed the on-screen mode, so the Excel figures are identical to what the user
sees — that is the whole point: a downloadable file to defend the numbers.

Brand: Alpha Navy #1D3270 header, Direct Orange #F47C20 accent.
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, Sequence, Tuple

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# report key -> human label (shared by the command + the download endpoint)
KEY_LABELS = {
    'broker-analysis': 'Broker Analysis',
    'top-50-dom': 'Top 50 Domestic',
    'claims-registry': 'Claims Registry',
}


def reports_dir() -> Path:
    d = Path(settings.MEDIA_ROOT) / 'aware_reports'
    d.mkdir(parents=True, exist_ok=True)
    return d


NAVY = '1D3270'
ORANGE = 'F47C20'
WHITE = 'FFFFFF'
GREY = 'F2F4F8'


def _wb(title: str, subtitle: str, kpis: Sequence[Tuple[str, str]],
        columns: Sequence[str], rows: Sequence[Sequence[Any]],
        notes: Sequence[str], money_cols: Sequence[int] = (),
        money_format: str = '#,##0') -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Report'
    ncol = max(len(columns), 2)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, title)
    c.font = Font(bold=True, size=15, color=NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c = ws.cell(2, 1, subtitle)
    c.font = Font(size=10, color='555555')

    r = 4
    for label, value in kpis:
        ws.cell(r, 1, label).font = Font(bold=True, color='555555')
        ws.cell(r, 2, value).font = Font(bold=True, color=NAVY)
        r += 1
    r += 1

    hdr = r
    for j, name in enumerate(columns, 1):
        cell = ws.cell(hdr, j, name)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r = hdr + 1
    for i, row in enumerate(rows):
        for j, val in enumerate(row, 1):
            cell = ws.cell(r, j, val)
            if (j - 1) in money_cols and isinstance(val, (int, float)):
                # Default unchanged ('#,##0'). A caller whose report turns on
                # the cents - the claims payment movement report's excl-VAT and
                # VAT columns have to add back to the gross to the cent - asks
                # for '#,##0.00' instead.
                cell.number_format = money_format
            if i % 2:
                cell.fill = PatternFill('solid', fgColor=GREY)
        r += 1

    r += 1
    for n in notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        cell = ws.cell(r, 1, f'• {n}')
        cell.font = Font(size=9, italic=True, color='666666')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1

    widths = [max(len(str(columns[j])) if j < len(columns) else 10,
                  *[len(str(row[j])) for row in rows if j < len(row)] or [0], 10) + 3
              for j in range(ncol)]
    widths[0] = max(widths[0], 30)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(w, 48)
    ws.freeze_panes = ws.cell(hdr + 1, 1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _P(n) -> float:
    try:
        return round(float(n or 0))
    except (TypeError, ValueError):
        return 0


def broker_xlsx(d: Dict[str, Any], stamp: str) -> io.BytesIO:
    t = d.get('totals', {})
    ch = d.get('channel', {})
    dirc, brk = ch.get('direct', {}), ch.get('broker', {})
    kpis = [
        ('In-force book (annualised)', f"P {_P(t.get('inforce_gwp')):,.0f}"),
        ('Active policies', f"{t.get('active_pol', 0):,}"),
        ('External brokers', f"P {_P(brk.get('gwp')):,.0f}  ({brk.get('pct')}%)"),
        ('Direct & retail (in-house)', f"P {_P(dirc.get('gwp')):,.0f}  ({dirc.get('pct')}%)"),
        ('Claims paid (12m)', f"P {_P(t.get('claims_paid_12m')):,.0f}"),
        ('Outstanding reserve', f"P {_P(t.get('outstanding')):,.0f}"),
    ]
    cols = ['Broker', 'Active policies', 'In-force GWP', '% of broker book',
            'New biz 12m (n)', 'New biz GWP', 'Claims paid 12m', 'Outstanding reserve']
    rows = [[b['broker'], b['active_pol'], _P(b['inforce_gwp']), b.get('pct'),
             b.get('nb_cnt'), _P(b.get('nb_gwp')), _P(b.get('net_paid_12m')),
             _P(b.get('outstanding'))] for b in d.get('brokers', [])]
    return _wb('Broker Analysis — production book by intermediary',
               f'Alpha Direct Insurance · snapshot {stamp} · BWP · read-only from Graphite',
               kpis, cols, rows, d.get('notes', []), money_cols=(2, 5, 6, 7))


def top_dom_xlsx(d: Dict[str, Any], stamp: str) -> io.BytesIO:
    kpis = [('Domestic policies shown', f"{d.get('count', 0)}"),
            ('Their annual premium', f"P {_P(d.get('top50_gwp')):,.0f}")]
    cols = ['#', 'Policy', 'Policyholder', 'Broker', 'Annual premium', 'Pays', 'Since']
    rows = [[b['rank'], b['policy'], b['customer'], b['broker'],
             _P(b['annual_premium']), b['freq'], b['since']] for b in d.get('rows', [])]
    return _wb('Top 50 Domestic Policies — by annual premium',
               f'Alpha Direct Insurance · snapshot {stamp} · BWP · individual names shown as initials',
               kpis, cols, rows, d.get('notes', []), money_cols=(4,))


def claims_xlsx(d: Dict[str, Any], stamp: str) -> io.BytesIO:
    by = d.get('by_status', {})
    kpis = [('Open & pending claims', f"{d.get('open_count', 0):,}"),
            ('Outstanding reserve', f"P {_P(d.get('total_outstanding')):,.0f}"),
            ('Paid to date (on open)', f"P {_P(d.get('total_paid')):,.0f}"),
            ('By status', ' · '.join(f'{k} {v}' for k, v in by.items()) or '—')]
    cols = ['Claim', 'Policy', 'Status', 'Reported', 'Broker', 'Paid', 'Outstanding reserve']
    rows = [[b['claim'], b['policy'], b['status'], b['reported'], b['broker'],
             _P(b['paid']), _P(b['outstanding'])] for b in d.get('rows', [])]
    return _wb('Claims Registry — open & pending',
               f'Alpha Direct Insurance · snapshot {stamp} · BWP · safe fields only, no personal data',
               kpis, cols, rows, d.get('notes', []), money_cols=(5, 6))
